"""Native creation layouts: no external assets, hidden content, or grouped objects."""

from __future__ import annotations

from collections.abc import Callable
from typing import Any, cast

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

TEMPLATES = {
    "classic": ("F8F9FC", "222737", "353C4F", "5D667C", "E9EDF5"),
    "midnight": ("172033", "FFFFFF", "E4EAF4", "7DD3FC", "24334D"),
    "paper": ("FAF6EE", "302D28", "49443C", "956337", "EEE5D6"),
}
LAYOUTS = {"points", "columns", "cards", "statement"}
_rgb = cast(Callable[[str], RGBColor], RGBColor.from_string)


def render_slide(
    slide: Any,
    *,
    deck_title: str,
    ordinal: int,
    heading: str,
    points: list[str],
    template: str,
    layout: str,
) -> None:
    background, title_color, body_color, accent_color, panel_color = TEMPLATES[template]
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(background)

    def text(
        name: str,
        value: str,
        x: float,
        y: float,
        w: float,
        h: float,
        size: int,
        color: str,
        *,
        bold: bool = False,
    ) -> None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.name = name
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.06)
        frame.margin_top = frame.margin_bottom = Inches(0.04)
        paragraph = frame.paragraphs[0]
        paragraph.text = value
        paragraph.font.name = "Arial"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = _rgb(color)

    text("Deck label", f"{deck_title}  /  {ordinal:02d}", 0.7, 0.4, 11.8, 0.35, 10, accent_color)
    # A native title placeholder keeps normal slide-title extraction intact.
    box = slide.shapes.title
    box.left, box.top = Inches(0.7), Inches(1)
    box.width, box.height = Inches(11.8), Inches(1.25)
    box.name = "Slide title"
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = heading
    paragraph.font.name = "Arial"
    paragraph.font.size = Pt(30 if len(heading) <= 70 else 24)
    paragraph.font.bold = True
    paragraph.font.color.rgb = _rgb(title_color)

    count = len(points)
    for index, point in enumerate(points):
        if layout == "columns":
            rows = max(1, (count + 1) // 2)
            column, row = divmod(index, rows)
            x, y, w, h = 0.8 + column * 6.0, 2.5 + row * (4.2 / rows), 5.6, 4.2 / rows
            size = 20 if len(point) < 100 else 16
        elif layout == "cards":
            columns = 2 if count <= 4 else 3
            rows = max(1, (count + columns - 1) // columns)
            row, column = divmod(index, columns)
            cell_width, cell_height = 11.8 / columns, 4.2 / rows
            x, y, w, h = (
                0.7 + column * cell_width,
                2.5 + row * cell_height,
                cell_width - 0.2,
                cell_height - 0.2,
            )
            panel = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
            )
            panel.name = f"Point {index + 1} panel"
            panel.fill.solid()
            panel.fill.fore_color.rgb = _rgb(panel_color)
            panel.line.fill.background()
            x, y, w, h = x + 0.18, y + 0.18, w - 0.36, h - 0.36
            size = 20 if count <= 4 and len(point) < 100 else 14
            if count > 4 and len(point) >= 100:
                size = 11
        elif layout == "statement":
            x, y, w, h = 1.0, 2.7, 11.2, 3.5
            size = 36 if len(point) < 100 else 26
        else:
            row_height = min(0.85, 4.4 / max(1, count))
            x, y, w, h = 0.8, 2.4 + index * row_height, 11.7, row_height
            size = 20 if count <= 5 else 16
        text(f"Point {index + 1}", point, x, y, w, h, size, body_color)
