"""Structured native diagrams with connected shapes and verified semantic editing."""

from __future__ import annotations

import copy
import hashlib
import json
import math
import re
import shutil
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast

from lxml import etree  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from presenter_core.errors import CoreDomainError, invalid_request, reject_unknown_fields
from presenter_core.presentation.fidelity import (
    _NS,
    DRAWING_NS,
    OBJECT_LIMIT,
    FidelityGeometry,
    _parse_xml_part,
    inspect_pptx_snapshot,
)
from presenter_core.presentation.group_geometry import read_geometry

if TYPE_CHECKING:
    from presenter_core.presentation.editing import PresentationEditService
    from presenter_core.presentation.workspace_source import PresentationWorkspaceSourceService

PREFIX = "Presenter Copilot diagram v1:"
PALETTES = {
    "classic": ("EAF0FA", "1B2A41", "5070A5"),
    "midnight": ("23375B", "FFFFFF", "7193CE"),
    "paper": ("FFF4D9", "3B2F1E", "B98B39"),
}


def validate_data(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise invalid_request("Provide a structured diagram.")
    reject_unknown_fields(value, {"name", "kind", "palette", "nodes"})
    name, kind, palette, nodes = (value.get(k) for k in ("name", "kind", "palette", "nodes"))
    if not isinstance(name, str) or not 1 <= len(name.strip()) <= 120:
        raise invalid_request("Use a diagram name of 1–120 characters.")
    if (
        not isinstance(kind, str)
        or not isinstance(palette, str)
        or kind not in ("process", "hierarchy", "cycle")
        or palette not in PALETTES
    ):
        raise invalid_request("Choose process, hierarchy or cycle and a supported palette.")
    if not isinstance(nodes, list) or not 2 <= len(nodes) <= 8:
        raise invalid_request("A diagram needs 2–8 nodes.")
    known: dict[str, int] = {}
    clean = []
    for index, node in enumerate(nodes):
        if not isinstance(node, dict):
            raise invalid_request("Each diagram node needs an identity and label.")
        reject_unknown_fields(node, {"id", "label", "parent"})
        identity, label, parent = (node.get(k) for k in ("id", "label", "parent"))
        if (
            not isinstance(identity, str)
            or not re.fullmatch(r"[A-Za-z0-9_-]{1,40}", identity)
            or identity in known
            or not isinstance(label, str)
            or not 1 <= len(label.strip()) <= 80
            or len(label.split("\n")) > 3
        ):
            raise invalid_request("Use distinct node identities and labels of 1–80 characters.")
        if any(
            not (
                ord(c) in (9, 10)
                or 0x20 <= ord(c) <= 0xD7FF
                or 0xE000 <= ord(c) <= 0xFFFD
                or 0x10000 <= ord(c) <= 0x10FFFF
            )
            for c in label + name
        ):
            raise invalid_request("Diagram text contains unsupported characters.")
        if kind == "hierarchy":
            if index == 0:
                if parent is not None:
                    raise invalid_request("The first hierarchy node is its root.")
                depth = 0
            else:
                if not isinstance(parent, str) or parent not in known:
                    raise invalid_request("Choose an earlier node as each hierarchy parent.")
                depth = known[parent] + 1
                if depth > 3:
                    raise invalid_request("Use at most four hierarchy levels.")
        else:
            if parent is not None:
                raise invalid_request("Only hierarchy nodes use parents.")
            depth = 0
        known[identity] = depth
        clean.append({"id": identity, "label": label.strip(), "parent": parent})
    return {"name": name.strip(), "kind": kind, "palette": palette, "nodes": clean}


def layout(
    data: dict[str, Any],
) -> tuple[list[tuple[float, float, float, float]], list[tuple[str, str]]]:
    nodes, kind = data["nodes"], data["kind"]
    count = len(nodes)
    if kind == "hierarchy":
        depth: dict[str, int] = {}
        for node in nodes:
            depth[node["id"]] = depth[node["parent"]] + 1 if node["parent"] else 0
        levels = max(depth.values()) + 1
        boxes = []
        for node in nodes:
            peers = [v for v in nodes if depth[v["id"]] == depth[node["id"]]]
            width, height = min(0.32, 0.9 / len(peers)), 0.7 / levels
            center = (peers.index(node) + 0.5) / len(peers)
            boxes.append((center - width / 2, (depth[node["id"]] + 0.15) / levels, width, height))
        edges = [(node["parent"], node["id"]) for node in nodes if node["parent"]]
    elif kind == "cycle":
        boxes = [
            (
                0.39 + 0.34 * math.cos(-math.pi / 2 + i * 2 * math.pi / count),
                0.42 + 0.33 * math.sin(-math.pi / 2 + i * 2 * math.pi / count),
                0.22,
                0.16,
            )
            for i in range(count)
        ]
        edges = [(node["id"], nodes[(i + 1) % count]["id"]) for i, node in enumerate(nodes)]
    else:
        columns = min(4, count)
        rows = math.ceil(count / columns)
        width, height = 0.76 / columns, 0.28 if rows == 1 else 0.24
        boxes = []
        for i in range(count):
            row, column = divmod(i, columns)
            if row == 1:
                column = columns - 1 - column
            boxes.append(
                (
                    (column + 0.5) / columns - width / 2,
                    0.36 if rows == 1 else 0.14 + row * 0.48,
                    width,
                    height,
                )
            )
        edges = [(nodes[i]["id"], nodes[i + 1]["id"]) for i in range(count - 1)]
    return boxes, edges


def keys(data: dict[str, Any]) -> list[str]:
    return [
        "group",
        *("node:" + v["id"] for v in data["nodes"]),
        *(f"edge:{a}>{b}" for a, b in layout(data)[1]),
    ]


def rgb(value: str) -> RGBColor:
    return cast(RGBColor, RGBColor.from_string(value))  # type: ignore[no-untyped-call]


def canonical(node: Any) -> bytes:
    return bytes(etree.tostring(node, method="c14n", exclusive=True))


def generated(data: dict[str, Any], geometry: FidelityGeometry, native: dict[str, str]) -> Any:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    identity = group._element.find("p:nvGrpSpPr/p:cNvPr", _NS)
    identity.set("id", native["group"])
    identity.set("name", f"{data['kind'].title()} diagram · {data['name']}")
    identity.set(
        "descr",
        PREFIX
        + json.dumps({"data": data, "native": native}, sort_keys=True, separators=(",", ":")),
    )
    fill, text, edge = PALETTES[data["palette"]]
    boxes, edges = layout(data)
    shapes = {}
    for node, box in zip(data["nodes"], boxes, strict=True):
        x, y, w, h = [
            round(v * scale)
            for v, scale in zip(box, [geometry.width_emu, geometry.height_emu] * 2, strict=True)
        ]
        shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        props = shape._element.find("p:nvSpPr/p:cNvPr", _NS)
        props.set("id", native["node:" + node["id"]])
        props.set("name", "Diagram node · " + node["label"])
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.line.color.rgb = rgb(edge)
        shape.line.width = Pt(1.5)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = Pt(6)
        frame.margin_top = frame.margin_bottom = Pt(4)
        frame.text = node["label"]
        for paragraph in frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(max(12, min(20, int(min(w / 12700 / 7, h / 12700 / 3)))))
            paragraph.font.color.rgb = rgb(text)
            paragraph.font.bold = True
        shapes[node["id"]] = shape
    for source, target in edges:
        a, b = shapes[source], shapes[target]
        if data["kind"] == "hierarchy":
            begin, end = 2, 0
        elif b.left >= a.left + a.width:
            begin, end = 3, 1
        elif a.left >= b.left + b.width:
            begin, end = 1, 3
        elif b.top >= a.top + a.height:
            begin, end = 2, 0
        else:
            begin, end = 0, 2
        if data["kind"] == "cycle" and len(data["nodes"]) == 2:
            begin = end = 3 if source == data["nodes"][0]["id"] else 1
        connector = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 1, 1)
        connector.begin_connect(a, begin)
        connector.end_connect(b, end)
        props = connector._element.find("p:nvCxnSpPr/p:cNvPr", _NS)
        props.set("id", native[f"edge:{source}>{target}"])
        props.set("name", f"Diagram connection · {source} to {target}")
        connector.line.color.rgb = rgb(edge)
        connector.line.width = Pt(2)
        line = connector._element.find("p:spPr/a:ln", _NS)
        etree.SubElement(line, f"{{{DRAWING_NS}}}tailEnd", type="triangle", w="med", len="med")
    transform = group._element.find("p:grpSpPr/a:xfrm", _NS)
    for tag, values in [
        ("off", {"x": geometry.x_emu, "y": geometry.y_emu}),
        ("ext", {"cx": geometry.width_emu, "cy": geometry.height_emu}),
        ("chOff", {"x": 0, "y": 0}),
        ("chExt", {"cx": geometry.width_emu, "cy": geometry.height_emu}),
    ]:
        child = transform.find("a:" + tag, _NS)
        for key, value in values.items():
            child.set(key, str(value))
    return etree.fromstring(etree.tostring(group._element))


def inspect_diagram(group: Any) -> dict[str, Any] | None:
    identity = group.find("p:nvGrpSpPr/p:cNvPr", _NS)
    description = identity.get("descr", "") if identity is not None else ""
    if not description.startswith(PREFIX):
        return None
    info: dict[str, Any] = {
        "editable": False,
        "reason": "This diagram has manual or unsupported native changes.",
        "fingerprint": hashlib.sha256(canonical(group)).hexdigest(),
    }
    try:
        if len(description) > 12000:
            return info
        metadata = json.loads(description[len(PREFIX) :])
        if set(metadata) != {"data", "native"}:
            return info
        data = validate_data(metadata["data"])
        native = metadata["native"]
        if (
            not isinstance(native, dict)
            or set(native) != set(keys(data))
            or any(
                not isinstance(v, str)
                or not re.fullmatch(r"[1-9][0-9]{0,9}", v)
                or int(v) > 2_147_483_647
                for v in native.values()
            )
            or len(set(native.values())) != len(native)
        ):
            return info
        geometry, _ = read_geometry(group)
        if geometry is None or geometry.width_emu <= 0 or geometry.height_emu <= 0:
            return info
        if canonical(generated(data, geometry, native)) != canonical(group):
            return info
        return {
            **info,
            "editable": True,
            "reason": None,
            "data": data,
            "geometry": {
                "x_emu": geometry.x_emu,
                "y_emu": geometry.y_emu,
                "width_emu": geometry.width_emu,
                "height_emu": geometry.height_emu,
            },
            "_native": native,
        }
    except (ValueError, TypeError, KeyError, CoreDomainError):
        return info


def external_dependencies(root: Any, group: Any) -> bool:
    ids = set(group.xpath(".//p:cNvPr/@id", namespaces=_NS))
    outside = copy.deepcopy(root)
    native_id = group.find("p:nvGrpSpPr/p:cNvPr", _NS).get("id")
    matches = outside.xpath(
        "./p:cSld/p:spTree/p:grpSp[p:nvGrpSpPr/p:cNvPr/@id=$id]", namespaces=_NS, id=native_id
    )
    if len(matches) != 1:
        return True
    victim = matches[0]
    victim.getparent().remove(victim)
    return bool(
        ids.intersection(
            outside.xpath("./p:timing//@spid | .//a:stCxn/@id | .//a:endCxn/@id", namespaces=_NS)
        )
    )


class PresentationDiagramService:
    def __init__(self, sources: PresentationWorkspaceSourceService) -> None:
        self._sources = sources

    def list_diagrams(self, params: dict[str, Any]) -> dict[str, Any]:
        from presenter_core.presentation.editing import _verified_source_archive
        from presenter_core.presentation.semantics import PresentationSemanticService

        reject_unknown_fields(params, {"project_id", "document_id", "revision_id", "slide_number"})
        project, document, revision, native = self._sources.read_revision_inventory(
            {k: v for k, v in params.items() if k != "slide_number"}
        )
        ordinal = params.get("slide_number")
        if type(ordinal) is not int or not 1 <= ordinal <= len(native.slides):
            raise invalid_request("Choose a slide number.")
        slide = native.slides[ordinal - 1]
        with self._sources._storage.project_database(project) as conn:
            row = conn.execute(
                "SELECT local_snapshot_path, sha256, byte_size "
                "FROM presentation_revisions WHERE id=?",
                (revision,),
            ).fetchone()
        path = self._sources._storage.paths.snapshot_path(
            project, row["local_snapshot_path"], require_exists=True
        )
        items = []
        with _verified_source_archive(
            path, expected_sha256=row["sha256"], expected_byte_size=row["byte_size"]
        ) as archive:
            root = _parse_xml_part(archive, archive.getinfo(slide.part_name))
            for group in root.findall("p:cSld/p:spTree/p:grpSp", _NS):
                info = inspect_diagram(group)
                if info is not None:
                    if external_dependencies(root, group):
                        info = {
                            **info,
                            "editable": False,
                            "reason": "External animation or connector dependencies are preserved.",
                        }
                    identity = group.find("p:nvGrpSpPr/p:cNvPr", _NS)
                    obj = next(v for v in slide.objects if v.ooxml_id == identity.get("id"))
                    items.append(
                        {
                            "object_id": obj.object_id,
                            "name": obj.name,
                            **{k: v for k, v in info.items() if not k.startswith("_")},
                        }
                    )
                    if len(items) == 20:
                        break
        with self._sources._storage.project_database(project) as conn:
            PresentationSemanticService._require_current_presentation(
                conn, project, document, revision
            )
        return {
            "revision_id": revision,
            "slide_id": slide.slide_id,
            "slide_width_emu": native.slide_width_emu,
            "slide_height_emu": native.slide_height_emu,
            "items": items,
        }


def edit_diagram(
    editor: PresentationEditService, params: dict[str, Any], *, create: bool, delete: bool = False
) -> dict[str, Any]:
    from presenter_core.presentation.editing import _next_ooxml_shape_id, _verified_source_archive

    reject_unknown_fields(
        params,
        {"project_id", "document_id", "expected_revision_id", "slide_id"}
        | (set() if delete else {"data", "geometry"})
        | (set() if create else {"object_id", "expected_fingerprint"}),
    )
    project, document = (
        editor._project_id(params.get("project_id")),
        editor._document_id(params.get("document_id")),
    )
    revision = editor._uuid_field(params.get("expected_revision_id"), "expected_revision_id")
    slide_id = editor._uuid_field(params.get("slide_id"), "slide_id")
    data = None if delete else validate_data(params.get("data"))
    geometry = None if delete else editor._geometry_field(params.get("geometry"), "geometry")
    identity = editor._uuid_field(params.get("object_id"), "object_id") if not create else None
    fingerprint = params.get("expected_fingerprint")
    if not create and (
        not isinstance(fingerprint, str) or not re.fullmatch(r"[0-9a-f]{64}", fingerprint)
    ):
        raise invalid_request("Provide the verified diagram fingerprint.")
    source = editor._current_revision_source(project, document)
    if source["revision_id"] != revision:
        raise editor._revision_conflict()
    created: list[str] = []

    def write(original: Path, target: Path) -> bool:
        with original.open("rb") as handle:
            before = inspect_pptx_snapshot(
                handle,
                document_id=document,
                expected_sha256=source["sha256"],
                byte_size=source["byte_size"],
            )
        slide = next((v for v in before.slides if v.slide_id == slide_id), None)
        if slide is None:
            raise CoreDomainError("PRESENTATION_EDIT_TARGET_NOT_FOUND", "The slide is unavailable.")
        if geometry is not None:
            editor._validate_geometry_bounds(
                geometry,
                slide_width_emu=before.slide_width_emu,
                slide_height_emu=before.slide_height_emu,
                field="geometry",
            )
        with _verified_source_archive(
            original, expected_sha256=source["sha256"], expected_byte_size=source["byte_size"]
        ) as archive:
            root = _parse_xml_part(archive, archive.getinfo(slide.part_name))
            tree = root.find("p:cSld/p:spTree", _NS)
            old = None
            old_native: dict[str, str] = {}
            if not create:
                obj = next((v for v in slide.objects if v.object_id == identity), None)
                if obj is None or obj.kind != "grpSp":
                    raise CoreDomainError(
                        "PRESENTATION_EDIT_TARGET_NOT_FOUND", "The diagram is unavailable."
                    )
                old = editor._find_object_node(root, ooxml_id=obj.ooxml_id, kind="grpSp")
                info = inspect_diagram(old)
                if old.getparent() is not tree or not info or not info["editable"]:
                    raise CoreDomainError(
                        "PRESENTATION_EDIT_UNSUPPORTED",
                        "This diagram has native changes that must be preserved.",
                    )
                if info["fingerprint"] != fingerprint:
                    raise CoreDomainError(
                        "PRESENTATION_EDIT_STALE_TARGET",
                        "The diagram changed. Reload it before editing.",
                    )
                if external_dependencies(root, old):
                    raise CoreDomainError(
                        "PRESENTATION_EDIT_UNSUPPORTED",
                        "External animation or connector dependencies "
                        "prevent diagram regeneration.",
                    )
                old_native = info["_native"]
            if delete:
                assert old is not None
                tree.remove(old)
            else:
                assert data is not None and geometry is not None
                next_id = int(_next_ooxml_shape_id(root))
                native = {}
                for key in keys(data):
                    if key in old_native:
                        native[key] = old_native[key]
                    else:
                        if next_id > 2_147_483_647:
                            raise CoreDomainError(
                                "PRESENTATION_EDIT_RESOURCE_LIMIT",
                                "Native shape identities are exhausted.",
                            )
                        native[key] = str(next_id)
                        next_id += 1
                group = generated(data, geometry, native)
                if old is not None and canonical(old) == canonical(group):
                    return False
                removed = len(old_native)
                if before.object_count - removed + len(native) > OBJECT_LIMIT:
                    raise CoreDomainError(
                        "PRESENTATION_EDIT_RESOURCE_LIMIT",
                        "The diagram would exceed the object limit.",
                    )
                if old is not None:
                    tree.replace(old, group)
                else:
                    extensions = tree.findall("p:extLst", _NS)
                    if len(extensions) > 1 or (extensions and tree[-1] is not extensions[0]):
                        raise CoreDomainError(
                            "PRESENTATION_EDIT_UNSUPPORTED", "The shape tree is unsupported."
                        )
                    tree.insert(tree.index(extensions[0]) if extensions else len(tree), group)
            expected = canonical(root)
            payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
            if len(payload) > 8 * 1024 * 1024:
                raise CoreDomainError(
                    "PRESENTATION_EDIT_RESOURCE_LIMIT", "The diagram slide is too large."
                )
            with zipfile.ZipFile(target, "w") as output:
                for member in archive.infolist():
                    if member.filename == slide.part_name:
                        output.writestr(copy.copy(member), payload)
                    else:
                        with (
                            archive.open(member) as reader,
                            output.open(copy.copy(member), "w") as writer,
                        ):
                            shutil.copyfileobj(reader, writer, length=1024 * 1024)
        with target.open("rb") as handle:
            after = inspect_pptx_snapshot(
                handle,
                document_id=document,
                expected_sha256=editor._sha256_path(target),
                byte_size=target.stat().st_size,
            )
        if [v.slide_id for v in after.slides] != [v.slide_id for v in before.slides] or any(
            a != b
            for a, b in zip(before.slides, after.slides, strict=True)
            if a.slide_id != slide_id
        ):
            raise CoreDomainError(
                "PRESENTATION_EDIT_VALIDATION_FAILED", "Diagram editing changed another slide."
            )
        with zipfile.ZipFile(target) as archive:
            if canonical(_parse_xml_part(archive, archive.getinfo(slide.part_name))) != expected:
                raise CoreDomainError(
                    "PRESENTATION_EDIT_VALIDATION_FAILED", "The diagram slide did not validate."
                )
        updated = next(v for v in after.slides if v.slide_id == slide_id)
        if create:
            group_id = next(v.object_id for v in updated.objects if v.ooxml_id == native["group"])
            created.append(group_id)
        return True

    result = editor._commit_revision(project, document, revision, source, write)
    return {
        **result,
        "slide_id": slide_id,
        **({"created_object_id": created[0]} if create and created else {}),
    }
