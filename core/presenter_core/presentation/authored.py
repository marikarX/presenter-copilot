"""Fill a verified new slide with generated native content, preserving its deck."""

from __future__ import annotations

import copy
import hashlib
import io
import posixpath
import shutil
import uuid
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from presenter_core.errors import CoreDomainError
from presenter_core.ingestion.security import MAX_SOURCE_BYTES, preflight_source
from presenter_core.presentation.editing import (
    _NOTES_MASTER_REL_TYPE,
    _NS,
    _verified_source_archive,
)
from presenter_core.presentation.fidelity import (
    NOTES_SLIDE_REL_TYPE,
    REL_NS,
    _parse_xml_part,
    _relationships,
    _resolve_internal_part,
    inspect_pptx_snapshot,
)
from presenter_core.presentation.templates import render_slide

PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


def _xml(root: Any) -> bytes:
    return bytes(etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True))


def fill_authored_slide(
    blank_path: Path,
    destination: Path,
    *,
    document_id: str,
    slide_id: str,
    heading: str,
    points: list[str],
    notes: str,
    template: str,
    layout: str,
) -> None:
    # Only our own small generated package is opened by python-pptx. Imported
    # packages remain under the bounded archive/fidelity reader and byte copier.
    with blank_path.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    with blank_path.open("rb") as handle:
        baseline = inspect_pptx_snapshot(
            handle,
            document_id=document_id,
            expected_sha256=digest,
            byte_size=blank_path.stat().st_size,
        )
    target = next(slide for slide in baseline.slides if slide.slide_id == slide_id)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    render_slide(
        slide,
        deck_title="",
        ordinal=target.ordinal,
        heading=heading,
        points=points,
        template=template,
        layout=layout,
    )
    # Keep title semantics for ingestion/retrieval, with explicit local styling
    # and geometry rather than an inherited foreign title placeholder index.
    slide.shapes.title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    for placeholder in slide._element.xpath(".//p:ph"):
        placeholder.set("type", "title")
        placeholder.attrib.pop("idx", None)
    slide._element.set("showMasterSp", "0")
    sx, sy = baseline.slide_width_emu / Inches(13.333), baseline.slide_height_emu / Inches(7.5)
    for transform in slide._element.xpath(".//a:xfrm"):
        for node in transform:
            for attribute in ("x", "cx"):
                if node.get(attribute) is not None:
                    node.set(attribute, str(round(int(node.get(attribute)) * sx)))
            for attribute in ("y", "cy"):
                if node.get(attribute) is not None:
                    node.set(attribute, str(round(int(node.get(attribute)) * sy)))
    for node in slide._element.xpath(".//*[@sz]"):
        node.set("sz", str(max(100, round(int(node.get("sz")) * min(sx, sy)))))
    slide.notes_slide.notes_text_frame.text = notes
    generated = io.BytesIO()
    deck.save(generated)
    _merge_generated_notes(
        blank_path, destination, digest, target.part_name, generated, _xml(slide._element)
    )
    if destination.stat().st_size > MAX_SOURCE_BYTES:
        raise CoreDomainError(
            "PRESENTATION_EDIT_RESOURCE_LIMIT", "The authored slide exceeds the safe package size."
        )
    preflight_source(destination, "pptx")
    with destination.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    with destination.open("rb") as handle:
        result = inspect_pptx_snapshot(
            handle,
            document_id=document_id,
            expected_sha256=digest,
            byte_size=destination.stat().st_size,
        )
    authored = next(item for item in result.slides if item.slide_id == slide_id)
    actual_text = [
        item.text
        for item in authored.objects
        if item.name.startswith("Point ") and not item.name.endswith(" panel")
    ]
    if (
        authored.speaker_notes != notes
        or actual_text != points
        or not any(item.text == heading for item in authored.objects)
        or [item.slide_id for item in result.slides] != [item.slide_id for item in baseline.slides]
        or any(item.geometry is None or not item.geometry_editable for item in authored.objects)
    ):
        raise CoreDomainError(
            "PRESENTATION_EDIT_VALIDATION_FAILED", "The authored content did not pass validation."
        )
    for before, after in zip(baseline.slides, result.slides, strict=True):
        if before.slide_id != slide_id and before != after:
            raise CoreDomainError(
                "PRESENTATION_EDIT_VALIDATION_FAILED",
                "Authored insertion changed an existing slide.",
            )


def _merge_generated_notes(
    blank_path: Path,
    destination: Path,
    digest: str,
    target_part: str,
    generated: io.BytesIO,
    slide_xml: bytes,
) -> None:
    prefix = uuid.uuid4().hex
    notes_part = f"ppt/notesSlides/authored-{prefix}.xml"
    notes_rels = f"ppt/notesSlides/_rels/authored-{prefix}.xml.rels"
    master_part = f"ppt/notesMasters/authored-{prefix}.xml"
    master_rels = f"ppt/notesMasters/_rels/authored-{prefix}.xml.rels"
    theme_part = f"ppt/theme/authored-{prefix}.xml"

    def relationships(values: list[tuple[str, str, str]]) -> bytes:
        root = etree.Element(f"{{{PACKAGE_REL_NS}}}Relationships", nsmap={None: PACKAGE_REL_NS})
        for identifier, kind, path in values:
            node = etree.SubElement(root, f"{{{PACKAGE_REL_NS}}}Relationship")
            node.set("Id", identifier)
            node.set("Type", kind)
            node.set("Target", path)
        return _xml(root)

    with (
        zipfile.ZipFile(generated) as native,
        _verified_source_archive(
            blank_path, expected_sha256=digest, expected_byte_size=blank_path.stat().st_size
        ) as original,
    ):
        members = {info.filename: info for info in original.infolist()}
        presentation_rels = etree.fromstring(original.read("ppt/_rels/presentation.xml.rels"))
        presentation = etree.fromstring(original.read("ppt/presentation.xml"))
        registered = _relationships(presentation_rels)
        existing = [v for v in registered.values() if v[0] == _NOTES_MASTER_REL_TYPE]
        if len(existing) > 1 or (existing and existing[0][2].casefold() != "internal"):
            raise CoreDomainError(
                "PRESENTATION_EDIT_UNSUPPORTED", "The notes master is ambiguous or external."
            )
        fresh_master = not existing
        if existing:
            master_part = _resolve_internal_part("ppt", existing[0][1])
            if (
                master_part not in members
                or _parse_xml_part(original, members[master_part]).tag
                != f"{{{_NS['p']}}}notesMaster"
            ):
                raise CoreDomainError(
                    "PRESENTATION_EDIT_UNSUPPORTED", "The native notes master is unavailable."
                )
            master_list = presentation.find("p:notesMasterIdLst", _NS)
            if master_list is not None and (
                len(master_list) != 1
                or registered.get(master_list[0].get(f"{{{REL_NS}}}id", "")) != existing[0]
            ):
                raise CoreDomainError(
                    "PRESENTATION_EDIT_UNSUPPORTED",
                    "The notes master declaration does not match its relationship.",
                )
        additions = {
            notes_part: native.read("ppt/notesSlides/notesSlide1.xml"),
            notes_rels: relationships(
                [
                    (
                        "rId1",
                        _NOTES_MASTER_REL_TYPE,
                        posixpath.relpath(master_part, "ppt/notesSlides"),
                    ),
                    ("rId2", f"{REL_NS}/slide", posixpath.relpath(target_part, "ppt/notesSlides")),
                ]
            ),
        }
        if fresh_master:
            additions.update(
                {
                    master_part: native.read("ppt/notesMasters/notesMaster1.xml"),
                    master_rels: relationships(
                        [("rId1", f"{REL_NS}/theme", f"../theme/authored-{prefix}.xml")]
                    ),
                    theme_part: native.read("ppt/theme/theme2.xml"),
                }
            )
        if set(additions).intersection(members):
            raise CoreDomainError(
                "PRESENTATION_EDIT_VALIDATION_FAILED", "Generated package paths are already in use."
            )
        slide_rels_part = f"ppt/slides/_rels/{posixpath.basename(target_part)}.rels"
        rels = etree.fromstring(original.read(slide_rels_part))
        relation = etree.SubElement(rels, f"{{{PACKAGE_REL_NS}}}Relationship")
        relation.set("Id", "authoredNotes")
        relation.set("Type", NOTES_SLIDE_REL_TYPE)
        relation.set("Target", f"../notesSlides/authored-{prefix}.xml")
        if fresh_master:
            relation = etree.SubElement(presentation_rels, f"{{{PACKAGE_REL_NS}}}Relationship")
            master_id = f"authoredMaster{prefix}"
            relation.set("Id", master_id)
            relation.set("Type", _NOTES_MASTER_REL_TYPE)
            relation.set("Target", f"notesMasters/authored-{prefix}.xml")
            master_list = presentation.find("p:notesMasterIdLst", _NS)
            if master_list is None:
                master_list = etree.Element(f"{{{_NS['p']}}}notesMasterIdLst")
                # OOXML order: sldMasterIdLst, notesMasterIdLst, handoutMasterIdLst,
                # sldIdLst, sldSz, notesSz, followed by optional presentation data.
                position = next(
                    (
                        index
                        for index, node in enumerate(presentation)
                        if node.tag != f"{{{_NS['p']}}}sldMasterIdLst"
                    ),
                    len(presentation),
                )
                presentation.insert(position, master_list)
            master = etree.SubElement(master_list, f"{{{_NS['p']}}}notesMasterId")
            master.set(f"{{{REL_NS}}}id", master_id)
        types = etree.fromstring(original.read("[Content_Types].xml"))
        for part, kind in [
            (notes_part, "notesSlide"),
            *([(master_part, "notesMaster"), (theme_part, "theme")] if fresh_master else []),
        ]:
            override = etree.SubElement(types, f"{{{CONTENT_NS}}}Override")
            override.set("PartName", f"/{part}")
            override.set(
                "ContentType",
                "application/vnd.openxmlformats-officedocument."
                + ("theme+xml" if kind == "theme" else f"presentationml.{kind}+xml"),
            )
        replacements = {
            target_part: slide_xml,
            slide_rels_part: _xml(rels),
            **(
                {
                    "ppt/presentation.xml": _xml(presentation),
                    "ppt/_rels/presentation.xml.rels": _xml(presentation_rels),
                }
                if fresh_master
                else {}
            ),
            "[Content_Types].xml": _xml(types),
        }
        with zipfile.ZipFile(destination, "w") as output:
            output.comment = original.comment
            for info in original.infolist():
                replacement = replacements.get(info.filename)
                if replacement is not None:
                    output.writestr(copy.copy(info), replacement)
                else:
                    with (
                        original.open(info) as input_member,
                        output.open(copy.copy(info), "w") as output_member,
                    ):
                        shutil.copyfileobj(input_member, output_member, length=1024 * 1024)
            for part, payload in additions.items():
                output.writestr(part, payload)


def fill_native_notes(
    blank_path: Path, destination: Path, *, document_id: str, slide_id: str, notes: str
) -> None:
    """Attach notes only to the new native slide without rewriting its objects."""
    with blank_path.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    with blank_path.open("rb") as handle:
        before = inspect_pptx_snapshot(
            handle,
            document_id=document_id,
            expected_sha256=digest,
            byte_size=blank_path.stat().st_size,
        )
    target = next(slide for slide in before.slides if slide.slide_id == slide_id)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = notes
    generated = io.BytesIO()
    deck.save(generated)
    with _verified_source_archive(
        blank_path, expected_sha256=digest, expected_byte_size=blank_path.stat().st_size
    ) as archive:
        slide_xml = archive.read(target.part_name)
    _merge_generated_notes(blank_path, destination, digest, target.part_name, generated, slide_xml)
    preflight_source(destination, "pptx")
    with destination.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    with destination.open("rb") as handle:
        after = inspect_pptx_snapshot(
            handle,
            document_id=document_id,
            expected_sha256=digest,
            byte_size=destination.stat().st_size,
        )
    if [slide.slide_id for slide in before.slides] != [slide.slide_id for slide in after.slides]:
        raise CoreDomainError(
            "PRESENTATION_EDIT_VALIDATION_FAILED", "Native creation changed slide order."
        )
    for old, new in zip(before.slides, after.slides, strict=True):
        if old.slide_id == slide_id:
            if (
                old.objects != new.objects
                or new.speaker_notes != notes
                or not new.speaker_notes_editable
            ):
                raise CoreDomainError(
                    "PRESENTATION_EDIT_VALIDATION_FAILED", "Native slide notes did not validate."
                )
        elif old != new:
            raise CoreDomainError(
                "PRESENTATION_EDIT_VALIDATION_FAILED", "Native creation changed an existing slide."
            )
