"""Materialize a user-reviewed outline as native editable slides."""

from __future__ import annotations

import json
import re
import sqlite3
import tempfile
import uuid
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from presenter_core.errors import CoreDomainError, invalid_request, reject_unknown_fields
from presenter_core.ingestion.service import IngestionService
from presenter_core.presentation.templates import LAYOUTS, TEMPLATES, render_slide
from presenter_core.storage.paths import normalize_project_id
from presenter_core.storage.service import StorageManager


class PresentationCreationService:
    def __init__(self, storage: StorageManager, ingestion: IngestionService) -> None:
        self._storage = storage
        self._ingestion = ingestion

    def create_from_outline(self, params: dict[str, Any]) -> dict[str, Any]:
        reject_unknown_fields(
            params, {"project_id", "title", "slides", "provider_run_id", "template"}
        )
        project_value = params.get("project_id")
        if not isinstance(project_value, str):
            raise invalid_request("project_id must be a UUID.", field="project_id")
        project_id = normalize_project_id(project_value)
        title = self._text(params.get("title"), "title", 120)
        template = params.get("template", "classic")
        if not isinstance(template, str) or template not in TEMPLATES:
            raise invalid_request("Choose a supported presentation template.", field="template")
        raw_slides = params.get("slides")
        if not isinstance(raw_slides, list) or not 1 <= len(raw_slides) <= 40:
            raise invalid_request("Provide between 1 and 40 slides.", field="slides")
        slides: list[tuple[str, list[str], str, str]] = []
        citations: list[list[str]] = []
        for raw in raw_slides:
            if not isinstance(raw, dict):
                raise invalid_request("Each slide must be an object.", field="slides")
            reject_unknown_fields(raw, {"title", "bullets", "notes", "layout", "evidence_ids"})
            references = raw.get("evidence_ids", [])
            if (
                not isinstance(references, list)
                or len(references) > 8
                or any(
                    not isinstance(item, str) or not 1 <= len(item) <= 120 for item in references
                )
                or len(set(references)) != len(references)
            ):
                raise invalid_request("Choose up to eight distinct outline evidence references.")
            citations.append(references.copy())
            heading = self._text(raw.get("title"), "slide.title", 120)
            bullets = raw.get("bullets", [])
            if not isinstance(bullets, list) or len(bullets) > 8:
                raise invalid_request("Each slide accepts up to 8 points.", field="slide.bullets")
            points = [self._text(item, "slide.bullets", 200) for item in bullets]
            notes = self._text(raw.get("notes", ""), "slide.notes", 4_000, allow_empty=True)
            layout = raw.get("layout", "points")
            if not isinstance(layout, str) or layout not in LAYOUTS:
                raise invalid_request("Choose a supported slide layout.", field="slide.layout")
            if layout == "statement" and len(points) > 1:
                raise invalid_request(
                    "Statement layout accepts at most one point.", field="slide.layout"
                )
            slides.append((heading, points, notes, layout))
        # Establish project ownership before writing temporary content. Import then
        # owns snapshotting, original revision identity, chunks and index events.
        provider_run_id = params.get("provider_run_id")
        if any(citations) and provider_run_id is None:
            raise invalid_request(
                "Outline evidence references require their reviewed planning run."
            )
        with self._storage.project_database(project_id) as connection:
            if provider_run_id is not None:
                try:
                    provider_run_id = str(uuid.UUID(provider_run_id))
                except (ValueError, TypeError, AttributeError) as error:
                    raise invalid_request(
                        "provider_run_id must be a UUID.", field="provider_run_id"
                    ) from error
                if not connection.execute(
                    "SELECT 1 FROM provider_runs WHERE id = ? "
                    "AND task_type = 'presentation_outline' AND status = 'success'",
                    (provider_run_id,),
                ).fetchone():
                    raise invalid_request(
                        "The reviewed outline run is unavailable in this project.",
                        field="provider_run_id",
                    )
                if any(citations):
                    proposal = connection.execute(
                        "SELECT proposal_json FROM presentation_outline_proposals "
                        "WHERE provider_run_id=?",
                        (provider_run_id,),
                    ).fetchone()
                    if proposal is None:
                        raise invalid_request("The original outline evidence is unavailable.")
                    supplied = {item["evidence_id"] for item in json.loads(proposal[0])["evidence"]}
                    if any(item not in supplied for references in citations for item in references):
                        raise invalid_request(
                            "The outline cites evidence not supplied to its planning run."
                        )
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = title
        presentation.core_properties.subject = "User-reviewed presentation outline"
        if provider_run_id is not None:
            presentation.core_properties.comments = (
                f"Reviewed outline provider run: {provider_run_id}"
            )
        for index, (heading, points, notes, layout) in enumerate(slides, 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            render_slide(
                slide,
                deck_title=title,
                ordinal=index,
                heading=heading,
                points=points,
                template=template,
                layout=layout,
            )
            slide.notes_slide.notes_text_frame.text = notes
        filename = re.sub(r"[^A-Za-z0-9._ -]", "_", title).strip(" .")[:80] or "Presentation"
        if filename.split(".")[0].upper() in {
            "CON",
            "PRN",
            "AUX",
            "NUL",
            *[f"COM{i}" for i in range(1, 10)],
            *[f"LPT{i}" for i in range(1, 10)],
        }:
            filename = f"Deck-{filename}"
        with tempfile.TemporaryDirectory(prefix="presenter-outline-") as directory:
            path = Path(directory) / f"{filename}.pptx"
            presentation.save(str(path))
            imported = self._ingestion.import_source(
                {"project_id": project_id, "kind": "presentation", "path": str(path)},
                initialize_document=(
                    lambda connection, document: self._record_evidence_receipt(
                        connection, project_id, document, provider_run_id, slides, citations
                    )
                )
                if any(citations)
                else None,
            )
        return imported

    def _record_evidence_receipt(
        self,
        connection: sqlite3.Connection,
        project: str,
        document: str,
        run: str | None,
        slides: list[tuple[str, list[str], str, str]],
        citations: list[list[str]],
    ) -> None:
        """Persist unsaved citation candidates; ingestion owns the native revision.

        Runs inside ingestion's document/original-revision registration transaction.
        Failure rolls back those rows and removes the new owned source snapshot.
        """
        row = connection.execute(
            "SELECT d.metadata_json,r.id,r.sha256 FROM documents d "
            "JOIN presentation_revisions r ON r.presentation_id=d.id "
            "WHERE d.id=? AND d.project_id=? AND r.revision_number=1",
            (document, project),
        ).fetchone()
        if row is None:
            raise CoreDomainError(
                "PRESENTATION_OUTLINE_EVIDENCE_UNAVAILABLE",
                "The created deck's original revision is unavailable.",
            )
        receipt = {
            "version": 1,
            "provider_run_id": run,
            "original_revision_id": row["id"],
            "original_sha256": row["sha256"],
            "slides": [
                {
                    "slide_number": index + 1,
                    "points": slide[1],
                    "evidence_ids": citations[index],
                }
                for index, slide in enumerate(slides)
            ],
        }
        metadata = json.loads(row["metadata_json"])
        previous = metadata.get("outline_evidence_receipt")
        if previous is not None and previous != receipt:
            raise CoreDomainError(
                "PRESENTATION_OUTLINE_EVIDENCE_CONFLICT",
                "This deck already has a different outline evidence receipt.",
            )
        metadata["outline_evidence_receipt"] = receipt
        connection.execute(
            "UPDATE documents SET metadata_json=? WHERE id=?", (json.dumps(metadata), document)
        )

    @staticmethod
    def _text(value: Any, field: str, limit: int, *, allow_empty: bool = False) -> str:
        if not isinstance(value, str):
            raise invalid_request("Expected text.", field=field)
        text = value.strip()
        if (not text and not allow_empty) or len(text) > limit:
            raise invalid_request(f"Text must contain at most {limit} characters.", field=field)
        if any(
            not (
                code in (9, 10, 13)
                or 0x20 <= code <= 0xD7FF
                or 0xE000 <= code <= 0xFFFD
                or 0x10000 <= code <= 0x10FFFF
            )
            for code in map(ord, text)
        ):
            raise invalid_request("Text contains unsupported characters.", field=field)
        return text
