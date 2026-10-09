"""Native chart authoring: own generated chart/workbook parts in a verified package."""

from __future__ import annotations

import copy
import io
import math
import posixpath
import shutil
import uuid
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any

from lxml import etree  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.xlsx import Workbook  # type: ignore[attr-defined]
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt

from presenter_core.errors import CoreDomainError, invalid_request, reject_unknown_fields
from presenter_core.presentation.creation import PresentationCreationService
from presenter_core.presentation.fidelity import (
    _NS,
    DRAWING_NS,
    OBJECT_LIMIT,
    REL_NS,
    _parse_xml_part,
    _relationships,
    _resolve_internal_part,
    inspect_pptx_snapshot,
)

if TYPE_CHECKING:
    from presenter_core.presentation.editing import PresentationEditService

CNS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NS = {**_NS, "c": CNS}
KINDS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}
COLORS = ("2563EB", "D97706", "059669", "9333EA", "DC2626", "0891B2")


def validate_data(params: dict[str, Any]) -> dict[str, Any]:
    title = PresentationCreationService._text(params.get("title"), "title", 120)
    kind = params.get("chart_type")
    if not isinstance(kind, str) or kind not in KINDS:
        raise invalid_request("Choose column, bar, line or pie.")
    categories = params.get("categories")
    if not isinstance(categories, list) or not 1 <= len(categories) <= 24:
        raise invalid_request("Provide 1–24 categories.")
    categories = [PresentationCreationService._text(v, "categories", 80) for v in categories]
    series = params.get("series")
    if not isinstance(series, list) or not 1 <= len(series) <= 6:
        raise invalid_request("Provide 1–6 series.")
    values = []
    for row in series:
        if not isinstance(row, dict) or set(row) != {"name", "values"}:
            raise invalid_request("Each series needs a name and values.")
        name = PresentationCreationService._text(row["name"], "name", 80)
        numbers = row["values"]
        if (
            not isinstance(numbers, list)
            or len(numbers) != len(categories)
            or any(
                type(n) not in (int, float) or abs(n) > 1e12 or not math.isfinite(n)
                for n in numbers
            )
        ):
            raise invalid_request("Provide one finite number per category, within ±1 trillion.")
        values.append({"name": name, "values": numbers})
    if kind == "pie" and (
        len(values) != 1 or any(n < 0 for n in values[0]["values"]) or sum(values[0]["values"]) <= 0
    ):
        raise invalid_request("Pie charts need one nonnegative series with a positive total.")
    return {"title": title, "chart_type": kind, "categories": categories, "series": values}


def summary(data: dict[str, Any]) -> str:
    return (
        "Chart: "
        + str(data["title"])
        + "\n"
        + "\n".join(
            series["name"]
            + ": "
            + "; ".join(
                f"{category}: {int(number) if float(number).is_integer() else repr(float(number))}"
                for category, number in zip(data["categories"], series["values"], strict=True)
            )
            for series in data["series"]
        )
    )


def read_chart(root: Any) -> dict[str, Any] | None:
    """Read only bounded cached chart data; never open or resolve a workbook."""
    plots = root.findall("c:chart/c:plotArea", NS)
    if len(plots) != 1:
        return None
    kinds = [
        (node, name)
        for node in plots[0]
        for name in ("barChart", "lineChart", "pieChart")
        if node.tag == f"{{{CNS}}}{name}"
    ]
    if len(kinds) != 1 or any(
        isinstance(child.tag, str)
        and child.tag.startswith(f"{{{CNS}}}")
        and child.tag.endswith("Chart")
        and child is not kinds[0][0]
        for child in plots[0]
    ):
        return None
    node, name = kinds[0]
    if name == "barChart":
        direction = node.find("c:barDir", NS)
        grouping = node.find("c:grouping", NS)
        if (
            direction is None
            or direction.get("val") not in {"bar", "col"}
            or grouping is None
            or grouping.get("val") != "clustered"
        ):
            return None
        kind = "bar" if direction.get("val") == "bar" else "column"
    else:
        if name == "lineChart":
            grouping = node.find("c:grouping", NS)
            if grouping is None or grouping.get("val") != "standard":
                return None
        kind = "line" if name == "lineChart" else "pie"
    title = (
        "\n".join(
            "".join(node.text or "" for node in paragraph.findall(".//a:t", NS))
            for paragraph in root.findall("c:chart/c:title/c:tx/c:rich/a:p", NS)
        )
        or "Untitled chart"
    )
    series = []
    categories = None

    def cached(parent: Any, selector: str, numeric: bool = False) -> list[Any] | None:
        caches = parent.findall(selector, NS)
        if len(caches) != 1:
            return None
        pts = caches[0].findall("c:pt", NS)
        count = caches[0].find("c:ptCount", NS)
        if (
            count is None
            or count.get("val") != str(len(pts))
            or not 1 <= len(pts) <= 24
            or [p.get("idx") for p in pts] != [str(i) for i in range(len(pts))]
        ):
            return None
        result = []
        for p in pts:
            v = p.find("c:v", NS)
            if v is None:
                return None
            try:
                result.append(float(v.text or "") if numeric else v.text or "")
            except ValueError:
                return None
        return result

    nodes = node.findall("c:ser", NS)
    if not 1 <= len(nodes) <= 6:
        return None
    for part in nodes:
        labels = cached(part, "c:cat/c:strRef/c:strCache")
        numbers = cached(part, "c:val/c:numRef/c:numCache", True)
        names = part.findall("c:tx/c:strRef/c:strCache/c:pt/c:v", NS)
        if (
            labels is None
            or numbers is None
            or len(names) != 1
            or (categories is not None and categories != labels)
        ):
            return None
        categories = labels
        series.append({"name": names[0].text or "", "values": numbers})
    try:
        return validate_data(
            {"title": title, "chart_type": kind, "categories": categories, "series": series}
        )
    except CoreDomainError:
        return None


def chart_texts(
    archive: zipfile.ZipFile, members: dict[str, Any], slide_root: Any, slide_part: str
) -> dict[str, str]:
    """Verified native cache text for semantic/retrieval use, without external I/O."""
    from presenter_core.presentation.fidelity import _iter_slide_objects

    tree = slide_root.find("p:cSld/p:spTree", NS)
    frames = (
        [
            node
            for node, alternate in _iter_slide_objects(tree)
            if node.tag == f"{{{NS['p']}}}graphicFrame" and not alternate
        ]
        if tree is not None
        else []
    )
    refs = [(frame, frame.findall("a:graphic/a:graphicData/c:chart", NS)) for frame in frames]
    if not any(nodes for _, nodes in refs):
        return {}
    relpart = posixpath.join(
        posixpath.dirname(slide_part), "_rels", posixpath.basename(slide_part) + ".rels"
    )
    if relpart not in members:
        return {}
    relations = _relationships(_parse_xml_part(archive, members[relpart]))
    result = {}
    for frame, nodes in refs:
        if len(nodes) != 1:
            continue
        relation = relations.get(nodes[0].get(f"{{{REL_NS}}}id", ""))
        if not relation or relation[0] != f"{REL_NS}/chart" or relation[2].casefold() != "internal":
            continue
        part = _resolve_internal_part(posixpath.dirname(slide_part), relation[1])
        if part not in members:
            continue
        data = read_chart(_parse_xml_part(archive, members[part]))
        identity = frame.find("p:nvGraphicFramePr/p:cNvPr", NS)
        if data is not None and identity is not None:
            result[identity.get("id", "")] = summary(data)
    return result


def rgb(value: str) -> RGBColor:
    return RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16))  # type: ignore[no-untyped-call]


def generated_chart(data: dict[str, Any], geometry: Any) -> tuple[Any, bytes, bytes]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    values = CategoryChartData()  # type: ignore[no-untyped-call]
    values.categories = data["categories"]
    for row in data["series"]:
        values.add_series(row["name"], row["values"])  # type: ignore[no-untyped-call]
    frame = slide.shapes.add_chart(
        KINDS[data["chart_type"]],
        geometry.x_emu,
        geometry.y_emu,
        geometry.width_emu,
        geometry.height_emu,
        values,
    )
    # App-authored charts support native grouping; preserve locks on imported charts.
    locks = frame._element.find("p:nvGraphicFramePr/p:cNvGraphicFramePr/a:graphicFrameLocks", NS)
    if locks is not None:
        locks.set("noGrp", "0")
    chart = frame.chart
    chart.font.name = "Arial"
    chart.font.size = Pt(12)
    chart.font.color.rgb = rgb("182132")
    space = chart._chartSpace
    properties = etree.Element(f"{{{CNS}}}spPr")
    fill = etree.SubElement(properties, f"{{{DRAWING_NS}}}solidFill")
    etree.SubElement(fill, f"{{{DRAWING_NS}}}srgbClr", val="FFFFFF")
    space.insert(space.index(space.find("c:chart", NS)) + 1, properties)
    chart.has_title = True
    chart.chart_title.text_frame.text = data["title"]
    if data["chart_type"] == "pie":
        title = chart.chart_title._element
        layout = title.find("c:layout", NS)
        if layout is None:
            layout = etree.Element(f"{{{CNS}}}layout")
            title.insert(1, layout)
        manual = etree.SubElement(layout, f"{{{CNS}}}manualLayout")
        etree.SubElement(manual, f"{{{CNS}}}yMode", val="factor")
        etree.SubElement(manual, f"{{{CNS}}}y", val="0")
    for run in chart.chart_title.text_frame.paragraphs[0].runs:
        run.font.size = Pt(20)
        run.font.color.rgb = rgb("182132")
    chart.has_legend = len(data["series"]) > 1 or data["chart_type"] == "pie"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    for index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(COLORS[index])
        series.format.line.color.rgb = rgb(COLORS[index])
        if data["chart_type"] == "pie":
            for offset, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(COLORS[offset % len(COLORS)])
    workbook_bytes = literal_workbook(data)
    chart.part.chart_workbook.update_from_xlsx_blob(workbook_bytes)
    xml = etree.tostring(chart._chartSpace, encoding="UTF-8", xml_declaration=True, standalone=True)
    return copy.deepcopy(frame._element), xml, workbook_bytes


def literal_workbook(data: dict[str, Any]) -> bytes:
    workbook_stream = io.BytesIO()
    with Workbook(
        workbook_stream, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}
    ) as workbook:
        sheet = workbook.add_worksheet("Sheet1")
        for index, category in enumerate(data["categories"], 1):
            sheet.write_string(index, 0, category)
        for column, row in enumerate(data["series"], 1):
            sheet.write_string(0, column, row["name"])
            for index, number in enumerate(row["values"], 1):
                sheet.write_number(index, column, number)
    return workbook_stream.getvalue()


def create_chart(editor: PresentationEditService, params: dict[str, Any]) -> dict[str, Any]:
    from presenter_core.presentation.editing import _next_ooxml_shape_id, _verified_source_archive

    reject_unknown_fields(
        params,
        {
            "project_id",
            "document_id",
            "expected_revision_id",
            "slide_id",
            "geometry",
            "title",
            "chart_type",
            "categories",
            "series",
        },
    )
    project = editor._project_id(params.get("project_id"))
    document = editor._document_id(params.get("document_id"))
    revision = editor._uuid_field(params.get("expected_revision_id"), "expected_revision_id")
    slide_id = editor._uuid_field(params.get("slide_id"), "slide_id")
    data = validate_data(params)
    geometry = editor._geometry_field(params.get("geometry"), "geometry")
    source = editor._current_revision_source(project, document)
    if source["revision_id"] != revision:
        raise editor._revision_conflict()
    created: list[str] = []

    def write(original: Path, target: Path) -> None:
        with original.open("rb") as handle:
            before = inspect_pptx_snapshot(
                handle,
                document_id=document,
                expected_sha256=source["sha256"],
                byte_size=source["byte_size"],
            )
        slide = next((v for v in before.slides if v.slide_id == slide_id), None)
        if slide is None:
            raise CoreDomainError("PRESENTATION_EDIT_TARGET_NOT_FOUND", "The slide is unavailable.")
        if before.object_count >= OBJECT_LIMIT:
            raise CoreDomainError(
                "PRESENTATION_EDIT_RESOURCE_LIMIT", "The object limit has been reached."
            )
        editor._validate_geometry_bounds(
            geometry,
            slide_width_emu=before.slide_width_emu,
            slide_height_emu=before.slide_height_emu,
            field="geometry",
        )
        frame, chart, workbook = generated_chart(data, geometry)
        prefix = uuid.uuid4().hex
        chartpart = f"ppt/charts/copilot-{prefix}.xml"
        bookpart = f"ppt/embeddings/copilot-{prefix}.xlsx"
        chartrels = f"ppt/charts/_rels/copilot-{prefix}.xml.rels"
        relns = "http://schemas.openxmlformats.org/package/2006/relationships"
        ctns = "http://schemas.openxmlformats.org/package/2006/content-types"
        with _verified_source_archive(
            original, expected_sha256=source["sha256"], expected_byte_size=source["byte_size"]
        ) as archive:
            members = {v.filename: v for v in archive.infolist()}
            root = _parse_xml_part(archive, members[slide.part_name])
            tree = root.find("p:cSld/p:spTree", NS)
            shape_id = _next_ooxml_shape_id(tree)
            identity = frame.find("p:nvGraphicFramePr/p:cNvPr", NS)
            identity.set("id", shape_id)
            identity.set("name", f"Chart {shape_id}")
            identity.set("descr", data["title"])
            relpart = editor._relationship_part_name(slide.part_name)
            relations = (
                _parse_xml_part(archive, members[relpart])
                if relpart in members
                else etree.Element(f"{{{relns}}}Relationships", nsmap={None: relns})
            )
            rid = editor._new_relationship_id(set(_relationships(relations)), "Chart")
            frame.find("a:graphic/a:graphicData/c:chart", NS).set(f"{{{REL_NS}}}id", rid)
            etree.SubElement(
                relations,
                f"{{{relns}}}Relationship",
                Id=rid,
                Type=f"{REL_NS}/chart",
                Target=posixpath.relpath(chartpart, posixpath.dirname(slide.part_name)),
            )
            extensions = tree.findall("p:extLst", NS)
            if len(extensions) > 1 or (extensions and tree[-1] is not extensions[0]):
                raise CoreDomainError(
                    "PRESENTATION_EDIT_UNSUPPORTED", "The shape tree is unsupported."
                )
            tree.insert(tree.index(extensions[0]) if extensions else len(tree), frame)
            chartroot = etree.fromstring(chart)
            external = chartroot.find("c:externalData", NS)
            external.set(f"{{{REL_NS}}}id", "rId1")
            external.find("c:autoUpdate", NS).set("val", "0")
            relroot = etree.Element(f"{{{relns}}}Relationships", nsmap={None: relns})
            etree.SubElement(
                relroot,
                f"{{{relns}}}Relationship",
                Id="rId1",
                Type=f"{REL_NS}/package",
                Target=posixpath.relpath(bookpart, "ppt/charts"),
            )
            additions = {
                chartpart: etree.tostring(
                    chartroot, encoding="UTF-8", xml_declaration=True, standalone=True
                ),
                bookpart: workbook,
                chartrels: etree.tostring(
                    relroot, encoding="UTF-8", xml_declaration=True, standalone=True
                ),
            }
            if {v.casefold() for v in additions} & {v.casefold() for v in members}:
                raise CoreDomainError(
                    "PRESENTATION_EDIT_VALIDATION_FAILED", "Chart part paths already exist."
                )
            types = _parse_xml_part(archive, members["[Content_Types].xml"])
            for part, kind in [
                (chartpart, "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"),
                (bookpart, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ]:
                etree.SubElement(
                    types, f"{{{ctns}}}Override", PartName="/" + part, ContentType=kind
                )
            replacements = {
                slide.part_name: etree.tostring(
                    root, encoding="UTF-8", xml_declaration=True, standalone=True
                ),
                relpart: etree.tostring(
                    relations, encoding="UTF-8", xml_declaration=True, standalone=True
                ),
                "[Content_Types].xml": etree.tostring(
                    types, encoding="UTF-8", xml_declaration=True, standalone=True
                ),
            }
            with zipfile.ZipFile(target, "w") as output:
                output.comment = archive.comment
                for member in archive.infolist():
                    if member.filename in replacements:
                        output.writestr(copy.copy(member), replacements.pop(member.filename))
                    else:
                        with (
                            archive.open(member) as reader,
                            output.open(copy.copy(member), "w") as writer,
                        ):
                            shutil.copyfileobj(reader, writer, length=1024 * 1024)
                for name, payload in {**replacements, **additions}.items():
                    output.writestr(name, payload)
        with target.open("rb") as handle:
            after = inspect_pptx_snapshot(
                handle,
                document_id=document,
                expected_sha256=editor._sha256_path(target),
                byte_size=target.stat().st_size,
            )
        updated = next(v for v in after.slides if v.slide_id == slide_id)
        if (
            len(before.slides) != len(after.slides)
            or any(
                old != new
                for old, new in zip(before.slides, after.slides, strict=True)
                if old.slide_id != slide_id
            )
            or updated.objects[:-1] != slide.objects
            or len(updated.objects) != len(slide.objects) + 1
            or updated.objects[-1].text != summary(data)
            or updated.speaker_notes != slide.speaker_notes
            or updated.speaker_notes_editable != slide.speaker_notes_editable
        ):
            raise CoreDomainError(
                "PRESENTATION_EDIT_VALIDATION_FAILED", "Chart creation changed unrelated content."
            )
        with zipfile.ZipFile(target) as archive:
            actual = read_chart(etree.fromstring(archive.read(chartpart)))
            if actual != data or archive.read(bookpart) != workbook:
                raise CoreDomainError(
                    "PRESENTATION_EDIT_VALIDATION_FAILED", "Chart cache/workbook validation failed."
                )
        created.append(updated.objects[-1].object_id)

    result = editor._commit_revision(project, document, revision, source, write)
    return {**result, "created_object_id": created[0], "slide_id": slide_id}
