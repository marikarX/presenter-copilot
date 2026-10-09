"""PPTX parser preserving slide order, visible text, and notes."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation

from presenter_core.errors import CoreDomainError
from presenter_core.ingestion.models import ParsedSourceUnit

from .utils import normalize_text


class PptxParser:
    parser_id = "pptx.python-pptx"

    def parse(self, path: Path) -> list[ParsedSourceUnit]:
        try:
            presentation = Presentation(str(path))
            units: list[ParsedSourceUnit] = []
            for ordinal, slide in enumerate(presentation.slides, start=1):
                title = self._title(slide)
                visible_text = self._visible_text(slide)
                notes = self._notes(slide)
                metadata: dict[str, Any] = {"parser": self.parser_id}
                if notes:
                    metadata["notes"] = notes
                    metadata["notes_searchable"] = True
                search_text = self._search_text(visible_text, notes)
                units.append(
                    ParsedSourceUnit(
                        unit_type="slide",
                        ordinal=ordinal,
                        title=title,
                        text=visible_text,
                        metadata=metadata,
                        search_text=search_text,
                    )
                )
            return units
        except CoreDomainError:
            raise
        except Exception as exc:
            raise CoreDomainError(
                "SOURCE_PARSE_FAILED",
                "The PowerPoint presentation could not be parsed.",
                details={"parser_id": self.parser_id, "source_type": "pptx"},
            ) from exc

    @staticmethod
    def _title(slide: Any) -> str | None:
        try:
            title_shape = slide.shapes.title
            if title_shape is None or not title_shape.has_text_frame:
                return None
            title = normalize_text(title_shape.text)
            return title or None
        except (AttributeError, ValueError):
            return None

    @staticmethod
    def _visible_text(slide: Any) -> str:
        parts: list[str] = []
        pending = [(shape, 0) for shape in reversed(list(slide.shapes))]
        visited = 0
        while pending:
            shape, depth = pending.pop()
            visited += 1
            if visited > 100_000 or depth > 32:
                raise CoreDomainError(
                    "SOURCE_PARSE_FAILED",
                    "The slide exceeds supported group nesting or object limits.",
                )
            try:
                if hasattr(shape, "shapes"):
                    pending.extend((child, depth + 1) for child in reversed(list(shape.shapes)))
                    continue
                if shape.has_text_frame:
                    value = normalize_text(shape.text)
                    if value:
                        parts.append(value)
                elif shape.has_table:
                    rows = [
                        " | ".join(normalize_text(cell.text) for cell in row.cells)
                        for row in shape.table.rows
                    ]
                    parts.extend(row for row in rows if row.strip())
                elif shape.has_chart:
                    from presenter_core.presentation.charts import read_chart, summary

                    data = read_chart(shape.chart._chartSpace)
                    if data is not None:
                        parts.append(summary(data))
            except (AttributeError, ValueError, KeyError):
                continue
        return "\n".join(parts)

    @staticmethod
    def _notes(slide: Any) -> str:
        try:
            notes_slide = slide.notes_slide
            text_frame = notes_slide.notes_text_frame
            return normalize_text(text_frame.text) if text_frame else ""
        except (AttributeError, ValueError):
            return ""

    @staticmethod
    def _search_text(visible_text: str, notes: str) -> str:
        if not notes:
            return visible_text
        if not visible_text:
            return f"Speaker notes:\n{notes}"
        return f"{visible_text}\n\nSpeaker notes:\n{notes}"
