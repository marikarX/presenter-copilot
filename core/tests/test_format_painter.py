"""Format painter: copy one object's fill, outline, effects and text formatting onto others."""

import io
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _fixture(path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Agenda"
    source = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1)
    )
    source.name = "Painter"
    source.fill.solid()
    source.fill.fore_color.rgb = RGBColor(0xC0, 0x1F, 0x2F)
    source.line.color.rgb = RGBColor(0x1F, 0x3F, 0xC0)
    source.line.width = Pt(3)
    run = source.text_frame.paragraphs[0].add_run()
    run.text = "Source"
    run.font.bold = True
    run.font.size = Pt(24)
    run.font.name = "Georgia"
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    plain = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5), Inches(2), Inches(2), Inches(1))
    plain.name = "Plain"
    linked = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    linked.name = "Linked"
    first = linked.text_frame.paragraphs[0].add_run()
    first.text = "Read more"
    first.hyperlink.address = "https://example.org/more"
    image = io.BytesIO()
    Image.new("RGB", (40, 30), "#2a7").save(image, format="PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(6), Inches(4), Inches(2), Inches(1.5)).name = "Photo"
    deck.save(path)


def test_format_painter_copies_shape_and_text_formatting(tmp_path: Path) -> None:
    _fixture(tmp_path / "paint.pptx")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, tmp_path / "paint.pptx")["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def inventory() -> dict:
            return _call(
                core,
                "l",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]

        before = inventory()
        named = {o["name"]: o for o in before["objects"]}

        def paint(source: str, targets: list[str]) -> dict:
            return _call(
                core,
                "p",
                "presentation.workspace.copy_object_format",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": before["slide_id"],
                    "source_object_id": named[source]["object_id"],
                    "target_object_ids": [named[name]["object_id"] for name in targets],
                },
            )

        reply = paint("Painter", ["Plain", "Linked", "Photo", "Title 1"])
        assert reply["ok"], reply["error"]
        revision = reply["result"]["revision_id"]
        aspects = {item["object_id"]: item["aspects"] for item in reply["result"]["copied"]}
        assert aspects[named["Plain"]["object_id"]] == [
            "fill",
            "theme style",
            "outline",
            "effects",
            "text",
        ]
        assert aspects[named["Photo"]["object_id"]] == ["outline", "effects"]
        # A placeholder keeps its layout's text formatting.
        assert "text" not in aspects[named["Title 1"]["object_id"]]

        after = {o["name"]: o for o in inventory()["objects"]}
        for name in ("Plain", "Linked"):
            style = after[name]["shape_style"]
            assert style["fill_color"] == "#C01F2F" and style["line_color"] == "#1F3FC0"
            assert style["line_width_pt"] == 3
            assert after[name]["geometry"] == named[name]["geometry"]
        linked_text = next(
            t for t in inventory()["targets"] if t["object_id"] == named["Linked"]["object_id"]
        )["text_style"]
        assert linked_text["bold"] is True and linked_text["font_family"] == "Georgia"
        assert linked_text["font_size_pt"] == 24
        assert after["Painter"] == named["Painter"]

        with core._storage.project_database(project) as connection:
            snapshot = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, snapshot)) as archive:
            xml = archive.read("ppt/slides/slide1.xml").decode()
        # The text box keeps its own link and the picture keeps its image.
        assert "hlinkClick" in xml and xml.count("<a:blip ") == 1
        assert xml.count('w="38100"') == 5  # 3 pt outline on the source and four targets

        unchanged = paint("Painter", ["Plain"])
        assert unchanged["ok"] and unchanged["result"]["changed"] is False

        for source, targets, field in (
            ("Painter", ["Painter"], "INVALID_REQUEST"),
            ("Painter", [], "INVALID_REQUEST"),
        ):
            bad = paint(source, targets)
            assert not bad["ok"] and bad["error"]["code"] == field
        stale = _call(
            core,
            "s",
            "presentation.workspace.copy_object_format",
            {
                **scope,
                "expected_revision_id": before["revision_id"],
                "slide_id": before["slide_id"],
                "source_object_id": named["Painter"]["object_id"],
                "target_object_ids": [named["Plain"]["object_id"]],
            },
        )
        assert not stale["ok"]
    finally:
        core.close()
