from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from test_graphic_geometry import fixture
from test_presentation_edit_planning import call
from test_presentation_native_layouts import path, scope
from test_presentation_transactions import context

from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def inventory(core, params, number):
    return call(
        core, "presentation.workspace.list_text_targets", {**scope(params), "slide_number": number}
    )


def paste_params(params, source, target, ids):
    return {
        "source_document_id": params["document_id"],
        "source_revision_id": params["expected_revision_id"],
        "source_slide_id": source["slide_id"],
        "slide_id": target["slide_id"],
        "object_ids": ids,
    }


def save_plan(core, params, operation):
    plan = call(
        core,
        "presentation.workspace.prepare_edit_plan",
        {**params, "summary": "Paste copied objects", "operations": [operation]},
    )
    call(
        core,
        "presentation.workspace.open_edit_plan_preview",
        {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": plan["plan_id"],
        },
    )
    return call(
        core,
        "presentation.workspace.apply_edit_plan",
        {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": plan["plan_id"],
            "reviewed_sha256": plan["sha256"],
            "confirm": True,
        },
    )


@pytest.mark.parametrize("kind", ["chart", "table"])
@pytest.mark.parametrize("same_slide", [False, True])
def test_native_graphic_paste_review_undo_and_independent_parts(tmp_path: Path, kind, same_slide):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first, obj = fixture(core, tmp_path, kind)
        second = inventory(core, params, 1 if same_slide else 2)
        original = path(core, params).read_bytes()
        copied = paste_params(params, first, second, [obj["object_id"]])
        saved = save_plan(core, params, {"type": "paste_objects", "params": copied})
        assert saved["changed"]
        params["expected_revision_id"] = saved["revision_id"]
        created_ids = saved["operations"][0]["created_object_ids"]
        assert len(created_ids) == 1 and created_ids[0] != obj["object_id"]
        fresh = inventory(core, params, 1 if same_slide else 2)
        created = next(o for o in fresh["objects"] if o["object_id"] == created_ids[0])
        assert (
            created["name"] != obj["name"]
            and created["geometry"]["width_emu"] == obj["geometry"]["width_emu"]
        )
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(path(core, params)) as after,
        ):
            assert all(
                before.read(n) == after.read(n)
                for n in before.namelist()
                if n
                not in {
                    "[Content_Types].xml",
                    first["slide_part"]
                    if "slide_part" in first
                    else "ppt/slides/slide1.xml"
                    if same_slide
                    else "ppt/slides/slide2.xml",
                    "ppt/slides/_rels/slide1.xml.rels"
                    if same_slide
                    else "ppt/slides/_rels/slide2.xml.rels",
                }
            )
            if kind == "chart":
                added = set(after.namelist()) - set(before.namelist())
                assert any(n.endswith(".xlsx") for n in added)
                workbook = next(n for n in added if n.endswith(".xlsx"))
                assert after.read(workbook) == before.read(
                    next(n for n in before.namelist() if n.endswith(".xlsx"))
                )
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


def test_copy_uses_pinned_revision_after_original_text_changes(tmp_path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, operations = context(core)
        first, second = inventory(core, params, 1), inventory(core, params, 2)
        obj = next(o for o in first["objects"] if o["delete_editable"] and o["geometry_editable"])
        text = next(
            t["current_text"] for t in first["targets"] if t["object_id"] == obj["object_id"]
        )
        operations[0]["params"].update(object_id=obj["object_id"], expected_text=text)
        copied = paste_params(params, first, second, [obj["object_id"]])
        changed = call(
            core,
            "presentation.workspace.replace_text_shape",
            {
                **params,
                **operations[0]["params"],
            },
        )
        params["expected_revision_id"] = changed["revision_id"]
        saved = save_plan(core, params, {"type": "paste_objects", "params": copied})
        params["expected_revision_id"] = saved["revision_id"]
        fresh = inventory(core, params, 2)
        assert any(t["current_text"] == text for t in fresh["targets"])
        assert any(
            o["current_text"] == "Decision requested" for o in inventory(core, params, 1)["targets"]
        )
    finally:
        core.close()


@pytest.mark.parametrize("invalid", ["duplicate", "missing", "wrong_revision", "oversized"])
def test_invalid_clipboard_is_atomic(tmp_path, invalid):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, _ = context(core)
        first, second = inventory(core, params, 1), inventory(core, params, 2)
        obj = first["objects"][0]
        values = paste_params(params, first, second, [obj["object_id"]])
        if invalid == "duplicate":
            values["object_ids"] *= 2
        if invalid == "missing":
            values["object_ids"] = ["00000000-0000-4000-8000-000000000001"]
        if invalid == "wrong_revision":
            values["source_revision_id"] = "00000000-0000-4000-8000-000000000001"
        if invalid == "oversized":
            values["object_ids"] *= 21
        original = path(core, params).read_bytes()
        with pytest.raises(CoreDomainError):
            core._presentation_edits.paste_objects({**params, **values})
        assert path(core, params).read_bytes() == original
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
    finally:
        core.close()


@pytest.mark.parametrize("kind", ["chart", "table"])
def test_paste_between_presentations_in_same_project(tmp_path, kind):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first, obj = fixture(core, tmp_path, kind)
        source_bytes = path(core, params).read_bytes()
        created = call(
            core,
            "presentation.workspace.create_from_outline",
            {
                "project_id": params["project_id"],
                "title": "Destination",
                "slides": [{"title": "Paste here", "bullets": [], "notes": "Keep these notes"}],
            },
        )
        doc = created["document"]["id"]
        head = core._presentation_edits._current_revision_source(params["project_id"], doc)
        target = {
            "project_id": params["project_id"],
            "document_id": doc,
            "expected_revision_id": head["revision_id"],
        }
        original = path(core, target).read_bytes()
        destination = inventory(core, target, 1)
        copied = paste_params(params, first, destination, [obj["object_id"]])
        saved = save_plan(core, target, {"type": "paste_objects", "params": copied})
        target["expected_revision_id"] = saved["revision_id"]
        assert inventory(core, target, 1)["objects"][-1]["kind"] == "graphicFrame"
        assert path(core, params).read_bytes() == source_bytes
        call(core, "presentation.workspace.step_revision", {**target, "direction": "undo"})
        assert path(core, target).read_bytes() == original
    finally:
        core.close()


def test_batch_preserves_source_layer_order_and_relative_geometry(tmp_path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, _ = context(core)
        first, second = inventory(core, params, 1), inventory(core, params, 2)
        objects = [o for o in first["objects"] if o["delete_editable"] and o["geometry_editable"]][
            :3
        ]
        assert len(objects) == 3
        values = paste_params(params, first, second, [o["object_id"] for o in reversed(objects)])
        saved = save_plan(core, params, {"type": "paste_objects", "params": values})
        params["expected_revision_id"] = saved["revision_id"]
        copies = inventory(core, params, 2)["objects"][-3:]
        dx = copies[0]["geometry"]["x_emu"] - objects[0]["geometry"]["x_emu"]
        dy = copies[0]["geometry"]["y_emu"] - objects[0]["geometry"]["y_emu"]
        for copied, source in zip(copies, objects, strict=True):
            assert copied["name"].startswith(source["name"])
            assert copied["geometry"] == {
                **source["geometry"],
                "x_emu": source["geometry"]["x_emu"] + dx,
                "y_emu": source["geometry"]["y_emu"] + dy,
            }
    finally:
        core.close()


def test_shared_picture_dependency_crop_and_external_link_are_remapped(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image = tmp_path / "picture.png"
    Image.new("RGB", (50, 40), "purple").save(image)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for i in range(2):
        picture = slide.shapes.add_picture(
            str(image), Inches(1 + 3 * i), Inches(1), width=Inches(2)
        )
        picture.crop_left = 0.1
        picture.click_action.hyperlink.address = "https://example.com/presentation-reference"
    fixture_path = tmp_path / "pictures.pptx"
    deck.save(fixture_path)
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        target, _ = context(core)
        doc = call(
            core,
            "source.import",
            {"project_id": target["project_id"], "path": str(fixture_path), "kind": "presentation"},
        )["document"]["id"]
        head = core._presentation_edits._current_revision_source(target["project_id"], doc)
        source = {**target, "document_id": doc, "expected_revision_id": head["revision_id"]}
        original = path(core, source).read_bytes()
        first, second = inventory(core, source, 1), inventory(core, target, 2)
        copied = paste_params(source, first, second, [o["object_id"] for o in first["objects"]])
        saved = save_plan(core, target, {"type": "paste_objects", "params": copied})
        target["expected_revision_id"] = saved["revision_id"]
        parsed = Presentation(path(core, target))
        pictures = [shape for shape in parsed.slides[1].shapes if shape.shape_type == 13]
        assert len(pictures) == 2
        assert all(p.image.blob == image.read_bytes() and p.crop_left == 0.1 for p in pictures)
        assert all(
            p.click_action.hyperlink.address == "https://example.com/presentation-reference"
            for p in pictures
        )
        with zipfile.ZipFile(path(core, target)) as archive:
            assert (
                len([n for n in archive.namelist() if n.startswith("ppt/media/copilot-paste-")])
                == 1
            )
        assert path(core, source).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize("catalog_part", ["ppt/tableStyles.xml", "ppt/catalog/custom.xml"])
def test_custom_table_style_cannot_silently_disappear_on_cross_deck_paste(tmp_path, catalog_part):
    from lxml import etree

    from presenter_core.presentation.fidelity import _NS, DRAWING_NS, PACKAGE_REL_NS, REL_NS
    from presenter_core.presentation.graphic_duplication import CT

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        target, _, _ = fixture(core, tmp_path, "table")
        original = path(core, target).read_bytes()
        with zipfile.ZipFile(io.BytesIO(original)) as archive:
            parts = {n: archive.read(n) for n in archive.namelist()}
        style_id = "{11111111-2222-3333-4444-555555555555}"
        slide = etree.fromstring(parts["ppt/slides/slide1.xml"])
        props = slide.find(".//a:tblPr", _NS)
        etree.SubElement(props, f"{{{DRAWING_NS}}}tableStyleId").text = style_id
        parts["ppt/slides/slide1.xml"] = etree.tostring(slide)
        styles = etree.Element(f"{{{DRAWING_NS}}}tblStyleLst", def_=style_id)
        etree.SubElement(
            styles, f"{{{DRAWING_NS}}}tblStyle", styleId=style_id, styleName="Custom source table"
        )
        parts[catalog_part] = etree.tostring(styles)
        rels = etree.fromstring(parts["ppt/_rels/presentation.xml.rels"])
        for n in list(rels):
            if n.get("Type") == f"{REL_NS}/tableStyles":
                rels.remove(n)
        etree.SubElement(
            rels,
            f"{{{PACKAGE_REL_NS}}}Relationship",
            Id="rIdTestStyles",
            Type=f"{REL_NS}/tableStyles",
            Target=catalog_part.removeprefix("ppt/"),
        )
        parts["ppt/_rels/presentation.xml.rels"] = etree.tostring(rels)
        types = etree.fromstring(parts["[Content_Types].xml"])
        if not any(n.get("PartName") == "/" + catalog_part for n in types):
            etree.SubElement(
                types,
                f"{{{CT}}}Override",
                PartName="/" + catalog_part,
                ContentType="application/vnd.openxmlformats-officedocument.presentationml.tableStyles+xml",
            )
        parts["[Content_Types].xml"] = etree.tostring(types)
        candidate = tmp_path / "custom-style.pptx"
        with zipfile.ZipFile(candidate, "w") as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
        doc = call(
            core,
            "source.import",
            {"project_id": target["project_id"], "path": str(candidate), "kind": "presentation"},
        )["document"]["id"]
        head = core._presentation_edits._current_revision_source(target["project_id"], doc)
        source = {**target, "document_id": doc, "expected_revision_id": head["revision_id"]}
        first = inventory(core, source, 1)
        obj = next(o for o in first["objects"] if o["kind"] == "graphicFrame")
        assert obj["delete_editable"] and obj["geometry_editable"]
        copied = paste_params(source, first, inventory(core, target, 2), [obj["object_id"]])
        with pytest.raises(CoreDomainError, match="custom style unavailable"):
            core._presentation_edits.paste_objects({**target, **copied})
        assert path(core, target).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize("cross_slide", [False, True])
def test_missing_chart_relationship_root_is_a_structured_rejection(cross_slide):
    from lxml import etree

    from presenter_core.presentation.fidelity import (
        DRAWING_NS,
        PACKAGE_REL_NS,
        PRESENTATION_NS,
        REL_NS,
    )
    from presenter_core.presentation.graphic_duplication import CNS, clone_chart_graph

    frame = etree.Element(f"{{{PRESENTATION_NS}}}graphicFrame")
    graphic = etree.SubElement(frame, f"{{{DRAWING_NS}}}graphic")
    data_node = etree.SubElement(graphic, f"{{{DRAWING_NS}}}graphicData")
    chart = etree.SubElement(data_node, f"{{{CNS}}}chart")
    chart.set(f"{{{REL_NS}}}id", "rId1")
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", b"<unused/>")
    with (
        zipfile.ZipFile(data) as archive,
        pytest.raises(CoreDomainError, match="package dependencies"),
    ):
        clone_chart_graph(
            archive,
            "ppt/slides/slide2.xml" if cross_slide else "ppt/slides/slide1.xml",
            frame,
            source_slide_part="ppt/slides/slide1.xml" if cross_slide else None,
            slide_relationships=etree.Element(f"{{{PACKAGE_REL_NS}}}Relationships")
            if cross_slide
            else None,
        )


def test_paste_never_overwrites_an_existing_destination_dependency(tmp_path, monkeypatch):
    import presenter_core.presentation.object_paste as module

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first, obj = fixture(core, tmp_path, "chart")
        original = path(core, params).read_bytes()
        monkeypatch.setattr(
            module,
            "clone_chart_graph",
            lambda *_args, **_kwargs: {"ppt/charts/chart1.xml": b"<corrupt/>"},
        )
        copied = paste_params(params, first, first, [obj["object_id"]])
        with pytest.raises(CoreDomainError, match="overwrite an existing part"):
            core._presentation_edits.paste_objects({**params, **copied})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


def test_batch_chart_graph_sharing_stays_inside_the_pasted_batch(tmp_path):
    import copy

    from lxml import etree

    from presenter_core.presentation.fidelity import _NS, REL_NS
    from presenter_core.presentation.graphic_duplication import CNS

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        target, _, _ = fixture(core, tmp_path, "chart")
        original = path(core, target).read_bytes()
        with zipfile.ZipFile(io.BytesIO(original)) as archive:
            parts = {n: archive.read(n) for n in archive.namelist()}
        slide = etree.fromstring(parts["ppt/slides/slide1.xml"])
        tree = slide.find("p:cSld/p:spTree", _NS)
        duplicate = copy.deepcopy(tree.find("p:graphicFrame", _NS))
        native = duplicate.find("p:nvGraphicFramePr/p:cNvPr", _NS)
        native.set("id", "3")
        native.set("name", "Shared graph second chart")
        duplicate.find("p:xfrm/a:off", _NS).set("x", "1828800")
        tree.append(duplicate)
        parts["ppt/slides/slide1.xml"] = etree.tostring(slide)
        candidate = tmp_path / "shared-graph.pptx"
        with zipfile.ZipFile(candidate, "w") as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
        doc = call(
            core,
            "source.import",
            {"project_id": target["project_id"], "path": str(candidate), "kind": "presentation"},
        )["document"]["id"]
        head = core._presentation_edits._current_revision_source(target["project_id"], doc)
        source = {**target, "document_id": doc, "expected_revision_id": head["revision_id"]}
        source_bytes = path(core, source).read_bytes()
        first, second = inventory(core, source, 1), inventory(core, target, 2)
        copied = paste_params(source, first, second, [o["object_id"] for o in first["objects"]])
        saved = save_plan(core, target, {"type": "paste_objects", "params": copied})
        target["expected_revision_id"] = saved["revision_id"]
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(path(core, target)) as after,
        ):
            slide = etree.fromstring(after.read("ppt/slides/slide2.xml"))
            rels = {
                n.get("Id"): n.get("Target")
                for n in etree.fromstring(after.read("ppt/slides/_rels/slide2.xml.rels"))
            }
            targets = [
                rels[n.get(f"{{{REL_NS}}}id")]
                for n in slide.findall(".//c:chart", {**_NS, "c": CNS})
            ]
            assert len(targets) == 3 and targets[1] == targets[2] and targets[1] != targets[0]
            assert (
                len(
                    [
                        n
                        for n in set(after.namelist()) - set(before.namelist())
                        if n.endswith(".xlsx")
                    ]
                )
                == 1
            )
            assert before.read("ppt/charts/chart2.xml") == after.read("ppt/charts/chart2.xml")
        assert path(core, source).read_bytes() == source_bytes
    finally:
        core.close()
