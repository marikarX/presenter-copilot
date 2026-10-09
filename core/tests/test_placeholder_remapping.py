from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches
from test_appendix_workflow import slides
from test_presentation_native_layouts import path, scope
from test_presentation_transactions import call
from test_selection_actions import save

from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.presentation.fidelity import _NS
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def fixture(core, directory: Path, dependent=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    title = slide.shapes.title
    title.text_frame.clear()
    paragraph = title.text_frame.paragraphs[0]
    for text, color in [("Reviewed ", "CC0000"), ("recommendation", "0000CC")]:
        run = paragraph.add_run()
        run.text = text
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(color)
    slide.placeholders[1].text = "Keep rich text and native dependencies"
    slide.notes_slide.notes_text_frame.text = "Preserved review notes"
    image = directory / "image.png"
    Image.new("RGB", (80, 50), "purple").save(image)
    slide.shapes.add_picture(str(image), Inches(7), Inches(5), width=Inches(1))
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Unchanged slide"
    if dependent:
        timing = etree.SubElement(slide._element, f"{{{_NS['p']}}}timing")
        etree.SubElement(timing, f"{{{_NS['p']}}}spTgt", spid=str(title.shape_id))
    source = directory / "remap.pptx"
    deck.save(source)
    project = call(core, "project.create", {"name": "Remapping acceptance"})["project"]["id"]
    document = call(
        core, "source.import", {"project_id": project, "path": str(source), "kind": "presentation"}
    )["document"]["id"]
    revision = core._presentation_edits._current_revision_source(project, document)
    params = {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": revision["revision_id"],
    }
    slide_id = slides(core, params)[0].slide_id
    catalog = call(
        core, "presentation.workspace.list_native_layouts", {**scope(params), "slide_id": slide_id}
    )
    return params, slide_id, catalog


def assignments(catalog, chosen):
    title = next(
        item for item in catalog["remap_candidates"] if item["type"] in {"title", "ctrTitle"}
    )
    body = next(item for item in catalog["remap_candidates"] if item["type"] in {"body", "obj"})
    first = next(
        slot for slot in chosen["creation_placeholders"] if slot["type"] in {"title", "ctrTitle"}
    )
    second = next(
        slot
        for kind in ("obj", "body", "subTitle")
        for slot in chosen["creation_placeholders"]
        if slot["type"] == kind
    )
    return [
        {"object_id": title["object_id"], "placeholder_id": first["placeholder_id"]},
        {"object_id": body["object_id"], "placeholder_id": second["placeholder_id"]},
    ]


def request(params, slide, layout, mapping):
    return {
        **params,
        "summary": "Remap reviewed native placeholders",
        "operations": [
            {
                "type": "apply_slide_layout",
                "params": {"slide_id": slide, "layout_id": layout, "placeholder_mapping": mapping},
            }
        ],
    }


@pytest.mark.parametrize("name", ["Title and Content", "Comparison", "Title Slide"])
def test_remapping_preserves_rich_text_ids_parts_notes_and_atomic_history(tmp_path: Path, name):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, slide, choices = fixture(core, tmp_path)
        original = path(core, params).read_bytes()
        before = slides(core, params)
        chosen = next(item for item in choices["items"] if item["name"] == name)
        mapping = assignments(choices, chosen)
        proposed = core._presentation_edit_plans.prepare(
            request(params, slide, chosen["layout_id"], mapping)
        )
        assert proposed["changed"] and path(core, params).read_bytes() == original
        applied = save(core, params, proposed)
        revised = path(core, params).read_bytes()
        after = slides(core, params)
        assert [obj.object_id for obj in before[0].objects] == [
            obj.object_id for obj in after[0].objects
        ]
        assert after[1] == before[1]
        assert after[0].speaker_notes == before[0].speaker_notes
        assert [obj.text for obj in before[0].objects] == [obj.text for obj in after[0].objects]
        assert after[0].objects[-1] == before[0].objects[-1]
        assert all(obj.geometry for obj in after[0].objects[:2])
        with (
            zipfile.ZipFile(io.BytesIO(original)) as old,
            zipfile.ZipFile(io.BytesIO(revised)) as new,
        ):
            assert old.namelist() == new.namelist()
            changed = {name for name in old.namelist() if old.read(name) != new.read(name)}
            assert changed <= {"ppt/slides/slide1.xml", "ppt/slides/_rels/slide1.xml.rels"}
            left = etree.fromstring(old.read("ppt/slides/slide1.xml"))
            right = etree.fromstring(new.read("ppt/slides/slide1.xml"))
            assert [
                etree.tostring(body, method="c14n")
                for body in left.findall("p:cSld/p:spTree/p:sp/p:txBody", _NS)
            ] == [
                etree.tostring(body, method="c14n")
                for body in right.findall("p:cSld/p:spTree/p:sp/p:txBody", _NS)
            ]
        assert applied["sha256"] == proposed["sha256"]
        repeat = core._presentation_edit_plans.prepare(
            request(params, slide, chosen["layout_id"], mapping)
        )
        assert not repeat["changed"]
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
        params["expected_revision_id"] = core._presentation_edits._current_revision_source(
            params["project_id"], params["document_id"]
        )["revision_id"]
        call(core, "presentation.workspace.step_revision", {**params, "direction": "redo"})
        assert path(core, params).read_bytes() == revised
    finally:
        core.close()


@pytest.mark.parametrize(
    "bad",
    [
        "duplicate_object",
        "duplicate_slot",
        "unknown_slot",
        "unknown_object",
        "collision",
        "empty",
        "late",
    ],
)
def test_invalid_mapping_never_publishes_partial_layout(tmp_path: Path, bad):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, slide, choices = fixture(core, tmp_path)
        original = path(core, params).read_bytes()
        chosen = next(item for item in choices["items"] if item["name"] == "Title and Content")
        mapping = assignments(choices, chosen)
        if bad == "duplicate_object":
            mapping[1]["object_id"] = mapping[0]["object_id"]
        if bad == "duplicate_slot":
            mapping[1]["placeholder_id"] = mapping[0]["placeholder_id"]
        if bad == "unknown_slot":
            mapping[0]["placeholder_id"] = "missing"
        if bad == "unknown_object":
            mapping[0]["object_id"] = "missing"
        if bad == "collision":
            mapping = [
                {
                    "object_id": mapping[0]["object_id"],
                    "placeholder_id": mapping[1]["placeholder_id"],
                }
            ]
        if bad == "empty":
            mapping = []
        req = request(params, slide, chosen["layout_id"], mapping)
        if bad == "late":
            req["operations"].append(
                {"type": "delete_object", "params": {"slide_id": slide, "object_id": "missing"}}
            )
        with pytest.raises(CoreDomainError):
            core._presentation_edit_plans.prepare(req)
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


def test_animated_placeholder_is_not_offered_or_remapped(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        _, _, choices = fixture(core, tmp_path, dependent=True)
        assert len(choices["remap_candidates"]) == 1 and choices["remap_candidates"][0]["type"] in {
            "obj",
            "body",
        }
    finally:
        core.close()
