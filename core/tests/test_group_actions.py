from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from test_object_groups import group_request, mixed_fixture, operation, prepare
from test_object_paste import inventory, paste_params
from test_presentation_native_layouts import path
from test_presentation_transactions import call
from test_selection_actions import save

from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.presentation.fidelity import _NS
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def grouped(core, directory, nested, picture_format):
    params, first = mixed_fixture(core, directory, picture_format)
    save(
        core,
        params,
        prepare(core, group_request(core, params, [first["objects"][0], *first["objects"][2:]])),
    )
    if nested:
        first = inventory(core, params, 1)
        save(
            core,
            params,
            prepare(
                core,
                group_request(
                    core,
                    params,
                    [
                        o
                        for o in first["objects"]
                        if o["geometry_editable"] and not o["group_child"]
                    ],
                ),
            ),
        )
    first = inventory(core, params, 1)
    return params, first, next(o for o in first["objects"] if o["kind"] == "grpSp")


@pytest.mark.parametrize("nested", [False, True])
@pytest.mark.parametrize("picture_format", ["PNG", "JPEG"])
@pytest.mark.parametrize("action", ["duplicate", "paste_slide", "paste_document"])
def test_group_copy_relationships_delete_and_whole_history(
    tmp_path: Path, nested, picture_format, action
):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first, group = grouped(core, tmp_path, nested, picture_format)
        original = path(core, params).read_bytes()
        original_revision = params["expected_revision_id"]
        original_ids = {o["object_id"] for o in first["objects"]}
        if action == "duplicate":
            request = operation(
                core,
                params,
                "duplicate_object",
                {
                    "slide_id": first["slide_id"],
                    "object_id": group["object_id"],
                    "expected_geometry": group["geometry"],
                },
            )
            target = first
        else:
            copied = paste_params(params, first, inventory(core, params, 2), [group["object_id"]])
            if action == "paste_document":
                source_directory = tmp_path / "target"
                source_directory.mkdir()
                mixed_fixture(core, source_directory, picture_format)
                target_source = source_directory / "mixed.pptx"
                target_deck = Presentation(target_source)
                target_deck.slides[0].shapes[0].text = "Destination evidence"
                target_deck.save(target_source)
                imported = call(
                    core,
                    "source.import",
                    {
                        "project_id": params["project_id"],
                        "path": str(source_directory / "mixed.pptx"),
                        "kind": "presentation",
                    },
                )["document"]["id"]
                current = core._presentation_edits._current_revision_source(
                    params["project_id"], imported
                )
                params = {
                    "project_id": params["project_id"],
                    "document_id": imported,
                    "expected_revision_id": current["revision_id"],
                }
                target = inventory(core, params, 1)
                original = path(core, params).read_bytes()
                original_revision = params["expected_revision_id"]
                original_ids = {o["object_id"] for o in target["objects"]}
                copied["slide_id"] = target["slide_id"]
            else:
                target = inventory(core, params, 2)
                original_ids = {o["object_id"] for o in target["objects"]}
            request = operation(core, params, "paste_objects", copied)
        result = save(core, params, prepare(core, request))
        values = result["operations"][0]
        new_id = values.get("created_object_id") or values["created_object_ids"][0]
        fresh = inventory(core, params, 1 if action != "paste_slide" else 2)
        new = next(o for o in fresh["objects"] if o["object_id"] == new_id)
        assert new["kind"] == "grpSp" and new["delete_editable"] and new["group_editable"]
        assert len({o["ooxml_shape_id"] for o in fresh["objects"]}) == len(fresh["objects"])
        assert all(any(f["object_id"] == old for f in fresh["objects"]) for old in original_ids)
        copied_bytes = path(core, params).read_bytes()
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(io.BytesIO(copied_bytes)) as after,
        ):
            # A group chart is copied with an independent native chart and workbook graph.
            charts = [
                n
                for n in after.namelist()
                if n.startswith("ppt/charts/") and n.endswith(".xml") and "/_rels/" not in n
            ]
            workbooks = [n for n in after.namelist() if n.startswith("ppt/embeddings/")]
            assert len(charts) == 2 and len(workbooks) == 2
            assert all(
                after.read(n) == before.read(n)
                for n in before.namelist()
                if n
                not in {
                    "[Content_Types].xml",
                    f"ppt/slides/slide{2 if action == 'paste_slide' else 1}.xml",
                    f"ppt/slides/_rels/slide{2 if action == 'paste_slide' else 1}.xml.rels",
                }
            )
            root = etree.fromstring(
                after.read(f"ppt/slides/slide{2 if action == 'paste_slide' else 1}.xml")
            )
            ids = root.xpath(".//p:cNvPr/@id", namespaces=_NS)
            assert len(ids) == len(set(ids))
        deletion = operation(
            core,
            params,
            "delete_object",
            {
                "slide_id": target["slide_id"],
                "object_id": new_id,
                "expected_geometry": new["geometry"],
            },
        )
        save(core, params, prepare(core, deletion))
        assert {
            o["object_id"]
            for o in inventory(core, params, 1 if action != "paste_slide" else 2)["objects"]
        } == original_ids
        stepped = call(
            core,
            "presentation.workspace.step_revision",
            {**params, "direction": "undo"},
        )
        params["expected_revision_id"] = stepped["revision_id"]
        assert path(core, params).read_bytes() == copied_bytes
        stepped = call(
            core,
            "presentation.workspace.step_revision",
            {**params, "direction": "undo"},
        )
        params["expected_revision_id"] = stepped["revision_id"]
        assert path(core, params).read_bytes() == original
        stepped = call(
            core,
            "presentation.workspace.step_revision",
            {**params, "direction": "redo"},
        )
        params["expected_revision_id"] = stepped["revision_id"]
        assert path(core, params).read_bytes() == copied_bytes
        assert original_revision != result["revision_id"]
    finally:
        core.close()


@pytest.mark.parametrize("direction", ["front", "back", "forward", "backward"])
def test_group_layer_changes_only_order(tmp_path: Path, direction):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first, group = grouped(core, tmp_path, False, "PNG")
        original = path(core, params).read_bytes()
        result = call(
            core,
            "presentation.workspace.change_object_layer",
            {
                **params,
                "slide_id": first["slide_id"],
                "object_id": group["object_id"],
                "expected_object_ids": [o["object_id"] for o in first["objects"]],
                "direction": direction,
            },
        )
        params["expected_revision_id"] = result["revision_id"]
        fresh = inventory(core, params, 1)
        assert {o["object_id"]: o for o in first["objects"]} == {
            o["object_id"]: o for o in fresh["objects"]
        }
        if direction in {"back", "backward"}:
            assert result["changed"]
            assert fresh["objects"][0]["object_id"] == group["object_id"]
            stepped = call(
                core,
                "presentation.workspace.step_revision",
                {**params, "direction": "undo"},
            )
            params["expected_revision_id"] = stepped["revision_id"]
            assert path(core, params).read_bytes() == original
        else:
            assert not result["changed"] and path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize(
    "failure", ["cancel", "stale_geometry", "stale_revision", "late_target", "object_limit"]
)
def test_group_action_discard_and_atomic_failures(tmp_path: Path, failure, monkeypatch):
    import uuid

    import presenter_core.presentation.object_paste as paste_module

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first, group = grouped(core, tmp_path, True, "PNG")
        original = path(core, params).read_bytes()
        head = params["expected_revision_id"]
        request = operation(
            core,
            params,
            "duplicate_object",
            {
                "slide_id": first["slide_id"],
                "object_id": group["object_id"],
                "expected_geometry": dict(group["geometry"]),
            },
        )
        if failure == "stale_geometry":
            request["operations"][0]["params"]["expected_geometry"]["x_emu"] += 1
        elif failure == "stale_revision":
            request["expected_revision_id"] = str(uuid.uuid4())
        elif failure == "late_target":
            request["operations"].append(
                {
                    "type": "delete_object",
                    "params": {
                        "slide_id": first["slide_id"],
                        "object_id": str(uuid.uuid4()),
                        "expected_geometry": group["geometry"],
                    },
                }
            )
        elif failure == "object_limit":
            monkeypatch.setattr(paste_module, "OBJECT_LIMIT", len(first["objects"]) + 1)
        if failure == "cancel":
            proposed = prepare(core, request)
            call(
                core,
                "presentation.workspace.reject_edit_plan",
                {
                    "project_id": params["project_id"],
                    "document_id": params["document_id"],
                    "plan_id": proposed["plan_id"],
                },
            )
        else:
            with pytest.raises(CoreDomainError):
                core._presentation_edit_plans.prepare(request)
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == head
        )
        assert path(core, params).read_bytes() == original
    finally:
        core.close()
