"""Deck-wide find and run-preserving replace."""

import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _deck(path: Path) -> Path:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    box = first.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.name = "Mixed runs"
    paragraph = box.text_frame.paragraphs[0]
    plain = paragraph.add_run()
    plain.text = "Retire the Legacy plan and the "
    bold = paragraph.add_run()
    bold.text = "legacy tools"
    bold.font.bold = True
    bold.font.size = Pt(30)
    split = first.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    split.name = "Split word"
    paragraph = split.text_frame.paragraphs[0]
    paragraph.add_run().text = "Leg"
    paragraph.add_run().text = "acy budget"
    second = deck.slides.add_slide(deck.slide_layouts[6])
    other = second.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    other.name = "Second slide"
    other.text_frame.text = "No match here, and LEGACY in capitals"
    deck.save(path)
    return path


class _Workspace:
    def __init__(self, tmp_path: Path) -> None:
        self.core = CoreService(data_root=tmp_path / "data")
        self.project = _project_id(self.core)
        source = _deck(tmp_path / "find.pptx")
        self.document = _import_presentation(self.core, self.project, source)["id"]
        self.revision = _open(self.core, self.project, self.document)["revision_id"]

    def call(self, method: str, params: dict) -> dict:
        return _call(
            self.core,
            method,
            f"presentation.workspace.{method}",
            {
                "project_id": self.project,
                "document_id": self.document,
                "expected_revision_id": self.revision,
                **params,
            },
        )

    def find(self, query: str, match_case: bool = False) -> dict:
        reply = self.call("find_text", {"query": query, "match_case": match_case})
        assert reply["ok"], reply
        return reply["result"]

    def replace(self, query: str, replacement: str, found: dict, **extra) -> dict:
        reply = self.call(
            "replace_text_all",
            {
                "query": query,
                "match_case": found["match_case"],
                "replacement": replacement,
                "expected_matches": [
                    {
                        key: item[key]
                        for key in ("target", "slide_id", "object_id", "replaceable_count")
                    }
                    for item in found["objects"]
                    if item["replaceable_count"]
                ],
                **extra,
            },
        )
        if reply["ok"]:
            self.revision = reply["result"]["revision_id"]
        return reply

    def slide_xml(self, number: int) -> str:
        with self.core._storage.project_database(self.project) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?",
                (self.revision,),
            ).fetchone()
        path = self.core._storage.paths.snapshot_path(
            self.project, row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(path) as archive:
            return archive.read(f"ppt/slides/slide{number}.xml").decode()


def test_find_counts_replaceable_and_cross_run_matches(tmp_path: Path) -> None:
    workspace = _Workspace(tmp_path)
    try:
        found = workspace.find("legacy")
        by_name = {item["name"]: item for item in found["objects"]}
        assert by_name["Mixed runs"]["replaceable_count"] == 2
        assert by_name["Split word"] == {
            **by_name["Split word"],
            "replaceable_count": 0,
            "skipped_count": 1,
        }
        assert by_name["Second slide"]["slide_number"] == 2
        assert (found["replaceable_matches"], found["skipped_matches"]) == (3, 1)
        assert workspace.find("legacy", match_case=True)["replaceable_matches"] == 1
    finally:
        workspace.core.close()


def test_replace_all_keeps_run_formatting_in_one_revision(tmp_path: Path) -> None:
    workspace = _Workspace(tmp_path)
    try:
        before = workspace.revision
        found = workspace.find("legacy")
        reply = workspace.replace("legacy", "current", found)
        assert reply["ok"], reply
        assert reply["result"]["replaced_matches"] == 3
        assert reply["result"]["objects_changed"] == 2
        assert reply["result"]["parent_revision_id"] == before
        first = workspace.slide_xml(1)
        assert "Retire the current plan and the " in first
        assert ">current tools<" in first and 'b="1"' in first and 'sz="3000"' in first
        assert ">Leg<" in first and ">acy budget<" in first
        assert "No match here, and current in capitals" in workspace.slide_xml(2)
        after = workspace.find("legacy")
        assert (after["replaceable_matches"], after["skipped_matches"]) == (0, 1)
    finally:
        workspace.core.close()


def test_stale_matches_and_invalid_queries_are_rejected(tmp_path: Path) -> None:
    workspace = _Workspace(tmp_path)
    try:
        found = workspace.find("legacy")
        found["objects"][0]["replaceable_count"] += 1
        stale = workspace.replace("legacy", "current", found)
        assert not stale["ok"] and stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        for query in ("", "   ", "a" * 201, "line\nbreak"):
            reply = workspace.call("find_text", {"query": query})
            assert not reply["ok"] and reply["error"]["code"] == "INVALID_REQUEST"
    finally:
        workspace.core.close()


def test_text_inside_group_shapes_is_found_and_replaced(tmp_path: Path) -> None:
    source = tmp_path / "grouped.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    inner.name = "Grouped text"
    inner.text_frame.text = "Grouped legacy note"
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document, "expected_revision_id": revision}
        found = _call(
            core, "find", "presentation.workspace.find_text", {**scope, "query": "legacy"}
        )
        assert found["ok"], found
        [match] = found["result"]["objects"]
        assert (match["name"], match["replaceable_count"]) == ("Grouped text", 1)
        reply = _call(
            core,
            "replace",
            "presentation.workspace.replace_text_all",
            {
                **scope,
                "query": "legacy",
                "replacement": "current",
                "expected_matches": [
                    {k: match[k] for k in ("target", "slide_id", "object_id", "replaceable_count")}
                ],
            },
        )
        assert reply["ok"] and reply["result"]["replaced_matches"] == 1, reply
    finally:
        core.close()


def test_speaker_notes_are_searched_and_replaced_with_slides(tmp_path: Path) -> None:
    source = tmp_path / "notes.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text_frame.text = "Legacy roadmap"
    slide.notes_slide.notes_text_frame.text = "Mention the legacy budget twice: legacy."
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document, "expected_revision_id": revision}
        found = _call(
            core, "find", "presentation.workspace.find_text", {**scope, "query": "legacy"}
        )
        assert found["ok"], found
        targets = {(m["target"], m["replaceable_count"]) for m in found["result"]["objects"]}
        assert targets == {("slide", 1), ("speaker_notes", 2)}
        reply = _call(
            core,
            "replace",
            "presentation.workspace.replace_text_all",
            {
                **scope,
                "query": "legacy",
                "replacement": "current",
                "expected_matches": [
                    {k: m[k] for k in ("target", "slide_id", "object_id", "replaceable_count")}
                    for m in found["result"]["objects"]
                ],
            },
        )
        assert reply["ok"] and reply["result"]["replaced_matches"] == 3, reply
        inventory = _call(
            core,
            "targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": document,
                "revision_id": reply["result"]["revision_id"],
                "slide_number": 1,
            },
        )["result"]
        assert inventory["speaker_notes"]["current_text"] == (
            "Mention the current budget twice: current."
        )
    finally:
        core.close()
