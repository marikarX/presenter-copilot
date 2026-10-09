"""Resolve native formatting without modifying imported styles or following external links."""

import io
import zipfile

import pytest
from lxml import etree  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.util import Inches, Pt

from presenter_core.errors import CoreDomainError
from presenter_core.presentation.fidelity import _NS
from presenter_core.presentation.text_inheritance import TextInheritance, empty


@pytest.fixture
def package():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Inherited title"
    slide.placeholders[1].text = "Inherited body"
    paragraph = slide.placeholders[1].text_frame.add_paragraph()
    paragraph.text, paragraph.level = "Second-level body", 1
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(0.7)).text = "Default box"
    box = slide.shapes.add_textbox(Inches(1), Inches(6), Inches(4), Inches(0.7))
    box.text = "Explicit box"
    font = box.text_frame.paragraphs[0].runs[0].font
    font.name, font.size, font.bold = "Arial", Pt(24), True
    file = io.BytesIO()
    deck.save(file)
    with zipfile.ZipFile(file) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def reader(package):
    file = io.BytesIO()
    with zipfile.ZipFile(file, "w") as archive:
        for name, data in package.items():
            archive.writestr(name, data)
    return TextInheritance(zipfile.ZipFile(file), "ppt/slides/slide1.xml")


def shape(resolver, number):
    return resolver.slide.findall("p:cSld/p:spTree/p:sp", _NS)[number]


def change(package, part, callback):
    root = etree.fromstring(package[part])
    callback(root)
    package[part] = etree.tostring(root)


@pytest.mark.parametrize(
    "number,size,family",
    [(0, 44, "Calibri"), (1, 32, "Calibri"), (2, 18, "Calibri"), (3, 24, "Arial")],
)
def test_native_defaults_and_preservation(package, number, size, family):
    resolver = reader(package)
    original = shape(resolver, number)
    before = etree.tostring(original)
    result = resolver.resolve(original)
    run = result.paragraphs[0]["runs"][0]
    assert run["family"] == family
    assert int(run["size_hundredths"]) / 100 == size
    assert len(result.fingerprint) == 64
    assert result.fingerprint == resolver.resolve(original).fingerprint
    result.transform.set("rot", "99")
    result.body.set("probe", "detached")
    assert etree.tostring(original) == before
    assert resolver.archive.read("ppt/slides/slide1.xml") == package["ppt/slides/slide1.xml"]


def test_outline_level_and_empty_paragraph(package):
    resolver = reader(package)
    assert (
        resolver.resolve(shape(resolver, 1)).paragraphs[1]["runs"][0]["size_hundredths"] == "2800"
    )
    paragraph = etree.SubElement(shape(resolver, 1).find("p:txBody", _NS), f"{{{_NS['a']}}}p")
    end = etree.SubElement(paragraph, f"{{{_NS['a']}}}endParaRPr", sz="2100")
    etree.SubElement(end, f"{{{_NS['a']}}}latin", typeface="Arial")
    result = resolver.resolve(shape(resolver, 1)).paragraphs[-1]
    assert not result["runs"]
    assert result["end"]["family"] == "Arial"
    assert result["end"]["size_hundredths"] == "2100"


def test_run_overrides_partial_paragraph_defaults(package):
    resolver = reader(package)
    target = shape(resolver, 3)
    paragraph = target.find("p:txBody/a:p", _NS)
    pp = empty("pPr")
    paragraph.insert(0, pp)
    etree.SubElement(pp, f"{{{_NS['a']}}}defRPr", sz="1700", i="1")
    defaults = target.find("p:txBody/a:lstStyle", _NS)
    level = etree.SubElement(defaults, f"{{{_NS['a']}}}lvl1pPr")
    props = etree.SubElement(level, f"{{{_NS['a']}}}defRPr", sz="1400", b="0")
    etree.SubElement(props, f"{{{_NS['a']}}}latin", typeface="Times New Roman")
    assert resolver.resolve(target).paragraphs[0]["runs"][0] == {
        "family": "Arial",
        "size_hundredths": "2400",
        "bold": "1",
        "italic": "1",
    }


def test_theme_changes_bind_fingerprint(package):
    first = reader(package)
    fingerprint = first.resolve(shape(first, 1)).fingerprint
    change(
        package,
        "ppt/theme/theme1.xml",
        lambda root: root.find("a:themeElements/a:fontScheme/a:minorFont/a:latin", _NS).set(
            "typeface", "Arial"
        ),
    )
    second = reader(package)
    result = second.resolve(shape(second, 1))
    assert result.paragraphs[0]["runs"][0]["family"] == "Arial"
    assert result.fingerprint != fingerprint
    assert second.resolve(shape(second, 0)).paragraphs[0]["runs"][0]["family"] == "Calibri"


def test_presentation_defaults_and_placeholder_geometry(package):
    change(
        package,
        "ppt/presentation.xml",
        lambda root: root.find("p:defaultTextStyle/a:lvl1pPr/a:defRPr", _NS).set("sz", "1900"),
    )
    resolver = reader(package)
    assert (
        resolver.resolve(shape(resolver, 2)).paragraphs[0]["runs"][0]["size_hundredths"] == "1900"
    )
    change(
        package,
        "ppt/slideMasters/slideMaster1.xml",
        lambda root: root.find("p:txStyles/p:otherStyle/a:lvl1pPr/a:defRPr", _NS).attrib.pop("sz"),
    )
    inherited = reader(package)
    assert (
        inherited.resolve(shape(inherited, 2)).paragraphs[0]["runs"][0]["size_hundredths"] == "1900"
    )
    assert resolver.resolve(shape(resolver, 0)).transform.find("a:ext", _NS).get("cx") == "8229600"


@pytest.mark.parametrize("source", ["sample_paragraph", "list_style"])
def test_layout_sample_text_is_not_an_inherited_style(package, source):
    def alter(root):
        target = next(
            s
            for s in root.findall("p:cSld/p:spTree/p:sp", _NS)
            if (ph := s.find("p:nvSpPr/p:nvPr/p:ph", _NS)) is not None and ph.get("idx") == "1"
        )
        if source == "list_style":
            pp = etree.SubElement(target.find("p:txBody/a:lstStyle", _NS), f"{{{_NS['a']}}}lvl1pPr")
        else:
            paragraph = target.find("p:txBody/a:p", _NS)
            pp = paragraph.find("a:pPr", _NS)
            if pp is None:
                pp = empty("pPr")
                paragraph.insert(0, pp)
        props = etree.SubElement(pp, f"{{{_NS['a']}}}defRPr", sz="2200", i="1")
        etree.SubElement(props, f"{{{_NS['a']}}}latin", typeface="Arial")

    change(package, "ppt/slideLayouts/slideLayout2.xml", alter)
    resolver = reader(package)
    result = resolver.resolve(shape(resolver, 1))
    expected = ("Arial", "2200") if source == "list_style" else ("Calibri", "3200")
    assert (
        result.paragraphs[0]["runs"][0]["family"],
        result.paragraphs[0]["runs"][0]["size_hundredths"],
    ) == expected
    assert result.paragraphs[1]["runs"][0]["size_hundredths"] == "2800"


def test_subtitle_uses_body_master_category(package):
    def alter(root):
        node = next(n for n in root if n.get("Type", "").endswith("/slideLayout"))
        node.set("Target", "../slideLayouts/slideLayout1.xml")

    change(package, "ppt/slides/_rels/slide1.xml.rels", alter)
    resolver = reader(package)
    result = resolver.resolve(shape(resolver, 1))
    assert result.paragraphs[0]["runs"][0]["size_hundredths"] == "3200"
    assert result.transform.find("a:ext", _NS).get("cy") == "1752600"


@pytest.mark.parametrize("case", ["external", "duplicate", "missing", "level", "script"])
def test_unresolved_inheritance_rejects_without_mutation(package, case):
    if case in {"external", "duplicate"}:

        def alter(root):
            node = next(n for n in root if n.get("Type", "").endswith("/theme"))
            if case == "external":
                node.set("TargetMode", "External")
                node.set("Target", "https://example.invalid/theme.xml")
            else:
                clone = etree.fromstring(etree.tostring(node))
                clone.set("Id", "anotherTheme")
                root.append(clone)

        change(package, "ppt/slideMasters/_rels/slideMaster1.xml.rels", alter)
    elif case == "missing":
        del package["ppt/theme/theme1.xml"]
    else:

        def alter(root):
            target = root.findall("p:cSld/p:spTree/p:sp", _NS)[1]
            paragraph = target.find("p:txBody/a:p", _NS)
            pp = empty("pPr")
            paragraph.insert(0, pp)
            if case == "level":
                pp.set("lvl", "9")
            else:
                rp = etree.SubElement(pp, f"{{{_NS['a']}}}defRPr")
                etree.SubElement(rp, f"{{{_NS['a']}}}latin", typeface="+mn-ea")

        change(package, "ppt/slides/slide1.xml", alter)
    original = dict(package)
    with pytest.raises(CoreDomainError):
        resolver = reader(package)
        resolver.resolve(shape(resolver, 1))
    assert package == original
