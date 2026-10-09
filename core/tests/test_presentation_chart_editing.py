from __future__ import annotations

import copy
import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from test_presentation_edit_planning import call
from test_presentation_native_layouts import path, scope

from presenter_core.ipc.core import CoreService
from presenter_core.presentation.charts import CNS, KINDS, NS, read_chart
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def imported(core, directory, kind="column", *, shared=False, mutation=None):
    deck = Presentation()
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        values = CategoryChartData()
        values.categories = ["Baseline", "Pilot", "Target"]
        values.add_series("Revenue", [10, 12.25, 15])
        if kind != "pie":
            values.add_series("Cost", [8, 9, 10])
        chart = slide.shapes.add_chart(
            KINDS[kind], Inches(1), Inches(2), Inches(7), Inches(4), values
        ).chart
        chart.has_title = True
        chart.chart_title.text_frame.text = "Imported pilot"
        font = chart.chart_title.text_frame.paragraphs[0].runs[0].font
        font.size = Pt(23)
        font.bold = True
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = RGBColor(110, 30, 160)
        slide.notes_slide.notes_text_frame.text = "Preserved imported notes"
    fixture = directory / "imported-chart.pptx"
    deck.save(fixture)
    if shared or mutation:
        with zipfile.ZipFile(fixture) as archive:
            parts = {v.filename: archive.read(v) for v in archive.infolist()}
        if shared:
            rel = etree.fromstring(parts["ppt/slides/_rels/slide2.xml.rels"])
            for node in rel:
                if node.get("Type", "").endswith("/chart"):
                    node.set("Target", "../charts/chart1.xml")
            parts["ppt/slides/_rels/slide2.xml.rels"] = etree.tostring(rel)
        if mutation:
            chart = etree.fromstring(parts["ppt/charts/chart1.xml"])
            if mutation == "extension":
                etree.SubElement(chart, f"{{{CNS}}}extLst")
            if mutation == "stacked":
                chart.find(".//c:grouping", NS).set("val", "stacked")
            if mutation == "missing_cache":
                cache = chart.find(".//c:numCache", NS)
                cache.getparent().remove(cache)
            if mutation == "mixed_title":
                paragraph = chart.find("c:chart/c:title/c:tx/c:rich/a:p", NS)
                paragraph.append(copy.deepcopy(paragraph.find("a:r", NS)))
            if mutation == "linked_title":
                tx = chart.find("c:chart/c:title/c:tx", NS)
                tx.remove(tx.find("c:rich", NS))
                ref = etree.SubElement(tx, f"{{{CNS}}}strRef")
                etree.SubElement(ref, f"{{{CNS}}}f").text = "OldSheet!$A$20"
                cache = etree.SubElement(ref, f"{{{CNS}}}strCache")
                etree.SubElement(cache, f"{{{CNS}}}ptCount", val="1")
                etree.SubElement(
                    etree.SubElement(cache, f"{{{CNS}}}pt", idx="0"), f"{{{CNS}}}v"
                ).text = "Linked title"
            if mutation == "multiline_title":
                body = chart.find("c:chart/c:title/c:tx/c:rich", NS)
                paragraph = copy.deepcopy(body.find("a:p", NS))
                paragraph.find("a:r/a:t", NS).text = "Second line"
                body.append(paragraph)
            if mutation == "linked":
                rel = etree.fromstring(parts["ppt/charts/_rels/chart1.xml.rels"])
                rel[0].set("TargetMode", "External")
                rel[0].set("Target", "https://example.test/data.xlsx")
                parts["ppt/charts/_rels/chart1.xml.rels"] = etree.tostring(rel)
            parts["ppt/charts/chart1.xml"] = etree.tostring(chart)
        with zipfile.ZipFile(fixture, "w") as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
    project = call(core, "project.create", {"name": "Imported chart review"})["project"]["id"]
    document = call(
        core, "source.import", {"project_id": project, "path": str(fixture), "kind": "presentation"}
    )["document"]["id"]
    source = core._presentation_edits._current_revision_source(project, document)
    params = {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": source["revision_id"],
    }
    discovered = call(
        core, "presentation.workspace.list_charts", {**scope(params), "slide_number": 1}
    )
    return params, discovered


def draft(params, discovered):
    chart = discovered["items"][0]
    return {
        **params,
        "slide_id": discovered["slide_id"],
        "object_id": chart["object_id"],
        "expected_fingerprint": chart.get("fingerprint", "0" * 64),
        **copy.deepcopy(
            chart.get(
                "data",
                {
                    "title": "Imported pilot",
                    "chart_type": "column",
                    "categories": ["A"],
                    "series": [{"name": "Revenue", "values": [1]}],
                },
            )
        ),
    }


@pytest.mark.parametrize("kind", ["column", "bar", "line", "pie"])
@pytest.mark.parametrize("shared", [False, True])
def test_imported_chart_review_preserves_styles_notes_shared_parts_and_undo(
    tmp_path: Path, kind: str, shared: bool
):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path, kind, shared=shared)
        assert discovered["items"][0]["editable"]
        assert discovered["items"][0]["title_editable"]
        changed = draft(params, discovered)
        original = path(core, params).read_bytes()
        changed["title"] = "Reviewed imported pilot"
        changed["categories"][1] = "=2+3"
        changed["series"][0]["name"] = "https://example.test/Reviewed revenue"
        changed["series"][0]["values"][1] = 17.125
        plan = call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {
                **params,
                "summary": "Review imported chart",
                "operations": [
                    {
                        "type": "replace_chart_data",
                        "params": {k: v for k, v in changed.items() if k not in params},
                    }
                ],
            },
        )
        assert path(core, params).read_bytes() == original
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
            },
        )
        saved = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
                "reviewed_sha256": plan["sha256"],
                "confirm": True,
            },
        )
        params["expected_revision_id"] = saved["revision_id"]
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(path(core, params)) as after,
        ):
            differences = {
                name for name in before.namelist() if before.read(name) != after.read(name)
            }
            assert differences == {
                "ppt/slides/slide1.xml",
                "ppt/slides/_rels/slide1.xml.rels",
                "[Content_Types].xml",
            }
            assert after.read("ppt/charts/chart1.xml") == before.read("ppt/charts/chart1.xml")
            new_part = next(
                name
                for name in after.namelist()
                if name not in before.namelist() and name.endswith(".xml") and "/charts/" in name
            )
            old_root = etree.fromstring(before.read("ppt/charts/chart1.xml"))
            new_root = etree.fromstring(after.read(new_part))
            assert read_chart(new_root)["series"][0]["values"][1] == 17.125
            for root in [old_root, new_root]:
                for text in (
                    root.findall(".//c:f", NS)
                    + root.findall(".//c:v", NS)
                    + root.findall("c:chart/c:title/c:tx/c:rich//a:t", NS)
                ):
                    text.text = ""
                for auto in root.findall("c:externalData/c:autoUpdate", NS):
                    auto.set("val", "0")
            assert etree.tostring(old_root, method="c14n") == etree.tostring(
                new_root, method="c14n"
            )
            old_slide = etree.fromstring(before.read("ppt/slides/slide1.xml"))
            new_slide = etree.fromstring(after.read("ppt/slides/slide1.xml"))
            for root in [old_slide, new_slide]:
                root.find(".//c:chart", NS).set(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
                    "rIdSame",
                )
            assert etree.tostring(old_slide, method="c14n") == etree.tostring(
                new_slide, method="c14n"
            )
            bookpart = next(
                name
                for name in after.namelist()
                if name not in before.namelist() and name.endswith(".xlsx")
            )
            with zipfile.ZipFile(io.BytesIO(after.read(bookpart))) as workbook:
                assert b"<f>" not in workbook.read("xl/worksheets/sheet1.xml")
                assert b"=2+3" in workbook.read("xl/sharedStrings.xml")
        fresh = call(
            core, "presentation.workspace.list_charts", {**scope(params), "slide_number": 1}
        )
        assert fresh["items"][0]["object_id"] == discovered["items"][0]["object_id"]
        assert fresh["items"][0]["data"]["title"] == changed["title"]
        second = call(
            core, "presentation.workspace.list_charts", {**scope(params), "slide_number": 2}
        )
        assert second["items"][0]["data"]["title"] == "Imported pilot"
        call(
            core,
            "source.reindex",
            {"project_id": params["project_id"], "document_id": params["document_id"]},
        )
        with core._storage.project_database(params["project_id"]) as conn:
            assert any(
                "Reviewed imported pilot" in row["text"] and "17.125" in row["text"]
                for row in conn.execute("SELECT text FROM source_units")
            )
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize(
    "case", ["fingerprint", "revision", "object", "slide", "type", "categories", "series", "noop"]
)
def test_chart_edit_stale_bounds_and_noop_preserve_revision(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path)
        changed = draft(params, discovered)
        original = path(core, params).read_bytes()
        if case == "fingerprint":
            changed["expected_fingerprint"] = "0" * 64
        if case == "revision":
            changed["expected_revision_id"] = "00000000-0000-4000-8000-000000000001"
        if case == "object":
            changed["object_id"] = "00000000-0000-4000-8000-000000000001"
        if case == "slide":
            changed["slide_id"] = "00000000-0000-4000-8000-000000000001"
        if case == "type":
            changed["chart_type"] = "line"
        if case == "categories":
            changed["categories"].append("Fourth")
            for series in changed["series"]:
                series["values"].append(20)
        if case == "series":
            changed["series"].append(copy.deepcopy(changed["series"][0]))
        result = call(core, "presentation.workspace.replace_chart_data", changed, ok=case == "noop")
        if case == "noop":
            assert (
                result["changed"] is False
                and result["revision_id"] == params["expected_revision_id"]
            )
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize(
    "case", ["linked", "stacked", "extension", "missing_cache", "linked_title"]
)
def test_unsupported_chart_is_discovered_readonly_and_cannot_edit(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path, mutation=case)
        assert not discovered["items"][0]["editable"]
        assert discovered["items"][0]["reason"]
        original = path(core, params).read_bytes()
        call(core, "presentation.workspace.replace_chart_data", draft(params, discovered), ok=False)
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


def test_mixed_title_is_preserved_while_values_remain_editable(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path, mutation="mixed_title")
        assert discovered["items"][0]["editable"] and not discovered["items"][0]["title_editable"]
        changed = draft(params, discovered)
        changed["title"] = "Unsafe title replacement"
        call(core, "presentation.workspace.replace_chart_data", changed, ok=False)
        changed["title"] = discovered["items"][0]["data"]["title"]
        changed["series"][0]["values"][1] = 14
        call(core, "presentation.workspace.replace_chart_data", changed)
    finally:
        core.close()


def test_multiline_title_roundtrips_native_review_and_style(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path, mutation="multiline_title")
        assert discovered["items"][0]["title_editable"]
        assert discovered["items"][0]["data"]["title"] == "Imported pilot\nSecond line"
        changed = draft(params, discovered)
        changed["title"] = "Reviewed pilot\nNew second line"
        saved = call(core, "presentation.workspace.replace_chart_data", changed)
        params["expected_revision_id"] = saved["revision_id"]
        fresh = call(
            core, "presentation.workspace.list_charts", {**scope(params), "slide_number": 1}
        )
        assert fresh["items"][0]["data"]["title"] == changed["title"]
        with zipfile.ZipFile(path(core, params)) as archive:
            part = next(
                name
                for name in archive.namelist()
                if name.startswith("ppt/charts/copilot-") and name.endswith(".xml")
            )
            chart = etree.fromstring(archive.read(part))
            fonts = chart.findall("c:chart/c:title/c:tx/c:rich/a:p/a:r/a:rPr", NS)
            assert len(fonts) == 2 and all(
                font.get("sz") == "2300" and font.get("b") == "1" for font in fonts
            )
    finally:
        core.close()


def test_failed_later_operation_does_not_publish_chart_edit(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path)
        changed = draft(params, discovered)
        changed["series"][0]["values"][0] = 30
        original = path(core, params).read_bytes()
        operation = {
            "type": "replace_chart_data",
            "params": {k: v for k, v in changed.items() if k not in params},
        }
        invalid = copy.deepcopy(operation)
        invalid["params"]["object_id"] = "00000000-0000-4000-8000-000000000001"
        call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {**params, "summary": "Atomic chart review", "operations": [operation, invalid]},
            ok=False,
        )
        assert path(core, params).read_bytes() == original
        assert not list(path(core, params).parent.glob(".edit-stage-*.pptx"))
    finally:
        core.close()


def test_chart_edit_rejects_corrupt_verified_source(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, discovered = imported(core, tmp_path)
        changed = draft(params, discovered)
        changed["series"][0]["values"][0] = 30
        snapshot = path(core, params)
        corrupt = snapshot.read_bytes() + b"corrupt"
        snapshot.write_bytes(corrupt)
        call(core, "presentation.workspace.replace_chart_data", changed, ok=False)
        assert snapshot.read_bytes() == corrupt
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
    finally:
        core.close()
