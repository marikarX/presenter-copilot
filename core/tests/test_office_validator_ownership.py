"""Office acceptance must not mutate or close user-owned application state."""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation

from presenter_core.tools.presentation_corpus_roundtrip import (
    AcceptanceFailure,
    PowerPointOpenValidator,
)


def _source(tmp_path: Path) -> Path:
    path = tmp_path / "original.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    return path


class SharedOffice:
    def __init__(
        self,
        original: Path,
        *,
        fail_export: bool = False,
        fail_close: bool = False,
        mutate_copy: bool = False,
    ) -> None:
        self.original = original
        self.opened: list[Path] = []
        self.closed: list[Path] = []
        self.fail_export = fail_export
        self.fail_close = fail_close
        self.mutate_copy = mutate_copy
        self.Presentations = SimpleNamespace(Open=self.open)

    def __setattr__(self, key: str, value: Any) -> None:
        if key in {"WindowState", "DisplayAlerts", "AutomationSecurity", "Visible"}:
            raise AssertionError("Do not change a shared Office setting")
        object.__setattr__(self, key, value)

    def Quit(self) -> None:
        raise AssertionError("Do not Quit a shared Office application")

    def open(self, path: str, **kwargs: Any) -> Any:
        snapshot = Path(path)
        assert snapshot != self.original
        assert snapshot.read_bytes() == self.original.read_bytes()
        assert kwargs == {"ReadOnly": True, "Untitled": False, "WithWindow": False}
        self.opened.append(snapshot)

        def export(destination: str, format: str, width: int, height: int) -> None:
            assert format == "PNG"
            if self.fail_export:
                raise RuntimeError("Synthetic export failure")
            Image.new("RGB", (width, height), "blue").save(destination)

        def close() -> None:
            self.closed.append(snapshot)
            if self.mutate_copy:
                snapshot.write_bytes(snapshot.read_bytes() + b"changed")
            if self.fail_close:
                raise RuntimeError("Synthetic close failure")

        return SimpleNamespace(
            Slides=SimpleNamespace(Count=1, Item=lambda _: SimpleNamespace(Export=export)),
            Close=close,
        )


def _validator(office: SharedOffice) -> PowerPointOpenValidator:
    validator = PowerPointOpenValidator.__new__(PowerPointOpenValidator)
    validator._application = office
    validator._com_apartment = None
    return validator


def test_construct_release_never_sets_global_properties_or_quits(
    tmp_path: Path, monkeypatch: Any
) -> None:
    office = SharedOffice(_source(tmp_path))
    calls: list[str] = []
    monkeypatch.setattr(sys, "platform", "win32")

    def initialize() -> Any:
        calls.append("init")
        return SimpleNamespace(CoUninitialize=lambda: calls.append("release"))

    monkeypatch.setattr(
        "presenter_core.tools.presentation_corpus_roundtrip._initialize_com_apartment",
        initialize,
    )
    monkeypatch.setitem(
        sys.modules, "win32com.client", SimpleNamespace(DispatchEx=lambda _: office)
    )
    validator = PowerPointOpenValidator()
    validator.close()
    validator.close()
    assert calls == ["init", "release"]
    assert office.opened == []


def test_open_and_render_use_distinct_owned_copies(tmp_path: Path) -> None:
    source = _source(tmp_path)
    before = source.read_bytes()
    office = SharedOffice(source)
    validator = _validator(office)
    assert validator.open_read_only(source, 1) == 1
    rendered = validator._render_slides(office, source, tmp_path / "renders", 1, [(12, 6)])
    assert Image.open(rendered[0]).size == (12, 6)
    validator.close()
    assert len(office.opened) == 2 and len(set(office.opened)) == 2
    assert office.closed == office.opened
    assert all(not path.exists() for path in office.opened)
    assert source.read_bytes() == before


@pytest.mark.parametrize(
    "failure,code",
    [
        ("export", "POWERPOINT_RENDER_FAILED"),
        ("close", "POWERPOINT_PRESENTATION_CLOSE_FAILED"),
        ("mutation", "POWERPOINT_CHANGED_READ_ONLY_EXPORT"),
    ],
)
def test_failure_closes_only_the_owned_copy(tmp_path: Path, failure: str, code: str) -> None:
    source = _source(tmp_path)
    before = source.read_bytes()
    office = SharedOffice(
        source,
        fail_export=failure == "export",
        fail_close=failure == "close",
        mutate_copy=failure == "mutation",
    )
    with pytest.raises(AcceptanceFailure) as caught:
        PowerPointOpenValidator._render_slides(office, source, tmp_path / "render", 1, [(12, 6)])
    assert caught.value.code == code
    assert office.closed == office.opened and len(office.closed) == 1
    assert source.read_bytes() == before


@pytest.mark.parametrize("part", ["ppt/vbaProject.bin", "ppt/activeX/activeX1.xml"])
def test_active_content_is_blocked_before_open_without_global_security_changes(
    tmp_path: Path, part: str
) -> None:
    source = _source(tmp_path)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr(part, b"untrusted")
    office = SharedOffice(source)
    with pytest.raises(AcceptanceFailure) as caught:
        _validator(office).open_read_only(source, 1)
    assert caught.value.code == "POWERPOINT_ACTIVE_CONTENT_BLOCKED"
    assert office.opened == []


def test_undeclared_uri_link_is_blocked_before_office_open(tmp_path: Path) -> None:
    source = _source(tmp_path)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr(
            "ppt/_rels/untrusted.xml.rels",
            """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
        <Relationship Id="unsafe" Type="image" Target="https://example.invalid/image.png" />
        </Relationships>""",
        )
    office = SharedOffice(source)
    with pytest.raises(AcceptanceFailure) as caught:
        _validator(office).open_read_only(source, 1)
    assert caught.value.code == "POWERPOINT_EXTERNAL_CONTENT_BLOCKED"
    assert office.opened == []


@pytest.mark.parametrize("result", [0, 1, -2147417850])
def test_com_initialization_owns_native_reference(result: int, monkeypatch: Any) -> None:
    from presenter_core.tools import presentation_corpus_roundtrip as tool

    calls: list[tuple[Any, ...]] = []

    def initialize(*args: Any) -> int:
        calls.append(args)
        return result

    def release() -> None:
        calls.append(("release",))

    apartment = SimpleNamespace(CoInitializeEx=initialize, CoUninitialize=release)
    monkeypatch.setattr(tool.ctypes, "WinDLL", lambda _: apartment, raising=False)
    if result < 0:
        with pytest.raises(OSError):
            tool._initialize_com_apartment()
        assert calls == [(None, 2)]
    else:
        assert tool._initialize_com_apartment() is apartment
        assert calls == [(None, 2)]
        apartment.CoUninitialize()
        assert calls == [(None, 2), ("release",)]


def test_failed_dispatch_releases_owned_apartment(tmp_path: Path, monkeypatch: Any) -> None:
    calls: list[str] = []
    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setattr(
        "presenter_core.tools.presentation_corpus_roundtrip._initialize_com_apartment",
        lambda: SimpleNamespace(CoUninitialize=lambda: calls.append("release")),
    )

    def fail_dispatch(_: str) -> Any:
        raise RuntimeError("Synthetic server unavailable")

    monkeypatch.setitem(sys.modules, "win32com.client", SimpleNamespace(DispatchEx=fail_dispatch))
    with pytest.raises(AcceptanceFailure, match="POWERPOINT_AUTOMATION_UNAVAILABLE"):
        PowerPointOpenValidator()
    assert calls == ["release"]
