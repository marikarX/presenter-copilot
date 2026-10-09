"""Reviewed shape fill transparency (solid fill a:alpha)."""

import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService
from presenter_core.presentation.fidelity import DRAWING_NS, PRESENTATION_NS

A = f"{{{DRAWING_NS}}}"
P = f"{{{PRESENTATION_NS}}}"


def _deck(path: Path, *, alpha: str | None = None) -> Path:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    filled = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(2))
    filled.name = "Filled"
    filled.fill.solid()
    filled.fill.fore_color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    if alpha is not None:
        color = filled._element.spPr.find(f"{A}solidFill/{A}srgbClr")
        etree.SubElement(color, f"{A}alpha", val=alpha)
    empty = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(1), Inches(3), Inches(2))
    empty.name = "Empty"
    empty.fill.background()
    deck.save(path)
    return path


class _Shapes:
    def __init__(self, tmp_path: Path, **options) -> None:
        self.core = CoreService(data_root=tmp_path / "data")
        self.project = _project_id(self.core)
        source = _deck(tmp_path / "shapes.pptx", **options)
        self.document = _import_presentation(self.core, self.project, source)["id"]
        self.revision = _open(self.core, self.project, self.document)["revision_id"]

    def shape(self, name: str) -> tuple[dict, str]:
        inventory = _call(
            self.core,
            "targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": self.project,
                "document_id": self.document,
                "revision_id": self.revision,
                "slide_number": 1,
            },
        )["result"]
        return next(o for o in inventory["objects"] if o["name"] == name), inventory["slide_id"]

    def style(self, name: str, replacement: dict) -> dict:
        target, slide_id = self.shape(name)
        reply = _call(
            self.core,
            "style",
            "presentation.workspace.format_shape_style",
            {
                "project_id": self.project,
                "document_id": self.document,
                "expected_revision_id": self.revision,
                "slide_id": slide_id,
                "object_id": target["object_id"],
                "expected_style_fingerprint": target["shape_style"]["fingerprint"],
                "replacement_style": replacement,
            },
        )
        if reply["ok"]:
            self.revision = reply["result"]["revision_id"]
        return reply

    def color(self, name: str) -> etree._Element:
        with self.core._storage.project_database(self.project) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?",
                (self.revision,),
            ).fetchone()
        path = self.core._storage.paths.snapshot_path(
            self.project, row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(path) as archive:
            root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        shape = next(
            sp for sp in root.iter(f"{P}sp") if sp.find(f"{P}nvSpPr/{P}cNvPr").get("name") == name
        )
        return shape.find(f"{P}spPr/{A}solidFill/{A}srgbClr")


def test_transparency_sets_survives_recolor_and_clears(tmp_path: Path) -> None:
    shapes = _Shapes(tmp_path)
    try:
        assert shapes.shape("Filled")[0]["shape_style"]["fill_transparency_percent"] == 0
        assert shapes.style("Filled", {"fill_transparency_percent": 40})["ok"]
        assert shapes.color("Filled").find(f"{A}alpha").get("val") == "60000"
        assert shapes.shape("Filled")[0]["shape_style"]["fill_transparency_percent"] == 40
        recolor = shapes.style("Filled", {"fill_color": "#AA2200"})
        assert recolor["ok"], recolor
        color = shapes.color("Filled")
        assert (color.get("val"), color.find(f"{A}alpha").get("val")) == ("AA2200", "60000")
        assert shapes.style("Filled", {"fill_transparency_percent": 0})["ok"]
        assert shapes.color("Filled").find(f"{A}alpha") is None
    finally:
        shapes.core.close()


def test_unfilled_shape_needs_a_color_first_but_both_can_be_set_together(tmp_path: Path) -> None:
    shapes = _Shapes(tmp_path)
    try:
        assert shapes.shape("Empty")[0]["shape_style"]["fill_transparency_percent"] is None
        refused = shapes.style("Empty", {"fill_transparency_percent": 50})
        assert not refused["ok"] and refused["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert shapes.style("Empty", {"fill_color": "#00AA00"})["ok"]
        assert shapes.shape("Empty")[0]["shape_style"]["fill_transparency_percent"] == 0
        both = shapes.style("Empty", {"fill_color": "#0000AA", "fill_transparency_percent": 25})
        assert both["ok"], both
        assert shapes.color("Empty").find(f"{A}alpha").get("val") == "75000"
    finally:
        shapes.core.close()


def test_imported_alpha_fill_is_editable_and_unusual_alpha_stays_read_only(
    tmp_path: Path,
) -> None:
    shapes = _Shapes(tmp_path, alpha="30000")
    try:
        style = shapes.shape("Filled")[0]["shape_style"]
        assert style["fill_editable"] and style["fill_transparency_percent"] == 70
    finally:
        shapes.core.close()
    odd = _Shapes(tmp_path / "odd", alpha="33333")
    try:
        assert odd.shape("Filled")[0]["shape_style"]["fill_editable"] is False
    finally:
        odd.core.close()
