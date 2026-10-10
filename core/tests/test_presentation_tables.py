from __future__ import annotations

import io
import uuid
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from test_presentation_edit_planning import call
from test_presentation_native_layouts import path, scope
from test_presentation_transactions import context

from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.presentation.tables import inspect_table, validate_cells
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def table_request(params, operations):
    return {
        **params,
        "slide_id": operations[0]["params"]["slide_id"],
        "geometry": {
            "x_emu": 457200,
            "y_emu": 2743200,
            "width_emu": 7315200,
            "height_emu": 2743200,
        },
        "cells": [
            ["Metric", "Current", "Target"],
            ["Revenue", "$10M", "$12M"],
            ["Risk", "Measured", "Bounded"],
        ],
        "header_row": True,
    }


def test_table_create_edit_review_preserves_parts_retrieval_and_undo(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, operations = context(core)
        original = path(core, params).read_bytes()
        request = table_request(params, operations)
        plan = call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {
                **params,
                "summary": "Insert native table",
                "operations": [
                    {
                        "type": "create_table",
                        "params": {k: v for k, v in request.items() if k not in params},
                    }
                ],
            },
        )
        assert path(core, params).read_bytes() == original
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
            },
        )
        saved = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
                "reviewed_sha256": plan["sha256"],
                "confirm": True,
            },
        )
        created_revision = saved["revision_id"]
        params["expected_revision_id"] = created_revision
        tables = call(
            core, "presentation.workspace.list_tables", {**scope(params), "slide_number": 1}
        )
        assert tables["slide_id"] == request["slide_id"]
        table = tables["items"][0]
        assert table["editable"] and table["cells"] == request["cells"]
        inserted_bytes = path(core, params).read_bytes()
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(io.BytesIO(inserted_bytes)) as after,
        ):
            assert before.namelist() == after.namelist()
            assert [n for n in before.namelist() if before.read(n) != after.read(n)] == [
                "ppt/slides/slide1.xml"
            ]
        new_cells = [list(row) for row in table["cells"]]
        new_cells[1][2] = "$13M\nPilot target"
        edit = {
            **params,
            "slide_id": request["slide_id"],
            "object_id": table["object_id"],
            "expected_fingerprint": table["fingerprint"],
            "cells": new_cells,
        }
        saved = call(core, "presentation.workspace.replace_table_cells", edit)
        params["expected_revision_id"] = saved["revision_id"]
        fresh = call(
            core, "presentation.workspace.list_tables", {**scope(params), "slide_number": 1}
        )["items"][0]
        assert fresh["object_id"] == table["object_id"] and fresh["cells"] == new_cells
        assert fresh["fingerprint"] != table["fingerprint"]
        call(
            core,
            "presentation.workspace.replace_table_cells",
            {**edit, "expected_revision_id": saved["revision_id"]},
            ok=False,
        )
        no_op = call(
            core,
            "presentation.workspace.replace_table_cells",
            {
                **edit,
                "expected_revision_id": saved["revision_id"],
                "expected_fingerprint": fresh["fingerprint"],
            },
        )
        assert not no_op["changed"]
        call(
            core,
            "source.reindex",
            {"project_id": params["project_id"], "document_id": params["document_id"]},
        )
        with core._storage.project_database(params["project_id"]) as connection:
            assert any(
                "Pilot target" in row[0]
                for row in connection.execute("SELECT text FROM source_units")
            )
        with (
            zipfile.ZipFile(io.BytesIO(inserted_bytes)) as before,
            zipfile.ZipFile(path(core, params)) as after,
        ):
            old = etree.fromstring(before.read("ppt/slides/slide1.xml"))
            new = etree.fromstring(after.read("ppt/slides/slide1.xml"))
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for root in [old, new]:
                for node in root.findall(".//a:tc/a:txBody", ns):
                    for para in node.findall("a:p", ns):
                        node.remove(para)
            assert etree.tostring(old, method="c14n") == etree.tostring(new, method="c14n")
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == inserted_bytes
        params["expected_revision_id"] = created_revision
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize(
    "cells",
    [
        [],
        [["x"], ["a", "b"]],
        [["x"]] * 13,
        [["x"] * 9],
        [["x" * 501]],
        [["x\n" * 11]],
        [["\x00"]],
        [[True]],
        [["\uffff"]],
        [["x" * 500] * 8] * 3,
    ],
)
def test_table_content_limits(cells):
    with pytest.raises(CoreDomainError):
        validate_cells(cells)


@pytest.mark.parametrize(
    "case", ["foreign_slide", "outside", "boolean_geometry", "bad_header", "unknown_field", "stale"]
)
def test_bad_table_creation_never_publishes(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, operations = context(core)
        request = table_request(params, operations)
        if case == "foreign_slide":
            request["slide_id"] = str(uuid.uuid4())
        if case == "outside":
            request["geometry"]["x_emu"] = 999999999
        if case == "boolean_geometry":
            request["geometry"]["x_emu"] = True
        if case == "bad_header":
            request["header_row"] = "yes"
        if case == "unknown_field":
            request["path"] = "foreign"
        if case == "stale":
            request["expected_revision_id"] = str(uuid.uuid4())
        original = path(core, params).read_bytes()
        call(core, "presentation.workspace.create_table", request, ok=False)
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize("case", ["merge", "hyperlink", "mixed_runs", "extension", "oversized"])
def test_complex_tables_are_read_only(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, ops = context(core)
        call(core, "presentation.workspace.create_table", table_request(params, ops))
        with zipfile.ZipFile(path(core, params)) as archive:
            root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        ns = {
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        }
        frame = root.find("p:cSld/p:spTree/p:graphicFrame", ns)
        cell = frame.find("a:graphic/a:graphicData/a:tbl/a:tr/a:tc", ns)
        if case == "merge":
            cell.set("gridSpan", "2")
        if case == "hyperlink":
            etree.SubElement(
                cell.find("a:txBody/a:p/a:r/a:rPr", ns), "{" + ns["a"] + "}hlinkClick"
            ).set(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
                "rIdForeign",
            )
        if case == "mixed_runs":
            import copy

            cell.find("a:txBody/a:p", ns).append(copy.deepcopy(cell.find("a:txBody/a:p/a:r", ns)))
        if case == "extension":
            etree.SubElement(frame, "{" + ns["p"] + "}extLst")
        if case == "oversized":
            cell.find("a:txBody/a:p/a:r/a:t", ns).text = "x" * 501
        info = inspect_table(frame)
        if case in {"mixed_runs", "hyperlink"}:
            # Rich-text and linked cells are locked individually; the rest stays editable.
            assert info["editable"] and info["locked_cells"] == [[0, 0]]
        else:
            assert not info["editable"]
    finally:
        core.close()


def test_imported_table_retains_custom_cell_style_and_dimensions(tmp_path: Path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        native = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(3)).table
        native.columns[0].width = Inches(2)
        native.columns[1].width = Inches(4)
        for row in range(2):
            for col in range(2):
                cell = native.cell(row, col)
                cell.text = f"Original {row} {col}"
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(20, 40, 80)
                cell.text_frame.paragraphs[0].runs[0].font.size = Pt(23)
                cell.text_frame.paragraphs[0].runs[0].font.italic = True
        fixture = tmp_path / "imported-table.pptx"
        deck.save(fixture)
        project = call(core, "project.create", {"name": "Imported native table"})["project"]["id"]
        document = call(
            core,
            "source.import",
            {"project_id": project, "path": str(fixture), "kind": "presentation"},
        )["document"]["id"]
        source = core._presentation_edits._current_revision_source(project, document)
        params = {
            "project_id": project,
            "document_id": document,
            "expected_revision_id": source["revision_id"],
        }
        result = call(
            core, "presentation.workspace.list_tables", {**scope(params), "slide_number": 1}
        )
        table = result["items"][0]
        assert table["editable"]
        original = path(core, params).read_bytes()
        changed = table["cells"]
        changed[1][1] = "Reviewed imported value"
        saved = call(
            core,
            "presentation.workspace.replace_table_cells",
            {
                **params,
                "slide_id": result["slide_id"],
                "object_id": table["object_id"],
                "expected_fingerprint": table["fingerprint"],
                "cells": changed,
            },
        )
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(path(core, params)) as after,
        ):
            old = etree.fromstring(before.read("ppt/slides/slide1.xml"))
            new = etree.fromstring(after.read("ppt/slides/slide1.xml"))
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for root in [old, new]:
                for text in root.findall(".//a:tc/a:txBody/a:p/a:r/a:t", ns):
                    text.text = ""
            assert etree.tostring(old, method="c14n") == etree.tostring(new, method="c14n")
        params["expected_revision_id"] = saved["revision_id"]
        fresh = call(
            core, "presentation.workspace.list_tables", {**scope(params), "slide_number": 1}
        )["items"][0]
        assert (
            fresh["cells"][1][1] == "Reviewed imported value"
            and fresh["object_id"] == table["object_id"]
        )
    finally:
        core.close()


@pytest.mark.parametrize("case", ["corrupt", "changed_head"])
def test_table_discovery_rejects_changed_snapshot_or_head(tmp_path: Path, monkeypatch, case: str):
    from presenter_core.presentation import tables

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, operations = context(core)
        saved = call(core, "presentation.workspace.create_table", table_request(params, operations))
        params["expected_revision_id"] = saved["revision_id"]
        if case == "corrupt":
            source = path(core, params)
            source.write_bytes(source.read_bytes() + b"tampered")
        else:
            original = tables.inspect_table
            changed = False

            def advance(frame):
                nonlocal changed
                result = original(frame)
                if not changed:
                    changed = True
                    call(
                        core,
                        "presentation.workspace.replace_speaker_notes",
                        {**params, **operations[1]["params"]},
                    )
                return result

            monkeypatch.setattr(tables, "inspect_table", advance)
        error = call(
            core,
            "presentation.workspace.list_tables",
            {**scope(params), "slide_number": 1},
            ok=False,
        )
        assert error["code"] in {"SOURCE_SNAPSHOT_CHANGED", "PRESENTATION_REVISION_CONFLICT"}
    finally:
        core.close()
