"""Insert video and audio: PowerPoint-style media pictures with a poster frame."""

import math
import struct
import wave
import zipfile
from pathlib import Path

import av
import numpy as np
from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _video(path: Path) -> None:
    with av.open(str(path), "w") as container:
        stream = container.add_stream("libx264", rate=10)
        stream.width, stream.height, stream.pix_fmt = 64, 48, "yuv420p"
        for index in range(10):
            pixels = np.zeros((48, 64, 3), dtype=np.uint8)
            pixels[:, :, 1] = 200
            pixels[:, : 6 * (index + 1), 0] = 255
            frame = av.VideoFrame.from_ndarray(pixels, format="rgb24")
            for packet in stream.encode(frame):
                container.mux(packet)
        for packet in stream.encode():
            container.mux(packet)


def _audio(path: Path) -> None:
    with wave.open(str(path), "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(2)
        output.setframerate(8000)
        output.writeframes(
            b"".join(
                struct.pack("<h", int(8000 * math.sin(2 * math.pi * 440 * n / 8000)))
                for n in range(8000)
            )
        )


def test_insert_video_and_audio_as_editable_media(tmp_path: Path) -> None:
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(tmp_path / "deck.pptx")
    _video(tmp_path / "clip.mp4")
    _audio(tmp_path / "tone.wav")
    (tmp_path / "fake.mp4").write_bytes(b"not a video")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, tmp_path / "deck.pptx")["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def inventory() -> dict:
            return _call(
                core,
                "l",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]

        slide_id = inventory()["slide_id"]

        def insert(path: Path) -> dict:
            return _call(
                core,
                "m",
                "presentation.workspace.insert_media",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": slide_id,
                    "path": str(path),
                },
            )

        video = insert(tmp_path / "clip.mp4")
        assert video["ok"], video.get("error")
        assert video["result"]["media_kind"] == "video"
        assert 0.5 < video["result"]["duration_seconds"] < 2
        revision = video["result"]["revision_id"]
        audio = insert(tmp_path / "tone.wav")
        assert audio["ok"], audio.get("error")
        revision = audio["result"]["revision_id"]

        objects = {o["name"]: o for o in inventory()["objects"]}
        clip = objects["Video 1"]
        # 64x48 scaled to 60% of the slide height, keeping its aspect.
        ratio = clip["geometry"]["width_emu"] / clip["geometry"]["height_emu"]
        assert abs(ratio - 64 / 48) < 0.001
        assert clip["geometry_editable"] and clip["delete_editable"]
        assert objects["Audio 1"]["geometry"]["width_emu"] == 457200

        with core._storage.project_database(project) as connection:
            snapshot = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, snapshot)) as archive:
            slide = archive.read("ppt/slides/slide1.xml").decode()
            names = archive.namelist()
            stored = next(i for i in archive.infolist() if i.filename.endswith(".mp4"))
        assert "<a:videoFile" in slide and "<a:audioFile" in slide and "p14:media" in slide
        assert stored.compress_type == zipfile.ZIP_STORED
        assert stored.file_size == (tmp_path / "clip.mp4").stat().st_size
        assert sum(n.startswith("ppt/media/copilot-poster-") for n in names) == 2

        moved = _call(
            core,
            "t",
            "presentation.workspace.transform_object",
            {
                **scope,
                "expected_revision_id": revision,
                "slide_id": slide_id,
                "object_id": clip["object_id"],
                "expected_geometry": clip["geometry"],
                "replacement_geometry": {**clip["geometry"], "x_emu": 0, "y_emu": 0},
            },
        )
        assert moved["ok"], moved.get("error")
        revision = moved["result"]["revision_id"]

        refused = insert(tmp_path / "fake.mp4")
        assert refused["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert insert(tmp_path / "deck.pptx")["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        relative = insert(Path("clip.mp4"))
        assert relative["error"]["code"] == "INVALID_REQUEST"
    finally:
        core.close()
