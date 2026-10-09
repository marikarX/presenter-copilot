"""Measured mixed-run fitting stages atomically and restores original packages."""

import io
import zipfile

import pytest
from lxml import etree  # type: ignore[import-untyped]
from pptx import Presentation
from pptx.util import Inches, Pt
from test_chart_legend import core_at
from test_presentation_edit_planning import call
from test_presentation_native_layouts import path, scope
from test_selection_actions import save

from presenter_core.errors import CoreDomainError
from presenter_core.presentation.fidelity import _NS
from presenter_core.presentation.text_fit import apply_fit, calculate, measure
from presenter_core.presentation.text_fit_editing import unmanaged
from presenter_core.presentation.text_inheritance import TextInheritance


@pytest.fixture(autouse=True)
def metrics(monkeypatch):
    class Font:
        def __init__(self, size):
            self.size = size

        def getlength(self, text, **kwargs):
            return len(text) * self.size / 2

    monkeypatch.setattr(
        "presenter_core.presentation.text_fit.installed_font",
        lambda family, bold, italic, size: Font(size),
    )


def fixture(core, directory, kind="mixed"):
    deck = Presentation()
    for index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.5))
        box.name = "Fit target"
        paragraph = box.text_frame.paragraphs[0]
        first = paragraph.add_run()
        first.text = "A reviewed pilot needs careful measurement "
        first.font.name, first.font.size, first.font.bold = "Arial", Pt(32), True
        second = paragraph.add_run()
        second.text = "before a decision."
        second.font.name, second.font.size, second.font.italic = "Arial", Pt(20), True
        if kind in {"bullet", "numbered"}:
            props = paragraph._p.get_or_add_pPr()
            props.set("marL", "457200")
            props.set("indent", "-228600")
            if kind == "numbered":
                etree.SubElement(props, f"{{{_NS['a']}}}buSzPct", val="50000")
            etree.SubElement(props, f"{{{_NS['a']}}}buFont", typeface="Arial")
            if kind == "bullet":
                etree.SubElement(props, f"{{{_NS['a']}}}buChar", char="•")
            else:
                etree.SubElement(props, f"{{{_NS['a']}}}buAutoNum", type="arabicPeriod")
        if kind == "cached":
            props = box._element.find("p:txBody/a:bodyPr", _NS)
            for node in list(props):
                props.remove(node)
            etree.SubElement(
                props, f"{{{_NS['a']}}}normAutofit", fontScale="60000", lnSpcReduction="20000"
            )
        slide.notes_slide.notes_text_frame.text = f"Retained notes {index}"
    file = directory / "text-fit.pptx"
    deck.save(file)
    project = call(core, "project.create", {"name": "Native fitting"})["project"]["id"]
    document = call(
        core, "source.import", {"project_id": project, "path": str(file), "kind": "presentation"}
    )["document"]["id"]
    revision = core._presentation_edits._current_revision_source(project, document)["revision_id"]
    return {"project_id": project, "document_id": document, "expected_revision_id": revision}


def target(core, params):
    info = call(
        core, "presentation.workspace.list_text_targets", {**scope(params), "slide_number": 1}
    )
    obj = info["targets"][0]
    with zipfile.ZipFile(path(core, params)) as archive:
        reader = TextInheritance(archive, "ppt/slides/slide1.xml")
        shape = reader.slide.find("p:cSld/p:spTree/p:sp", _NS)
        result = reader.resolve(shape)
    return (
        {
            "slide_id": info["slide_id"],
            "object_id": obj["object_id"],
            "expected_fingerprint": result.fingerprint,
        },
        result,
        shape,
    )


def prepare(core, params, strategy="grow", minimum=6):
    picked, _, _ = target(core, params)
    return core._presentation_edit_plans.prepare(
        {
            **params,
            "summary": "Review native fit",
            "operations": [
                {
                    "type": "fit_text_shape",
                    "params": {**picked, "strategy": strategy, "minimum_size_pt": minimum},
                }
            ],
        }
    )


@pytest.mark.parametrize("strategy", ["grow", "shrink"])
def test_mixed_run_fitting_preserves_parts_and_exact_undo(tmp_path, strategy):
    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path)
        original = path(core, params).read_bytes()
        _, before, shape = target(core, params)
        prior = unmanaged(shape)
        plan = prepare(core, params, strategy)
        assert path(core, params).read_bytes() == original
        save(core, params, plan)
        _, after, fitted = target(core, params)
        assert unmanaged(fitted) == prior
        assert measure(after) <= int(after.transform.find("a:ext", _NS).get("cy")) / 12700 + 0.02
        if strategy == "grow":
            assert int(after.transform.find("a:ext", _NS).get("cy")) > int(
                before.transform.find("a:ext", _NS).get("cy")
            )
            assert after.paragraphs[0]["runs"] == before.paragraphs[0]["runs"]
        else:
            assert etree.tostring(after.transform) == etree.tostring(before.transform)
            sizes = [int(r["size_hundredths"]) for r in after.paragraphs[0]["runs"]]
            assert abs(sizes[0] / sizes[1] - 32 / 20) < 0.01
        with (
            zipfile.ZipFile(io.BytesIO(original)) as old,
            zipfile.ZipFile(path(core, params)) as new,
        ):
            assert [n for n in old.namelist() if old.read(n) != new.read(n)] == [
                "ppt/slides/slide1.xml"
            ]
        undo = call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        params["expected_revision_id"] = undo["revision_id"]
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize(
    "case", ["bounds", "minimum", "strategy", "font", "vertical", "field", "stale"]
)
def test_unsupported_or_unfittable_inputs_reject_before_mutation(tmp_path, case):
    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path)
        picked, resolved, shape = target(core, params)
        original = path(core, params).read_bytes()
        if case == "stale":
            picked["expected_fingerprint"] = "0" * 64
            with pytest.raises(CoreDomainError):
                core._presentation_edit_plans.prepare(
                    {
                        **params,
                        "summary": "Stale fit",
                        "operations": [
                            {
                                "type": "fit_text_shape",
                                "params": {**picked, "strategy": "grow", "minimum_size_pt": 6},
                            }
                        ],
                    }
                )
        else:
            if case == "bounds":
                resolved.transform.find("a:off", _NS).set("y", "6500000")
            if case == "font":
                resolved.body.find("a:p/a:r/a:rPr", _NS).set("sz", "invalid")
            if case == "vertical":
                resolved.body.find("a:bodyPr", _NS).set("vert", "vert")
            if case == "field":
                etree.SubElement(resolved.body.find("a:p", _NS), f"{{{_NS['a']}}}fld")
            with pytest.raises(CoreDomainError):
                calculate(
                    resolved,
                    [] if case == "strategy" else "shrink" if case == "minimum" else "grow",
                    19 if case == "minimum" else 6,
                    6858000,
                )
        assert path(core, params).read_bytes() == original
        assert etree.tostring(shape) == etree.tostring(target(core, params)[2])
    finally:
        core.close()


@pytest.mark.parametrize("case", ["bullet", "numbered", "cached"])
def test_bullets_and_cached_scaling_are_measured_then_normalized(tmp_path, case):
    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path)
        _, resolved, shape = target(core, params)
        paragraph = resolved.body.find("a:p", _NS)
        pp = paragraph.find("a:pPr", _NS)
        if case in {"bullet", "numbered"}:
            pp.set("marL", "457200")
            pp.set("indent", "-228600")
            for node in list(pp):
                if etree.QName(node).localname.startswith("bu"):
                    pp.remove(node)
            if case == "bullet":
                etree.SubElement(pp, f"{{{_NS['a']}}}buChar", char="•")
            else:
                etree.SubElement(pp, f"{{{_NS['a']}}}buAutoNum", type="arabicPeriod")
                etree.SubElement(pp, f"{{{_NS['a']}}}buSzPct", val="50000")
        else:
            bp = resolved.body.find("a:bodyPr", _NS)
            for node in list(bp):
                bp.remove(node)
            etree.SubElement(
                bp, f"{{{_NS['a']}}}normAutofit", fontScale="60000", lnSpcReduction="20000"
            )
        original = etree.tostring(shape)
        result = calculate(resolved, "shrink", 6, 6858000)
        assert result.font_scale <= (0.6 if case == "cached" else 1)
        assert etree.tostring(shape) == original
        apply_fit(shape, resolved, result)
        assert shape.find("p:txBody/a:bodyPr/a:noAutofit", _NS) is not None
        assert shape.find("p:txBody/a:bodyPr", _NS).get("wrap") == "square"
    finally:
        core.close()


@pytest.mark.parametrize("kind", ["mixed", "bullet", "numbered", "cached"])
def test_fit_catalog_is_read_only_and_bound_to_revision(tmp_path, kind):
    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path, kind)
        original = path(core, params).read_bytes()
        info = call(
            core,
            "presentation.workspace.list_text_fit_targets",
            {**scope(params), "slide_number": 1},
        )
        item = info["targets"][0]
        assert item["editable"] and item["name"] == "Fit target"
        assert item["geometry_pt"]["height"] == 36
        assert item["cached_font_scale"] == (0.6 if kind == "cached" else 1)
        assert [r["size_hundredths"] for r in item["paragraphs"][0]["runs"]] == ["3200", "2000"]
        assert item["fingerprint"] == target(core, params)[0]["expected_fingerprint"]
        assert path(core, params).read_bytes() == original
        old_scope = scope(params)
        save(core, params, prepare(core, params))
        call(
            core,
            "presentation.workspace.list_text_fit_targets",
            {**old_scope, "slide_number": 1},
            ok=False,
        )
    finally:
        core.close()


def test_fit_catalog_readonly_for_unsupported_script_and_numbering(tmp_path):
    from presenter_core.presentation.text_fit_editing import inspect_fit

    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path, "numbered")
        with zipfile.ZipFile(path(core, params)) as archive:
            reader = TextInheritance(archive, "ppt/slides/slide1.xml")
            shape = reader.slide.find("p:cSld/p:spTree/p:sp", _NS)
            shape.find("p:txBody/a:p/a:pPr/a:buAutoNum", _NS).set("type", "romanUcPeriod")
            before = etree.tostring(shape)
            assert not inspect_fit(reader, shape, text="text")["editable"]
            assert etree.tostring(shape) == before
            shape.find("p:txBody/a:p/a:pPr/a:buAutoNum", _NS).set("type", "arabicPeriod")
            shape.find("p:txBody/a:p/a:r/a:t", _NS).text = "文字"
            assert not inspect_fit(reader, shape, text="文字")["editable"]
    finally:
        core.close()


@pytest.mark.parametrize("failure", [None, "minimum", "private_fingerprint"])
def test_copilot_fit_binds_actual_staged_fonts_and_undo(tmp_path, failure):
    from test_copilot_data import DataProvider
    from test_copilot_diagrams import operations

    provider = DataProvider(locality="local")
    core = core_at(tmp_path, provider)
    try:
        params = fixture(core, tmp_path)
        original = path(core, params).read_bytes()

        def build(snapshot):
            slide = snapshot["slides"][0]
            fit = slide["text_fits"][0]
            assert fit["editable"] and "fingerprint" not in fit
            target_params = {"slide_id": slide["slide_id"], "object_id": fit["object_id"]}
            shrink = {
                **target_params,
                "strategy": "shrink",
                "minimum_size_pt": 19 if failure == "minimum" else 6,
            }
            if failure == "private_fingerprint":
                shrink["expected_fingerprint"] = "0" * 64
            return {
                "summary": "Fit mixed text",
                "operations": operations(
                    [
                        ("fit_text_shape", shrink),
                        (
                            "fit_text_shape",
                            {**target_params, "strategy": "grow", "minimum_size_pt": 6},
                        ),
                    ]
                ),
            }

        provider.builder = build
        proposal = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {**params, "instruction": "Shrink then grow the box", "slide_numbers": [1]},
            ok=failure is None,
        )
        assert path(core, params).read_bytes() == original
        if failure:
            assert not core._presentation_edit_plans._plans
            return
        reviewed = proposal["review_operations"]
        old = reviewed[0]["previous_data"]["paragraphs"][0]["runs"]
        staged = reviewed[1]["previous_data"]["paragraphs"][0]["runs"]
        assert int(old[0]["size_hundredths"]) == 3200
        assert int(staged[0]["size_hundredths"]) < 3200
        assert (
            proposal["operations"][0]["params"]["expected_fingerprint"]
            != proposal["operations"][1]["params"]["expected_fingerprint"]
        )
        save(core, params, proposal)
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


def test_native_kerning_threshold_does_not_require_optional_shaping():
    from presenter_core.presentation.text_fit import FontMetrics

    class BasicFont:
        def getlength(self, text):
            return len(text) * 10 - (3 if text == "AV" else 0)

    font = BasicFont()
    assert FontMetrics(font, False).getlength("AV") == 17
    assert FontMetrics(font, True).getlength("AV") == 20


def test_font_rounding_at_selected_minimum_does_not_cross_boundary(tmp_path):
    from presenter_core.presentation.text_fit import TextFitResult

    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path)
        _, resolved, shape = target(core, params)
        apply_fit(shape, resolved, TextFitResult(457200, 0.3 - 1e-14, 28))
        assert shape.find("p:txBody/a:p/a:r[2]/a:rPr", _NS).get("sz") == "600"
    finally:
        core.close()


def test_repeated_fit_is_a_package_preserving_noop(tmp_path):
    core = core_at(tmp_path)
    try:
        params = fixture(core, tmp_path)
        save(core, params, prepare(core, params, "shrink"))
        before = path(core, params).read_bytes()
        plan = prepare(core, params, "shrink")
        assert not plan["changed"]
        assert path(core, params).read_bytes() == before
    finally:
        core.close()
