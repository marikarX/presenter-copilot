"""Objects extending past the slide edge can be moved back onto it."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def test_offslide_object_moves_back_but_never_off(tmp_path: Path) -> None:
    source = tmp_path / "bleed.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(-Inches(1)), Inches(1), Inches(3), Inches(2)
    )
    shape.name = "Bleed"
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}
        inventory = _call(
            core,
            "t",
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": revision, "slide_number": 1},
        )["result"]
        bleed = next(o for o in inventory["objects"] if o["name"] == "Bleed")
        assert bleed["geometry_editable"] and bleed["geometry"]["x_emu"] < 0

        def move(x: int) -> dict:
            return _call(
                core,
                "m",
                "presentation.workspace.transform_object",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": inventory["slide_id"],
                    "object_id": bleed["object_id"],
                    "expected_geometry": bleed["geometry"],
                    "replacement_geometry": {**bleed["geometry"], "x_emu": x},
                },
            )

        still_off = move(-Inches(1) + 100)
        assert not still_off["ok"] and still_off["error"]["code"] == "INVALID_REQUEST"
        back = move(0)
        assert back["ok"], back.get("error")
    finally:
        core.close()
