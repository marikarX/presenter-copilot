"""Reviewed hide/show of slides in the slide show."""

from pathlib import Path

from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id
from test_slide_background import P, _Deck

from presenter_core.ipc.core import CoreService


def _hide(core, project, document, revision, slide_id, expected, hidden):
    return _call(
        core,
        "hide",
        "presentation.workspace.set_slide_hidden",
        {
            "project_id": project,
            "document_id": document,
            "expected_revision_id": revision,
            "slide_id": slide_id,
            "expected_hidden": expected,
            "hidden": hidden,
        },
    )


def test_slide_hides_and_shows_again(tmp_path: Path) -> None:
    deck = _Deck(tmp_path)
    try:
        inventory = deck.inventory()
        assert inventory["hidden"] is False
        reply = _hide(
            deck.core,
            deck.project,
            deck.document,
            deck.revision,
            inventory["slide_id"],
            False,
            True,
        )
        assert reply["ok"] and reply["result"]["changed"], reply
        deck.revision = reply["result"]["revision_id"]
        assert deck.slide().get("show") == "0"
        assert deck.inventory()["hidden"] is True
        stale = _hide(
            deck.core,
            deck.project,
            deck.document,
            deck.revision,
            inventory["slide_id"],
            False,
            True,
        )
        assert not stale["ok"] and stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        shown = _hide(
            deck.core,
            deck.project,
            deck.document,
            deck.revision,
            inventory["slide_id"],
            True,
            False,
        )
        assert shown["ok"], shown
        deck.revision = shown["result"]["revision_id"]
        assert deck.slide().get("show") is None and deck.slide().tag == f"{P}sld"
    finally:
        deck.core.close()


def test_last_visible_slide_cannot_be_hidden(tmp_path: Path) -> None:
    source = tmp_path / "one.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        slide_id = _call(
            core,
            "targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": document,
                "revision_id": revision,
                "slide_number": 1,
            },
        )["result"]["slide_id"]
        reply = _hide(core, project, document, revision, slide_id, False, True)
        assert not reply["ok"] and reply["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
    finally:
        core.close()
