from __future__ import annotations

import uuid
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from test_presentation_edit_planning import call
from test_presentation_transactions import context

from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.presentation.layouts import catalog
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def scope(params):
    return {
        "project_id": params["project_id"],
        "document_id": params["document_id"],
        "revision_id": params["expected_revision_id"],
    }


def path(core, params):
    source = core._presentation_edits._current_revision_source(
        params["project_id"], params["document_id"]
    )
    return core._storage.paths.snapshot_path(
        params["project_id"], source["local_snapshot_path"], require_exists=True
    )


def test_unused_native_layout_is_reviewed_atomic_preserves_parts_and_undo(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, targets = context(core)
        choices = call(core, "presentation.workspace.list_native_layouts", scope(params))
        assert len(choices["items"]) == 11
        assert all(
            set(item)
            == {
                "layout_id",
                "name",
                "master_name",
                "placeholder_count",
                "used_by_slides",
                "creation_supported",
                "creation_placeholders",
            }
            for item in choices["items"]
        )
        chosen = next(item for item in choices["items"] if item["name"] == "Comparison")
        assert chosen["used_by_slides"] == []
        original = path(core, params).read_bytes()
        plan = call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "expected_revision_id": params["expected_revision_id"],
                "summary": "Review unused native layout",
                "operations": [
                    {
                        "type": "apply_slide_layout",
                        "params": {
                            "slide_id": targets[0]["params"]["slide_id"],
                            "layout_id": chosen["layout_id"],
                        },
                    }
                ],
            },
        )
        assert plan["changed"] and path(core, params).read_bytes() == original
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
            },
        )
        saved = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
                "reviewed_sha256": plan["sha256"],
                "confirm": True,
            },
        )
        assert saved["sha256"] == plan["sha256"]
        import io

        with (
            zipfile.ZipFile(io.BytesIO(original)) as old,
            zipfile.ZipFile(path(core, params)) as new,
        ):
            assert old.namelist() == new.namelist()
            changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
            assert changed == ["ppt/slides/_rels/slide1.xml.rels"]
        params["expected_revision_id"] = saved["revision_id"]
        fresh = call(core, "presentation.workspace.list_native_layouts", scope(params))
        assert next(item for item in fresh["items"] if item["layout_id"] == chosen["layout_id"])[
            "used_by_slides"
        ] == [1]
        no_op = call(
            core,
            "presentation.workspace.apply_slide_layout",
            {
                **params,
                "slide_id": targets[0]["params"]["slide_id"],
                "layout_id": chosen["layout_id"],
            },
        )
        assert not no_op["changed"] and no_op["revision_id"] == saved["revision_id"]
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize("case", ["fabricated", "both", "neither", "path", "wrong_slide"])
def test_native_layout_selection_rejects_untrusted_ids_without_revision(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, targets = context(core)
        selection = {
            **params,
            "slide_id": targets[0]["params"]["slide_id"],
            "layout_id": str(uuid.uuid4()),
        }
        if case == "both":
            selection["layout_source_slide_id"] = targets[0]["params"]["slide_id"]
        if case == "neither":
            selection.pop("layout_id")
        if case == "path":
            selection["layout_id"] = "../../other.xml"
        if case == "wrong_slide":
            selection["slide_id"] = str(uuid.uuid4())
        before = path(core, params).read_bytes()
        call(core, "presentation.workspace.apply_slide_layout", selection, ok=False)
        assert path(core, params).read_bytes() == before
        with core._storage.project_database(params["project_id"]) as connection:
            assert (
                connection.execute("SELECT COUNT(*) FROM presentation_revisions").fetchone()[0] == 1
            )
    finally:
        core.close()


@pytest.mark.parametrize("case", ["external", "wrong_master", "duplicate"])
def test_native_catalog_rejects_broken_master_relationships(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, _ = context(core)
        damaged = tmp_path / "invalid-layout.pptx"
        with zipfile.ZipFile(path(core, params)) as source, zipfile.ZipFile(damaged, "w") as output:
            for info in source.infolist():
                data = source.read(info.filename)
                if (
                    case in {"external", "wrong_master"}
                    and info.filename == "ppt/slideLayouts/_rels/slideLayout1.xml.rels"
                ):
                    root = etree.fromstring(data)
                    parent = next(v for v in root if v.get("Type", "").endswith("/slideMaster"))
                    if case == "external":
                        parent.set("TargetMode", "External")
                    else:
                        parent.set("Target", "../slideMasters/other.xml")
                    data = etree.tostring(root)
                if case == "duplicate" and info.filename == "ppt/slideMasters/slideMaster1.xml":
                    root = etree.fromstring(data)
                    nodes = root.find(
                        "{http://schemas.openxmlformats.org/presentationml/2006/main}sldLayoutIdLst"
                    )
                    import copy

                    nodes.append(copy.deepcopy(nodes[0]))
                    data = etree.tostring(root)
                output.writestr(info, data)
        with zipfile.ZipFile(damaged) as archive, pytest.raises(CoreDomainError):
            catalog(archive, params["document_id"])
    finally:
        core.close()


@pytest.mark.parametrize("case", ["corrupt", "changed_head"])
def test_native_catalog_rejects_changed_snapshot_or_revision(
    tmp_path: Path, monkeypatch, case: str
):
    from presenter_core.presentation import layouts

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, operations = context(core)
        if case == "corrupt":
            source = path(core, params)
            source.write_bytes(source.read_bytes() + b"corrupt")
        else:
            original = layouts.catalog

            def change_head(archive, document):
                result = original(archive, document)
                call(
                    core,
                    "presentation.workspace.replace_speaker_notes",
                    {**params, **operations[1]["params"]},
                )
                return result

            monkeypatch.setattr(layouts, "catalog", change_head)
        error = call(core, "presentation.workspace.list_native_layouts", scope(params), ok=False)
        assert error["code"] in {"SOURCE_SNAPSHOT_CHANGED", "PRESENTATION_REVISION_CONFLICT"}
    finally:
        core.close()


def test_native_creation_populates_comparison_slots_notes_and_preserves_package(tmp_path: Path):
    import io

    from presenter_core.presentation.fidelity import inspect_pptx_snapshot

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, targets = context(core)
        choices = call(core, "presentation.workspace.list_native_layouts", scope(params))
        layout = next(item for item in choices["items"] if item["name"] == "Comparison")
        assert layout["creation_supported"]
        assert len(layout["creation_placeholders"]) == 5
        texts = [
            "Approve the pilot",
            "Option A",
            "Measured benefit\nKnown risk",
            "Option B",
            "Lower cost\nSlower rollout",
        ]
        original = path(core, params).read_bytes()
        plan = call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {
                **params,
                "summary": "Insert native comparison",
                "operations": [
                    {
                        "type": "create_slide",
                        "params": {
                            "slide_id": targets[0]["params"]["slide_id"],
                            "layout_id": layout["layout_id"],
                            "placeholder_text": [
                                {"placeholder_id": slot["placeholder_id"], "text": text}
                                for slot, text in zip(
                                    layout["creation_placeholders"], texts, strict=True
                                )
                            ],
                            "notes": "Compare the two alternatives.\nAsk for a pilot decision.",
                        },
                    }
                ],
            },
        )
        assert path(core, params).read_bytes() == original
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
            },
        )
        saved = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
                "reviewed_sha256": plan["sha256"],
                "confirm": True,
            },
        )
        with path(core, params).open("rb") as handle:
            native = inspect_pptx_snapshot(
                handle,
                document_id=params["document_id"],
                expected_sha256=saved["sha256"],
                byte_size=saved["byte_size"],
            )
        assert len(native.slides) == 3
        created = native.slides[1]
        assert [obj.text for obj in created.objects] == texts
        assert all(obj.geometry_editable for obj in created.objects)
        assert created.speaker_notes == "Compare the two alternatives.\nAsk for a pilot decision."
        with (
            zipfile.ZipFile(io.BytesIO(original)) as old,
            zipfile.ZipFile(path(core, params)) as new,
        ):
            changed = {name for name in old.namelist() if old.read(name) != new.read(name)}
            assert changed == {
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
                "[Content_Types].xml",
            }
            assert b'idx="1"' in new.read(created.part_name)
            assert [name for name in new.namelist() if name.startswith("ppt/notesMasters/")] == [
                name for name in old.namelist() if name.startswith("ppt/notesMasters/")
            ]
            assert [name for name in new.namelist() if name.startswith("ppt/theme/")] == [
                name for name in old.namelist() if name.startswith("ppt/theme/")
            ]
        params["expected_revision_id"] = saved["revision_id"]
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize(
    "case",
    [
        "foreign_slot",
        "duplicate_slot",
        "wrong_layout",
        "too_long",
        "too_many_lines",
        "no_layout",
        "both_sources",
    ],
)
def test_native_creation_bad_drafts_leave_revision_untouched(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, targets = context(core)
        layout = next(
            v
            for v in call(core, "presentation.workspace.list_native_layouts", scope(params))[
                "items"
            ]
            if v["name"] == "Title and Content"
        )
        slot = {
            "placeholder_id": layout["creation_placeholders"][0]["placeholder_id"],
            "text": "Native title",
        }
        draft = {
            **params,
            "slide_id": targets[0]["params"]["slide_id"],
            "layout_id": layout["layout_id"],
            "placeholder_text": [slot],
            "notes": "New notes",
        }
        if case == "foreign_slot":
            slot["placeholder_id"] = str(uuid.uuid4())
        if case == "duplicate_slot":
            draft["placeholder_text"].append(dict(slot))
        if case == "wrong_layout":
            draft["layout_id"] = str(uuid.uuid4())
        if case == "too_long":
            slot["text"] = "x" * 2001
        if case == "too_many_lines":
            slot["text"] = "\n".join(["x"] * 21)
        if case == "no_layout":
            draft.pop("layout_id")
        if case == "both_sources":
            draft["layout_source_slide_id"] = draft["slide_id"]
        original = path(core, params).read_bytes()
        call(core, "presentation.workspace.create_slide", draft, ok=False)
        assert path(core, params).read_bytes() == original
        with core._storage.project_database(params["project_id"]) as connection:
            assert (
                connection.execute("SELECT COUNT(*) FROM presentation_revisions").fetchone()[0] == 1
            )
    finally:
        core.close()


@pytest.mark.parametrize(
    "case",
    ["missing", "rotated", "outside", "duplicate", "grouped", "bad_index", "foreign_transform"],
)
def test_native_creation_disables_ambiguous_placeholder_geometry(tmp_path: Path, case: str):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, _ = context(core)
        damaged = tmp_path / "unsupported.pptx"
        with zipfile.ZipFile(path(core, params)) as source, zipfile.ZipFile(damaged, "w") as output:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "ppt/slideLayouts/slideLayout4.xml":
                    root = etree.fromstring(data)
                    ns = {
                        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
                        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                    }
                    tree = root.find("p:cSld/p:spTree", ns)
                    shape = tree.findall("p:sp", ns)[1]
                    transform = shape.find("p:spPr/a:xfrm", ns)
                    if case == "missing":
                        shape.find("p:nvSpPr/p:nvPr/p:ph", ns).set("type", "subTitle")
                        transform.getparent().remove(transform)
                    if case == "bad_index":
                        shape.find("p:nvSpPr/p:nvPr/p:ph", ns).set("idx", "../invalid")
                    if case == "foreign_transform":
                        transform.set(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
                            "foreign",
                        )
                    if case == "rotated":
                        transform.set("rot", "60000")
                    if case == "outside":
                        transform.find("a:off", ns).set("x", "999999999")
                    if case == "duplicate":
                        shape.find("p:nvSpPr/p:nvPr/p:ph", ns).set("idx", "0")
                    if case == "grouped":
                        tree.remove(shape)
                        group = etree.SubElement(tree, "{" + ns["p"] + "}grpSp")
                        group.append(shape)
                    data = etree.tostring(root)
                output.writestr(info, data)
        with zipfile.ZipFile(damaged) as archive:
            item = next(
                v for v in catalog(archive, params["document_id"]) if v["name"] == "Two Content"
            )
            assert not item["creation_supported"]
    finally:
        core.close()


def test_native_creation_empty_slots_and_followup_edit_share_one_plan(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, targets = context(core)
        layout = next(
            v
            for v in call(core, "presentation.workspace.list_native_layouts", scope(params))[
                "items"
            ]
            if v["name"] == "Title and Content"
        )
        plan = call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {
                **params,
                "summary": "Native title with empty body",
                "operations": [
                    {
                        "type": "create_slide",
                        "params": {
                            "slide_id": targets[0]["params"]["slide_id"],
                            "layout_id": layout["layout_id"],
                            "placeholder_text": [
                                {
                                    "placeholder_id": layout["creation_placeholders"][0][
                                        "placeholder_id"
                                    ],
                                    "text": "Native title",
                                }
                            ],
                            "notes": "",
                        },
                    },
                    {
                        "type": "replace_speaker_notes",
                        "params": {
                            "slide_id": {"created_by": 1},
                            "expected_text": "",
                            "replacement_text": "Notes added to created slide",
                        },
                    },
                ],
            },
        )
        assert plan["changed"]
        reader = call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
            },
        )
        from presenter_core.presentation.fidelity import inspect_pptx_snapshot

        staged = core._presentation_edit_plans._plans[plan["plan_id"]].package
        with staged.path.open("rb") as handle:
            native = inspect_pptx_snapshot(
                handle,
                document_id=params["document_id"],
                expected_sha256=staged.sha256,
                byte_size=staged.byte_size,
            )
        created = native.slides[1]
        assert created.slide_id == reader["fidelity_report"]["slides"][1]["slide_id"]
        assert created.speaker_notes == "Notes added to created slide"
        assert [obj.text for obj in created.objects] == ["Native title", ""]
        call(
            core,
            "presentation.workspace.reject_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "plan_id": plan["plan_id"],
            },
        )
        with core._storage.project_database(params["project_id"]) as connection:
            assert (
                connection.execute("SELECT COUNT(*) FROM presentation_revisions").fetchone()[0] == 1
            )
    finally:
        core.close()


@pytest.mark.parametrize("case", ["none", "duplicate", "external", "wrong_declaration"])
def test_native_creation_notes_master_allocation_and_ambiguity(tmp_path: Path, case: str):
    from pptx import Presentation

    from presenter_core.presentation.fidelity import inspect_pptx_snapshot

    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        fixture = tmp_path / "source.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "Opening"
        if case != "none":
            deck.slides[0].notes_slide.notes_text_frame.text = "Existing notes"
        deck.save(fixture)
        if case != "none":
            altered = tmp_path / "altered.pptx"
            with zipfile.ZipFile(fixture) as archive, zipfile.ZipFile(altered, "w") as output:
                for info in archive.infolist():
                    data = archive.read(info.filename)
                    if info.filename == "ppt/_rels/presentation.xml.rels":
                        root = etree.fromstring(data)
                        master = next(
                            node for node in root if node.get("Type", "").endswith("/notesMaster")
                        )
                        if case == "duplicate":
                            import copy

                            clone = copy.deepcopy(master)
                            clone.set("Id", "otherNotesMaster")
                            root.append(clone)
                        if case == "external":
                            master.set("TargetMode", "External")
                        data = etree.tostring(root)
                    if case == "wrong_declaration" and info.filename == "ppt/presentation.xml":
                        root = etree.fromstring(data)
                        nodes = etree.Element(
                            "{http://schemas.openxmlformats.org/presentationml/2006/main}notesMasterIdLst"
                        )
                        child = etree.SubElement(
                            nodes,
                            "{http://schemas.openxmlformats.org/presentationml/2006/main}notesMasterId",
                        )
                        child.set(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
                            "unknownMaster",
                        )
                        root.insert(1, nodes)
                        data = etree.tostring(root)
                    output.writestr(info, data)
            fixture = altered
        project = call(core, "project.create", {"name": "Notes master dependency"})["project"]["id"]
        document = call(
            core,
            "source.import",
            {"project_id": project, "path": str(fixture), "kind": "presentation"},
        )["document"]["id"]
        source = core._presentation_edits._current_revision_source(project, document)
        params = {
            "project_id": project,
            "document_id": document,
            "expected_revision_id": source["revision_id"],
        }
        inventory = call(
            core, "presentation.workspace.list_text_targets", {**scope(params), "slide_number": 1}
        )
        layout = next(
            v
            for v in call(core, "presentation.workspace.list_native_layouts", scope(params))[
                "items"
            ]
            if v["name"] == "Title and Content"
        )
        original = path(core, params).read_bytes()
        saved = call(
            core,
            "presentation.workspace.create_slide",
            {
                **params,
                "slide_id": inventory["slide_id"],
                "layout_id": layout["layout_id"],
                "placeholder_text": [],
                "notes": "New notes",
            },
            ok=case == "none",
        )
        if case == "none":
            with path(core, params).open("rb") as handle:
                native = inspect_pptx_snapshot(
                    handle,
                    document_id=document,
                    expected_sha256=saved["sha256"],
                    byte_size=saved["byte_size"],
                )
            assert native.slides[1].speaker_notes == "New notes"
            with zipfile.ZipFile(path(core, params)) as archive:
                assert (
                    len(
                        [
                            name
                            for name in archive.namelist()
                            if name.startswith("ppt/notesMasters/") and name.endswith(".xml")
                        ]
                    )
                    == 1
                )
        else:
            assert saved["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
            assert path(core, params).read_bytes() == original
    finally:
        core.close()
