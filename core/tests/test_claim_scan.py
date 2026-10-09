"""Deck scan preflight and dense text windows preserve existing claim authority."""

import copy
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from test_presentation_edit_planning import PlanProvider, call

from presenter_core.ipc.core import CoreService
from presenter_core.presentation import workspace_source
from presenter_core.presentation.claim_extraction import claim_context_size, claim_request
from presenter_core.presentation.claims import digest
from presenter_core.providers.context import MAX_TOTAL_CONTEXT_CHARS
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def setup(tmp_path: Path, *, slides: int = 3):
    source = tmp_path / "scan.pptx"
    deck = Presentation()
    dense = "Gross margin increased to 28%.\n" + " ".join(
        f"Supporting observation {i}." for i in range(250)
    )
    assert 6000 < len(dense) < 8000
    for number in range(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = (
            dense
            if number == 1
            else "Unavailable text " * 600
            if number == 2
            else f"Metric improved on slide {number + 1}."
        )
    deck.save(source)
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    project = call(core, "project.create", {"name": "Deck scan", "privacy_mode": "local_only"})[
        "project"
    ]["id"]
    document = call(
        core, "source.import", {"project_id": project, "path": str(source), "kind": "presentation"}
    )["document"]["id"]
    base = {"project_id": project, "document_id": document}
    opened = call(core, "presentation.workspace.open_source", base)
    call(core, "presentation.workspace.close_source", {**base, "session_id": opened["session_id"]})
    return core, provider, {**base, "revision_id": opened["revision_id"]}, dense


def test_full_deck_local_preflight_pages_once_and_reports_unavailable_text(
    tmp_path: Path, monkeypatch
):
    core, provider, scope, dense = setup(tmp_path, slides=26)
    original = workspace_source.inspect_pptx_snapshot
    inspections = []

    def inspect(*args, **kwargs):
        inspections.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(workspace_source, "inspect_pptx_snapshot", inspect)
    try:
        first = call(core, "presentation.workspace.prepare_claim_scan", {**scope, "limit": 25})
        assert len(inspections) == 1 and first["has_more"] and len(first["items"]) == 25
        last = call(core, "presentation.workspace.prepare_claim_scan", {**scope, "offset": 25})
        assert not last["has_more"] and last["items"][0]["slide_number"] == 26
        assert first["items"][2]["excluded_objects"] == 1 and not first["items"][2]["batches"]
        assert first["items"][1]["eligible_objects"] == 1 and len(first["items"][1]["batches"]) >= 2
        assert not provider.requests and dense not in json.dumps(first)
        with core._storage.project_database(scope["project_id"]) as connection:
            assert connection.execute("SELECT COUNT(*) FROM provider_runs").fetchone()[0] == 0
            assert connection.execute("SELECT COUNT(*) FROM presentation_claims").fetchone()[0] == 0
    finally:
        core.close()


def test_dense_windows_cover_native_text_and_bind_quotes_to_full_object(tmp_path: Path):
    core, provider, scope, dense = setup(tmp_path)
    try:
        plan = call(core, "presentation.workspace.prepare_claim_scan", scope)
        slide = plan["items"][1]
        native = call(
            core, "presentation.workspace.list_text_targets", {**scope, "slide_number": 2}
        )
        assert native["targets"][0]["current_text"] is None
        last_end = 0
        for batch in slide["batches"]:
            span = batch["spans"][0]
            assert span["start"] <= last_end and span["end"] > last_end
            last_end = span["end"]
            text = dense[span["start"] : span["end"]]
            request = claim_request(
                {span["object_id"]: text}, "local_only", scope["document_id"], scope["revision_id"]
            )
            assert claim_context_size(request) <= MAX_TOTAL_CONTEXT_CHARS and span[
                "sha256"
            ] == digest(dense)
            result = call(
                core,
                "presentation.workspace.extract_claims",
                {
                    **scope,
                    "slide_id": slide["slide_id"],
                    "slide_number": 2,
                    "object_ids": batch["object_ids"],
                    "spans": batch["spans"],
                },
            )
            assert (
                json.loads(provider.request_objects[-1].current_slide_summary)["objects"][0]["text"]
                == text
            )
            for candidate in result["candidates"]:
                assert candidate["object_sha256"] == digest(dense)
                assert dense[candidate["start"] : candidate["end"]] == candidate["quote"]
        assert last_end == len(dense)
        assert not call(
            core,
            "presentation.workspace.list_claims",
            {**scope, "slide_id": slide["slide_id"], "slide_number": 2},
        )["items"]
    finally:
        core.close()


@pytest.mark.parametrize(
    "bad", ["hash", "negative", "outside", "bool", "missing", "duplicate", "unknown"]
)
def test_invalid_prepared_windows_never_reach_provider(tmp_path: Path, bad: str):
    core, provider, scope, _ = setup(tmp_path)
    try:
        slide = call(core, "presentation.workspace.prepare_claim_scan", scope)["items"][1]
        batch = copy.deepcopy(slide["batches"][0])
        span = batch["spans"][0]
        if bad == "hash":
            span["sha256"] = "0" * 64
        if bad == "negative":
            span["start"] = -1
        if bad == "outside":
            span["end"] = 9000
        if bad == "bool":
            span["start"] = True
        if bad == "missing":
            batch["spans"] = []
        if bad == "duplicate":
            batch["spans"].append(copy.deepcopy(span))
        if bad == "unknown":
            span["object_id"] = "00000000-0000-0000-0000-000000000000"
        call(
            core,
            "presentation.workspace.extract_claims",
            {
                **scope,
                "slide_id": slide["slide_id"],
                "slide_number": 2,
                "object_ids": batch["object_ids"],
                "spans": batch["spans"],
            },
            ok=False,
        )
        assert not provider.requests
    finally:
        core.close()


def test_escaped_text_plan_respects_actual_serialized_budget(tmp_path: Path):
    # A quote-heavy supported object expands through nested JSON. Never relax the 12K limit.
    from presenter_core.presentation.claim_scan import PresentationClaimScanService

    core, _, scope, _ = setup(tmp_path)
    try:
        original = core._presentation_semantics._workspace_sources.read_revision_inventory
        project, document, revision, inventory = original(scope)
        native = inventory.text_targets(1)
        text = '"' * 7990 + " a claim"
        native["targets"][0]["decision_text"] = text
        native["semantic_targets"][0]["decision_text"] = text

        class Inventory:
            slides = inventory.slides[:1]

            def text_targets(self, _number):
                return native

        core._presentation_semantics._workspace_sources.read_revision_inventory = lambda _params: (
            project,
            document,
            revision,
            Inventory(),
        )
        result = PresentationClaimScanService(core._presentation_semantics).prepare_claim_scan(
            scope
        )
        windows = result["items"][0]["batches"]
        assert len(windows) > 2
        for batch in windows:
            selected = {s["object_id"]: text[s["start"] : s["end"]] for s in batch["spans"]}
            assert (
                claim_context_size(
                    claim_request(selected, "selected_context_cloud", document, revision)
                )
                <= MAX_TOTAL_CONTEXT_CHARS - 100
            )
    finally:
        core.close()


def test_quote_outside_disclosed_window_is_rejected_even_if_present_in_full_object(tmp_path: Path):
    core, provider, scope, _ = setup(tmp_path)
    try:
        slide = call(core, "presentation.workspace.prepare_claim_scan", scope)["items"][1]
        batch = slide["batches"][-1]
        assert batch["spans"][0]["start"] > 0
        provider.modifier = lambda output, _invocation: output["candidates"][0].update(
            quote="Gross margin increased to 28%."
        )
        error = call(
            core,
            "presentation.workspace.extract_claims",
            {
                **scope,
                "slide_id": slide["slide_id"],
                "slide_number": 2,
                "object_ids": batch["object_ids"],
                "spans": batch["spans"],
            },
            ok=False,
        )
        assert error["code"] == "PROVIDER_MALFORMED_OUTPUT"
        assert not call(
            core,
            "presentation.workspace.list_claims",
            {**scope, "slide_id": slide["slide_id"], "slide_number": 2},
        )["items"]
    finally:
        core.close()


def test_preflight_and_prepared_batch_reject_a_real_new_revision(tmp_path: Path):
    from test_presentation_narrative import change

    core, provider, scope, _ = setup(tmp_path)
    try:
        plan = call(core, "presentation.workspace.prepare_claim_scan", scope)
        native = call(
            core, "presentation.workspace.list_text_targets", {**scope, "slide_number": 1}
        )
        slide_scope = {**scope, "slide_id": native["slide_id"], "slide_number": 1}
        change(core, slide_scope, native["targets"][0])
        assert (
            call(core, "presentation.workspace.prepare_claim_scan", scope, ok=False)["code"]
            == "PRESENTATION_REVISION_CONFLICT"
        )
        batch = plan["items"][0]["batches"][0]
        error = call(
            core,
            "presentation.workspace.extract_claims",
            {**slide_scope, "object_ids": batch["object_ids"], "spans": batch["spans"]},
            ok=False,
        )
        assert error["code"] == "PRESENTATION_REVISION_CONFLICT" and not provider.requests
    finally:
        core.close()
