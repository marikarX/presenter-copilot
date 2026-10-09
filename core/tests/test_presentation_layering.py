"""Layer changes preserve native content and relationships while changing paint order."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from presenter_core.ipc.core import CoreService
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter
from tests.test_presentation_planning import call


@pytest.mark.parametrize(
    "direction,order",
    [("forward", [0, 2, 1]), ("front", [0, 2, 1]), ("backward", [1, 0, 2]), ("back", [1, 0, 2])],
)
def test_layer_change_preserves_payloads_and_other_parts(
    tmp_path: Path, direction: str, order: list[int]
) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for name in ("Bottom", "Middle", "Top"):
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        shape.text = name
    slide.notes_slide.notes_text_frame.text = "Keep these notes"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "layers.pptx"
    presentation.save(str(path))
    original_bytes = path.read_bytes()
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        project = call(core, "project.create", {"name": "Layers"})["project"]["id"]
        document = call(
            core,
            "source.import",
            {"project_id": project, "kind": "presentation", "path": str(path)},
        )["document"]["id"]
        original = core._presentation_edits._current_revision_source(project, document)
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": document,
                "revision_id": original["revision_id"],
                "slide_number": 1,
            },
        )
        identities = [obj["object_id"] for obj in inventory["objects"]]
        params = {
            "project_id": project,
            "document_id": document,
            "expected_revision_id": original["revision_id"],
            "slide_id": inventory["slide_id"],
            "object_id": identities[1],
            "expected_object_ids": identities,
            "direction": direction,
        }
        assert (
            call(
                core,
                "presentation.workspace.change_object_layer",
                {**params, "expected_object_ids": list(reversed(identities))},
                ok=False,
            )["code"]
            == "PRESENTATION_EDIT_STALE_TARGET"
        )
        result = call(core, "presentation.workspace.change_object_layer", params)
        assert result["changed"] is True and result["revision_number"] == 2
        current = core._presentation_edits._current_revision_source(project, document)
        current_path = core._storage.paths.snapshot_path(
            project, current["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(path) as before, zipfile.ZipFile(current_path) as after:
            assert before.namelist() == after.namelist()
            for part in before.namelist():
                if part != "ppt/slides/slide1.xml":
                    assert before.read(part) == after.read(part)
            ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            old_nodes = etree.fromstring(before.read("ppt/slides/slide1.xml")).xpath(
                "./p:cSld/p:spTree/p:sp", namespaces=ns
            )
            new_nodes = etree.fromstring(after.read("ppt/slides/slide1.xml")).xpath(
                "./p:cSld/p:spTree/p:sp", namespaces=ns
            )
            assert [etree.tostring(node) for node in new_nodes] == [
                etree.tostring(old_nodes[index]) for index in order
            ]
        updated = call(
            core,
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": document,
                "revision_id": result["revision_id"],
                "slide_number": 1,
            },
        )
        assert [obj["object_id"] for obj in updated["objects"]] == [
            identities[index] for index in order
        ]
        boundary = "front" if direction in ("front", "forward") else "back"
        noop = call(
            core,
            "presentation.workspace.change_object_layer",
            {
                **params,
                "expected_revision_id": result["revision_id"],
                "expected_object_ids": [identities[index] for index in order],
                "direction": boundary,
            },
        )
        assert noop["changed"] is False and noop["revision_id"] == result["revision_id"]
        assert (
            call(core, "presentation.workspace.change_object_layer", params, ok=False)["code"]
            == "PRESENTATION_REVISION_CONFLICT"
        )
        call(
            core,
            "presentation.workspace.step_revision",
            {
                "project_id": project,
                "document_id": document,
                "expected_revision_id": result["revision_id"],
                "direction": "undo",
            },
        )
        assert (
            core._presentation_edits._current_revision_source(project, document)["revision_id"]
            == original["revision_id"]
        )
        assert path.read_bytes() == original_bytes
    finally:
        core.close()
