from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from test_chart_fills import GRADIENT, SOLID
from test_chart_legend import call, core_at, path, save, scope
from test_copilot_data import DataProvider
from test_copilot_diagrams import operations

from presenter_core.errors import CoreDomainError
from presenter_core.presentation.shape_fills import NS, inspect_fill


def imported(core, directory):
    deck = Presentation()
    for _ in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for index in range(2):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + index * 3), Inches(1), Inches(2), Inches(2)
            )
            shape.text = f"Preserved text {index}"
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(20, 40, 60)
            shape.line.color.rgb = RGBColor(100, 120, 140)
            shape.line.width = Pt(3)
        slide.notes_slide.notes_text_frame.text = "Preserved notes"
    fixture = directory / "shape-fill.pptx"
    deck.save(fixture)
    project = call(core, "project.create", {"name": "Shape fill review"})["project"]["id"]
    document = call(
        core, "source.import", {"project_id": project, "path": str(fixture), "kind": "presentation"}
    )["document"]["id"]
    source = core._presentation_edits._current_revision_source(project, document)
    params = {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": source["revision_id"],
    }
    return params, catalog(core, params)


def catalog(core, params):
    return call(
        core, "presentation.workspace.list_shape_fills", {**scope(params), "slide_number": 1}
    )


def prepare(core, params, inventory, fill, fingerprint=None):
    item = inventory["items"][0]
    return core._presentation_edit_plans.prepare(
        {
            **params,
            "summary": "Review shape fill",
            "operations": [
                {
                    "type": "set_shape_fill",
                    "params": {
                        "slide_id": inventory["slide_id"],
                        "object_id": item["object_id"],
                        "expected_fingerprint": fingerprint or item["fingerprint"],
                        "fill": fill,
                    },
                }
            ],
        }
    )


@pytest.mark.parametrize("fill", [SOLID, GRADIENT])
def test_shape_fill_exact_package_preservation_save_and_undo(tmp_path: Path, fill):
    core = core_at(tmp_path)
    try:
        params, before = imported(core, tmp_path)
        original = path(core, params).read_bytes()
        with zipfile.ZipFile(path(core, params)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        plan = prepare(core, params, before, fill)
        assert path(core, params).read_bytes() == original
        save(core, params, plan)
        after = catalog(core, params)
        assert after["items"][0]["fill"] == fill
        assert after["items"][1] == before["items"][1]
        with zipfile.ZipFile(path(core, params)) as archive:
            assert set(archive.namelist()) == set(parts)
            assert {name for name, data in parts.items() if archive.read(name) != data} == {
                "ppt/slides/slide1.xml"
            }
            old = etree.fromstring(parts["ppt/slides/slide1.xml"])
            new = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
            for root in [old, new]:
                props = root.find("p:cSld/p:spTree/p:sp/p:spPr", NS)
                for child in list(props):
                    if etree.QName(child).localname in {"solidFill", "gradFill"}:
                        props.remove(child)
            assert etree.tostring(old, method="c14n") == etree.tostring(new, method="c14n")
        assert not prepare(core, params, after, fill)["changed"]
        call(core, "presentation.workspace.step_revision", {**params, "direction": "undo"})
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


def test_stale_shape_fill_and_invalid_fill_preserve_head(tmp_path: Path):
    core = core_at(tmp_path)
    try:
        params, before = imported(core, tmp_path)
        original = path(core, params).read_bytes()
        for fill, fingerprint in [(SOLID, "0" * 64), ({**SOLID, "opacity_percent": 101}, None)]:
            with pytest.raises(CoreDomainError):
                prepare(core, params, before, fill, fingerprint)
            assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize("fill_tag", ["blipFill", "pattFill", "grpFill"])
def test_unsupported_shape_fill_is_readonly(fill_tag):
    shape = etree.fromstring(
        f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}">'
        '<p:nvSpPr><p:cNvPr id="2"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:prstGeom prst="rect"/><a:{fill_tag}/></p:spPr></p:sp>'
    )
    assert not inspect_fill(shape)["editable"]
    assert inspect_fill(shape)["fingerprint"] is None


def test_copilot_shape_fill_uses_staged_previous_fill(tmp_path: Path):
    provider = DataProvider(locality="local")
    core = core_at(tmp_path, provider)
    try:
        params, _ = imported(core, tmp_path)

        def build(snapshot):
            slide = snapshot["slides"][0]
            target = slide["shape_fills"][0]
            assert target["editable"]
            return {
                "summary": "Review shape fills",
                "operations": operations(
                    [
                        (
                            "set_shape_fill",
                            {
                                "slide_id": slide["slide_id"],
                                "object_id": target["object_id"],
                                "fill": value,
                            },
                        )
                        for value in (SOLID, GRADIENT)
                    ]
                ),
            }

        provider.builder = build
        original = path(core, params).read_bytes()
        plan = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {**params, "instruction": "Style shape fills", "slide_numbers": [1]},
        )
        assert plan["review_operations"][1]["previous_data"]["fill"] == SOLID
        assert path(core, params).read_bytes() == original
        save(core, params, plan)
        assert catalog(core, params)["items"][0]["fill"] == GRADIENT
    finally:
        core.close()


def test_copilot_fill_binds_shape_created_earlier_in_plan(tmp_path: Path):
    provider = DataProvider(locality="local")
    core = core_at(tmp_path, provider)
    try:
        params, _ = imported(core, tmp_path)

        def build(snapshot):
            slide = snapshot["slides"][0]
            return {
                "summary": "Create and fill a text box",
                "operations": operations(
                    [
                        ("create_text_box", {"slide_id": slide["slide_id"], "text": "New shape"}),
                        (
                            "set_shape_fill",
                            {
                                "slide_id": slide["slide_id"],
                                "object_id": {"created_by": 1},
                                "fill": GRADIENT,
                            },
                        ),
                    ]
                ),
            }

        provider.builder = build
        plan = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {**params, "instruction": "Create a gradient text box", "slide_numbers": [1]},
        )
        save(core, params, plan)
        assert any(item["fill"] == GRADIENT for item in catalog(core, params)["items"])
    finally:
        core.close()
