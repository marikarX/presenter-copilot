"""Native templates preserve content and enter the normal revision/editing path."""

from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation

from presenter_core.ipc.core import CoreService
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter
from tests.test_presentation_planning import call, context


@pytest.mark.parametrize("template", ["classic", "midnight", "paper"])
def test_native_layouts_preserve_content_bounds_and_editability(
    tmp_path: Path, template: str
) -> None:
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        project_id, _ = context(core)
        slides = [
            {
                "title": "A" * 120,
                "bullets": [f"{index}: " + "B" * 195 for index in range(8)],
                "notes": "Speaker detail",
                "layout": layout,
            }
            for layout in ("points", "columns", "cards")
        ] + [
            {
                "title": "Decision",
                "bullets": ["Approve the pilot"],
                "notes": "Ask aloud",
                "layout": "statement",
            }
        ]
        document = call(
            core,
            "presentation.workspace.create_from_outline",
            {
                "project_id": project_id,
                "title": "Template deck",
                "template": template,
                "slides": slides,
            },
        )["document"]
        with core._storage.project_database(project_id) as connection:
            snapshot = connection.execute(
                "SELECT local_snapshot_path FROM documents WHERE id = ?", (document["id"],)
            ).fetchone()[0]
        deck = Presentation(
            str(core._storage.paths.snapshot_path(project_id, snapshot, require_exists=True))
        )
        assert len(deck.slides) == 4
        for slide, expected in zip(deck.slides, slides, strict=True):
            assert slide.shapes.title.text == expected["title"]
            assert slide.notes_slide.notes_text_frame.text == expected["notes"]
            assert [
                shape.text
                for shape in slide.shapes
                if shape.name.startswith("Point ")
                and shape.has_text_frame
                and not shape.name.endswith(" panel")
            ] == expected["bullets"]
            for shape in slide.shapes:
                assert 0 <= shape.left < shape.left + shape.width <= deck.slide_width
                assert 0 <= shape.top < shape.top + shape.height <= deck.slide_height
        # All content is native OOXML rather than a slide-sized bitmap.
        assert all(shape.shape_type != 13 for slide in deck.slides for shape in slide.shapes)
    finally:
        core.close()


@pytest.mark.parametrize(
    "extra",
    [
        {"template": "unknown"},
        {"template": []},
        {"slides": [{"title": "Decision", "layout": "unknown"}]},
        {"slides": [{"title": "Decision", "layout": "statement", "bullets": ["one", "two"]}]},
    ],
)
def test_invalid_design_never_imports_a_deck(tmp_path: Path, extra: dict[str, Any]) -> None:
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        project_id, _ = context(core)
        error = call(
            core,
            "presentation.workspace.create_from_outline",
            {"project_id": project_id, "title": "Deck", "slides": [{"title": "Decision"}], **extra},
            ok=False,
        )
        assert error["code"] == "INVALID_REQUEST"
        assert call(core, "source.list", {"project_id": project_id})["sources"] == []
    finally:
        core.close()
