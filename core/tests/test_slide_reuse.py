"""Reuse Slides: copy a slide from another deck in the project into this one."""

import io
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _source(path: Path) -> None:
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[0])
    first.shapes.title.text = "Library"
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Quarterly results"
    body = slide.placeholders[1].text_frame
    body.text = "Revenue grew"
    run = body.paragraphs[0].add_run()
    run.text = " (details)"
    run.hyperlink.address = "https://example.org/results"
    image = io.BytesIO()
    Image.new("RGB", (40, 30), "#2a7").save(image, format="PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(6), Inches(4.5), Inches(2), Inches(1.5)).name = "Logo"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", (3, 5))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4.5), Inches(4), Inches(2.5), data
    ).name = "Revenue chart"
    slide.notes_slide.notes_text_frame.text = "Mention the pilot."
    linked = deck.slides.add_slide(deck.slide_layouts[5])
    linked.shapes.title.text = "Back"
    linked.shapes.title.click_action.target_slide = first
    deck.save(path)


def _destination(path: Path, wide: bool = False) -> None:
    deck = Presentation()
    if wide:
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for title in ("Opening", "Closing"):
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = title
    deck.save(path)


def test_reuse_slide_brings_content_notes_and_takes_destination_layout(tmp_path: Path) -> None:
    _source(tmp_path / "library.pptx")
    _destination(tmp_path / "talk.pptx")
    _destination(tmp_path / "wide.pptx", wide=True)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        library = _import_presentation(core, project, tmp_path / "library.pptx")["id"]
        talk = _import_presentation(core, project, tmp_path / "talk.pptx")["id"]
        wide = _import_presentation(core, project, tmp_path / "wide.pptx")["id"]
        library_revision = _open(core, project, library)["revision_id"]
        revision = _open(core, project, talk)["revision_id"]

        def slide(document: str, revision_id: str, number: int) -> dict:
            return _call(
                core,
                "l",
                "presentation.workspace.list_text_targets",
                {
                    "project_id": project,
                    "document_id": document,
                    "revision_id": revision_id,
                    "slide_number": number,
                },
            )["result"]

        source = slide(library, library_revision, 2)

        def reuse(document: str, baseline: str, source_slide: str, position: int) -> dict:
            return _call(
                core,
                "r",
                "presentation.workspace.reuse_slide",
                {
                    "project_id": project,
                    "document_id": document,
                    "expected_revision_id": baseline,
                    "source_document_id": library,
                    "source_revision_id": library_revision,
                    "source_slide_id": source_slide,
                    "position": position,
                },
            )

        reply = reuse(talk, revision, source["slide_id"], 2)
        assert reply["ok"], reply.get("error")
        assert reply["result"]["layout_name"] == "Title and Content"
        assert reply["result"]["notes_copied"] is True
        revision = reply["result"]["revision_id"]

        reused = slide(talk, revision, 2)
        assert reused["slide_id"] == reply["result"]["created_slide_id"]
        assert reused["speaker_notes"]["current_text"] == "Mention the pilot."
        names = {o["name"]: o for o in reused["objects"]}
        original = {o["name"]: o for o in source["objects"]}
        assert set(names) == set(original)
        for name, obj in original.items():
            assert names[name]["geometry"] == obj["geometry"], name
        assert slide(talk, revision, 1)["objects"][0]["name"]  # neighbours intact
        titles = [
            next(
                t["current_text"] for t in slide(talk, revision, n)["targets"] if t["current_text"]
            )
            for n in (1, 2, 3)
        ]
        assert titles == ["Opening", "Quarterly results", "Closing"]

        with core._storage.project_database(project) as connection:
            snapshot = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, snapshot)) as archive:
            names_in_package = archive.namelist()
            types = archive.read("[Content_Types].xml").decode()
            rels = archive.read(
                next(n for n in names_in_package if n.startswith("ppt/slides/_rels/slide3"))
            ).decode()
        assert any(n.startswith("ppt/media/copilot-reuse-") for n in names_in_package)
        assert any(n.startswith("ppt/charts/copilot-copy-") for n in names_in_package)
        assert "https://example.org/results" in rels and "slideLayout2.xml" in rels
        assert "notesSlide+xml" in types

        linked = slide(library, library_revision, 3)
        refused = reuse(talk, revision, linked["slide_id"], 1)
        assert refused["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        wide_revision = _open(core, project, wide)["revision_id"]
        sized = reuse(wide, wide_revision, source["slide_id"], 1)
        assert sized["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert reuse(talk, revision, source["slide_id"], 9)["error"]["code"] == "INVALID_REQUEST"
    finally:
        core.close()
