from __future__ import annotations

import base64
import hashlib
import io
import uuid
import zipfile
from pathlib import Path
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def call(core: CoreService, method: str, params: dict[str, Any]) -> Any:
    response = core.handle_message(
        {
            "protocol_version": 1,
            "type": "request",
            "request_id": str(uuid.uuid4()),
            "method": method,
            "params": params,
        }
    )
    assert response["ok"], response
    return response["result"]


def fixture(core: CoreService, tmp_path: Path, count: int = 2) -> tuple[dict[str, Any], Path]:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    for index in range(count):
        image = io.BytesIO()
        Image.new("RGB", (10 + index, 8), (index, 100, 180)).save(image, format="PNG")
        slide.shapes.add_picture(
            io.BytesIO(image.getvalue()), Inches(1), Inches(1), width=Inches(1)
        )
    original = tmp_path / "pictures.pptx"
    deck.save(original)
    with zipfile.ZipFile(original, "a") as archive:
        archive.writestr("ppt/media/unsupported.svg", "<svg/>")
        archive.writestr("ppt/media/broken.png", b"not a picture")
    project = call(core, "project.create", {"name": "Deck pictures"})["project"]["id"]
    document = call(core, "source.import", {"project_id": project, "path": str(original)})[
        "document"
    ]["id"]
    head = core._presentation_edits._current_revision_source(project, document)
    return {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": head["revision_id"],
    }, original


def test_deck_picture_inventory_and_reuse_preserve_bytes(tmp_path: Path) -> None:
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        scope, original = fixture(core, tmp_path)
        original_bytes = original.read_bytes()
        listing = call(core, "presentation.workspace.list_picture_assets", scope)
        assert len(listing["assets"]) == 2
        assert listing["skipped"] == 2
        assert listing["next_cursor"] is None and not listing["has_more"]
        asset = listing["assets"][0]
        with Image.open(io.BytesIO(base64.b64decode(asset["thumbnail_base64"]))) as thumb:
            assert thumb.width <= 128 and thumb.height <= 128
        read = call(
            core,
            "presentation.workspace.read_picture_asset",
            {**scope, "asset_id": asset["asset_id"]},
        )
        data = base64.b64decode(read["image_base64"])
        assert hashlib.sha256(data).hexdigest() == asset["sha256"]
        assert len(data) == asset["byte_size"]
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {
                "project_id": scope["project_id"],
                "document_id": scope["document_id"],
                "revision_id": scope["expected_revision_id"],
                "slide_number": 2,
            },
        )
        result = call(
            core,
            "presentation.workspace.create_picture",
            {
                **scope,
                "slide_id": inventory["slide_id"],
                "geometry": {
                    "x_emu": 10000,
                    "y_emu": 10000,
                    "width_emu": 100000,
                    "height_emu": 100000,
                },
                "alt_text": "Reused visual",
                "image_base64": read["image_base64"],
            },
        )
        assert result["revision_number"] == 2
        head = core._presentation_edits._current_revision_source(
            scope["project_id"], scope["document_id"]
        )
        saved = core._storage.paths.snapshot_path(
            scope["project_id"], head["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original) as before, zipfile.ZipFile(saved) as after:
            assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")
            assert (
                sum(
                    after.read(name) == data
                    for name in after.namelist()
                    if name.startswith("ppt/media/")
                )
                >= 1
            )
        assert original.read_bytes() == original_bytes
        with pytest.raises(CoreDomainError):
            core._presentation_edits.read_picture_asset({**scope, "asset_id": asset["asset_id"]})
    finally:
        core.close()


def test_asset_pages_are_bounded_and_progress(tmp_path: Path) -> None:
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        scope, _ = fixture(core, tmp_path, 21)
        first = call(core, "presentation.workspace.list_picture_assets", scope)
        second = call(
            core,
            "presentation.workspace.list_picture_assets",
            {**scope, "cursor": first["next_cursor"]},
        )
        assert len(first["assets"]) == 20 and first["has_more"]
        assert len(second["assets"]) == 1 and not second["has_more"]
        assert len({a["asset_id"] for page in (first, second) for a in page["assets"]}) == 21
        for cursor in (True, -1, 999):
            with pytest.raises(CoreDomainError):
                core._presentation_edits.list_picture_assets({**scope, "cursor": cursor})
    finally:
        core.close()


def test_asset_scope_identity_and_snapshot_checks(tmp_path: Path) -> None:
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        scope, _ = fixture(core, tmp_path)
        asset = call(core, "presentation.workspace.list_picture_assets", scope)["assets"][0]
        other = call(core, "project.create", {"name": "Other"})["project"]["id"]
        for overrides in (
            {"project_id": other},
            {"expected_revision_id": str(uuid.uuid4())},
            {"asset_id": "0" * 64},
            {"asset_id": "../../image.png"},
            {"path": "C:/private.png"},
        ):
            with pytest.raises(CoreDomainError):
                core._presentation_edits.read_picture_asset(
                    {**scope, "asset_id": asset["asset_id"], **overrides}
                )
        head = core._presentation_edits._current_revision_source(
            scope["project_id"], scope["document_id"]
        )
        path = core._storage.paths.snapshot_path(
            scope["project_id"], head["local_snapshot_path"], require_exists=True
        )
        path.write_bytes(path.read_bytes() + b"changed")
        with pytest.raises(CoreDomainError):
            core._presentation_edits.list_picture_assets(scope)
    finally:
        core.close()
