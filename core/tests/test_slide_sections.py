"""Slide sections and their upkeep across slide-list edits."""

from pathlib import Path

from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService
from presenter_core.presentation.slide_sections import _stable


class _Deck:
    def __init__(self, tmp_path: Path, count: int = 6) -> None:
        source = tmp_path / "sections.pptx"
        deck = Presentation()
        for number in range(1, count + 1):
            deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = f"S{number}"
        deck.save(source)
        self.core = CoreService(data_root=tmp_path / "data")
        self.project = _project_id(self.core)
        self.document = _import_presentation(self.core, self.project, source)["id"]
        self.revision = _open(self.core, self.project, self.document)["revision_id"]

    def scope(self) -> dict:
        return {
            "project_id": self.project,
            "document_id": self.document,
            "expected_revision_id": self.revision,
        }

    def call(self, method: str, **params) -> dict:
        reply = _call(
            self.core, "x", f"presentation.workspace.{method}", {**self.scope(), **params}
        )
        if (
            reply["ok"]
            and "revision_id" in reply["result"]
            and method.startswith(("set_", "delete_", "duplicate_", "create_", "reorder_"))
        ):
            self.revision = reply["result"]["revision_id"]
        return reply

    def titles(self) -> dict[str, str]:
        result, number = {}, 1
        while True:
            reply = _call(
                self.core,
                "t",
                "presentation.workspace.list_text_targets",
                {
                    "project_id": self.project,
                    "document_id": self.document,
                    "revision_id": self.revision,
                    "slide_number": number,
                },
            )
            if not reply["ok"]:
                return result
            inventory = reply["result"]
            title = next(
                (o["current_text"] for o in inventory["targets"] if o.get("current_text")), ""
            )
            result[inventory["slide_id"]] = title or f"new{number}"
            number += 1

    def sections(self) -> list[tuple[str, list[str]]]:
        titles = self.titles()
        state = self.call("get_slide_sections")["result"]
        assert state["reason"] is None, state
        return [(s["name"], [titles[i] for i in s["slide_ids"]]) for s in state["sections"]]

    def by_title(self) -> dict[str, str]:
        return {title: slide_id for slide_id, title in self.titles().items()}


def test_stable_subsequence() -> None:
    assert sorted(_stable([4, 0, 1, 2, 3, 5])) == [1, 2, 3, 4, 5]
    assert _stable([]) == set()


def test_sections_split_and_follow_slide_edits(tmp_path: Path) -> None:
    deck = _Deck(tmp_path)
    try:
        initial = deck.call("get_slide_sections")["result"]
        assert initial["sections"] == [] and initial["editable"]
        ids = deck.by_title()
        reply = deck.call(
            "set_slide_sections",
            expected_fingerprint=initial["fingerprint"],
            sections=[
                {"section_id": None, "name": "Intro", "first_slide_id": ids["S1"]},
                {"section_id": None, "name": "Body", "first_slide_id": ids["S3"]},
                {"section_id": None, "name": "Close", "first_slide_id": ids["S5"]},
            ],
        )
        assert reply["ok"], reply
        assert deck.sections() == [
            ("Intro", ["S1", "S2"]),
            ("Body", ["S3", "S4"]),
            ("Close", ["S5", "S6"]),
        ]
        # Moving S5 to the front joins the section it lands in.
        order = list(deck.titles())
        ids = deck.by_title()
        moved = [ids["S5"]] + [slide for slide in order if slide != ids["S5"]]
        assert deck.call("reorder_slides", ordered_slide_ids=moved)["ok"]
        assert deck.sections() == [
            ("Intro", ["S5", "S1", "S2"]),
            ("Body", ["S3", "S4"]),
            ("Close", ["S6"]),
        ]
        ids = deck.by_title()
        assert deck.call("delete_slide", slide_id=ids["S3"])["ok"]
        assert deck.call("duplicate_slide", slide_id=deck.by_title()["S6"])["ok"]
        assert deck.sections() == [
            ("Intro", ["S5", "S1", "S2"]),
            ("Body", ["S4"]),
            ("Close", ["S6", "S6"]),
        ]
        # Rename and merge: keep Intro's identity, fold Body into Close.
        state = deck.call("get_slide_sections")["result"]
        ids = deck.by_title()
        reply = deck.call(
            "set_slide_sections",
            expected_fingerprint=state["fingerprint"],
            sections=[
                {
                    "section_id": state["sections"][0]["section_id"],
                    "name": "Opening",
                    "first_slide_id": ids["S5"],
                },
                {"section_id": None, "name": "Rest", "first_slide_id": ids["S4"]},
            ],
        )
        assert reply["ok"], reply
        after = deck.call("get_slide_sections")["result"]
        assert after["sections"][0]["section_id"] == state["sections"][0]["section_id"]
        assert deck.sections() == [("Opening", ["S5", "S1", "S2"]), ("Rest", ["S4", "S6", "S6"])]
        stale = deck.call(
            "set_slide_sections", expected_fingerprint=state["fingerprint"], sections=[]
        )
        assert not stale["ok"] and stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        bad = deck.call(
            "set_slide_sections",
            expected_fingerprint=after["fingerprint"],
            sections=[{"section_id": None, "name": "Late", "first_slide_id": ids["S4"]}],
        )
        assert not bad["ok"] and bad["error"]["code"] == "INVALID_REQUEST"
        assert deck.call(
            "set_slide_sections", expected_fingerprint=after["fingerprint"], sections=[]
        )["ok"]
        assert deck.sections() == []
    finally:
        deck.core.close()
