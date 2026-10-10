"""Deck-wide slide numbers and footers."""

import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService
from presenter_core.presentation.fidelity import DRAWING_NS, PRESENTATION_NS

P = f"{{{PRESENTATION_NS}}}"
A = f"{{{DRAWING_NS}}}"


class _Footers:
    def __init__(self, tmp_path: Path) -> None:
        source = tmp_path / "footers.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = "Title"
        for title in ("Agenda", "Details"):
            deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = title
        deck.save(source)
        self.core = CoreService(data_root=tmp_path / "data")
        self.project = _project_id(self.core)
        self.document = _import_presentation(self.core, self.project, source)["id"]
        self.revision = _open(self.core, self.project, self.document)["revision_id"]

    def scope(self) -> dict:
        return {
            "project_id": self.project,
            "document_id": self.document,
            "expected_revision_id": self.revision,
        }

    def state(self) -> dict:
        reply = _call(self.core, "get", "presentation.workspace.get_slide_footers", self.scope())
        assert reply["ok"], reply
        return reply["result"]

    def apply(self, **values) -> dict:
        fingerprint = values.pop("fingerprint", None) or self.state()["fingerprint"]
        reply = _call(
            self.core,
            "set",
            "presentation.workspace.set_slide_footers",
            {**self.scope(), "expected_fingerprint": fingerprint, **values},
        )
        if reply["ok"]:
            self.revision = reply["result"]["revision_id"]
        return reply

    def slide(self, number: int) -> etree._Element:
        with self.core._storage.project_database(self.project) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?",
                (self.revision,),
            ).fetchone()
        path = self.core._storage.paths.snapshot_path(
            self.project, row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(path) as archive:
            return etree.fromstring(archive.read(f"ppt/slides/slide{number}.xml"))


def _kinds(root: etree._Element) -> list[str]:
    return [ph.get("type") for ph in root.iter(f"{P}ph") if ph.get("type") in {"sldNum", "ftr"}]


def test_numbers_and_footer_apply_skip_title_and_remove(tmp_path: Path) -> None:
    deck = _Footers(tmp_path)
    try:
        initial = deck.state()
        assert (initial["slide_numbers"], initial["footers"], initial["slide_count"]) == (
            "none",
            "none",
            3,
        )
        reply = deck.apply(slide_numbers=True, footer_text="Confidential")
        assert reply["ok"] and reply["result"]["slides_changed"] == 2, reply
        assert _kinds(deck.slide(1)) == []
        second = deck.slide(2)
        assert sorted(_kinds(second)) == ["ftr", "sldNum"]
        assert second.find(f".//{A}fld[@type='slidenum']") is not None
        assert "Confidential" in "".join(t.text or "" for t in second.iter(f"{A}t"))
        state = deck.state()
        assert (state["slide_numbers"], state["footer_text"]) == ("some", "Confidential")
        assert deck.apply(slide_numbers=True, footer_text="Internal")["ok"]
        assert deck.state()["footer_text"] == "Internal"
        assert deck.apply(slide_numbers=False, footer_text=None)["ok"]
        assert all(_kinds(deck.slide(n)) == [] for n in (1, 2, 3))
        same = deck.apply(slide_numbers=False, footer_text=None)
        assert same["ok"] and same["result"]["changed"] is False
    finally:
        deck.core.close()


def test_title_slide_can_be_included_and_stale_requests_fail(tmp_path: Path) -> None:
    deck = _Footers(tmp_path)
    try:
        stale = deck.apply(slide_numbers=True, footer_text=None, fingerprint="0" * 64)
        assert not stale["ok"] and stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        reply = deck.apply(slide_numbers=True, footer_text=None, skip_title_slides=False)
        assert reply["ok"] and reply["result"]["slides_changed"] == 3, reply
        assert deck.state()["slide_numbers"] == "all"
        bad = deck.apply(slide_numbers=True, footer_text="   ")
        assert not bad["ok"] and bad["error"]["code"] == "INVALID_REQUEST"
    finally:
        deck.core.close()
