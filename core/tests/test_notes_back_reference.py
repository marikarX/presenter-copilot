"""Slides whose notes part omits the notes-to-slide back reference."""

import re
import zipfile
from pathlib import Path

from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _deck(tmp_path: Path) -> Path:
    source = tmp_path / "notes.pptx"
    deck = Presentation()
    for title in ("One", "Two"):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = f"Say {title}"
    deck.save(source)
    stripped = tmp_path / "stripped.pptx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(stripped, "w") as output:
        for info in original.infolist():
            data = original.read(info)
            if re.fullmatch(r"ppt/notesSlides/_rels/notesSlide\d+\.xml\.rels", info.filename):
                data = re.sub(rb'<Relationship [^>]*relationships/slide"[^>]*/>', b"", data)
                data = re.sub(rb'<Relationship [^>]*Target="\.\./slides/[^"]*"[^>]*/>', b"", data)
                assert b"/slides/" not in data
            output.writestr(info, data)
    return stripped


def test_duplicate_and_delete_without_notes_back_reference(tmp_path: Path) -> None:
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, _deck(tmp_path))["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def slide(number: int) -> dict:
            return _call(
                core,
                "t",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": number},
            )["result"]

        first = slide(1)
        duplicated = _call(
            core,
            "d",
            "presentation.workspace.duplicate_slide",
            {**scope, "expected_revision_id": revision, "slide_id": first["slide_id"]},
        )
        assert duplicated["ok"], duplicated
        revision = duplicated["result"]["revision_id"]
        copy = slide(2)
        assert copy["speaker_notes"]["current_text"] == "Say One"
        with core._storage.project_database(project) as connection:
            path = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, path)) as archive:
            back = [
                name
                for name in archive.namelist()
                if name.startswith("ppt/notesSlides/_rels/") and b"/slides/" in archive.read(name)
            ]
        assert len(back) == 1
        deleted = _call(
            core,
            "x",
            "presentation.workspace.delete_slide",
            {**scope, "expected_revision_id": revision, "slide_id": first["slide_id"]},
        )
        assert deleted["ok"], deleted
    finally:
        core.close()
