"""Give untitled slides a title placeholder."""

from pathlib import Path

from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def test_untitled_slides_get_a_title_and_the_check_clears(tmp_path: Path) -> None:
    source = tmp_path / "untitled.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])  # Blank layout: no title placeholder.
    title_only = deck.slides.add_slide(deck.slide_layouts[5])
    title_only.shapes._spTree.remove(title_only.shapes.title._element)
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def check() -> list:
            return _call(
                core,
                "a",
                "presentation.workspace.check_accessibility",
                {**scope, "revision_id": revision},
            )["result"]["issues"]

        missing = check()
        assert [item["slide_number"] for item in missing] == [1, 2]
        for item in missing:
            reply = _call(
                core,
                "t",
                "presentation.workspace.set_slide_title",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": item["slide_id"],
                    "title": f"Slide {item['slide_number']} summary",
                },
            )
            assert reply["ok"], reply.get("error")
            revision = reply["result"]["revision_id"]
        assert check() == []
        again = _call(
            core,
            "t",
            "presentation.workspace.set_slide_title",
            {
                **scope,
                "expected_revision_id": revision,
                "slide_id": missing[0]["slide_id"],
                "title": "Another",
            },
        )
        assert not again["ok"] and again["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        title = next(
            o
            for o in _call(
                core,
                "l",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]["objects"]
            if o["name"] == "Title"
        )
        # The blank layout has no title placeholder: the master title position is written.
        assert title["geometry_editable"] and not title["geometry_inherited"]
    finally:
        core.close()
