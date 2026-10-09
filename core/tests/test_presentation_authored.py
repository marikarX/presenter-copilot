"""Authored insertion reviews exact bytes without reserializing imported content."""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from presenter_core.ipc.core import CoreService
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter
from tests.test_presentation_planning import call


@pytest.mark.parametrize("layout", ["points", "columns", "cards", "statement"])
@pytest.mark.parametrize("with_notes", [False, True])
def test_authored_insert_review_apply_and_undo_preserves_existing_parts(
    tmp_path: Path, layout: str, with_notes: bool
) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(7.5)
    original = deck.slides.add_slide(deck.slide_layouts[5])
    original.shapes.title.text = "Original title"
    if with_notes:
        original.notes_slide.notes_text_frame.text = "Original notes"
    path = tmp_path / "source.pptx"
    deck.save(str(path))
    source_digest = hashlib.sha256(path.read_bytes()).hexdigest()
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        project = call(core, "project.create", {"name": "Authored"})["project"]["id"]
        document = call(
            core,
            "source.import",
            {"project_id": project, "kind": "presentation", "path": str(path)},
        )["document"]["id"]
        scope = {"project_id": project, "document_id": document}
        baseline = core._presentation_edits._current_revision_source(project, document)
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": baseline["revision_id"], "slide_number": 1},
        )
        params = {
            "slide_id": inventory["slide_id"],
            "title": "New decision",
            "bullets": ["Approve a measured pilot"]
            if layout == "statement"
            else ["Approve a measured pilot", "Review the outcome"],
            "notes": "Ask for approval\nExplain the tradeoffs",
            "template": "midnight",
            "layout": layout,
        }
        plan = call(
            core,
            "presentation.workspace.prepare_edit_plan",
            {
                **scope,
                "expected_revision_id": baseline["revision_id"],
                "summary": "Insert authored content",
                "operations": [{"type": "create_authored_slide", "params": params}],
            },
        )
        assert (
            core._presentation_edits._current_revision_source(project, document)["revision_id"]
            == baseline["revision_id"]
        )
        preview = call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": plan["plan_id"]},
        )
        assert preview["fidelity_report"]["slide_count"] == 2
        call(
            core,
            "presentation.workspace.close_source",
            {**scope, "session_id": preview["session_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": plan["plan_id"],
                "reviewed_sha256": plan["sha256"],
                "confirm": True,
            },
        )
        assert result["revision_number"] == 2
        current = core._presentation_edits._current_revision_source(project, document)
        assert current["sha256"] == plan["sha256"]
        current_path = core._storage.paths.snapshot_path(
            project, current["local_snapshot_path"], require_exists=True
        )
        revised = Presentation(str(current_path))
        assert len(revised.slides) == 2
        assert revised.slides[1].notes_slide.notes_text_frame.text == params["notes"]
        assert [
            shape.text
            for shape in revised.slides[1].shapes
            if shape.name.startswith("Point ")
            and shape.has_text_frame
            and not shape.name.endswith(" panel")
        ] == params["bullets"]
        with zipfile.ZipFile(path) as before, zipfile.ZipFile(current_path) as after:
            allowed = {
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
                "[Content_Types].xml",
            }
            for member in before.namelist():
                if member not in allowed:
                    assert before.read(member) == after.read(member), member
        authored = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": current["revision_id"], "slide_number": 2},
        )
        assert authored["speaker_notes"]["editable"] is True
        target = next(
            item for item in authored["targets"] if item["current_text"] == "New decision"
        )
        edited = call(
            core,
            "presentation.workspace.replace_text_shape",
            {
                **scope,
                "expected_revision_id": current["revision_id"],
                "slide_id": authored["slide_id"],
                "object_id": target["object_id"],
                "operation_type": "replace_text_shape",
                "expected_text": "New decision",
                "replacement_text": "Reviewed decision",
            },
        )
        assert edited["revision_number"] == 3
        assert hashlib.sha256(path.read_bytes()).hexdigest() == source_digest
        call(
            core,
            "presentation.workspace.restore_revision",
            {
                **scope,
                "expected_revision_id": edited["revision_id"],
                "revision_id": baseline["revision_id"],
            },
        )
        assert (
            core._presentation_edits._current_revision_source(project, document)["sha256"]
            == baseline["sha256"]
        )
    finally:
        core.close()
