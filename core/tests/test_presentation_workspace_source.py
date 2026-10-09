from __future__ import annotations

import base64
import copy
import hashlib
import io
import json
import posixpath
import sqlite3
import uuid
import zipfile
from pathlib import Path
from typing import Any

import pytest
from lxml import etree
from PIL import Image

from presenter_core.credentials import InMemoryCredentialStore
from presenter_core.errors import CoreDomainError
from presenter_core.ipc.core import CoreService
from presenter_core.presentation.editing import PresentationEditService
from presenter_core.presentation.fidelity import (
    DRAWING_NS,
    MAX_TEXT_FIRST_LINE_INDENT_EMU,
    MAX_TEXT_LEFT_INDENT_EMU,
    PRESENTATION_NS,
    _delete_geometry_candidate,
    _paragraph_bullet_mode_for_shape,
    _paragraph_property_insert_index,
    inspect_pptx_snapshot,
)
from presenter_core.presentation.workspace_source import READER_TTL_SECONDS

FIXTURE_ROOT = Path(__file__).parents[2] / "samples" / "synthetic-deck"


def _call(
    core: CoreService, request_id: str, method: str, params: dict[str, Any]
) -> dict[str, Any]:
    return core.handle_message(
        {
            "protocol_version": 1,
            "type": "request",
            "request_id": request_id,
            "method": method,
            "params": params,
        }
    )


def _project_id(core: CoreService) -> str:
    result = _call(core, "project", "project.create", {"name": "Workspace preview"})
    assert result["ok"] is True
    return result["result"]["project"]["id"]


def _import_presentation(
    core: CoreService, project_id: str, source_path: Path | None = None
) -> dict[str, Any]:
    result = _call(
        core,
        "import-pptx",
        "source.import",
        {
            "project_id": project_id,
            "path": str((source_path or FIXTURE_ROOT / "deck" / "presentation.pptx").resolve()),
            "kind": "presentation",
        },
    )
    assert result["ok"] is True
    return result["result"]["document"]


def _write_multirun_text_fixture(
    source_path: Path,
    destination_path: Path,
    *,
    uniform_runs: bool = False,
) -> tuple[int, str]:
    with source_path.open("rb") as source_handle:
        source_bytes = source_handle.read()
    digest = hashlib.sha256(source_bytes).hexdigest()
    with io.BytesIO(source_bytes) as source_handle:
        fidelity = inspect_pptx_snapshot(
            source_handle,
            document_id=str(uuid.uuid4()),
            expected_sha256=digest,
            byte_size=len(source_bytes),
        )
    slide = next(
        slide
        for slide in fidelity.slides
        if any(obj.single_run_text_edit_candidate and len(obj.text) >= 4 for obj in slide.objects)
    )
    target = next(
        obj for obj in slide.objects if obj.single_run_text_edit_candidate and len(obj.text) >= 4
    )
    namespaces = {"p": PRESENTATION_NS, "a": DRAWING_NS}
    with zipfile.ZipFile(source_path) as source_archive:
        slide_root = etree.fromstring(source_archive.read(slide.part_name))
        shapes = slide_root.xpath(
            ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
            namespaces=namespaces,
            shape_id=target.ooxml_id,
        )
        assert len(shapes) == 1
        runs = shapes[0].xpath("./p:txBody/a:p/a:r", namespaces=namespaces)
        assert len(runs) == 1
        run = runs[0]
        text_node = run.find(f"{{{DRAWING_NS}}}t")
        assert text_node is not None and text_node.text is not None
        original_text = text_node.text
        split_at = len(original_text) // 2
        assert split_at > 0
        text_node.text = original_text[:split_at]
        second_run = copy.deepcopy(run)
        second_text = second_run.find(f"{{{DRAWING_NS}}}t")
        assert second_text is not None
        second_text.text = original_text[split_at:]
        paragraph = run.getparent()
        assert paragraph is not None
        paragraph_properties = paragraph.find(f"{{{DRAWING_NS}}}pPr")
        if paragraph_properties is None:
            paragraph_properties = etree.Element(f"{{{DRAWING_NS}}}pPr")
            paragraph.insert(0, paragraph_properties)
        paragraph_properties.set("marL", "228600")
        paragraph_properties.set("indent", "-114300")
        etree.SubElement(paragraph_properties, f"{{{DRAWING_NS}}}buChar", char="•")
        first_properties = run.find(f"{{{DRAWING_NS}}}rPr")
        if first_properties is None:
            first_properties = etree.Element(f"{{{DRAWING_NS}}}rPr")
            run.insert(0, first_properties)
        second_properties = second_run.find(f"{{{DRAWING_NS}}}rPr")
        if second_properties is None:
            second_properties = etree.Element(f"{{{DRAWING_NS}}}rPr")
            second_run.insert(0, second_properties)
        first_properties.set("b", "0")
        first_properties.set("i", "1")
        if uniform_runs:
            second_run.replace(second_properties, copy.deepcopy(first_properties))
        else:
            second_properties.set("b", "1")
            second_properties.set("i", "0")
            second_properties.set("u", "sng")
        run.addnext(second_run)
        rewritten_slide = etree.tostring(
            slide_root,
            encoding="UTF-8",
            xml_declaration=True,
            standalone=True,
        )
        with zipfile.ZipFile(destination_path, "w", allowZip64=True) as output_archive:
            for info in source_archive.infolist():
                output_archive.writestr(
                    info,
                    rewritten_slide
                    if info.filename == slide.part_name
                    else source_archive.read(info),
                )
    return slide.ordinal, target.ooxml_id


def _open(core: CoreService, project_id: str, document_id: str) -> dict[str, Any]:
    result = _call(
        core,
        "open-presentation",
        "presentation.workspace.open_source",
        {"project_id": project_id, "document_id": document_id},
    )
    assert result["ok"] is True, result
    return result["result"]


def _list_revisions(
    core: CoreService,
    project_id: str,
    document_id: str,
    **options: Any,
) -> dict[str, Any]:
    result = _call(
        core,
        "list-presentation-revisions",
        "presentation.workspace.list_revisions",
        {"project_id": project_id, "document_id": document_id, **options},
    )
    assert result["ok"] is True, result
    return result["result"]


def _add_test_revision(
    core: CoreService,
    project_id: str,
    document: dict[str, Any],
    revision_id: str,
    snapshot_path: Path,
) -> None:
    with core._storage.project_database(project_id) as connection:
        current = connection.execute(
            "SELECT revision_id FROM presentation_revision_heads WHERE presentation_id = ?",
            (document["id"],),
        ).fetchone()
        assert current is not None
        parent_revision_id = str(current["revision_id"])
        connection.execute(
            """
            INSERT INTO presentation_revisions (
                id, presentation_id, parent_revision_id, revision_number,
                revision_kind, local_snapshot_path, sha256, byte_size, created_at
            ) VALUES (?, ?, ?, 2, 'edited', ?, ?, ?, '2026-09-29T00:00:00.000Z')
            """,
            (
                revision_id,
                document["id"],
                parent_revision_id,
                f"sources/{snapshot_path.name}",
                document["sha256"],
                document["byte_size"],
            ),
        )
        connection.execute(
            "UPDATE presentation_revision_heads SET revision_id = ? WHERE presentation_id = ?",
            (revision_id, document["id"]),
        )
        connection.commit()
    snapshot_path.write_bytes((FIXTURE_ROOT / "deck" / "presentation.pptx").read_bytes())


def test_workspace_source_stream_is_chunked_ordered_and_hash_verified(tmp_path: Path) -> None:
    original = (FIXTURE_ROOT / "deck" / "presentation.pptx").read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        snapshot_path = (
            tmp_path / "data" / "projects" / project_id / "sources" / document["snapshot_name"]
        )
        opened = _open(core, project_id, document["id"])
        assert opened["byte_size"] == len(original)
        assert opened["chunk_bytes"] == 512 * 1024
        report = opened["fidelity_report"]
        assert report["status"] == "inventory_only"
        assert report["slide_count"] == 22
        assert report["object_count"] > 0
        assert set(report["object_kinds"]) == {
            "shapes",
            "pictures",
            "graphic_frames",
            "connectors",
            "groups",
        }
        assert sum(report["object_kinds"].values()) == report["object_count"]
        assert report["text_edit_candidates"] > 0
        assert report["nontext_or_unmodeled_objects"] >= 0
        assert str((tmp_path / "data").resolve()) not in str(opened)
        assert "What would you ask it to do" not in json.dumps(report)
        history = _list_revisions(core, project_id, document["id"])
        assert history["presentation_id"] == document["id"]
        assert history["current_revision_id"] == history["revisions"][0]["revision_id"]
        assert history["revisions"] == [
            {
                "revision_id": history["current_revision_id"],
                "parent_revision_id": None,
                "revision_number": 1,
                "revision_kind": "original",
                "sha256": document["sha256"],
                "byte_size": len(original),
                "created_at": document["imported_at"],
                "is_current": True,
            }
        ]
        invalid_history = _call(
            core,
            "too-many-revisions",
            "presentation.workspace.list_revisions",
            {"project_id": project_id, "document_id": document["id"], "limit": 101},
        )
        assert invalid_history["ok"] is False
        assert invalid_history["error"]["code"] == "INVALID_REQUEST"

        chunks: list[bytes] = []
        offset = 0
        request_id = 0
        while offset < opened["byte_size"]:
            request_id += 1
            result = _call(
                core,
                f"read-{request_id}",
                "presentation.workspace.read_source",
                {
                    "project_id": project_id,
                    "document_id": document["id"],
                    "session_id": opened["session_id"],
                    "offset": offset,
                    "max_bytes": 4096,
                },
            )
            assert result["ok"] is True, result
            payload = result["result"]
            assert payload["offset"] == offset
            chunk = base64.b64decode(payload["bytes_b64"], validate=True)
            assert 0 < len(chunk) <= 4096
            assert payload["complete"] is (offset + len(chunk) == len(original))
            if offset == 0:
                bad_offset = _call(
                    core,
                    "read-out-of-order",
                    "presentation.workspace.read_source",
                    {
                        "project_id": project_id,
                        "document_id": document["id"],
                        "session_id": opened["session_id"],
                        "offset": 1,
                        "max_bytes": 4096,
                    },
                )
                assert bad_offset["error"]["code"] == "INVALID_REQUEST"
            chunks.append(chunk)
            offset += len(chunk)

        rendered_input = b"".join(chunks)
        assert rendered_input == original
        assert hashlib.sha256(rendered_input).hexdigest() == document["sha256"]
        assert not core._presentation_workspace_sources._readers

        reopened = _open(core, project_id, document["id"])
        assert reopened["fidelity_report"]["slides"] == report["slides"]
        core._presentation_workspace_sources.close_source(
            {
                "project_id": project_id,
                "document_id": document["id"],
                "session_id": reopened["session_id"],
            }
        )
        assert snapshot_path.read_bytes() == original
    finally:
        core.close()


def test_fidelity_inventory_uses_repeatable_ids_and_rejects_dtd(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    digest = hashlib.sha256(original_bytes).hexdigest()
    with original_path.open("rb") as handle:
        first = inspect_pptx_snapshot(
            handle,
            document_id="00000000-0000-4000-8000-000000000001",
            expected_sha256=digest,
            byte_size=len(original_bytes),
        )
    with original_path.open("rb") as handle:
        second = inspect_pptx_snapshot(
            handle,
            document_id="00000000-0000-4000-8000-000000000001",
            expected_sha256=digest,
            byte_size=len(original_bytes),
        )
    ids = [slide.slide_id for slide in first.slides]
    assert ids == [slide.slide_id for slide in second.slides]
    assert len(ids) == len(set(ids))
    assert (
        len({obj.object_id for slide in first.slides for obj in slide.objects})
        == first.object_count
    )
    assert original_path.read_bytes() == original_bytes

    # Revision content changes the source hash, while stable OOXML slide and
    # shape identifiers retain their document-lineage identities.
    revision_path = tmp_path / "revision.pptx"
    with (
        zipfile.ZipFile(original_path, "r") as source,
        zipfile.ZipFile(revision_path, "w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == first.slides[0].part_name:
                original_text = b"Northstar migration decision"
                assert payload.count(original_text) == 1
                payload = payload.replace(original_text, b"Northstar migration recommendation", 1)
            destination.writestr(info, payload)
    revision_bytes = revision_path.read_bytes()
    revision_digest = hashlib.sha256(revision_bytes).hexdigest()
    assert revision_digest != digest
    with revision_path.open("rb") as handle:
        revision = inspect_pptx_snapshot(
            handle,
            document_id="00000000-0000-4000-8000-000000000001",
            expected_sha256=revision_digest,
            byte_size=len(revision_bytes),
        )
    assert [slide.slide_id for slide in revision.slides] == [
        slide.slide_id for slide in first.slides
    ]
    assert [[item.object_id for item in slide.objects] for slide in revision.slides] == [
        [item.object_id for item in slide.objects] for slide in first.slides
    ]
    assert revision.slides[0].objects != first.slides[0].objects

    with original_path.open("rb") as handle:
        another_document = inspect_pptx_snapshot(
            handle,
            document_id="00000000-0000-4000-8000-000000000002",
            expected_sha256=digest,
            byte_size=len(original_bytes),
        )
    assert [slide.slide_id for slide in another_document.slides] != ids

    malicious_path = tmp_path / "dtd.pptx"
    with (
        zipfile.ZipFile(original_path, "r") as source,
        zipfile.ZipFile(malicious_path, "w", compression=zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == first.slides[0].part_name:
                declaration_end = payload.find(b"?>")
                assert declaration_end >= 0
                payload = (
                    payload[: declaration_end + 2]
                    + b'<!DOCTYPE p:sld [<!ENTITY injected "blocked">]>'
                    + payload[declaration_end + 2 :]
                )
            destination.writestr(info, payload)

    malicious_bytes = malicious_path.read_bytes()
    with malicious_path.open("rb") as handle, pytest.raises(CoreDomainError) as error:
        inspect_pptx_snapshot(
            handle,
            document_id="00000000-0000-4000-8000-000000000001",
            expected_sha256=hashlib.sha256(malicious_bytes).hexdigest(),
            byte_size=len(malicious_bytes),
        )
    assert error.value.code == "PRESENTATION_FIDELITY_STRUCTURE_INVALID"


def test_project_v7_migration_backfills_original_presentation_revision(tmp_path: Path) -> None:
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        project_database = tmp_path / "data" / "projects" / project_id / "project.db"
        with sqlite3.connect(project_database) as connection:
            connection.execute("PRAGMA foreign_keys = OFF")
            connection.execute("ALTER TABLE sessions DROP COLUMN presentation_revision_id")
            connection.execute("DROP INDEX source_units_presentation_revision_idx")
            connection.execute("ALTER TABLE source_units DROP COLUMN presentation_revision_id")
            connection.execute("DROP TABLE presentation_rehearsal_objection_findings")
            connection.execute("DROP TABLE presentation_rehearsal_objection_analyses")
            connection.execute("DROP TABLE presentation_rehearsal_analyses")
            connection.execute("DROP TABLE presentation_revision_checkpoints")
            connection.execute("DROP TABLE presentation_outline_sessions")
            connection.execute("DROP TABLE presentation_outline_proposals")
            connection.execute("DROP TABLE presentation_slide_takeaways")
            connection.execute("DROP TABLE presentation_narratives")
            connection.execute("DROP TABLE presentation_claims")
            connection.execute("DROP TABLE project_picture_assets")
            connection.execute("DROP TABLE presentation_edit_transactions")
            connection.execute("DROP TABLE presentation_rehearsal_reviews")
            connection.execute("DROP TABLE presentation_slide_purposes")
            connection.execute("DROP TABLE presentation_revision_heads")
            connection.execute("DROP TABLE presentation_revisions")
            connection.execute("DROP TABLE presentations")
            connection.execute("UPDATE project SET schema_version = 7 WHERE id = ?", (project_id,))
            connection.execute("PRAGMA user_version = 7")

        history = _list_revisions(core, project_id, document["id"])
        assert history["presentation_id"] == document["id"]
        assert len(history["revisions"]) == 1
        original = history["revisions"][0]
        assert original["revision_kind"] == "original"
        assert original["revision_number"] == 1
        assert original["sha256"] == document["sha256"]
        assert original["is_current"] is True
        with core._storage.project_database(project_id) as connection:
            scoped_units = connection.execute(
                "SELECT COUNT(*) FROM source_units WHERE document_id = ? "
                "AND presentation_revision_id = ?",
                (document["id"], original["revision_id"]),
            ).fetchone()[0]
        assert scoped_units > 0
    finally:
        core.close()


def test_source_delete_closes_workspace_reader_before_removing_snapshot(
    tmp_path: Path,
) -> None:
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        snapshot = (
            tmp_path / "data" / "projects" / project_id / "sources" / document["snapshot_name"]
        )
        revision_snapshot = snapshot.parent / f"revision-{uuid.uuid4()}.pptx"
        _add_test_revision(
            core,
            project_id,
            document,
            revision_snapshot.stem.removeprefix("revision-"),
            revision_snapshot,
        )
        started = _call(
            core,
            "start-session-before-source-delete",
            "session.start",
            {"project_id": project_id, "mode": "run"},
        )
        assert started["ok"] is True, started
        session_id = started["result"]["session"]["id"]
        first_page = _list_revisions(core, project_id, document["id"], limit=1)
        assert first_page["has_more"] is True
        assert first_page["next_before_revision_number"] == 2
        assert first_page["revisions"][0]["revision_number"] == 2
        second_page = _list_revisions(
            core,
            project_id,
            document["id"],
            limit=1,
            before_revision_number=first_page["next_before_revision_number"],
        )
        assert second_page["has_more"] is False
        assert second_page["revisions"][0]["revision_number"] == 1
        assert str(tmp_path) not in json.dumps(first_page)
        opened = _open(core, project_id, document["id"])
        first_chunk = _call(
            core,
            "read-first",
            "presentation.workspace.read_source",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "session_id": opened["session_id"],
                "offset": 0,
                "max_bytes": 32,
            },
        )
        assert first_chunk["ok"] is True
        assert snapshot.is_file()

        blocked_delete = _call(
            core,
            "delete-presentation-during-session",
            "source.delete",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert blocked_delete["ok"] is False
        assert blocked_delete["error"]["code"] == "PRESENTATION_SESSION_ACTIVE"
        assert snapshot.is_file()
        assert revision_snapshot.is_file()
        assert opened["session_id"] in core._presentation_workspace_sources._readers
        stopped = _call(
            core,
            "stop-session-before-source-delete",
            "session.stop",
            {"project_id": project_id, "session_id": session_id, "status": "aborted"},
        )
        assert stopped["ok"] is True, stopped

        deleted = _call(
            core,
            "delete-source",
            "source.delete",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert deleted["ok"] is True, deleted
        assert deleted["result"]["deleted"] is True
        assert not snapshot.exists()
        assert not revision_snapshot.exists()
        assert not core._presentation_workspace_sources._readers
        session_after_delete = _call(
            core,
            "session-after-presentation-delete",
            "session.get",
            {"project_id": project_id, "session_id": session_id},
        )
        assert session_after_delete["ok"] is True, session_after_delete
        assert session_after_delete["result"]["session"]["presentation_revision_id"] is None
        with core._storage.project_database(project_id) as connection:
            current_presentation_id = connection.execute(
                "SELECT current_presentation_id FROM project WHERE id = ?",
                (project_id,),
            ).fetchone()[0]
        assert current_presentation_id is None
    finally:
        core.close()


def test_project_delete_and_shutdown_close_workspace_readers(tmp_path: Path) -> None:
    core = CoreService(data_root=tmp_path / "data")
    project_id = _project_id(core)
    document = _import_presentation(core, project_id)
    opened = _open(core, project_id, document["id"])
    assert opened["session_id"] in core._presentation_workspace_sources._readers

    deleted = _call(core, "delete-project", "project.delete", {"project_id": project_id})
    assert deleted["ok"] is True, deleted
    assert not core._presentation_workspace_sources._readers

    second_project_id = _project_id(core)
    second_document = _import_presentation(core, second_project_id)
    second_opened = _open(core, second_project_id, second_document["id"])
    assert second_opened["session_id"] in core._presentation_workspace_sources._readers
    assert core.close() is True
    assert not core._presentation_workspace_sources._readers


def test_workspace_export_saves_a_verified_byte_identical_copy(tmp_path: Path) -> None:
    original = (FIXTURE_ROOT / "deck" / "presentation.pptx").read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        snapshot_path = (
            tmp_path / "data" / "projects" / project_id / "sources" / document["snapshot_name"]
        )
        destination_dir = tmp_path / "exports"
        destination_dir.mkdir()
        destination = destination_dir / "saved-copy.pptx"
        destination.write_bytes(b"previous file contents")

        result = _call(
            core,
            "export-original-copy",
            "presentation.workspace.export_original_copy",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "destination_path": str(destination),
            },
        )

        assert result["ok"] is True, result
        assert result["result"] == {
            "saved": True,
            "file_name": destination.name,
            "byte_size": len(original),
            "sha256": document["sha256"],
            "copy_kind": "original",
        }
        assert destination.read_bytes() == original
        assert hashlib.sha256(destination.read_bytes()).hexdigest() == document["sha256"]
        assert snapshot_path.read_bytes() == original
        assert not list(destination_dir.glob(".presenter-copilot-*.tmp"))
    finally:
        core.close()


def test_workspace_export_rejects_unsafe_destinations_without_touching_source(
    tmp_path: Path,
) -> None:
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        destination_dir = tmp_path / "exports"
        destination_dir.mkdir()
        destination = destination_dir / "preserve-existing.pptx"
        destination.write_bytes(b"keep this file on failure")

        for candidate in (
            str(destination_dir / "wrong-extension.pdf"),
            str(tmp_path / "data" / "copy.pptx"),
            str(
                tmp_path / "data" / "projects" / project_id / "sources" / document["snapshot_name"]
            ),
        ):
            result = _call(
                core,
                f"reject-export-{len(candidate)}",
                "presentation.workspace.export_original_copy",
                {
                    "project_id": project_id,
                    "document_id": document["id"],
                    "destination_path": candidate,
                },
            )
            assert result["ok"] is False
            assert result["error"]["code"] == "PRESENTATION_EXPORT_DESTINATION_INVALID"

        source = tmp_path / "data" / "projects" / project_id / "sources" / document["snapshot_name"]
        changed = bytearray(source.read_bytes())
        changed[-1] ^= 1
        source.write_bytes(changed)
        failed_copy = _call(
            core,
            "export-changed-source",
            "presentation.workspace.export_original_copy",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "destination_path": str(destination),
            },
        )

        assert failed_copy["ok"] is False
        assert failed_copy["error"]["code"] == "SOURCE_SNAPSHOT_CHANGED"
        assert destination.read_bytes() == b"keep this file on failure"
        assert not list(destination_dir.glob(".presenter-copilot-*.tmp"))
    finally:
        core.close()


def test_workspace_source_rejects_supporting_files_and_changed_snapshots(
    tmp_path: Path,
) -> None:
    text_source = tmp_path / "notes.md"
    text_source.write_text("Supporting notes", encoding="utf-8")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        imported = _call(
            core,
            "import-notes",
            "source.import",
            {"project_id": project_id, "path": str(text_source), "kind": "supporting"},
        )
        assert imported["ok"] is True
        rejected = _call(
            core,
            "open-notes",
            "presentation.workspace.open_source",
            {"project_id": project_id, "document_id": imported["result"]["document"]["id"]},
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_WORKSPACE_SOURCE_INVALID"

        document = _import_presentation(core, project_id)
        snapshot = (
            tmp_path / "data" / "projects" / project_id / "sources" / document["snapshot_name"]
        )
        snapshot.write_bytes(snapshot.read_bytes() + b"tamper")
        changed = _call(
            core,
            "open-tampered",
            "presentation.workspace.open_source",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert changed["ok"] is False
        assert changed["error"]["code"] == "SOURCE_SNAPSHOT_CHANGED"
    finally:
        core.close()


def test_expired_workspace_readers_and_local_reset_release_file_handles(
    tmp_path: Path,
) -> None:
    core = CoreService(
        data_root=tmp_path / "data",
        credential_store=InMemoryCredentialStore(),
    )
    clock = [10.0]
    core._presentation_workspace_sources._clock = lambda: clock[0]
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        clock[0] += READER_TTL_SECONDS + 1
        expired = _call(
            core,
            "read-expired",
            "presentation.workspace.read_source",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "session_id": opened["session_id"],
                "offset": 0,
                "max_bytes": 32,
            },
        )
        assert expired["error"]["code"] == "PRESENTATION_WORKSPACE_READER_EXPIRED"
        assert not core._presentation_workspace_sources._readers

        second_open = _open(core, project_id, document["id"])
        reset = _call(core, "reset", "app.reset_local_data", {"confirm": True})
        assert reset["ok"] is True, reset
        assert not core._presentation_workspace_sources._readers
        assert not (tmp_path / "data" / "projects" / project_id).exists()
        assert second_open["session_id"] not in core._presentation_workspace_sources._readers
    finally:
        core.close()


def test_speaker_notes_edit_creates_revision_preserves_package_and_reindexes(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    original_sha256 = hashlib.sha256(original_bytes).hexdigest()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        original_revision = _list_revisions(core, project_id, document["id"])["revisions"][0]
        opened = _open(core, project_id, document["id"])
        target_result = _call(
            core,
            "speaker-notes-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert target_result["ok"] is True, target_result
        target = target_result["result"]["speaker_notes"]
        assert target == {
            "current_text": "Synthetic presenter note for slide 1.",
            "text_truncated": False,
            "editable": True,
        }

        replacement = "Updated speaker note marker\nSecond note paragraph\n"
        edited = _call(
            core,
            "replace-speaker-notes",
            "presentation.workspace.replace_speaker_notes",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": target_result["result"]["slide_id"],
                "expected_text": target["current_text"],
                "replacement_text": replacement,
            },
        )
        assert edited["ok"] is True, edited
        assert edited["result"]["changed"] is True
        assert edited["result"]["revision_number"] == 2
        assert edited["result"]["parent_revision_id"] == opened["revision_id"]

        stale_revision = _call(
            core,
            "stale-speaker-notes-revision",
            "presentation.workspace.replace_speaker_notes",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": target_result["result"]["slide_id"],
                "expected_text": target["current_text"],
                "replacement_text": "stale edit",
            },
        )
        assert stale_revision["ok"] is False
        assert stale_revision["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"
        stale_text = _call(
            core,
            "stale-speaker-notes-target",
            "presentation.workspace.replace_speaker_notes",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "slide_id": target_result["result"]["slide_id"],
                "expected_text": target["current_text"],
                "replacement_text": "stale edit",
            },
        )
        assert stale_text["ok"] is False
        assert stale_text["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"

        revised_targets = _call(
            core,
            "revised-speaker-notes-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": edited["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        assert revised_targets["ok"] is True, revised_targets
        assert revised_targets["result"]["speaker_notes"] == {
            "current_text": replacement,
            "text_truncated": False,
            "editable": True,
        }

        revisions = _list_revisions(core, project_id, document["id"])["revisions"]
        assert revisions[0]["revision_id"] == edited["result"]["revision_id"]
        assert revisions[1]["revision_id"] == original_revision["revision_id"]
        with core._storage.project_database(project_id) as connection:
            edited_row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (edited["result"]["revision_id"],),
            ).fetchone()
        assert edited_row is not None
        edited_path = core._storage.paths.snapshot_path(
            project_id, edited_row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(edited_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
        assert changed_parts == {"ppt/notesSlides/notesSlide1.xml"}
        assert original_path.read_bytes() == original_bytes
        assert hashlib.sha256(original_bytes).hexdigest() == original_sha256

        reindexed = _call(
            core,
            "reindex-speaker-notes-revision",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        current_search = _call(
            core,
            "search-edited-speaker-notes",
            "search.lexical",
            {"project_id": project_id, "query": "Updated speaker note marker"},
        )
        assert current_search["ok"] is True, current_search
        assert current_search["result"]["results"]
        historical_search = _call(
            core,
            "search-original-speaker-notes",
            "search.lexical",
            {
                "project_id": project_id,
                "query": "Synthetic presenter note for slide 1",
                "presentation_revision_id": original_revision["revision_id"],
            },
        )
        assert historical_search["ok"] is True, historical_search
        assert historical_search["result"]["results"]
    finally:
        core.close()


def test_object_transform_commits_safe_slide_geometry_revision(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        target_result = _call(
            core,
            "geometry-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert target_result["ok"] is True, target_result
        inventory = target_result["result"]
        assert inventory["slide_width_emu"] == 9_144_000
        assert inventory["slide_height_emu"] == 6_858_000
        target = next(item for item in inventory["objects"] if item["geometry_editable"])
        assert target["name"] == "Workspace transform test shape"
        assert target["delete_editable"] is True
        original_geometry = target["geometry"]
        replacement_geometry = {
            "x_emu": original_geometry["x_emu"] + 228_600,
            "y_emu": original_geometry["y_emu"] + 45_720,
            "width_emu": original_geometry["width_emu"] + 45_720,
            "height_emu": original_geometry["height_emu"],
        }
        edited = _call(
            core,
            "transform-object",
            "presentation.workspace.transform_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": target["object_id"],
                "expected_geometry": original_geometry,
                "replacement_geometry": replacement_geometry,
            },
        )
        assert edited["ok"] is True, edited
        assert edited["result"]["changed"] is True
        assert edited["result"]["revision_number"] == 2
        assert edited["result"]["parent_revision_id"] == opened["revision_id"]

        current_inventory = _call(
            core,
            "geometry-targets-after-edit",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": edited["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        assert current_inventory["ok"] is True, current_inventory
        current_target = next(
            item
            for item in current_inventory["result"]["objects"]
            if item["object_id"] == target["object_id"]
        )
        assert current_target["geometry"] == replacement_geometry

        unchanged = _call(
            core,
            "unchanged-object-geometry",
            "presentation.workspace.transform_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": target["object_id"],
                "expected_geometry": replacement_geometry,
                "replacement_geometry": replacement_geometry,
            },
        )
        assert unchanged["ok"] is True, unchanged
        assert unchanged["result"]["changed"] is False

        stale_geometry = _call(
            core,
            "stale-object-geometry",
            "presentation.workspace.transform_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": target["object_id"],
                "expected_geometry": original_geometry,
                "replacement_geometry": replacement_geometry,
            },
        )
        assert stale_geometry["ok"] is False
        assert stale_geometry["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"

        out_of_bounds = {
            **replacement_geometry,
            "x_emu": inventory["slide_width_emu"] - 100,
        }
        rejected_bounds = _call(
            core,
            "reject-out-of-bounds-transform",
            "presentation.workspace.transform_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": target["object_id"],
                "expected_geometry": replacement_geometry,
                "replacement_geometry": out_of_bounds,
            },
        )
        assert rejected_bounds["ok"] is False
        assert rejected_bounds["error"]["code"] == "INVALID_REQUEST"

        with core._storage.project_database(project_id) as connection:
            edited_row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (edited["result"]["revision_id"],),
            ).fetchone()
        assert edited_row is not None
        edited_path = core._storage.paths.snapshot_path(
            project_id, edited_row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(edited_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
        with original_path.open("rb") as handle:
            original_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
                byte_size=len(original_bytes),
            )
        edited_slide_part = next(
            slide.part_name
            for slide in original_fidelity.slides
            if slide.slide_id == inventory["slide_id"]
        )
        assert changed_parts == {edited_slide_part}
        assert original_path.read_bytes() == original_bytes
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_picture_crop_commits_slide_only_revision_and_keeps_media_relationships(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        inventory_response = _call(
            core,
            "picture-crop-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory_response["ok"] is True, inventory_response
        inventory = inventory_response["result"]
        picture = next(
            item
            for item in inventory["objects"]
            if item["name"] == "Workspace duplicate test picture"
        )
        assert picture["kind"] == "pic"
        assert picture["geometry_editable"] is True
        assert picture["picture_crop"] == {
            "left_percent": 0,
            "right_percent": 0,
            "top_percent": 0,
            "bottom_percent": 0,
        }

        invalid_crop = _call(
            core,
            "reject-empty-crop",
            "presentation.workspace.crop_picture",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": picture["object_id"],
                "expected_geometry": picture["geometry"],
                "expected_crop": picture["picture_crop"],
                "replacement_crop": {
                    "left_percent": 75,
                    "right_percent": 25,
                    "top_percent": 0,
                    "bottom_percent": 0,
                },
            },
        )
        assert invalid_crop["ok"] is False
        assert invalid_crop["error"]["code"] == "INVALID_REQUEST"

        for request_id, invalid_left_crop in (
            ("reject-nonfinite-crop", float("nan")),
            ("reject-oversized-crop-number", 10**1_000),
        ):
            invalid_number = _call(
                core,
                request_id,
                "presentation.workspace.crop_picture",
                {
                    "project_id": project_id,
                    "document_id": document["id"],
                    "expected_revision_id": opened["revision_id"],
                    "slide_id": inventory["slide_id"],
                    "object_id": picture["object_id"],
                    "expected_geometry": picture["geometry"],
                    "expected_crop": picture["picture_crop"],
                    "replacement_crop": {
                        "left_percent": invalid_left_crop,
                        "right_percent": 0,
                        "top_percent": 0,
                        "bottom_percent": 0,
                    },
                },
            )
            assert invalid_number["ok"] is False
            assert invalid_number["error"]["code"] == "INVALID_REQUEST"

        edited = _call(
            core,
            "crop-picture",
            "presentation.workspace.crop_picture",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": picture["object_id"],
                "expected_geometry": picture["geometry"],
                "expected_crop": picture["picture_crop"],
                "replacement_crop": {
                    "left_percent": 12.5,
                    "right_percent": 5.25,
                    "top_percent": 3.125,
                    "bottom_percent": 0,
                },
            },
        )
        assert edited["ok"] is True, edited
        assert edited["result"]["changed"] is True
        assert edited["result"]["revision_number"] == 2
        assert edited["result"]["parent_revision_id"] == opened["revision_id"]

        refreshed = _call(
            core,
            "picture-crop-inventory-after-edit",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": edited["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        assert refreshed["ok"] is True, refreshed
        updated_picture = next(
            item
            for item in refreshed["result"]["objects"]
            if item["object_id"] == picture["object_id"]
        )
        assert updated_picture["picture_crop"] == {
            "left_percent": 12.5,
            "right_percent": 5.25,
            "top_percent": 3.125,
            "bottom_percent": 0,
        }
        assert updated_picture["geometry"] == picture["geometry"]

        stale_crop = _call(
            core,
            "reject-stale-picture-crop",
            "presentation.workspace.crop_picture",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": picture["object_id"],
                "expected_geometry": picture["geometry"],
                "expected_crop": picture["picture_crop"],
                "replacement_crop": picture["picture_crop"],
            },
        )
        assert stale_crop["ok"] is False
        assert stale_crop["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (edited["result"]["revision_id"],),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
            assert before.read("ppt/slides/_rels/slide1.xml.rels") == after.read(
                "ppt/slides/_rels/slide1.xml.rels"
            )
            assert before.read("ppt/media/image1.png") == after.read("ppt/media/image1.png")
            slide_xml = etree.fromstring(after.read("ppt/slides/slide1.xml"))
        assert changed_parts == {"ppt/slides/slide1.xml"}
        src_rect = slide_xml.xpath(
            "//p:pic[p:nvPicPr/p:cNvPr[@name='Workspace duplicate test picture']]"
            "/p:blipFill/a:srcRect",
            namespaces={"p": PRESENTATION_NS, "a": DRAWING_NS},
        )
        assert len(src_rect) == 1
        assert src_rect[0].attrib == {"l": "12500", "r": "5250", "t": "3125"}
        assert original_path.read_bytes() == original_bytes

        reindexed = _call(
            core,
            "reindex-picture-crop",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


@pytest.mark.parametrize(
    ("shape_type", "expected_name"),
    [("rect", "Rectangle"), ("ellipse", "Ellipse")],
)
def test_shape_creation_adds_bounded_direct_color_shape(
    tmp_path: Path,
    shape_type: str,
    expected_name: str,
) -> None:
    from pptx import Presentation

    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "shape-creation-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory["ok"] is True, inventory
        slide_id = inventory["result"]["slide_id"]
        geometry = {
            "x_emu": 228_600,
            "y_emu": 182_880,
            "width_emu": 1_828_800,
            "height_emu": 1_371_600,
        }
        out_of_bounds = {
            **geometry,
            "x_emu": inventory["result"]["slide_width_emu"] - 10,
            "width_emu": 20,
        }
        rejected = _call(
            core,
            "reject-out-of-bounds-shape",
            "presentation.workspace.create_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
                "shape_type": shape_type,
                "geometry": out_of_bounds,
                "fill_color": "#D7E6FF",
                "line_color": "#204A87",
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "INVALID_REQUEST"

        result = _call(
            core,
            f"create-{shape_type}",
            "presentation.workspace.create_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
                "shape_type": shape_type,
                "geometry": geometry,
                "fill_color": "#D7E6FF",
                "line_color": "#204A87",
            },
        )
        assert result["ok"] is True, result
        assert result["result"]["changed"] is True
        assert result["result"]["parent_revision_id"] == opened["revision_id"]

        updated = _call(
            core,
            f"shape-{shape_type}-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": result["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        assert updated["ok"] is True, updated
        created_object_id = result["result"]["created_object_id"]
        shape = next(
            item for item in updated["result"]["objects"] if item["object_id"] == created_object_id
        )
        assert shape["name"] == expected_name
        assert shape["kind"] == "sp"
        assert shape["geometry"] == geometry
        assert shape["geometry_editable"] is True
        assert shape["delete_editable"] is True
        assert shape["shape_style"]["editable"] is True
        assert shape["shape_style"]["fill_color"] == "#D7E6FF"
        assert shape["shape_style"]["line_color"] == "#204A87"

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (result["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
            assert changed_parts == {"ppt/slides/slide1.xml"}
            for name in before_names - changed_parts:
                assert before.read(name) == after.read(name)
            shape_node = PresentationEditService._find_object_node(
                etree.fromstring(after.read("ppt/slides/slide1.xml")),
                ooxml_id=shape["ooxml_shape_id"],
                kind="sp",
            )
            assert shape_node.xpath(
                "./p:spPr/a:prstGeom/@prst",
                namespaces={"p": PRESENTATION_NS, "a": DRAWING_NS},
            ) == [shape_type]

        reopened = Presentation(revised_path)
        original = Presentation(original_path)
        assert len(reopened.slides) == len(original.slides)
        assert any(item.name == expected_name for item in reopened.slides[0].shapes)
        assert original_path.read_bytes() == original_bytes
    finally:
        core.close()


@pytest.mark.parametrize(
    ("image_format", "extension", "media_type", "updates_content_types"),
    [("PNG", "png", "image/png", False), ("JPEG", "jpg", "image/jpeg", True)],
)
def test_picture_creation_adds_private_media_at_requested_bounded_geometry(
    tmp_path: Path,
    image_format: str,
    extension: str,
    media_type: str,
    updates_content_types: bool,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    image_buffer = io.BytesIO()
    Image.new("RGB", (32, 24), color=(18, 52, 86)).save(image_buffer, format=image_format)
    image_bytes = image_buffer.getvalue()
    image_base64 = base64.b64encode(image_bytes).decode("ascii")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "create-picture-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory["ok"] is True, inventory
        slide_id = inventory["result"]["slide_id"]
        geometry = {
            "x_emu": 228_600,
            "y_emu": 182_880,
            "width_emu": 1_828_800,
            "height_emu": 1_371_600,
        }
        out_of_bounds = {
            **geometry,
            "x_emu": inventory["result"]["slide_width_emu"] - 10,
            "width_emu": 20,
        }
        rejected = _call(
            core,
            "reject-out-of-bounds-picture",
            "presentation.workspace.create_picture",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
                "geometry": out_of_bounds,
                "alt_text": "A new picture",
                "image_base64": image_base64,
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "INVALID_REQUEST"

        created = _call(
            core,
            "create-picture",
            "presentation.workspace.create_picture",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
                "geometry": geometry,
                "alt_text": "A new picture",
                "image_base64": image_base64,
            },
        )
        assert created["ok"] is True, created
        assert created["result"]["changed"] is True
        assert created["result"]["parent_revision_id"] == opened["revision_id"]
        created_object_id = created["result"]["created_object_id"]
        assert isinstance(created_object_id, str)

        updated = _call(
            core,
            "create-picture-updated-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": created["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        assert updated["ok"] is True, updated
        picture = next(
            item for item in updated["result"]["objects"] if item["object_id"] == created_object_id
        )
        assert picture["kind"] == "pic"
        assert picture["name"].startswith("Picture")
        assert picture["geometry"] == geometry
        assert picture["geometry_editable"] is True
        assert picture["delete_editable"] is True

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (created["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            added_parts = after_names - before_names
            assert len(added_parts) == 1
            media_part = next(iter(added_parts))
            assert media_part.startswith("ppt/media/presenter-copilot-")
            assert media_part.endswith(f".{extension}")
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
            expected_changed_parts = {
                "ppt/slides/slide1.xml",
                "ppt/slides/_rels/slide1.xml.rels",
            }
            if updates_content_types:
                expected_changed_parts.add("[Content_Types].xml")
            assert changed_parts == expected_changed_parts
            assert after.read(media_part) == image_bytes
            slide_root = etree.fromstring(after.read("ppt/slides/slide1.xml"))
            picture_node = PresentationEditService._find_object_node(
                slide_root,
                ooxml_id=picture["ooxml_shape_id"],
                kind="pic",
            )
            assert (
                picture_node.find("p:nvPicPr/p:cNvPr", namespaces={"p": PRESENTATION_NS}).get(
                    "descr"
                )
                == "A new picture"
            )
            image_relationship_id = picture_node.xpath(
                "./p:blipFill/a:blip/@r:embed",
                namespaces={
                    "p": PRESENTATION_NS,
                    "a": DRAWING_NS,
                    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                },
            )[0]
            relationships_root = etree.fromstring(after.read("ppt/slides/_rels/slide1.xml.rels"))
            image_relationship = next(
                item for item in relationships_root if item.get("Id") == image_relationship_id
            )
            assert image_relationship.get("Type").endswith("/image")
            assert (
                posixpath.normpath(
                    posixpath.join("ppt/slides", image_relationship.get("Target", ""))
                )
                == media_part
            )
            content_types = etree.fromstring(after.read("[Content_Types].xml"))
            image_defaults = [
                item
                for item in content_types
                if item.tag.endswith("}Default") and item.get("Extension") == extension
            ]
            assert len(image_defaults) == 1
            assert image_defaults[0].get("ContentType") == media_type

        assert original_path.read_bytes() == original_bytes
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


@pytest.mark.parametrize(
    ("image_format", "extension", "media_type", "updates_content_types"),
    [("PNG", "png", "image/png", False), ("JPEG", "jpg", "image/jpeg", True)],
)
def test_picture_image_replacement_preserves_shared_media_and_picture_geometry(
    tmp_path: Path,
    image_format: str,
    extension: str,
    media_type: str,
    updates_content_types: bool,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    replacement_buffer = io.BytesIO()
    Image.new("RGB", (32, 24), color=(18, 52, 86)).save(replacement_buffer, format=image_format)
    replacement_bytes = replacement_buffer.getvalue()
    replacement_base64 = base64.b64encode(replacement_bytes).decode("ascii")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "replace-picture-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory["ok"] is True, inventory
        source_picture = next(
            item
            for item in inventory["result"]["objects"]
            if item["name"] == "Workspace duplicate test picture"
        )
        duplicate = _call(
            core,
            "duplicate-picture-for-shared-media-check",
            "presentation.workspace.duplicate_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": source_picture["object_id"],
                "expected_geometry": source_picture["geometry"],
            },
        )
        assert duplicate["ok"] is True, duplicate
        current_revision_id = duplicate["result"]["revision_id"]
        duplicate_inventory = _call(
            core,
            "replace-picture-duplicate-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": current_revision_id,
                "slide_number": 1,
            },
        )
        assert duplicate_inventory["ok"] is True, duplicate_inventory
        objects = duplicate_inventory["result"]["objects"]
        original_picture = next(
            item for item in objects if item["object_id"] == source_picture["object_id"]
        )
        copied_picture = next(
            item
            for item in objects
            if item["object_id"] == duplicate["result"]["created_object_id"]
        )
        with core._storage.project_database(project_id) as connection:
            before_revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (current_revision_id,),
            ).fetchone()
        assert before_revision is not None
        before_path = core._storage.paths.snapshot_path(
            project_id, before_revision["local_snapshot_path"], require_exists=True
        )
        slide_part = "ppt/slides/slide1.xml"
        relationships_part = "ppt/slides/_rels/slide1.xml.rels"
        with zipfile.ZipFile(before_path) as before:
            slide_root = etree.fromstring(before.read(slide_part))
            relationships_root = etree.fromstring(before.read(relationships_part))
            original_node = PresentationEditService._find_object_node(
                slide_root,
                ooxml_id=original_picture["ooxml_shape_id"],
                kind="pic",
            )
            copied_node = PresentationEditService._find_object_node(
                slide_root,
                ooxml_id=copied_picture["ooxml_shape_id"],
                kind="pic",
            )
            original_embed = original_node.xpath(
                "./p:blipFill/a:blip/@r:embed",
                namespaces={
                    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
                    "a": DRAWING_NS,
                    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                },
            )[0]
            copied_embed = copied_node.xpath(
                "./p:blipFill/a:blip/@r:embed",
                namespaces={
                    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
                    "a": DRAWING_NS,
                    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                },
            )[0]
            assert copied_embed == original_embed
            relationship = next(
                item for item in relationships_root if item.get("Id") == original_embed
            )
            old_media_part = posixpath.normpath(
                posixpath.join("ppt/slides", relationship.get("Target", ""))
            )
            old_media_bytes = before.read(old_media_part)

        for request_id, invalid_image in (
            ("reject-picture-invalid-base64", "%"),
            ("reject-picture-invalid-image", base64.b64encode(b"not an image").decode("ascii")),
        ):
            rejected = _call(
                core,
                request_id,
                "presentation.workspace.replace_picture_image",
                {
                    "project_id": project_id,
                    "document_id": document["id"],
                    "expected_revision_id": current_revision_id,
                    "slide_id": duplicate_inventory["result"]["slide_id"],
                    "object_id": original_picture["object_id"],
                    "expected_geometry": original_picture["geometry"],
                    "expected_crop": original_picture["picture_crop"],
                    "image_base64": invalid_image,
                },
            )
            assert rejected["ok"] is False
            assert rejected["error"]["code"] == "INVALID_REQUEST"

        edited = _call(
            core,
            "replace-picture-image",
            "presentation.workspace.replace_picture_image",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": current_revision_id,
                "slide_id": duplicate_inventory["result"]["slide_id"],
                "object_id": original_picture["object_id"],
                "expected_geometry": original_picture["geometry"],
                "expected_crop": original_picture["picture_crop"],
                "image_base64": replacement_base64,
            },
        )
        assert edited["ok"] is True, edited
        assert edited["result"]["changed"] is True
        assert edited["result"]["revision_number"] == 3
        assert edited["result"]["parent_revision_id"] == current_revision_id

        with core._storage.project_database(project_id) as connection:
            edited_revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (edited["result"]["revision_id"],),
            ).fetchone()
        assert edited_revision is not None
        edited_path = core._storage.paths.snapshot_path(
            project_id, edited_revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(before_path) as before, zipfile.ZipFile(edited_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            added_parts = after_names - before_names
            assert len(added_parts) == 1
            replacement_part = next(iter(added_parts))
            assert replacement_part.startswith("ppt/media/presenter-copilot-")
            assert replacement_part.endswith(f".{extension}")
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
            expected_changed_parts = {
                slide_part,
                relationships_part,
            }
            if updates_content_types:
                expected_changed_parts.add("[Content_Types].xml")
            assert changed_parts == expected_changed_parts
            assert after.read(old_media_part) == old_media_bytes
            assert after.read(replacement_part) == replacement_bytes
            edited_slide = etree.fromstring(after.read(slide_part))
            edited_relationships = etree.fromstring(after.read(relationships_part))
            edited_content_types = etree.fromstring(after.read("[Content_Types].xml"))

        namespaces = {
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
            "a": DRAWING_NS,
            "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        }
        edited_original = PresentationEditService._find_object_node(
            edited_slide,
            ooxml_id=original_picture["ooxml_shape_id"],
            kind="pic",
        )
        edited_copy = PresentationEditService._find_object_node(
            edited_slide,
            ooxml_id=copied_picture["ooxml_shape_id"],
            kind="pic",
        )
        replacement_embed = edited_original.xpath(
            "./p:blipFill/a:blip/@r:embed", namespaces=namespaces
        )[0]
        copied_embed_after = edited_copy.xpath(
            "./p:blipFill/a:blip/@r:embed", namespaces=namespaces
        )[0]
        assert replacement_embed != original_embed
        assert copied_embed_after == original_embed
        replacement_relationship = next(
            item for item in edited_relationships if item.get("Id") == replacement_embed
        )
        assert replacement_relationship.get("Type").endswith("/image")
        assert replacement_relationship.get("TargetMode") is None
        assert (
            posixpath.normpath(
                posixpath.join("ppt/slides", replacement_relationship.get("Target", ""))
            )
            == replacement_part
        )
        image_defaults = [
            item
            for item in edited_content_types
            if item.tag.endswith("}Default") and item.get("Extension") == extension
        ]
        assert len(image_defaults) == 1
        assert image_defaults[0].get("ContentType") == media_type

        refreshed = _call(
            core,
            "replace-picture-after-image-edit",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": edited["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        updated_picture = next(
            item
            for item in refreshed["result"]["objects"]
            if item["object_id"] == original_picture["object_id"]
        )
        assert updated_picture["geometry"] == original_picture["geometry"]
        assert updated_picture["picture_crop"] == original_picture["picture_crop"]
        no_op = _call(
            core,
            "replace-picture-image-no-op",
            "presentation.workspace.replace_picture_image",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "slide_id": refreshed["result"]["slide_id"],
                "object_id": updated_picture["object_id"],
                "expected_geometry": updated_picture["geometry"],
                "expected_crop": updated_picture["picture_crop"],
                "image_base64": replacement_base64,
            },
        )
        assert no_op["ok"] is True, no_op
        assert no_op["result"]["changed"] is False
        assert no_op["result"]["revision_id"] == edited["result"]["revision_id"]
        assert original_path.read_bytes() == original_bytes
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 3
    finally:
        core.close()


def test_slide_reorder_creates_safe_revision_and_preserves_slide_ids(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        original_order = [slide["slide_id"] for slide in opened["fidelity_report"]["slides"]]
        reordered = [original_order[1], original_order[0], *original_order[2:]]

        invalid_order = _call(
            core,
            "reject-incomplete-slide-order",
            "presentation.workspace.reorder_slides",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "ordered_slide_ids": original_order[:-1],
            },
        )
        assert invalid_order["ok"] is False
        assert invalid_order["error"]["code"] == "INVALID_REQUEST"

        result = _call(
            core,
            "reorder-first-two-slides",
            "presentation.workspace.reorder_slides",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "ordered_slide_ids": reordered,
            },
        )
        assert result["ok"] is True, result
        assert result["result"]["changed"] is True
        assert result["result"]["revision_number"] == 2
        assert result["result"]["parent_revision_id"] == opened["revision_id"]

        current = _open(core, project_id, document["id"])
        current_order = [slide["slide_id"] for slide in current["fidelity_report"]["slides"]]
        assert current_order == reordered
        assert set(current_order) == set(original_order)

        unchanged = _call(
            core,
            "unchanged-slide-order",
            "presentation.workspace.reorder_slides",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": result["result"]["revision_id"],
                "ordered_slide_ids": reordered,
            },
        )
        assert unchanged["ok"] is True, unchanged
        assert unchanged["result"]["changed"] is False

        stale = _call(
            core,
            "stale-slide-order",
            "presentation.workspace.reorder_slides",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "ordered_slide_ids": original_order,
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (result["result"]["revision_id"],),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
        assert changed_parts == {"ppt/presentation.xml"}
        assert original_path.read_bytes() == original_bytes
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_safe_object_delete_commits_revision_and_preserves_other_package_parts(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        inventory_response = _call(
            core,
            "object-delete-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory_response["ok"] is True, inventory_response
        inventory = inventory_response["result"]
        target = next(
            item
            for item in inventory["objects"]
            if item["name"] == "Workspace transform test shape"
        )
        assert target["delete_editable"] is True
        title_placeholder = next(item for item in inventory["objects"] if item["name"] == "Title 1")
        assert title_placeholder["delete_editable"] is False
        rejected_placeholder = _call(
            core,
            "reject-placeholder-delete",
            "presentation.workspace.delete_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": title_placeholder["object_id"],
                "expected_geometry": {
                    "x_emu": 0,
                    "y_emu": 0,
                    "width_emu": 1,
                    "height_emu": 1,
                },
            },
        )
        assert rejected_placeholder["ok"] is False
        assert rejected_placeholder["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 1

        deleted = _call(
            core,
            "delete-safe-shape",
            "presentation.workspace.delete_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": target["object_id"],
                "expected_geometry": target["geometry"],
            },
        )
        assert deleted["ok"] is True, deleted
        assert deleted["result"]["changed"] is True
        assert deleted["result"]["revision_number"] == 2
        assert deleted["result"]["parent_revision_id"] == opened["revision_id"]

        current = _call(
            core,
            "objects-after-delete",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": deleted["result"]["revision_id"],
                "slide_number": 1,
            },
        )
        assert current["ok"] is True, current
        assert all(
            item["object_id"] != target["object_id"] for item in current["result"]["objects"]
        )
        reindexed = _call(
            core,
            "reindex-after-delete",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (deleted["result"]["revision_id"],),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
        with original_path.open("rb") as handle:
            original_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
                byte_size=len(original_bytes),
            )
        slide_part = next(
            slide.part_name
            for slide in original_fidelity.slides
            if slide.slide_id == inventory["slide_id"]
        )
        assert changed_parts == {slide_part}
        assert original_path.read_bytes() == original_bytes
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_object_delete_rejects_connector_and_animation_references() -> None:
    namespace_declarations = f'xmlns:p="{PRESENTATION_NS}" xmlns:a="{DRAWING_NS}"'
    slide_open = f"<p:sld {namespace_declarations}>"
    shape = (
        '<p:sp><p:nvSpPr><p:cNvPr id="2" name="Target"/>'
        "<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>"
        '<a:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/>'
        "</a:xfrm></p:spPr></p:sp>"
    )
    slides = (
        f"{slide_open}<p:cSld><p:spTree>{shape}"
        '</p:spTree></p:cSld><p:timing><p:tnLst><p:spTgt spid="2"/>'
        "</p:tnLst></p:timing></p:sld>",
        f"{slide_open}<p:cSld><p:spTree>{shape}"
        "<p:cxnSp><p:nvCxnSpPr><p:cNvCxnSpPr>"
        '<a:stCxn id="2" idx="0"/></p:cNvCxnSpPr></p:nvCxnSpPr>'
        "</p:cxnSp></p:spTree></p:cSld></p:sld>",
    )

    for slide_xml in slides:
        root = etree.fromstring(slide_xml.encode())
        candidate = root.xpath("./p:cSld/p:spTree/p:sp", namespaces={"p": PRESENTATION_NS})[0]
        assert (
            _delete_geometry_candidate(
                root,
                candidate,
                kind="sp",
                ooxml_id="2",
                geometry_editable=True,
            )
            is False
        )


def test_safe_object_duplicate_supports_shapes_and_pictures(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        initial_inventory = _call(
            core,
            "duplicate-object-initial-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert initial_inventory["ok"] is True, initial_inventory
        inventory = initial_inventory["result"]
        source_object_ids = [item["object_id"] for item in inventory["objects"]]
        expected_object_ids = list(source_object_ids)
        targets = [
            next(
                item
                for item in inventory["objects"]
                if item["name"] == name and item["kind"] == kind
            )
            for name, kind in (
                ("Workspace transform test shape", "sp"),
                ("Workspace duplicate test picture", "pic"),
            )
        ]
        assert all(item["geometry_editable"] and item["delete_editable"] for item in targets)
        assert all(item["geometry"] is not None for item in targets)

        placeholder = next(item for item in inventory["objects"] if item["name"] == "Title 1")
        rejected_placeholder = _call(
            core,
            "reject-placeholder-duplicate",
            "presentation.workspace.duplicate_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": placeholder["object_id"],
                "expected_geometry": {
                    "x_emu": 0,
                    "y_emu": 0,
                    "width_emu": 1,
                    "height_emu": 1,
                },
            },
        )
        assert rejected_placeholder["ok"] is False
        assert rejected_placeholder["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"

        current_revision_id = opened["revision_id"]
        created_ids: list[str] = []
        for target in targets:
            duplicate = _call(
                core,
                f"duplicate-{target['kind']}",
                "presentation.workspace.duplicate_object",
                {
                    "project_id": project_id,
                    "document_id": document["id"],
                    "expected_revision_id": current_revision_id,
                    "slide_id": inventory["slide_id"],
                    "object_id": target["object_id"],
                    "expected_geometry": target["geometry"],
                },
            )
            assert duplicate["ok"] is True, duplicate
            result = duplicate["result"]
            assert result["changed"] is True
            assert result["revision_number"] == len(created_ids) + 2
            assert result["parent_revision_id"] == current_revision_id
            created_object_id = result["created_object_id"]
            uuid.UUID(created_object_id)
            created_ids.append(created_object_id)

            current_revision_id = result["revision_id"]
            refreshed = _call(
                core,
                f"inventory-after-{target['kind']}-duplicate",
                "presentation.workspace.list_text_targets",
                {
                    "project_id": project_id,
                    "document_id": document["id"],
                    "revision_id": current_revision_id,
                    "slide_number": 1,
                },
            )
            assert refreshed["ok"] is True, refreshed
            current_objects = refreshed["result"]["objects"]
            original = next(
                item for item in current_objects if item["object_id"] == target["object_id"]
            )
            copied = next(
                item for item in current_objects if item["object_id"] == created_object_id
            )
            assert original["geometry"] == target["geometry"]
            assert copied["kind"] == target["kind"]
            assert copied["name"] == f"{target['name']} copy"
            assert copied["geometry_editable"] is True
            assert copied["delete_editable"] is True
            assert copied["geometry"]["width_emu"] == target["geometry"]["width_emu"]
            assert copied["geometry"]["height_emu"] == target["geometry"]["height_emu"]
            assert copied["geometry"] != target["geometry"]
            target_index = expected_object_ids.index(target["object_id"])
            expected_object_ids.insert(target_index + 1, created_object_id)
            assert [item["object_id"] for item in current_objects] == expected_object_ids

        stale = _call(
            core,
            "reject-stale-object-duplicate",
            "presentation.workspace.duplicate_object",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["slide_id"],
                "object_id": targets[0]["object_id"],
                "expected_geometry": targets[0]["geometry"],
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"
        reindexed = _call(
            core,
            "reindex-after-object-duplicate",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (current_revision_id,),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
        assert changed_parts == {"ppt/slides/slide1.xml"}
        assert original_path.read_bytes() == original_bytes
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 3
    finally:
        core.close()


def test_safe_slide_delete_removes_only_target_slide_parts(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        original_slide_ids = [slide["slide_id"] for slide in opened["fidelity_report"]["slides"]]
        deleted_slide_id = original_slide_ids[0]
        deleted = _call(
            core,
            "delete-slide",
            "presentation.workspace.delete_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": deleted_slide_id,
            },
        )
        assert deleted["ok"] is True, deleted
        result = deleted["result"]
        assert result["changed"] is True
        assert result["revision_number"] == 2
        assert result["parent_revision_id"] == opened["revision_id"]

        current = _open(core, project_id, document["id"])
        current_slide_ids = [slide["slide_id"] for slide in current["fidelity_report"]["slides"]]
        assert current_slide_ids == original_slide_ids[1:]
        assert len(current_slide_ids) == len(original_slide_ids) - 1

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions "
                "WHERE id = ? AND presentation_id = ?",
                (result["revision_id"], document["id"]),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            removed_parts = {
                "ppt/slides/slide1.xml",
                "ppt/slides/_rels/slide1.xml.rels",
                "ppt/notesSlides/notesSlide1.xml",
                "ppt/notesSlides/_rels/notesSlide1.xml.rels",
            }
            assert after_names == before_names - removed_parts
            changed_parts = {name for name in after_names if before.read(name) != after.read(name)}
            assert changed_parts == {
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
                "[Content_Types].xml",
            }
        assert original_path.read_bytes() == original_bytes

        stale = _call(
            core,
            "reject-stale-slide-delete",
            "presentation.workspace.delete_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": deleted_slide_id,
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"
        reindexed = _call(
            core,
            "reindex-after-slide-delete",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_slide_delete_fails_closed_for_custom_show_reference(tmp_path: Path) -> None:
    source_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    custom_show_path = tmp_path / "custom-show-reference.pptx"
    with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(custom_show_path, "w") as after:
        for info in before.infolist():
            payload = before.read(info.filename)
            if info.filename == "ppt/presentation.xml":
                root = etree.fromstring(payload)
                slide_reference = root.xpath(
                    "./p:sldIdLst/p:sldId[1]", namespaces={"p": PRESENTATION_NS}
                )[0]
                relation_id = slide_reference.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                )
                show_list = etree.SubElement(root, f"{{{PRESENTATION_NS}}}custShowLst")
                show = etree.SubElement(show_list, f"{{{PRESENTATION_NS}}}custShow")
                show.set("id", "1")
                show.set("name", "Safety test")
                slide_list = etree.SubElement(show, f"{{{PRESENTATION_NS}}}sldLst")
                show_slide = etree.SubElement(slide_list, f"{{{PRESENTATION_NS}}}sld")
                show_slide.set(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
                    str(relation_id),
                )
                payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            after.writestr(info, payload)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, custom_show_path)
        opened = _open(core, project_id, document["id"])
        target_slide_id = opened["fidelity_report"]["slides"][0]["slide_id"]
        rejected = _call(
            core,
            "reject-custom-show-slide-delete",
            "presentation.workspace.delete_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": target_slide_id,
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        history = _list_revisions(core, project_id, document["id"])
        assert history["current_revision_id"] == opened["revision_id"]
        assert len(history["revisions"]) == 1
    finally:
        core.close()


def test_slide_delete_keeps_at_least_one_slide(tmp_path: Path) -> None:
    from pptx import Presentation

    one_slide_path = tmp_path / "one-slide.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(one_slide_path)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, one_slide_path)
        opened = _open(core, project_id, document["id"])
        slide_id = opened["fidelity_report"]["slides"][0]["slide_id"]
        rejected = _call(
            core,
            "reject-last-slide-delete",
            "presentation.workspace.delete_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        history = _list_revisions(core, project_id, document["id"])
        assert history["current_revision_id"] == opened["revision_id"]
        assert len(history["revisions"]) == 1
    finally:
        core.close()


def test_safe_slide_duplicate_preserves_notes_images_and_original_parts(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        source_slide_id = opened["fidelity_report"]["slides"][0]["slide_id"]
        original_inventory_result = _call(
            core,
            "duplicate-slide-original-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert original_inventory_result["ok"] is True, original_inventory_result
        original_inventory = original_inventory_result["result"]
        duplicate = _call(
            core,
            "duplicate-slide",
            "presentation.workspace.duplicate_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": source_slide_id,
            },
        )
        assert duplicate["ok"] is True, duplicate
        result = duplicate["result"]
        assert result["changed"] is True
        assert result["revision_number"] == 2
        assert result["parent_revision_id"] == opened["revision_id"]
        uuid.UUID(result["created_slide_id"])

        revised_open = _open(core, project_id, document["id"])
        revised_slides = revised_open["fidelity_report"]["slides"]
        assert len(revised_slides) == 23
        assert revised_slides[0]["slide_id"] == source_slide_id
        assert revised_slides[1]["slide_id"] == result["created_slide_id"]
        assert revised_slides[1]["ordinal"] == 2

        duplicate_inventory_result = _call(
            core,
            "duplicate-slide-copy-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": result["revision_id"],
                "slide_number": 2,
            },
        )
        assert duplicate_inventory_result["ok"] is True, duplicate_inventory_result
        duplicate_inventory = duplicate_inventory_result["result"]
        assert duplicate_inventory["slide_id"] == result["created_slide_id"]
        assert duplicate_inventory["speaker_notes"] == original_inventory["speaker_notes"]
        assert len(duplicate_inventory["objects"]) == len(original_inventory["objects"])
        assert [item["name"] for item in duplicate_inventory["objects"]] == [
            item["name"] for item in original_inventory["objects"]
        ]
        assert all(
            copied["object_id"] != source["object_id"]
            and copied["geometry"] == source["geometry"]
            and copied["kind"] == source["kind"]
            for source, copied in zip(
                original_inventory["objects"], duplicate_inventory["objects"], strict=True
            )
        )

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions "
                "WHERE id = ? AND presentation_id = ?",
                (result["revision_id"], document["id"]),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with revised_path.open("rb") as revised_handle:
            revised_fidelity = inspect_pptx_snapshot(
                revised_handle,
                document_id=document["id"],
                expected_sha256=result["sha256"],
                byte_size=result["byte_size"],
            )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names < after_names
            changed_parts = {name for name in before_names if before.read(name) != after.read(name)}
            assert changed_parts == {
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
                "[Content_Types].xml",
            }
            new_slide_part = next(
                slide.part_name
                for slide in revised_fidelity.slides
                if slide.slide_id == result["created_slide_id"]
            )
            new_slide_rels_part = f"ppt/slides/_rels/{Path(new_slide_part).name}.rels"
            new_slide_rels = etree.fromstring(after.read(new_slide_rels_part))
            package_namespaces = {
                "rel": "http://schemas.openxmlformats.org/package/2006/relationships"
            }
            notes_relation = new_slide_rels.xpath(
                "./rel:Relationship[contains(@Type, '/notesSlide')]",
                namespaces=package_namespaces,
            )
            image_relations = new_slide_rels.xpath(
                "./rel:Relationship[contains(@Type, '/image')]",
                namespaces=package_namespaces,
            )
            assert len(notes_relation) == 1
            assert len(image_relations) == 1
            notes_target = notes_relation[0].get("Target", "")
            assert notes_target.startswith("../notesSlides/")
            notes_part = f"ppt/notesSlides/{Path(notes_target).name}"
            new_notes_rels_part = f"ppt/notesSlides/_rels/{Path(notes_part).name}.rels"
            notes_rels = etree.fromstring(after.read(new_notes_rels_part))
            notes_slide_relation = notes_rels.xpath(
                "./rel:Relationship[contains(@Type, '/slide')]",
                namespaces=package_namespaces,
            )
            assert len(notes_slide_relation) == 1
            assert notes_slide_relation[0].get("Target") == f"../slides/{Path(new_slide_part).name}"
            original_notes_part = "ppt/notesSlides/notesSlide1.xml"
            assert after.read(notes_part) == before.read(original_notes_part)
            assert after.read("ppt/media/image1.png") == before.read("ppt/media/image1.png")
        assert original_path.read_bytes() == original_bytes

        stale = _call(
            core,
            "reject-stale-slide-duplicate",
            "presentation.workspace.duplicate_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": source_slide_id,
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"
        reindexed = _call(
            core,
            "reindex-after-slide-duplicate",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_safe_blank_slide_creation_uses_layout_without_copying_slide_content(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        source_slide_id = opened["fidelity_report"]["slides"][0]["slide_id"]
        created = _call(
            core,
            "create-blank-slide",
            "presentation.workspace.create_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": source_slide_id,
            },
        )
        assert created["ok"] is True, created
        result = created["result"]
        assert result["changed"] is True
        assert result["revision_number"] == 2
        assert result["parent_revision_id"] == opened["revision_id"]
        uuid.UUID(result["created_slide_id"])

        revised_open = _open(core, project_id, document["id"])
        revised_slides = revised_open["fidelity_report"]["slides"]
        assert len(revised_slides) == 23
        assert revised_slides[0]["slide_id"] == source_slide_id
        assert revised_slides[1]["slide_id"] == result["created_slide_id"]
        assert revised_slides[1]["ordinal"] == 2

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions "
                "WHERE id = ? AND presentation_id = ?",
                (result["revision_id"], document["id"]),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with revised_path.open("rb") as revised_handle:
            revised_fidelity = inspect_pptx_snapshot(
                revised_handle,
                document_id=document["id"],
                expected_sha256=result["sha256"],
                byte_size=result["byte_size"],
            )
        blank_slide = next(
            slide
            for slide in revised_fidelity.slides
            if slide.slide_id == result["created_slide_id"]
        )
        assert blank_slide.objects == ()
        assert blank_slide.speaker_notes is None
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            new_parts = after_names - before_names
            assert len(new_parts) == 2
            assert {name for name in before_names if before.read(name) != after.read(name)} == {
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
                "[Content_Types].xml",
            }
            blank_slide_rels_part = f"ppt/slides/_rels/{Path(blank_slide.part_name).name}.rels"
            blank_root = etree.fromstring(after.read(blank_slide.part_name))
            blank_shape_tree = blank_root.xpath(
                "./p:cSld/p:spTree",
                namespaces={"p": PRESENTATION_NS},
            )
            assert len(blank_shape_tree) == 1
            assert [child.tag.rsplit("}", 1)[-1] for child in blank_shape_tree[0]] == [
                "nvGrpSpPr",
                "grpSpPr",
            ]
            package_namespaces = {
                "rel": "http://schemas.openxmlformats.org/package/2006/relationships"
            }
            source_slide_part = revised_fidelity.slides[0].part_name
            source_slide_rels_part = f"ppt/slides/_rels/{Path(source_slide_part).name}.rels"
            source_slide_rels = etree.fromstring(after.read(source_slide_rels_part))
            blank_slide_rels = etree.fromstring(after.read(blank_slide_rels_part))
            source_layout = source_slide_rels.xpath(
                "./rel:Relationship[contains(@Type, '/slideLayout')]",
                namespaces=package_namespaces,
            )
            blank_relationships = blank_slide_rels.xpath(
                "./rel:Relationship",
                namespaces=package_namespaces,
            )
            assert len(source_layout) == 1
            assert len(blank_relationships) == 1
            assert blank_relationships[0].get("Type") == source_layout[0].get("Type")
            assert posixpath.normpath(
                posixpath.join(
                    posixpath.dirname(blank_slide.part_name),
                    blank_relationships[0].get("Target", ""),
                )
            ) == posixpath.normpath(
                posixpath.join(
                    posixpath.dirname(source_slide_part),
                    source_layout[0].get("Target", ""),
                )
            )
        assert original_path.read_bytes() == original_bytes

        stale = _call(
            core,
            "reject-stale-blank-slide",
            "presentation.workspace.create_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": source_slide_id,
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"
        reindexed = _call(
            core,
            "reindex-after-blank-slide",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_blank_slide_creation_can_reuse_another_slide_layout(tmp_path: Path) -> None:
    from pptx import Presentation

    source_path = tmp_path / "alternate-layouts.pptx"
    presentation = Presentation()
    first_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[2])
    assert first_slide.shapes.title is not None
    assert second_slide.shapes.title is not None
    first_slide.shapes.title.text = "Active slide content stays put"
    second_slide.shapes.title.text = "Alternate layout source content stays put"
    presentation.save(source_path)
    original_bytes = source_path.read_bytes()

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, source_path)
        opened = _open(core, project_id, document["id"])
        source_slides = opened["fidelity_report"]["slides"]
        active_slide_id = source_slides[0]["slide_id"]
        layout_source_slide_id = source_slides[1]["slide_id"]
        created = _call(
            core,
            "create-blank-slide-with-other-layout",
            "presentation.workspace.create_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": active_slide_id,
                "layout_source_slide_id": layout_source_slide_id,
            },
        )
        assert created["ok"] is True, created
        result = created["result"]
        assert result["layout_source_slide_id"] == layout_source_slide_id

        revised_open = _open(core, project_id, document["id"])
        revised_slides = revised_open["fidelity_report"]["slides"]
        assert len(revised_slides) == 3
        assert revised_slides[0]["slide_id"] == active_slide_id
        assert revised_slides[1]["slide_id"] == result["created_slide_id"]
        assert revised_slides[2]["slide_id"] == layout_source_slide_id

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions "
                "WHERE id = ? AND presentation_id = ?",
                (result["revision_id"], document["id"]),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with revised_path.open("rb") as revised_handle:
            revised_fidelity = inspect_pptx_snapshot(
                revised_handle,
                document_id=document["id"],
                expected_sha256=result["sha256"],
                byte_size=result["byte_size"],
            )
        blank_slide = revised_fidelity.slides[1]
        assert blank_slide.objects == ()
        assert blank_slide.speaker_notes is None
        assert any(
            "Active slide content stays put" in obj.text
            for obj in revised_fidelity.slides[0].objects
        )
        assert any(
            "Alternate layout source content stays put" in obj.text
            for obj in revised_fidelity.slides[2].objects
        )

        package_namespaces = {"rel": "http://schemas.openxmlformats.org/package/2006/relationships"}

        def resolved_layout_part(archive: zipfile.ZipFile, part_name: str) -> str:
            relationships_part = f"{posixpath.dirname(part_name)}/_rels/{Path(part_name).name}.rels"
            relationships = etree.fromstring(archive.read(relationships_part)).xpath(
                "./rel:Relationship[contains(@Type, '/slideLayout')]",
                namespaces=package_namespaces,
            )
            assert len(relationships) == 1
            return posixpath.normpath(
                posixpath.join(posixpath.dirname(part_name), relationships[0].get("Target", ""))
            )

        with zipfile.ZipFile(revised_path) as after:
            active_layout_part = resolved_layout_part(after, revised_fidelity.slides[0].part_name)
            blank_layout_part = resolved_layout_part(after, blank_slide.part_name)
            layout_source_part = resolved_layout_part(after, revised_fidelity.slides[2].part_name)
            assert blank_layout_part == layout_source_part
            assert blank_layout_part != active_layout_part
        assert source_path.read_bytes() == original_bytes

        invalid_layout = _call(
            core,
            "reject-layout-from-another-presentation",
            "presentation.workspace.create_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": revised_open["revision_id"],
                "slide_id": revised_slides[0]["slide_id"],
                "layout_source_slide_id": str(uuid.uuid4()),
            },
        )
        assert invalid_layout["ok"] is False
        assert invalid_layout["error"]["code"] == "PRESENTATION_EDIT_LAYOUT_NOT_FOUND"
        assert (
            _list_revisions(core, project_id, document["id"])["current_revision_id"]
            == revised_open["revision_id"]
        )
    finally:
        core.close()


def test_applying_existing_slide_layout_preserves_objects_and_is_revisioned(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    source_path = tmp_path / "apply-existing-layout.pptx"
    presentation = Presentation()
    first_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    second_slide = presentation.slides.add_slide(presentation.slide_layouts[2])
    assert first_slide.shapes.title is not None
    assert second_slide.shapes.title is not None
    first_slide.shapes.title.text = "Keep target slide content"
    second_slide.shapes.title.text = "Use this slide's layout"
    presentation.save(source_path)
    original_bytes = source_path.read_bytes()

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, source_path)
        opened = _open(core, project_id, document["id"])
        original_slides = opened["fidelity_report"]["slides"]
        target_slide_id = original_slides[0]["slide_id"]
        layout_source_slide_id = original_slides[1]["slide_id"]

        applied = _call(
            core,
            "apply-layout-from-another-slide",
            "presentation.workspace.apply_slide_layout",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": target_slide_id,
                "layout_source_slide_id": layout_source_slide_id,
            },
        )
        assert applied["ok"] is True, applied
        result = applied["result"]
        assert result["changed"] is True
        assert result["revision_number"] == 2
        assert result["parent_revision_id"] == opened["revision_id"]
        assert result["layout_source_slide_id"] == layout_source_slide_id

        revised_open = _open(core, project_id, document["id"])
        revised_slides = revised_open["fidelity_report"]["slides"]
        assert [slide["slide_id"] for slide in revised_slides] == [
            slide["slide_id"] for slide in original_slides
        ]

        with core._storage.project_database(project_id) as connection:
            revision = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions "
                "WHERE id = ? AND presentation_id = ?",
                (result["revision_id"], document["id"]),
            ).fetchone()
        assert revision is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, revision["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(revised_path) as after:
            assert set(before.namelist()) == set(after.namelist())
            changed_parts = {
                name for name in before.namelist() if before.read(name) != after.read(name)
            }
            target_relationships_part = "ppt/slides/_rels/slide1.xml.rels"
            assert changed_parts == {target_relationships_part}

            package_namespaces = {
                "rel": "http://schemas.openxmlformats.org/package/2006/relationships"
            }

            def layout_part(archive: zipfile.ZipFile, part_name: str) -> str:
                rels_part = f"{posixpath.dirname(part_name)}/_rels/{Path(part_name).name}.rels"
                relationships = etree.fromstring(archive.read(rels_part)).xpath(
                    "./rel:Relationship[contains(@Type, '/slideLayout')]",
                    namespaces=package_namespaces,
                )
                assert len(relationships) == 1
                return posixpath.normpath(
                    posixpath.join(posixpath.dirname(part_name), relationships[0].get("Target", ""))
                )

            target_part = "ppt/slides/slide1.xml"
            source_part = "ppt/slides/slide2.xml"
            assert layout_part(after, target_part) == layout_part(after, source_part)
            assert layout_part(before, target_part) != layout_part(before, source_part)
        assert source_path.read_bytes() == original_bytes

        unchanged = _call(
            core,
            "apply-same-layout-again",
            "presentation.workspace.apply_slide_layout",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": revised_open["revision_id"],
                "slide_id": target_slide_id,
                "layout_source_slide_id": layout_source_slide_id,
            },
        )
        assert unchanged["ok"] is True, unchanged
        assert unchanged["result"]["changed"] is False
        assert unchanged["result"]["revision_id"] == revised_open["revision_id"]

        rejected = _call(
            core,
            "reject-layout-from-another-presentation",
            "presentation.workspace.apply_slide_layout",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": revised_open["revision_id"],
                "slide_id": target_slide_id,
                "layout_source_slide_id": str(uuid.uuid4()),
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_LAYOUT_NOT_FOUND"
        assert (
            _list_revisions(core, project_id, document["id"])["current_revision_id"]
            == revised_open["revision_id"]
        )
    finally:
        core.close()


def test_slide_duplicate_fails_closed_on_unhandled_relationship(tmp_path: Path) -> None:
    source_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    unsafe_path = tmp_path / "unsupported-slide-relation.pptx"
    with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(unsafe_path, "w") as after:
        for info in before.infolist():
            payload = before.read(info.filename)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                root = etree.fromstring(payload)
                relationship = etree.SubElement(
                    root,
                    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                )
                relationship.set("Id", "rIdUnsupported")
                relationship.set(
                    "Type",
                    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
                )
                relationship.set("Target", "../embeddings/oleObject1.bin")
                payload = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            after.writestr(info, payload)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, unsafe_path)
        opened = _open(core, project_id, document["id"])
        target_slide_id = opened["fidelity_report"]["slides"][0]["slide_id"]
        rejected = _call(
            core,
            "reject-slide-duplicate-unsupported-relation",
            "presentation.workspace.duplicate_slide",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": target_slide_id,
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        history = _list_revisions(core, project_id, document["id"])
        assert history["current_revision_id"] == opened["revision_id"]
        assert len(history["revisions"]) == 1
    finally:
        core.close()


def test_rich_speaker_notes_are_readable_but_rejected_for_editing(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    rich_path = tmp_path / "rich-notes.pptx"
    run = b"<a:r><a:t>Synthetic presenter note for slide 1.</a:t></a:r>"
    split_runs = (
        b"<a:r><a:t>Synthetic presenter note</a:t></a:r><a:r><a:t> for slide 1.</a:t></a:r>"
    )
    with zipfile.ZipFile(original_path) as source, zipfile.ZipFile(rich_path, "w") as output:
        for info in source.infolist():
            content = source.read(info)
            if info.filename == "ppt/notesSlides/notesSlide1.xml":
                assert run in content
                content = content.replace(run, split_runs, 1)
            output.writestr(info, content)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, rich_path)
        opened = _open(core, project_id, document["id"])
        target_result = _call(
            core,
            "rich-speaker-notes-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert target_result["ok"] is True, target_result
        notes = target_result["result"]["speaker_notes"]
        assert notes["current_text"] == "Synthetic presenter note for slide 1."
        assert notes["text_truncated"] is False
        assert notes["editable"] is False

        rejected = _call(
            core,
            "reject-rich-speaker-notes-edit",
            "presentation.workspace.replace_speaker_notes",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": target_result["result"]["slide_id"],
                "expected_text": notes["current_text"],
                "replacement_text": "Do not flatten these formatted notes.",
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        history = _list_revisions(core, project_id, document["id"])
        assert history["current_revision_id"] == opened["revision_id"]
        assert len(history["revisions"]) == 1
    finally:
        core.close()


def test_text_box_creation_is_revisioned_and_changes_only_selected_slide(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        targets = _call(
            core,
            "text-box-creation-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert targets["ok"] is True, targets
        slide_id = targets["result"]["slide_id"]
        text = f"Created workspace text {uuid.uuid4().hex}"
        created = _call(
            core,
            "create-text-box",
            "presentation.workspace.create_text_box",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
                "text": text,
            },
        )
        assert created["ok"] is True, created
        assert created["result"]["changed"] is True
        created_object_id = created["result"]["created_object_id"]
        assert isinstance(created_object_id, str)
        assert created["result"]["parent_revision_id"] == opened["revision_id"]

        latest = _open(core, project_id, document["id"])
        updated_targets = _call(
            core,
            "text-box-creation-inventory",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": latest["revision_id"],
                "slide_number": 1,
            },
        )
        assert updated_targets["ok"] is True, updated_targets
        new_target = next(
            item
            for item in updated_targets["result"]["targets"]
            if item["object_id"] == created_object_id
        )
        assert new_target["current_text"] == text
        created_object = next(
            item
            for item in updated_targets["result"]["objects"]
            if item["object_id"] == created_object_id
        )
        assert created_object["name"].startswith("Text Box")
        assert created_object["geometry_editable"] is True
        assert created_object["delete_editable"] is True
        assert created_object["geometry"]["x_emu"] >= 0
        assert created_object["geometry"]["y_emu"] >= 0

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (latest["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        assert hashlib.sha256(revised_path.read_bytes()).hexdigest() == latest["sha256"]
        with original_path.open("rb") as original_handle:
            original_fidelity = inspect_pptx_snapshot(
                original_handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
                byte_size=len(original_bytes),
            )
        with revised_path.open("rb") as revised_handle:
            revised_fidelity = inspect_pptx_snapshot(
                revised_handle,
                document_id=document["id"],
                expected_sha256=latest["sha256"],
                byte_size=latest["byte_size"],
            )
        original_slide = next(
            slide for slide in original_fidelity.slides if slide.slide_id == slide_id
        )
        revised_slide = next(
            slide for slide in revised_fidelity.slides if slide.slide_id == slide_id
        )
        assert revised_slide.objects[:-1] == original_slide.objects
        assert revised_slide.objects[-1].object_id == created_object_id
        assert [slide.slide_id for slide in revised_fidelity.slides] == [
            slide.slide_id for slide in original_fidelity.slides
        ]
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            original_names = {info.filename for info in before.infolist()}
            revised_names = {info.filename for info in after.infolist()}
            assert original_names == revised_names
            for part_name in original_names - {original_slide.part_name}:
                assert before.read(part_name) == after.read(part_name)
        assert original_path.read_bytes() == original_bytes

        reindexed = _call(
            core,
            "reindex-created-text-box",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        search = _call(
            core,
            "search-created-text-box",
            "search.lexical",
            {"project_id": project_id, "query": text},
        )
        assert search["ok"] is True, search
        assert search["result"]["results"]

        stale = _call(
            core,
            "stale-create-text-box",
            "presentation.workspace.create_text_box",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": slide_id,
                "text": "stale",
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"
    finally:
        core.close()


def test_single_run_format_creates_revision_and_preserves_other_package_parts(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        targets_result = _call(
            core,
            "format-text-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert targets_result["ok"] is True, targets_result
        target = next(
            item for item in targets_result["result"]["targets"] if item["single_run_editable"]
        )
        expected_text = target["current_text"]
        style = target["text_style"]
        assert isinstance(expected_text, str) and expected_text
        assert style["editable"] is True
        assert isinstance(style["fingerprint"], str)
        replacement_style = {
            "bold": True,
            "italic": False,
            "font_size_pt": 20,
            "font_color": "#123456",
        }
        changed = _call(
            core,
            "format-single-run-text",
            "presentation.workspace.format_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": targets_result["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "expected_style_fingerprint": style["fingerprint"],
                "replacement_style": replacement_style,
            },
        )
        assert changed["ok"] is True, changed
        assert changed["result"]["changed"] is True
        assert changed["result"]["parent_revision_id"] == opened["revision_id"]

        history = _list_revisions(core, project_id, document["id"])
        assert history["current_revision_id"] == changed["result"]["revision_id"]
        assert history["revisions"][0]["revision_number"] == 2

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (changed["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        assert hashlib.sha256(revised_path.read_bytes()).hexdigest() == changed["result"]["sha256"]
        with revised_path.open("rb") as handle:
            revised_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=changed["result"]["sha256"],
                byte_size=changed["result"]["byte_size"],
            )
        revised_object = next(
            obj
            for slide in revised_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert revised_object.text == expected_text
        assert revised_object.text_style is not None
        assert revised_object.text_style.bold is True
        assert revised_object.text_style.italic is False
        assert revised_object.text_style.font_size_pt == 20
        assert revised_object.text_style.font_color == "#123456"

        with original_path.open("rb") as handle:
            original_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
                byte_size=len(original_bytes),
            )
        slide_part = next(
            slide.part_name
            for slide in original_fidelity.slides
            if slide.slide_id == targets_result["result"]["slide_id"]
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            for part_name in before_names - {slide_part}:
                assert before.read(part_name) == after.read(part_name)
        assert original_path.read_bytes() == original_bytes

        latest = _open(core, project_id, document["id"])
        latest_targets = _call(
            core,
            "format-text-targets-after-edit",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": latest["revision_id"],
                "slide_number": 1,
            },
        )
        updated_target = next(
            item
            for item in latest_targets["result"]["targets"]
            if item["object_id"] == target["object_id"]
        )
        no_op = _call(
            core,
            "format-single-run-text-no-op",
            "presentation.workspace.format_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": latest["revision_id"],
                "slide_id": latest_targets["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "expected_style_fingerprint": updated_target["text_style"]["fingerprint"],
                "replacement_style": replacement_style,
            },
        )
        assert no_op["ok"] is True, no_op
        assert no_op["result"]["changed"] is False
        assert no_op["result"]["revision_id"] == latest["revision_id"]

        stale = _call(
            core,
            "format-single-run-text-stale-style",
            "presentation.workspace.format_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": latest["revision_id"],
                "slide_id": latest_targets["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "expected_style_fingerprint": "0" * 64,
                "replacement_style": {"bold": False},
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"

        invalid = _call(
            core,
            "format-single-run-text-invalid-size",
            "presentation.workspace.format_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": latest["revision_id"],
                "slide_id": latest_targets["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "expected_style_fingerprint": updated_target["text_style"]["fingerprint"],
                "replacement_style": {"font_size_pt": 100},
            },
        )
        assert invalid["ok"] is False
        assert invalid["error"]["code"] == "INVALID_REQUEST"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_multirun_text_format_applies_to_each_run_and_preserves_unrelated_style(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "multirun.pptx"
    slide_number, target_shape_id = _write_multirun_text_fixture(
        FIXTURE_ROOT / "deck" / "presentation.pptx",
        source_path,
    )
    original_bytes = source_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, source_path)
        opened = _open(core, project_id, document["id"])
        inspected = inspect_pptx_snapshot(
            io.BytesIO(original_bytes),
            document_id=document["id"],
            expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
            byte_size=len(original_bytes),
        )
        assert any(
            obj.ooxml_id == target_shape_id
            and obj.text_format_edit_candidate
            and not obj.single_run_text_edit_candidate
            and not obj.uniform_run_text_edit_candidate
            for slide in inspected.slides
            if slide.ordinal == slide_number
            for obj in slide.objects
        )
        inventory = _call(
            core,
            "multirun-style-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": slide_number,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["targets"]
            if item["ooxml_shape_id"] == target_shape_id
            and item["text_format_editable"]
            and not item["single_run_editable"]
            and not item["uniform_run_editable"]
        )
        assert target["current_text"]
        assert target["text_style"]["bold"] is None
        assert target["text_style"]["italic"] is None
        assert target["text_style"]["alignment"] is None
        assert target["text_style"]["alignment_editable"] is True
        assert target["text_style"]["bullet_state"] == "bullet"
        assert target["text_style"]["bullet_editable"] is True
        assert target["text_style"]["left_indent_pt"] == 18
        assert target["text_style"]["left_indent_editable"] is True
        assert target["text_style"]["first_line_indent_pt"] == -9
        assert target["text_style"]["first_line_indent_editable"] is True
        assert target["text_style"]["line_spacing_pt"] is None
        assert target["text_style"]["line_spacing_editable"] is True

        legacy = _call(
            core,
            "multirun-single-run-alias-rejected",
            "presentation.workspace.format_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"bold": True},
            },
        )
        assert legacy["ok"] is False
        assert legacy["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"

        mixed_replacement = _call(
            core,
            "multirun-replacement-rejected",
            "presentation.workspace.replace_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "operation_type": "replace_text_shape",
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "replacement_text": "unsafe style collapse",
            },
        )
        assert mixed_replacement["ok"] is False
        assert mixed_replacement["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 1

        changed = _call(
            core,
            "multirun-format-text-shape",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {
                    "bold": True,
                    "font_color": "#A1B2C3",
                    "alignment": "center",
                    "bullet_mode": "none",
                    "left_indent_pt": 36,
                    "line_spacing_pt": 24,
                    "space_before_pt": 12,
                    "space_after_pt": 6,
                },
            },
        )
        assert changed["ok"] is True, changed
        assert changed["result"]["changed"] is True

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (changed["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        revised_bytes = revised_path.read_bytes()
        with revised_path.open("rb") as handle:
            revised_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(revised_bytes).hexdigest(),
                byte_size=len(revised_bytes),
            )
        revised_object = next(
            obj
            for slide in revised_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert revised_object.text == target["current_text"]
        assert revised_object.text_format_edit_candidate is True
        assert revised_object.text_style is not None
        assert revised_object.text_style.bold is True
        assert revised_object.text_style.italic is None
        assert revised_object.text_style.font_color == "#A1B2C3"
        assert revised_object.text_style.alignment == "center"
        assert revised_object.text_style.bullet_state == "none"
        assert revised_object.text_style.bullet_editable is True
        assert revised_object.text_style.left_indent_pt == 36
        assert revised_object.text_style.left_indent_editable is True
        assert revised_object.text_style.line_spacing_pt == 24
        assert revised_object.text_style.line_spacing_editable is True
        assert revised_object.text_style.space_before_pt == 12
        assert revised_object.text_style.space_before_editable is True
        assert revised_object.text_style.space_after_pt == 6
        assert revised_object.text_style.space_after_editable is True

        revised_slide = next(
            slide for slide in revised_fidelity.slides if slide.ordinal == slide_number
        )
        namespaces = {"p": PRESENTATION_NS, "a": DRAWING_NS}
        with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(revised_path) as after:
            assert source_path.read_bytes() == original_bytes
            assert {info.filename for info in before.infolist()} == {
                info.filename for info in after.infolist()
            }
            for part_name in {info.filename for info in before.infolist()} - {
                revised_slide.part_name
            }:
                assert before.read(part_name) == after.read(part_name)
            revised_root = etree.fromstring(after.read(revised_slide.part_name))
        revised_shapes = revised_root.xpath(
            ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
            namespaces=namespaces,
            shape_id=target["ooxml_shape_id"],
        )
        revised_runs = revised_shapes[0].xpath("./p:txBody/a:p/a:r", namespaces=namespaces)
        paragraph_properties = revised_shapes[0].xpath(
            "./p:txBody/a:p/a:pPr", namespaces=namespaces
        )[0]
        assert paragraph_properties.get("algn") == "ctr"
        assert paragraph_properties.get("marL") == "457200"
        assert paragraph_properties.get("indent") == "-114300"
        line_spacing = paragraph_properties.find(f"{{{DRAWING_NS}}}lnSpc")
        spacing_before = paragraph_properties.find(f"{{{DRAWING_NS}}}spcBef")
        spacing_after = paragraph_properties.find(f"{{{DRAWING_NS}}}spcAft")
        assert line_spacing is not None and spacing_before is not None
        assert spacing_after is not None
        assert paragraph_properties.index(line_spacing) < paragraph_properties.index(spacing_before)
        assert paragraph_properties.index(spacing_before) < paragraph_properties.index(
            spacing_after
        )
        assert line_spacing.xpath("./a:spcPts/@val", namespaces=namespaces) == ["2400"]
        assert paragraph_properties.xpath("./a:spcBef/a:spcPts/@val", namespaces=namespaces) == [
            "1200"
        ]
        assert paragraph_properties.xpath("./a:spcAft/a:spcPts/@val", namespaces=namespaces) == [
            "600"
        ]
        assert revised_shapes[0].xpath("./p:txBody/a:p/a:pPr/a:buNone", namespaces=namespaces)
        assert revised_shapes[0].xpath("./p:txBody/a:p/a:pPr/a:buChar", namespaces=namespaces) == []
        assert len(revised_runs) == 2
        run_styles = [run.find(f"{{{DRAWING_NS}}}rPr") for run in revised_runs]
        assert all(style is not None and style.get("b") == "1" for style in run_styles)
        assert [style.get("i") for style in run_styles if style is not None] == ["1", "0"]
        assert run_styles[1] is not None and run_styles[1].get("u") == "sng"

        updated_spacing = _call(
            core,
            "multirun-update-spacing",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": changed["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": revised_object.text_style.fingerprint,
                "replacement_style": {
                    "bullet_mode": "bullet",
                    "line_spacing_pt": 30,
                    "space_before_pt": 18,
                    "space_after_pt": 9,
                    "first_line_indent_pt": 18,
                },
            },
        )
        assert updated_spacing["ok"] is True, updated_spacing
        with core._storage.project_database(project_id) as connection:
            updated_row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (updated_spacing["result"]["revision_id"],),
            ).fetchone()
        assert updated_row is not None
        updated_path = core._storage.paths.snapshot_path(
            project_id, updated_row["local_snapshot_path"], require_exists=True
        )
        updated_bytes = updated_path.read_bytes()
        with updated_path.open("rb") as handle:
            updated_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(updated_bytes).hexdigest(),
                byte_size=len(updated_bytes),
            )
        updated_object = next(
            obj
            for slide in updated_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert updated_object.text_style is not None
        assert updated_object.text_style.line_spacing_pt == 30
        assert updated_object.text_style.line_spacing_editable is True
        assert updated_object.text_style.space_before_pt == 18
        assert updated_object.text_style.space_before_editable is True
        assert updated_object.text_style.space_after_pt == 9
        assert updated_object.text_style.space_after_editable is True
        assert updated_object.text_style.first_line_indent_pt == 18
        assert updated_object.text_style.first_line_indent_editable is True
        assert updated_object.text_style.bullet_state == "bullet"
        assert updated_object.text_style.bullet_editable is True
        with zipfile.ZipFile(updated_path) as updated_archive:
            updated_slide = next(
                slide for slide in updated_fidelity.slides if slide.ordinal == slide_number
            )
            updated_root = etree.fromstring(updated_archive.read(updated_slide.part_name))
        updated_paragraph_properties = updated_root.xpath(
            ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]/p:txBody/a:p/a:pPr",
            namespaces={"p": PRESENTATION_NS, "a": DRAWING_NS},
            shape_id=target["ooxml_shape_id"],
        )[0]
        assert updated_paragraph_properties.get("marL") == "457200"
        assert updated_paragraph_properties.get("indent") == "228600"
        assert updated_paragraph_properties.xpath(
            "./a:buChar/@char", namespaces={"a": DRAWING_NS}
        ) == ["•"]

        invalid_alignment = _call(
            core,
            "multirun-invalid-alignment",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": updated_spacing["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": updated_object.text_style.fingerprint,
                "replacement_style": {"alignment": "distributed"},
            },
        )
        assert invalid_alignment["ok"] is False
        assert invalid_alignment["error"]["code"] == "INVALID_REQUEST"
        invalid_bullet_mode = _call(
            core,
            "multirun-invalid-bullet-mode",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": updated_spacing["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": updated_object.text_style.fingerprint,
                "replacement_style": {"bullet_mode": "ordinal"},
            },
        )
        assert invalid_bullet_mode["ok"] is False
        assert invalid_bullet_mode["error"]["code"] == "INVALID_REQUEST"
        invalid_indent = _call(
            core,
            "multirun-invalid-indent",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": updated_spacing["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": updated_object.text_style.fingerprint,
                "replacement_style": {"left_indent_pt": 145},
            },
        )
        assert invalid_indent["ok"] is False
        assert invalid_indent["error"]["code"] == "INVALID_REQUEST"
        invalid_first_line_indent = _call(
            core,
            "multirun-invalid-first-line-indent",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": updated_spacing["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": updated_object.text_style.fingerprint,
                "replacement_style": {"first_line_indent_pt": 145},
            },
        )
        assert invalid_first_line_indent["ok"] is False
        assert invalid_first_line_indent["error"]["code"] == "INVALID_REQUEST"
        invalid_line_spacing = _call(
            core,
            "multirun-invalid-line-spacing",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": updated_spacing["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": updated_object.text_style.fingerprint,
                "replacement_style": {"line_spacing_pt": 145},
            },
        )
        assert invalid_line_spacing["ok"] is False
        assert invalid_line_spacing["error"]["code"] == "INVALID_REQUEST"
        invalid_spacing = _call(
            core,
            "multirun-invalid-spacing",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": updated_spacing["result"]["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": updated_object.text_style.fingerprint,
                "replacement_style": {"space_after_pt": 145},
            },
        )
        assert invalid_spacing["ok"] is False
        assert invalid_spacing["error"]["code"] == "INVALID_REQUEST"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 3
    finally:
        core.close()


@pytest.mark.parametrize(
    ("character", "expected_mode", "expected_editable"),
    (
        ("•", "bullet", True),
        ("◦", "bullet_hollow", True),
        ("▪", "bullet_square", True),
        ("–", "bullet_dash", True),
        ("★", None, False),
    ),
)
def test_only_predefined_bullet_glyphs_are_editable(
    character: str,
    expected_mode: str | None,
    expected_editable: bool,
) -> None:
    properties = etree.Element(f"{{{DRAWING_NS}}}pPr")
    etree.SubElement(properties, f"{{{DRAWING_NS}}}buChar", {"char": character})

    assert _paragraph_bullet_mode_for_shape(properties) == (
        expected_mode,
        expected_editable,
    )


def test_plain_paragraph_supports_automatic_numbering(tmp_path: Path) -> None:
    source_path = tmp_path / "numbered-paragraph.pptx"
    slide_number, target_shape_id = _write_multirun_text_fixture(
        FIXTURE_ROOT / "deck" / "presentation.pptx",
        source_path,
    )
    original_bytes = source_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, source_path)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "numbered-paragraph-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": slide_number,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["targets"]
            if item["ooxml_shape_id"] == target_shape_id
        )
        assert target["text_format_editable"] is True
        assert target["text_style"]["bullet_state"] == "bullet"
        assert target["text_style"]["bullet_editable"] is True

        changed = _call(
            core,
            "numbered-paragraph-format",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"bullet_mode": "numbered"},
            },
        )
        assert changed["ok"] is True, changed
        assert changed["result"]["changed"] is True
        assert changed["result"]["parent_revision_id"] == opened["revision_id"]
        assert changed["result"]["revision_number"] == 2

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (changed["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        revised_bytes = revised_path.read_bytes()
        with revised_path.open("rb") as handle:
            revised_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(revised_bytes).hexdigest(),
                byte_size=len(revised_bytes),
            )
        revised_object = next(
            obj
            for slide in revised_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert revised_object.text == target["current_text"]
        assert revised_object.text_format_edit_candidate is True
        assert revised_object.text_style is not None
        assert revised_object.text_style.bullet_state == "numbered"
        assert revised_object.text_style.bullet_editable is True

        revised_slide = next(
            slide for slide in revised_fidelity.slides if slide.ordinal == slide_number
        )
        namespaces = {"p": PRESENTATION_NS, "a": DRAWING_NS}
        with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(revised_path) as after:
            assert source_path.read_bytes() == original_bytes
            assert {info.filename for info in before.infolist()} == {
                info.filename for info in after.infolist()
            }
            for part_name in {info.filename for info in before.infolist()} - {
                revised_slide.part_name
            }:
                assert before.read(part_name) == after.read(part_name)
            original_slide = next(
                slide
                for slide in inspect_pptx_snapshot(
                    io.BytesIO(original_bytes),
                    document_id=document["id"],
                    expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
                    byte_size=len(original_bytes),
                ).slides
                if slide.ordinal == slide_number
            )
            original_root = etree.fromstring(before.read(original_slide.part_name))
            revised_root = etree.fromstring(after.read(revised_slide.part_name))

        original_properties = original_root.xpath(
            ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]/p:txBody/a:p/a:pPr",
            namespaces=namespaces,
            shape_id=target["ooxml_shape_id"],
        )[0]
        revised_properties = revised_root.xpath(
            ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]/p:txBody/a:p/a:pPr",
            namespaces=namespaces,
            shape_id=target["ooxml_shape_id"],
        )[0]
        assert original_properties.xpath("./a:buChar/@char", namespaces=namespaces) == ["•"]
        assert revised_properties.xpath("./a:buAutoNum/@type", namespaces=namespaces) == [
            "arabicPeriod"
        ]
        assert revised_properties.xpath("./a:buChar | ./a:buNone", namespaces=namespaces) == []
    finally:
        core.close()


@pytest.mark.parametrize(
    ("bullet_mode", "bullet_tag", "bullet_attribute", "bullet_value"),
    (
        ("numbered", "buAutoNum", "type", "arabicPeriod"),
        ("bullet_hollow", "buChar", "char", "◦"),
        ("bullet_square", "buChar", "char", "▪"),
        ("bullet_dash", "buChar", "char", "–"),
    ),
)
def test_simple_multiparagraph_list_supports_predefined_styles(
    tmp_path: Path,
    bullet_mode: str,
    bullet_tag: str,
    bullet_attribute: str,
    bullet_value: str,
) -> None:
    base_path = tmp_path / "mixed-numbered-paragraphs-base.pptx"
    slide_number, target_shape_id = _write_multirun_text_fixture(
        FIXTURE_ROOT / "deck" / "presentation.pptx",
        base_path,
    )
    base_bytes = base_path.read_bytes()
    base_fidelity = inspect_pptx_snapshot(
        io.BytesIO(base_bytes),
        document_id=str(uuid.uuid4()),
        expected_sha256=hashlib.sha256(base_bytes).hexdigest(),
        byte_size=len(base_bytes),
    )
    slide_part_name = next(
        slide.part_name for slide in base_fidelity.slides if slide.ordinal == slide_number
    )
    source_path = tmp_path / "mixed-numbered-paragraphs.pptx"
    with zipfile.ZipFile(base_path) as source_archive:
        with zipfile.ZipFile(source_path, "w", allowZip64=True) as output_archive:
            output_archive.comment = source_archive.comment
            for info in source_archive.infolist():
                content = source_archive.read(info.filename)
                if info.filename == slide_part_name:
                    root = etree.fromstring(content)
                    shapes = root.xpath(
                        ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
                        namespaces={"p": PRESENTATION_NS},
                        shape_id=target_shape_id,
                    )
                    text_body = shapes[0].find(f"{{{PRESENTATION_NS}}}txBody")
                    assert text_body is not None
                    paragraphs = text_body.findall(f"{{{DRAWING_NS}}}p")
                    assert len(paragraphs) == 1
                    for label, source_bullet_mode in (
                        ("Second item", None),
                        ("Third item", "none"),
                    ):
                        paragraph = copy.deepcopy(paragraphs[0])
                        runs = paragraph.findall(f"{{{DRAWING_NS}}}r")
                        assert len(runs) == 2
                        for run in runs[1:]:
                            paragraph.remove(run)
                        text_node = runs[0].find(f"{{{DRAWING_NS}}}t")
                        assert text_node is not None
                        text_node.text = label
                        text_node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
                        properties = paragraph.find(f"{{{DRAWING_NS}}}pPr")
                        assert properties is not None
                        existing_bullets = [
                            child
                            for child in list(properties)
                            if etree.QName(child).localname
                            in {"buChar", "buAutoNum", "buNone", "buBlip"}
                        ]
                        assert len(existing_bullets) == 1
                        old_bullet = existing_bullets[0]
                        properties.remove(old_bullet)
                        if source_bullet_mode == "none":
                            replacement = etree.Element(f"{{{DRAWING_NS}}}buNone")
                            insertion = _paragraph_property_insert_index(properties, "buNone")
                            assert insertion is not None
                            properties.insert(insertion, replacement)
                        text_body.append(paragraph)
                    content = etree.tostring(
                        root,
                        encoding="UTF-8",
                        xml_declaration=True,
                        standalone=True,
                    )
                output_archive.writestr(info, content)

    original_bytes = source_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, source_path)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "multiparagraph-numbering-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": slide_number,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["targets"]
            if item["ooxml_shape_id"] == target_shape_id
        )
        assert target["text_style"]["editable"] is True, target["text_style"]
        assert target["text_format_editable"] is True, target
        assert target["current_text"].splitlines()[-2:] == ["Second item", "Third item"]
        assert target["text_format_editable"] is True
        assert target["single_run_editable"] is False
        assert target["uniform_run_editable"] is False
        assert target["text_style"]["editable"] is True
        assert target["text_style"]["bullet_state"] is None
        assert target["text_style"]["bullet_editable"] is True

        changed = _call(
            core,
            "multiparagraph-numbering-format",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"bullet_mode": bullet_mode},
            },
        )
        assert changed["ok"] is True, changed
        assert changed["result"]["changed"] is True

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (changed["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        revised_bytes = revised_path.read_bytes()
        with revised_path.open("rb") as handle:
            revised_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(revised_bytes).hexdigest(),
                byte_size=len(revised_bytes),
            )
        revised_object = next(
            obj
            for slide in revised_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert revised_object.text == target["current_text"]
        assert revised_object.text_format_edit_candidate is True
        assert revised_object.single_run_text_edit_candidate is False
        assert revised_object.uniform_run_text_edit_candidate is False
        assert revised_object.text_style is not None
        assert revised_object.text_style.bullet_state == bullet_mode
        assert revised_object.text_style.bullet_editable is True

        revised_slide = next(
            slide for slide in revised_fidelity.slides if slide.ordinal == slide_number
        )
        namespaces = {"p": PRESENTATION_NS, "a": DRAWING_NS}
        with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(revised_path) as after:
            assert source_path.read_bytes() == original_bytes
            for part_name in {info.filename for info in before.infolist()} - {
                revised_slide.part_name
            }:
                assert before.read(part_name) == after.read(part_name)
            root = etree.fromstring(after.read(revised_slide.part_name))
        revised_paragraphs = root.xpath(
            ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]/p:txBody/a:p",
            namespaces=namespaces,
            shape_id=target["ooxml_shape_id"],
        )
        assert len(revised_paragraphs) == 3
        assert [
            paragraph.xpath(f"./a:pPr/a:{bullet_tag}/@{bullet_attribute}", namespaces=namespaces)
            for paragraph in revised_paragraphs
        ] == [[bullet_value], [bullet_value], [bullet_value]]
        assert all(
            paragraph.xpath("./a:pPr/@marL", namespaces=namespaces) == ["228600"]
            and paragraph.xpath("./a:pPr/@indent", namespaces=namespaces) == ["-114300"]
            for paragraph in revised_paragraphs
        )
        other_bullet_tag = "buChar" if bullet_tag == "buAutoNum" else "buAutoNum"
        assert all(
            not paragraph.xpath(
                f"./a:pPr/a:buNone | ./a:pPr/a:{other_bullet_tag}",
                namespaces=namespaces,
            )
            for paragraph in revised_paragraphs
        )
    finally:
        core.close()


def test_custom_numbering_style_remains_read_only(tmp_path: Path) -> None:
    source_path = tmp_path / "custom-numbered-paragraph.pptx"
    slide_number, target_shape_id = _write_multirun_text_fixture(
        FIXTURE_ROOT / "deck" / "presentation.pptx",
        source_path,
    )
    original_bytes = source_path.read_bytes()
    document_id = str(uuid.uuid4())
    original_fidelity = inspect_pptx_snapshot(
        io.BytesIO(original_bytes),
        document_id=document_id,
        expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
        byte_size=len(original_bytes),
    )
    slide_part_name = next(
        slide.part_name for slide in original_fidelity.slides if slide.ordinal == slide_number
    )
    custom_path = tmp_path / "custom-numbered-paragraph-source.pptx"
    with zipfile.ZipFile(source_path) as source_archive:
        with zipfile.ZipFile(custom_path, "w", allowZip64=True) as output_archive:
            output_archive.comment = source_archive.comment
            for info in source_archive.infolist():
                content = source_archive.read(info.filename)
                if info.filename == slide_part_name:
                    root = etree.fromstring(content)
                    shapes = root.xpath(
                        ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
                        namespaces={"p": PRESENTATION_NS},
                        shape_id=target_shape_id,
                    )
                    properties = shapes[0].xpath(
                        "./p:txBody/a:p/a:pPr",
                        namespaces={"p": PRESENTATION_NS, "a": DRAWING_NS},
                    )[0]
                    plain_bullet = properties.find(f"{{{DRAWING_NS}}}buChar")
                    assert plain_bullet is not None
                    custom_number = etree.Element(f"{{{DRAWING_NS}}}buAutoNum")
                    custom_number.set("type", "romanUcPeriod")
                    properties.replace(plain_bullet, custom_number)
                    content = etree.tostring(
                        root,
                        encoding="UTF-8",
                        xml_declaration=True,
                        standalone=True,
                    )
                output_archive.writestr(info, content)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, custom_path)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "custom-numbered-paragraph-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": slide_number,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["targets"]
            if item["ooxml_shape_id"] == target_shape_id
        )
        assert target["text_style"]["bullet_state"] is None
        assert target["text_style"]["bullet_editable"] is False

        rejected = _call(
            core,
            "custom-numbered-paragraph-write-rejected",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"bullet_mode": "none"},
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 1
    finally:
        core.close()


def test_unsupported_source_left_indent_stays_read_only(tmp_path: Path) -> None:
    source_path = tmp_path / "unsupported-left-indent-source.pptx"
    slide_number, target_shape_id = _write_multirun_text_fixture(
        FIXTURE_ROOT / "deck" / "presentation.pptx",
        source_path,
    )
    source_bytes = source_path.read_bytes()
    source_digest = hashlib.sha256(source_bytes).hexdigest()
    fidelity = inspect_pptx_snapshot(
        io.BytesIO(source_bytes),
        document_id=str(uuid.uuid4()),
        expected_sha256=source_digest,
        byte_size=len(source_bytes),
    )
    slide = next(item for item in fidelity.slides if item.ordinal == slide_number)
    rewritten_path = tmp_path / "unsupported-left-indent-edited.pptx"
    with zipfile.ZipFile(source_path) as source_archive:
        source_parts = []
        for info in source_archive.infolist():
            content = source_archive.read(info.filename)
            if info.filename == slide.part_name:
                root = etree.fromstring(content)
                shapes = root.xpath(
                    ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
                    namespaces={"p": PRESENTATION_NS},
                    shape_id=target_shape_id,
                )
                paragraph_properties = shapes[0].xpath(
                    "./p:txBody/a:p/a:pPr",
                    namespaces={"p": PRESENTATION_NS, "a": DRAWING_NS},
                )[0]
                paragraph_properties.set("marL", str(MAX_TEXT_LEFT_INDENT_EMU + 1))
                paragraph_properties.set("indent", str(MAX_TEXT_FIRST_LINE_INDENT_EMU + 1))
                bullet_chars = paragraph_properties.findall(f"{{{DRAWING_NS}}}buChar")
                assert len(bullet_chars) == 1
                paragraph_properties.remove(bullet_chars[0])
                auto_number = etree.Element(f"{{{DRAWING_NS}}}buAutoNum", {"type": "romanUcPeriod"})
                bullet_insert_index = _paragraph_property_insert_index(
                    paragraph_properties, "buAutoNum"
                )
                assert bullet_insert_index is not None
                paragraph_properties.insert(bullet_insert_index, auto_number)
                line_spacing = etree.Element(f"{{{DRAWING_NS}}}lnSpc")
                line_spacing_percentage = etree.SubElement(line_spacing, f"{{{DRAWING_NS}}}spcPct")
                line_spacing_percentage.set("val", "100000")
                paragraph_properties.insert(0, line_spacing)
                spacing_before = etree.Element(f"{{{DRAWING_NS}}}spcBef")
                spacing_percentage = etree.SubElement(spacing_before, f"{{{DRAWING_NS}}}spcPct")
                spacing_percentage.set("val", "100000")
                paragraph_properties.insert(1, spacing_before)
                content = etree.tostring(
                    root,
                    encoding="UTF-8",
                    xml_declaration=True,
                    standalone=True,
                )
            source_parts.append((copy.copy(info), content))
        with zipfile.ZipFile(rewritten_path, "w", allowZip64=True) as output_archive:
            output_archive.comment = source_archive.comment
            for info, content in source_parts:
                output_archive.writestr(info, content)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, rewritten_path)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "unsupported-left-indent-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": slide_number,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["targets"]
            if item["ooxml_shape_id"] == target_shape_id
        )
        assert target["text_style"]["editable"] is True
        assert target["text_style"]["left_indent_pt"] is None
        assert target["text_style"]["left_indent_editable"] is False
        assert target["text_style"]["first_line_indent_pt"] is None
        assert target["text_style"]["first_line_indent_editable"] is False
        assert target["text_style"]["line_spacing_pt"] is None
        assert target["text_style"]["line_spacing_editable"] is False
        assert target["text_style"]["bullet_state"] is None
        assert target["text_style"]["bullet_editable"] is False
        assert target["text_style"]["space_before_pt"] is None
        assert target["text_style"]["space_before_editable"] is False
        assert target["text_style"]["space_after_pt"] is None
        assert target["text_style"]["space_after_editable"] is True

        rejected = _call(
            core,
            "unsupported-left-indent-edit",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"left_indent_pt": 36},
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        first_line_rejected = _call(
            core,
            "unsupported-first-line-indent-edit",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"first_line_indent_pt": 36},
            },
        )
        assert first_line_rejected["ok"] is False
        assert first_line_rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        line_spacing_rejected = _call(
            core,
            "unsupported-line-spacing-edit",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"line_spacing_pt": 24},
            },
        )
        assert line_spacing_rejected["ok"] is False
        assert line_spacing_rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        bullet_rejected = _call(
            core,
            "unsupported-bullet-mode-edit",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"bullet_mode": "none"},
            },
        )
        assert bullet_rejected["ok"] is False
        assert bullet_rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        spacing_rejected = _call(
            core,
            "unsupported-before-spacing-edit",
            "presentation.workspace.format_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "expected_style_fingerprint": target["text_style"]["fingerprint"],
                "replacement_style": {"space_before_pt": 36},
            },
        )
        assert spacing_rejected["ok"] is False
        assert spacing_rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 1
    finally:
        core.close()


def test_uniform_multirun_text_replacement_preserves_style_and_paragraph(
    tmp_path: Path,
) -> None:
    source_path = tmp_path / "uniform-multirun.pptx"
    slide_number, target_shape_id = _write_multirun_text_fixture(
        FIXTURE_ROOT / "deck" / "presentation.pptx",
        source_path,
        uniform_runs=True,
    )
    original_bytes = source_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, source_path)
        opened = _open(core, project_id, document["id"])
        original_fidelity = inspect_pptx_snapshot(
            io.BytesIO(original_bytes),
            document_id=document["id"],
            expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
            byte_size=len(original_bytes),
        )
        slide_part_name = next(
            slide.part_name for slide in original_fidelity.slides if slide.ordinal == slide_number
        )
        inventory = _call(
            core,
            "uniform-multirun-text-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": slide_number,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["targets"]
            if item["ooxml_shape_id"] == target_shape_id
        )
        assert target["uniform_run_editable"] is True
        assert target["single_run_editable"] is False
        assert target["text_format_editable"] is True
        assert target["text_style"]["fingerprint"]

        legacy = _call(
            core,
            "uniform-multirun-legacy-replacement-rejected",
            "presentation.workspace.replace_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "operation_type": "replace_single_run_text",
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "replacement_text": "legacy route",
            },
        )
        assert legacy["ok"] is False
        assert legacy["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"

        replacement = "  Revised uniform text  "
        changed = _call(
            core,
            "uniform-multirun-replacement",
            "presentation.workspace.replace_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "operation_type": "replace_text_shape",
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "replacement_text": replacement,
            },
        )
        assert changed["ok"] is True, changed
        assert changed["result"]["changed"] is True
        assert changed["result"]["parent_revision_id"] == opened["revision_id"]
        assert changed["result"]["revision_number"] == 2

        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (changed["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        namespaces = {"p": PRESENTATION_NS, "a": DRAWING_NS}
        with zipfile.ZipFile(source_path) as before, zipfile.ZipFile(revised_path) as after:
            names = {info.filename for info in before.infolist()}
            assert names == {info.filename for info in after.infolist()}
            changed_parts = {name for name in names if before.read(name) != after.read(name)}
            original_root = etree.fromstring(before.read(slide_part_name))
            revised_root = etree.fromstring(after.read(slide_part_name))
            original_shape = original_root.xpath(
                ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
                namespaces=namespaces,
                shape_id=target_shape_id,
            )[0]
            revised_shape = revised_root.xpath(
                ".//p:sp[p:nvSpPr/p:cNvPr[@id=$shape_id]]",
                namespaces=namespaces,
                shape_id=target_shape_id,
            )[0]
            original_run = original_shape.xpath("./p:txBody/a:p/a:r", namespaces=namespaces)[0]
            revised_runs = revised_shape.xpath("./p:txBody/a:p/a:r", namespaces=namespaces)
            original_properties = original_run.find(f"{{{DRAWING_NS}}}rPr")
            revised_properties = revised_runs[0].find(f"{{{DRAWING_NS}}}rPr")
            revised_text = revised_runs[0].find(f"{{{DRAWING_NS}}}t")
            paragraph_properties = revised_shape.xpath(
                "./p:txBody/a:p/a:pPr", namespaces=namespaces
            )[0]
        assert changed_parts == {slide_part_name}
        assert len(revised_runs) == 1
        assert revised_text is not None and revised_text.text == replacement
        assert revised_text.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
        assert original_properties is not None and revised_properties is not None
        assert etree.tostring(original_properties, method="c14n") == etree.tostring(
            revised_properties, method="c14n"
        )
        assert paragraph_properties.get("marL") == "228600"
        assert source_path.read_bytes() == original_bytes

        stale = _call(
            core,
            "uniform-multirun-stale-target-rejected",
            "presentation.workspace.replace_text_shape",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": changed["result"]["revision_id"],
                "operation_type": "replace_text_shape",
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": target["current_text"],
                "replacement_text": "stale edit",
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 2
    finally:
        core.close()


def test_shape_fill_and_outline_colors_create_verified_revisions(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "shape-style-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["objects"]
            if item["name"] == "Workspace transform test shape"
        )
        style = target["shape_style"]
        assert style["editable"] is True
        assert style["fill_editable"] is True
        assert style["line_editable"] is True
        assert style["fill_color"] == "#E0EDFA"
        assert style["line_color"] == "#396EA7"

        changed = _call(
            core,
            "format-shape-fill",
            "presentation.workspace.format_shape_style",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_style_fingerprint": style["fingerprint"],
                "replacement_style": {"fill_color": "#123456"},
            },
        )
        assert changed["ok"] is True, changed
        assert changed["result"]["changed"] is True
        assert changed["result"]["parent_revision_id"] == opened["revision_id"]

        history = _list_revisions(core, project_id, document["id"])
        assert history["current_revision_id"] == changed["result"]["revision_id"]
        with core._storage.project_database(project_id) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id = ?",
                (changed["result"]["revision_id"],),
            ).fetchone()
        assert row is not None
        revised_path = core._storage.paths.snapshot_path(
            project_id, row["local_snapshot_path"], require_exists=True
        )
        with revised_path.open("rb") as handle:
            revised_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=changed["result"]["sha256"],
                byte_size=changed["result"]["byte_size"],
            )
        revised_object = next(
            obj
            for slide in revised_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert revised_object.shape_style is not None
        assert revised_object.shape_style.fill_color == "#123456"
        assert revised_object.shape_style.line_color == "#396EA7"
        with original_path.open("rb") as handle:
            original_revision_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=hashlib.sha256(original_bytes).hexdigest(),
                byte_size=len(original_bytes),
            )
        original_object = next(
            obj
            for slide in original_revision_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert revised_object.geometry == original_object.geometry
        edited_slide_part = next(
            slide.part_name
            for slide in original_revision_fidelity.slides
            if slide.slide_id == inventory["result"]["slide_id"]
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(revised_path) as after:
            assert {info.filename for info in before.infolist()} == {
                info.filename for info in after.infolist()
            }
            for part_name in {info.filename for info in before.infolist()} - {edited_slide_part}:
                assert before.read(part_name) == after.read(part_name)
        assert original_path.read_bytes() == original_bytes

        refreshed = _open(core, project_id, document["id"])
        refreshed_inventory = _call(
            core,
            "shape-style-targets-after-fill",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": refreshed["revision_id"],
                "slide_number": 1,
            },
        )
        refreshed_target = next(
            item
            for item in refreshed_inventory["result"]["objects"]
            if item["object_id"] == target["object_id"]
        )
        line_changed = _call(
            core,
            "format-shape-outline",
            "presentation.workspace.format_shape_style",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": refreshed["revision_id"],
                "slide_id": refreshed_inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_style_fingerprint": refreshed_target["shape_style"]["fingerprint"],
                "replacement_style": {"line_color": "#654321"},
            },
        )
        assert line_changed["ok"] is True, line_changed
        assert line_changed["result"]["changed"] is True

        latest = _open(core, project_id, document["id"])
        latest_inventory = _call(
            core,
            "shape-style-targets-after-outline",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": latest["revision_id"],
                "slide_number": 1,
            },
        )
        latest_target = next(
            item
            for item in latest_inventory["result"]["objects"]
            if item["object_id"] == target["object_id"]
        )
        assert latest_target["shape_style"]["fill_color"] == "#123456"
        assert latest_target["shape_style"]["line_color"] == "#654321"

        no_op = _call(
            core,
            "format-shape-color-no-op",
            "presentation.workspace.format_shape_style",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": latest["revision_id"],
                "slide_id": latest_inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_style_fingerprint": latest_target["shape_style"]["fingerprint"],
                "replacement_style": {"line_color": "#654321"},
            },
        )
        assert no_op["ok"] is True, no_op
        assert no_op["result"]["changed"] is False
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 3

        stale = _call(
            core,
            "format-shape-color-stale",
            "presentation.workspace.format_shape_style",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": latest["revision_id"],
                "slide_id": latest_inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_style_fingerprint": "0" * 64,
                "replacement_style": {"fill_color": "#FFFFFF"},
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"

        invalid = _call(
            core,
            "format-shape-color-invalid",
            "presentation.workspace.format_shape_style",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": latest["revision_id"],
                "slide_id": latest_inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_style_fingerprint": latest_target["shape_style"]["fingerprint"],
                "replacement_style": {"fill_color": "red"},
            },
        )
        assert invalid["ok"] is False
        assert invalid["error"]["code"] == "INVALID_REQUEST"
    finally:
        core.close()


def test_shape_fill_editor_refuses_gradient_styles_without_a_revision(tmp_path: Path) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    gradient_path = tmp_path / "gradient-fill.pptx"
    with (
        zipfile.ZipFile(original_path) as source,
        zipfile.ZipFile(gradient_path, "w", allowZip64=True) as output,
    ):
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml":
                root = etree.fromstring(data)
                shape = next(
                    item
                    for item in root.xpath(
                        "./p:cSld/p:spTree/p:sp", namespaces={"p": PRESENTATION_NS}
                    )
                    if item.xpath(
                        "./p:nvSpPr/p:cNvPr/@name",
                        namespaces={"p": PRESENTATION_NS},
                    )
                    == ["Workspace transform test shape"]
                )
                shape_properties = shape.find(f"{{{PRESENTATION_NS}}}spPr")
                assert shape_properties is not None
                solid = shape_properties.find(f"{{{DRAWING_NS}}}solidFill")
                assert solid is not None
                fill_position = shape_properties.index(solid)
                shape_properties.remove(solid)
                gradient = etree.Element(f"{{{DRAWING_NS}}}gradFill", rotWithShape="1")
                stops = etree.SubElement(gradient, f"{{{DRAWING_NS}}}gsLst")
                first_stop = etree.SubElement(stops, f"{{{DRAWING_NS}}}gs", pos="0")
                etree.SubElement(first_stop, f"{{{DRAWING_NS}}}srgbClr", val="E0EDFA")
                last_stop = etree.SubElement(stops, f"{{{DRAWING_NS}}}gs", pos="100000")
                etree.SubElement(last_stop, f"{{{DRAWING_NS}}}srgbClr", val="396EA7")
                etree.SubElement(gradient, f"{{{DRAWING_NS}}}lin", ang="0", scaled="1")
                shape_properties.insert(fill_position, gradient)
                data = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
            output.writestr(info, data)

    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id, gradient_path)
        opened = _open(core, project_id, document["id"])
        inventory = _call(
            core,
            "gradient-shape-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert inventory["ok"] is True, inventory
        target = next(
            item
            for item in inventory["result"]["objects"]
            if item["name"] == "Workspace transform test shape"
        )
        assert target["shape_style"]["editable"] is True
        assert target["shape_style"]["fill_editable"] is False
        assert target["shape_style"]["line_editable"] is True

        rejected = _call(
            core,
            "gradient-shape-fill-rejected",
            "presentation.workspace.format_shape_style",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "slide_id": inventory["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_style_fingerprint": target["shape_style"]["fingerprint"],
                "replacement_style": {"fill_color": "#123456"},
            },
        )
        assert rejected["ok"] is False
        assert rejected["error"]["code"] == "PRESENTATION_EDIT_UNSUPPORTED"
        assert len(_list_revisions(core, project_id, document["id"])["revisions"]) == 1
    finally:
        core.close()


def test_single_run_edit_creates_revision_and_preserves_other_package_parts(
    tmp_path: Path,
) -> None:
    original_path = FIXTURE_ROOT / "deck" / "presentation.pptx"
    original_bytes = original_path.read_bytes()
    original_sha256 = hashlib.sha256(original_bytes).hexdigest()
    core = CoreService(data_root=tmp_path / "data")
    try:
        project_id = _project_id(core)
        document = _import_presentation(core, project_id)
        original_revision = _list_revisions(core, project_id, document["id"])["revisions"][0]
        started_run = _call(
            core,
            "start-bound-run",
            "session.start",
            {"project_id": project_id, "mode": "run"},
        )
        assert started_run["ok"] is True, started_run
        run_session = started_run["result"]["session"]
        assert run_session["presentation_revision_id"] == original_revision["revision_id"]
        opened = _open(core, project_id, document["id"])
        assert opened["revision_id"] == original_revision["revision_id"]
        targets_result = _call(
            core,
            "text-targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "revision_id": opened["revision_id"],
                "slide_number": 1,
            },
        )
        assert targets_result["ok"] is True, targets_result
        targets = targets_result["result"]["targets"]
        target = next(item for item in targets if item["single_run_editable"])
        expected_text = target["current_text"]
        assert isinstance(expected_text, str) and expected_text
        replacement_text = f"New revision evidence marker {uuid.uuid4().hex}"
        edited = _call(
            core,
            "replace-text",
            "presentation.workspace.replace_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "operation_type": "replace_single_run_text",
                "slide_id": targets_result["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "replacement_text": replacement_text,
            },
        )
        assert edited["ok"] is True, edited
        assert edited["result"]["changed"] is True
        assert edited["result"]["revision_number"] == 2
        assert edited["result"]["parent_revision_id"] == opened["revision_id"]

        history = _list_revisions(core, project_id, document["id"])
        current_revision = history["revisions"][0]
        assert current_revision["revision_id"] == edited["result"]["revision_id"]
        assert current_revision["revision_kind"] == "edited"
        assert current_revision["is_current"] is True
        assert history["revisions"][1]["revision_id"] == original_revision["revision_id"]

        with core._storage.project_database(project_id) as connection:
            original_unit_ids = {
                str(row["id"])
                for row in connection.execute(
                    "SELECT id FROM source_units WHERE document_id = ? "
                    "AND presentation_revision_id = ?",
                    (document["id"], original_revision["revision_id"]),
                ).fetchall()
            }
        original_search = _call(
            core,
            "search-original-revision-before-edit-index",
            "search.lexical",
            {
                "project_id": project_id,
                "query": expected_text,
                "presentation_revision_id": original_revision["revision_id"],
            },
        )
        assert original_search["ok"] is True, original_search
        assert any(
            result["source_unit_id"] in original_unit_ids
            for result in original_search["result"]["results"]
        )
        current_before_reindex = _call(
            core,
            "search-current-revision-before-index",
            "search.lexical",
            {
                "project_id": project_id,
                "query": replacement_text,
            },
        )
        assert current_before_reindex["ok"] is True, current_before_reindex
        assert current_before_reindex["result"]["results"] == []
        stale_original = _call(
            core,
            "search-current-revision-old-text",
            "search.lexical",
            {
                "project_id": project_id,
                "query": expected_text,
            },
        )
        assert stale_original["ok"] is True, stale_original
        assert all(
            result["source_unit_id"] not in original_unit_ids
            for result in stale_original["result"]["results"]
        )

        reindexed = _call(
            core,
            "reindex-edited-presentation",
            "source.reindex",
            {"project_id": project_id, "document_id": document["id"]},
        )
        assert reindexed["ok"] is True, reindexed
        current_after_reindex = _call(
            core,
            "search-current-revision-after-index",
            "search.lexical",
            {
                "project_id": project_id,
                "query": replacement_text,
            },
        )
        assert current_after_reindex["ok"] is True, current_after_reindex
        assert current_after_reindex["result"]["results"]
        assert all(
            result["source_unit_id"] not in original_unit_ids
            for result in current_after_reindex["result"]["results"]
        )
        historical_after_reindex = _call(
            core,
            "search-historical-revision-after-index",
            "search.lexical",
            {
                "project_id": project_id,
                "query": expected_text,
                "presentation_revision_id": original_revision["revision_id"],
            },
        )
        assert historical_after_reindex["ok"] is True, historical_after_reindex
        assert any(
            result["source_unit_id"] in original_unit_ids
            for result in historical_after_reindex["result"]["results"]
        )

        with core._storage.project_database(project_id) as connection:
            edited_row = connection.execute(
                """
                SELECT local_snapshot_path FROM presentation_revisions
                WHERE id = ? AND presentation_id = ?
                """,
                (edited["result"]["revision_id"], document["id"]),
            ).fetchone()
        assert edited_row is not None
        edited_path = core._storage.paths.snapshot_path(
            project_id, edited_row["local_snapshot_path"], require_exists=True
        )
        assert hashlib.sha256(edited_path.read_bytes()).hexdigest() == edited["result"]["sha256"]
        with edited_path.open("rb") as handle:
            edited_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=edited["result"]["sha256"],
                byte_size=edited["result"]["byte_size"],
            )
        edited_object = next(
            obj
            for slide in edited_fidelity.slides
            for obj in slide.objects
            if obj.object_id == target["object_id"]
        )
        assert edited_object.text == replacement_text

        with original_path.open("rb") as handle:
            original_fidelity = inspect_pptx_snapshot(
                handle,
                document_id=document["id"],
                expected_sha256=original_sha256,
                byte_size=len(original_bytes),
            )
        touched_part = next(
            slide.part_name
            for slide in original_fidelity.slides
            if slide.slide_id == targets_result["result"]["slide_id"]
        )
        with zipfile.ZipFile(original_path) as before, zipfile.ZipFile(edited_path) as after:
            before_names = {info.filename for info in before.infolist()}
            after_names = {info.filename for info in after.infolist()}
            assert before_names == after_names
            for part_name in before_names - {touched_part}:
                assert before.read(part_name) == after.read(part_name)
        assert original_path.read_bytes() == original_bytes
        session_after_edit = _call(
            core,
            "get-bound-run",
            "session.get",
            {"project_id": project_id, "session_id": run_session["id"]},
        )
        assert session_after_edit["ok"] is True, session_after_edit
        assert (
            session_after_edit["result"]["session"]["presentation_revision_id"]
            == original_revision["revision_id"]
        )

        stale = _call(
            core,
            "stale-replace-text",
            "presentation.workspace.replace_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": opened["revision_id"],
                "operation_type": "replace_single_run_text",
                "slide_id": targets_result["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "replacement_text": "stale edit",
            },
        )
        assert stale["ok"] is False
        assert stale["error"]["code"] == "PRESENTATION_REVISION_CONFLICT"

        destination = tmp_path / "current-copy.pptx"
        current_copy = _call(
            core,
            "export-current-copy",
            "presentation.workspace.export_current_copy",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "destination_path": str(destination),
            },
        )
        assert current_copy["ok"] is True, current_copy
        assert current_copy["result"]["copy_kind"] == "current_revision"
        assert destination.read_bytes() == edited_path.read_bytes()

        restored = _call(
            core,
            "restore-original-revision",
            "presentation.workspace.restore_revision",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "revision_id": original_revision["revision_id"],
            },
        )
        assert restored["ok"] is True, restored
        assert restored["result"]["changed"] is True
        assert (
            _list_revisions(core, project_id, document["id"])["current_revision_id"]
            == (original_revision["revision_id"])
        )

        redone = _call(
            core,
            "redo-edited-revision",
            "presentation.workspace.step_revision",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": original_revision["revision_id"],
                "direction": "redo",
            },
        )
        assert redone["ok"] is True, redone
        assert redone["result"]["revision_id"] == edited["result"]["revision_id"]
        undone = _call(
            core,
            "undo-edited-revision",
            "presentation.workspace.step_revision",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": edited["result"]["revision_id"],
                "direction": "undo",
            },
        )
        assert undone["ok"] is True, undone
        assert undone["result"]["revision_id"] == original_revision["revision_id"]

        restored_path = tmp_path / "restored-original.pptx"
        restored_copy = _call(
            core,
            "export-restored-original",
            "presentation.workspace.export_current_copy",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "destination_path": str(restored_path),
            },
        )
        assert restored_copy["ok"] is True, restored_copy
        assert restored_path.read_bytes() == original_bytes

        branched = _call(
            core,
            "edit-from-restored-revision",
            "presentation.workspace.replace_single_run_text",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": original_revision["revision_id"],
                "operation_type": "replace_single_run_text",
                "slide_id": targets_result["result"]["slide_id"],
                "object_id": target["object_id"],
                "expected_text": expected_text,
                "replacement_text": "A new branch after restore",
            },
        )
        assert branched["ok"] is True, branched
        restore_again = _call(
            core,
            "restore-before-branch-redo",
            "presentation.workspace.restore_revision",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": branched["result"]["revision_id"],
                "revision_id": original_revision["revision_id"],
            },
        )
        assert restore_again["ok"] is True, restore_again
        ambiguous_redo = _call(
            core,
            "reject-ambiguous-redo",
            "presentation.workspace.step_revision",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": original_revision["revision_id"],
                "direction": "redo",
            },
        )
        assert ambiguous_redo["ok"] is False
        assert ambiguous_redo["error"]["code"] == "PRESENTATION_REVISION_BRANCH_AMBIGUOUS"
        no_earlier_revision = _call(
            core,
            "undo-at-original-revision",
            "presentation.workspace.step_revision",
            {
                "project_id": project_id,
                "document_id": document["id"],
                "expected_revision_id": original_revision["revision_id"],
                "direction": "undo",
            },
        )
        assert no_earlier_revision["ok"] is True, no_earlier_revision
        assert no_earlier_revision["result"]["changed"] is False
        stopped_run = _call(
            core,
            "stop-bound-run",
            "session.stop",
            {
                "project_id": project_id,
                "session_id": run_session["id"],
                "status": "aborted",
            },
        )
        assert stopped_run["ok"] is True, stopped_run
    finally:
        core.close()


@pytest.mark.parametrize("expected", ["stale", "invalid", "matching"])
def test_current_copy_revision_guard_preserves_destination(tmp_path: Path, expected: str) -> None:
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project)
        revision = _open(core, project, document["id"])["revision_id"]
        destination = tmp_path / "copy.pptx"
        destination.write_bytes(b"keep existing destination")
        result = _call(
            core,
            "guard-copy",
            "presentation.workspace.export_current_copy",
            {
                "project_id": project,
                "document_id": document["id"],
                "expected_revision_id": revision
                if expected == "matching"
                else (str(uuid.uuid4()) if expected == "stale" else "invalid"),
                "destination_path": str(destination),
            },
        )
        if expected == "matching":
            assert result["ok"] is True, result
            assert result["result"]["revision_id"] == revision
            assert result["result"]["sha256"] == document["sha256"]
        else:
            assert result["ok"] is False, result
            assert result["error"]["code"] == (
                "PRESENTATION_REVISION_CONFLICT" if expected == "stale" else "INVALID_REQUEST"
            )
            assert destination.read_bytes() == b"keep existing destination"
        assert not list(tmp_path.glob(".presenter-copilot-*.tmp"))
    finally:
        core.close()
