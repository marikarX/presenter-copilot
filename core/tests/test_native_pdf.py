from __future__ import annotations

import hashlib
import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from presenter_core.errors import CoreDomainError
from presenter_core.presentation.native_pdf import (
    export_native_pdf,
    run_pdf_worker,
    validate_native_source,
    validate_pdf,
)


def _job(tmp_path: Path, hidden: bool = False) -> dict[str, Any]:
    source = tmp_path / "revision.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    return {
        "snapshot_path": str(source),
        "output_path": str(tmp_path / "revision.pdf"),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "include_hidden": hidden,
    }


class OfficeFixture:
    def __init__(self, *, fail: bool = False, wrong_pages: bool = False) -> None:
        self.closed = False
        self.opened: list[Any] = []
        self.export_args: tuple[Any, ...] = ()
        self.fail = fail
        self.wrong_pages = wrong_pages
        self.Presentations = SimpleNamespace(Open=self.open)
        self.Slides = SimpleNamespace(
            Count=2,
            Item=lambda index: SimpleNamespace(
                SlideShowTransition=SimpleNamespace(Hidden=-1 if index == 2 else 0)
            ),
        )
        self.PageSetup = SimpleNamespace(SlideWidth=720, SlideHeight=405)

    def open(self, path: str, **kwargs: Any) -> OfficeFixture:
        self.opened.append((path, kwargs))
        return self

    def ExportAsFixedFormat(self, *args: Any) -> None:
        self.export_args = args
        if self.fail:
            raise RuntimeError("private-path/token should never reach the renderer")
        writer = PdfWriter()
        for _ in range(2 if args[6] == 1 and not self.wrong_pages else 1):
            writer.add_blank_page(width=720, height=405)
        writer.write(args[0])

    def Close(self) -> None:
        self.closed = True

    def Quit(self) -> None:
        raise AssertionError("Never quit the shared PowerPoint application")


@pytest.mark.parametrize("hidden,pages", [(False, 1), (True, 2)])
def test_native_export_reads_only_owned_copy_closes_it_and_validates_pages(
    tmp_path: Path, hidden: bool, pages: int
) -> None:
    job = _job(tmp_path, hidden)
    office = OfficeFixture()
    before = Path(job["snapshot_path"]).read_bytes()
    result = export_native_pdf(job, application=office)
    assert result["page_count"] == pages
    assert result["include_hidden"] is hidden
    assert result["source_sha256"] == job["source_sha256"]
    assert office.closed
    assert office.opened == [
        (job["snapshot_path"], {"ReadOnly": True, "Untitled": False, "WithWindow": False})
    ]
    assert office.export_args[1:7] == (2, 2, 0, 1, 1, 1 if hidden else 0)
    assert office.export_args[11:15] == (True, True, True, False)
    assert Path(job["snapshot_path"]).read_bytes() == before
    assert set(vars(office)) == {
        "closed",
        "opened",
        "export_args",
        "fail",
        "wrong_pages",
        "Presentations",
        "Slides",
        "PageSetup",
    }


@pytest.mark.parametrize("fail,wrong_pages", [(True, False), (False, True)])
def test_native_export_closes_only_its_deck_on_office_or_validation_failure(
    tmp_path: Path, fail: bool, wrong_pages: bool
) -> None:
    office = OfficeFixture(fail=fail, wrong_pages=wrong_pages)
    with pytest.raises((RuntimeError, CoreDomainError)):
        export_native_pdf(_job(tmp_path, True), application=office)
    assert office.closed


@pytest.mark.parametrize(
    "delta, accepted", [(0.025, True), (0.0625, True), (0.063, False), (1, False)]
)
def test_bounded_office_page_size_rounding(tmp_path: Path, delta: float, accepted: bool) -> None:
    office = OfficeFixture()
    office.PageSetup.SlideWidth += delta
    if accepted:
        assert export_native_pdf(_job(tmp_path, True), application=office)["page_count"] == 2
    else:
        with pytest.raises(CoreDomainError, match="geometry"):
            export_native_pdf(_job(tmp_path, True), application=office)
        assert not office.export_args
    assert office.closed


def test_all_hidden_deck_rejects_empty_pdf_without_export(tmp_path: Path) -> None:
    office = OfficeFixture()
    office.Slides.Item = lambda _: SimpleNamespace(SlideShowTransition=SimpleNamespace(Hidden=-1))
    with pytest.raises(CoreDomainError, match="No visible slides"):
        export_native_pdf(_job(tmp_path), application=office)
    assert office.closed and not office.export_args


@pytest.mark.parametrize(
    "change", ["changed_hash", "existing_output", "other_parent", "invalid_hidden"]
)
def test_invalid_pdf_jobs_never_open_office(tmp_path: Path, change: str) -> None:
    job = _job(tmp_path)
    if change == "changed_hash":
        job["source_sha256"] = "a" * 64
    elif change == "existing_output":
        Path(job["output_path"]).write_bytes(b"keep")
    elif change == "other_parent":
        job["output_path"] = str(tmp_path.parent / "output.pdf")
    else:
        job["include_hidden"] = 1
    office = OfficeFixture()
    with pytest.raises(CoreDomainError):
        export_native_pdf(job, application=office)
    assert not office.opened
    if change == "existing_output":
        assert Path(job["output_path"]).read_bytes() == b"keep"


@pytest.mark.parametrize(
    "kind",
    [
        "vba",
        "activex",
        "external_image",
        "external_without_mode",
        "macro_action",
        "macro_type",
        "vba_type_only",
        "ole_type",
    ],
)
def test_active_or_external_dependencies_reject_before_office(tmp_path: Path, kind: str) -> None:
    job = _job(tmp_path)
    source = Path(job["snapshot_path"])
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    if kind == "vba":
        parts["ppt/vbaProject.bin"] = b"macro"
    elif kind == "activex":
        parts["ppt/activeX/activeX1.xml"] = b"active"
    elif kind == "macro_action":
        parts["ppt/slides/extra.xml"] = b'<root action="ppaction://macro?name=Test"/>'
    elif kind in {"vba_type_only", "ole_type"}:
        marker = b"vbaProject" if kind == "vba_type_only" else b"oleObject"
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
            b"</Types>",
            b'<Override PartName="/ppt/unknown.bin" ContentType="application/vnd.ms-office.'
            + marker
            + b'"/></Types>',
        )
        parts["ppt/unknown.bin"] = b"active content"
    elif kind == "macro_type":
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
            b"presentationml.presentation.main+xml", b"presentation.macroEnabled.main+xml"
        )
    else:
        parts["ppt/_rels/external.rels"] = (
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Type="image" TargetMode="External" '
            b'Target="https://example.com/image.png"/></Relationships>'
        )
    if kind == "external_without_mode":
        parts["ppt/_rels/external.rels"] = parts["ppt/_rels/external.rels"].replace(
            b'TargetMode="External" ', b""
        )
    with zipfile.ZipFile(source, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    job["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    office = OfficeFixture()
    with pytest.raises(CoreDomainError):
        export_native_pdf(job, application=office)
    assert not office.opened


def test_pdf_dimensions_and_encryption_are_not_silently_accepted(tmp_path: Path) -> None:
    path = tmp_path / "output.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=720, height=405)
    writer.write(path)
    with pytest.raises(CoreDomainError, match="size differs"):
        validate_pdf(path, 1, 960, 540)
    writer.encrypt("secret")
    writer.write(path)
    with pytest.raises(CoreDomainError, match="protected PDF"):
        validate_pdf(path, 1, 720, 405)


def test_worker_sanitizes_raw_failures_and_bounds_input(monkeypatch: pytest.MonkeyPatch) -> None:
    import presenter_core.presentation.native_pdf as module

    def failed(_: Any) -> Any:
        raise RuntimeError("private content and credentials")

    monkeypatch.setattr(module, "export_native_pdf", failed)
    output = io.StringIO()
    assert run_pdf_worker(io.StringIO("{}"), output) == 1
    receipt = json.loads(output.getvalue())
    assert receipt["error"]["code"] == "PRESENTATION_PDF_FAILED"
    assert "private" not in output.getvalue()
    output = io.StringIO()
    assert run_pdf_worker(io.StringIO("x" * 8193), output) == 1
    assert json.loads(output.getvalue())["error"]["code"] == "INVALID_REQUEST"


def test_safe_hyperlinks_remain_allowed(tmp_path: Path) -> None:
    job = _job(tmp_path)
    path = Path(job["snapshot_path"])
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr(
            "ppt/_rels/test.rels",
            (
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                'relationships/hyperlink" TargetMode="External" '
                'Target="https://example.com"/></Relationships>'
            ),
        )
    validate_native_source(path, hashlib.sha256(path.read_bytes()).hexdigest())


@pytest.mark.parametrize("changed", ["count", "geometry"])
def test_office_cannot_silently_repair_slide_inventory(tmp_path: Path, changed: str) -> None:
    office = OfficeFixture()
    if changed == "count":
        office.Slides.Count = 1
    else:
        office.PageSetup.SlideWidth = 960
    with pytest.raises(CoreDomainError):
        export_native_pdf(_job(tmp_path), application=office)
    assert office.closed and not office.export_args
