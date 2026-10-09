"""Original native selection and explicit persisted inference through actual protocol."""

import json
import uuid

import pytest
import test_rehearsal_coverage as coverage_tests
from test_presentation_planning import call
from test_rehearsal_semantic_analysis import CoverageProvider
from test_selection_actions import save

from presenter_core.presentation.rehearsal_review_inventory import RehearsalReviewInventory

rehearsal = coverage_tests.rehearsal
PREFIX = "presentation.workspace."


def scope_only(scope):
    return {key: scope[key] for key in ("project_id", "document_id", "expected_revision_id")}


def messages(core, scope, **extra):
    return call(core, PREFIX + "list_rehearsal_review_messages", {**scope, **extra})


def point(core, scope):
    result = messages(core, scope)
    item = next(item for item in result["messages"] if item["name"] == "Point 1")
    return {**scope, "slide_id": item["slide_id"], "object_id": item["object_id"]}


def analyze(core, scope, monkeypatch):
    provider = CoverageProvider()
    monkeypatch.setattr(core._providers, "_injected_provider", provider)
    chosen = point(core, scope)
    prepared = call(core, PREFIX + "prepare_rehearsal_coverage_review", chosen)
    result = call(
        core,
        PREFIX + "analyze_rehearsal_coverage",
        {
            **chosen,
            "expected_context_sha256": prepared["context_sha256"],
            "window_index": 0,
        },
    )
    assert provider.call_count == 1
    return result


def test_sessions_page_beyond_ten_and_no_debrief_requirement(rehearsal):
    core, scope = rehearsal
    with core._storage.project_database(scope["project_id"]) as db:
        row = dict(
            db.execute("SELECT * FROM sessions WHERE id=?", (scope["session_id"],)).fetchone()
        )
        fields = list(row)
        for index in range(25):
            later = row | {"id": str(uuid.uuid4()), "started_at": f"2099-10-09T12:{index:02d}:00Z"}
            db.execute(
                f"INSERT INTO sessions ({','.join(fields)}) "
                f"VALUES ({','.join('?' for _ in fields)})",
                [later[key] for key in fields],
            )
        db.commit()
    first = call(core, PREFIX + "list_rehearsal_review_sessions", scope_only(scope))
    second = call(
        core, PREFIX + "list_rehearsal_review_sessions", scope_only(scope) | {"offset": 20}
    )
    assert first["total"] == second["total"] == 26 and first["has_more"] and not second["has_more"]
    assert len(first["sessions"]) == 20 and len(second["sessions"]) == 6
    assert len({item["session_id"] for item in first["sessions"] + second["sessions"]}) == 26
    assert not first["sessions"][0]["debrief_available"]
    chosen = scope | {"session_id": first["sessions"][0]["session_id"]}
    assert messages(core, chosen)["messages"]


def test_original_message_pages_are_complete_read_only_and_no_notes(rehearsal):
    core, scope = rehearsal
    full = messages(core, scope)
    gathered = []
    for offset in range(0, full["total"], 2):
        page = messages(core, scope, offset=offset, limit=2)
        assert page["total"] == full["total"]
        gathered += page["messages"]
    assert gathered == full["messages"]
    assert "private notes excluded" not in json.dumps(full)
    assert all(item["current_span_status"] == "unchanged" for item in gathered)
    assert full["requires_user_selection"]
    with core._storage.project_database(scope["project_id"]) as db:
        assert db.execute("SELECT COUNT(*) FROM provider_runs").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM presentation_rehearsal_analyses").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM presentation_rehearsal_reviews").fetchone()[0] == 0


def test_original_messages_survive_current_reorder_edit_and_delete(rehearsal):
    core, scope = rehearsal
    original = messages(core, scope)
    chosen = next(item for item in original["messages"] if item["name"] == "Point 1")
    coverage = call(core, PREFIX + "rehearsal_coverage", scope)
    edited = scope_only(scope)
    revision = call(
        core,
        PREFIX + "reorder_slides",
        edited | {"ordered_slide_ids": [item["slide_id"] for item in coverage["slides"]][::-1]},
    )["revision_id"]
    edited["expected_revision_id"] = revision
    revision = call(
        core,
        PREFIX + "replace_text_shape",
        edited
        | {
            "slide_id": chosen["slide_id"],
            "object_id": chosen["object_id"],
            "operation_type": "replace_text_shape",
            "expected_text": chosen["text"],
            "replacement_text": "Changed later wording",
        },
    )["revision_id"]
    edited["expected_revision_id"] = revision
    current = messages(core, scope | edited)
    item = next(item for item in current["messages"] if item["object_id"] == chosen["object_id"])
    assert item["text"] == "First point" and item["current_span_status"] == "changed"
    assert item["source_slide_number"] == 1 and item["current_slide_number"] == 3
    inventory = call(
        core,
        PREFIX + "list_text_targets",
        {
            "project_id": scope["project_id"],
            "document_id": scope["document_id"],
            "revision_id": revision,
            "slide_number": 3,
        },
    )
    obj = next(item for item in inventory["objects"] if item["object_id"] == chosen["object_id"])
    plan = core._presentation_edit_plans.prepare(
        {
            **edited,
            "summary": "Remove current point",
            "operations": [
                {
                    "type": "delete_object",
                    "params": {
                        "slide_id": chosen["slide_id"],
                        "object_id": chosen["object_id"],
                        "expected_geometry": obj["geometry"],
                    },
                }
            ],
        }
    )
    save(core, edited, plan)
    removed = next(
        item
        for item in messages(core, scope | edited)["messages"]
        if item["object_id"] == chosen["object_id"]
    )
    assert removed["text"] == "First point" and removed["current_span_status"] == "missing"
    assert not removed["current_text_editable"]


def test_actual_analyze_list_detail_and_save_keep_safe_finding_identity(rehearsal, monkeypatch):
    core, scope = rehearsal
    result = analyze(core, scope, monkeypatch)
    listed = call(core, PREFIX + "list_rehearsal_analyses", scope)
    assert listed["total"] == 1 and listed["analyses"][0]["finding_id"] == result["finding_id"]
    assert listed["analyses"][0]["requires_verified_detail"]
    request = scope_only(scope) | {"finding_id": result["finding_id"]}
    detail = call(core, PREFIX + "get_rehearsal_analysis", request)
    assert detail["finding"]["kind"] == "semantic_coverage" and detail["can_adopt"]
    assert "_semantic_guard" not in json.dumps(detail) and "recording_sha256" not in detail
    assert detail["snapshot"]["hint"] == result["hint"]
    applied = call(
        core,
        PREFIX + "apply_rehearsal_edit",
        request
        | {
            "slide_id": detail["finding"]["slide_id"],
            "edit_kind": "notes",
            "expected_text": "private notes excluded",
            "replacement_text": "User reviewed this point.",
        },
    )
    current = request | {"expected_revision_id": applied["revision_id"]}
    assert (
        call(core, PREFIX + "get_rehearsal_analysis", current)["finding"]["disposition"]
        == "applied"
    )


def test_analysis_summaries_page_and_dismiss_without_implicit_inference(rehearsal, monkeypatch):
    core, scope = rehearsal
    result = analyze(core, scope, monkeypatch)
    result2 = analyze(core, scope, monkeypatch)
    provider = core._providers._injected_provider
    call(
        core,
        PREFIX + "reject_rehearsal_finding",
        scope_only(scope) | {"finding_id": result["finding_id"]},
    )
    page = call(core, PREFIX + "list_rehearsal_analyses", scope | {"limit": 1})
    other = call(core, PREFIX + "list_rehearsal_analyses", scope | {"limit": 1, "offset": 1})
    assert page["total"] == other["total"] == 2 and page["has_more"] and not other["has_more"]
    items = page["analyses"] + other["analyses"]
    assert {item["finding_id"] for item in items} == {result["finding_id"], result2["finding_id"]}
    assert (
        next(item for item in items if item["finding_id"] == result["finding_id"])["disposition"]
        == "rejected"
    )
    assert provider.call_count == 1


@pytest.mark.parametrize(
    "method,params",
    [
        ("list_rehearsal_review_sessions", {"limit": True}),
        ("list_rehearsal_review_sessions", {"offset": -1}),
        ("list_rehearsal_review_messages", {"limit": 21}),
        ("list_rehearsal_review_messages", {"slide_id": str(uuid.uuid4())}),
        ("list_rehearsal_review_messages", {"session_id": str(uuid.uuid4())}),
        ("list_rehearsal_analyses", {"offset": True}),
        ("get_rehearsal_analysis", {"finding_id": "bad"}),
        ("list_rehearsal_review_messages", {"speech": "caller authority"}),
    ],
)
def test_strict_scope_and_page_parameters(rehearsal, method, params):
    core, scope = rehearsal
    given = (
        scope_only(scope)
        if method in {"list_rehearsal_review_sessions", "get_rehearsal_analysis"}
        else scope
    )
    call(core, PREFIX + method, given | params, ok=False)


def test_current_head_guard_and_foreign_project_session(rehearsal):
    core, scope = rehearsal
    foreign = call(core, "project.create", {"name": "Foreign selector"})["project"]["id"]
    call(core, PREFIX + "list_rehearsal_review_messages", scope | {"project_id": foreign}, ok=False)
    with core._storage.project_database(scope["project_id"]) as db:
        db.execute("UPDATE sessions SET status='active' WHERE id=?", (scope["session_id"],))
        db.commit()
    call(core, PREFIX + "list_rehearsal_review_messages", scope, ok=False)
    assert not call(core, PREFIX + "list_rehearsal_review_sessions", scope_only(scope))["sessions"]


def test_tampered_analysis_list_uses_the_same_strict_row_parser(rehearsal, monkeypatch):
    core, scope = rehearsal
    analyze(core, scope, monkeypatch)
    with core._storage.project_database(scope["project_id"]) as db:
        db.execute("UPDATE presentation_rehearsal_analyses SET snapshot_json='{}'")
        db.commit()
    assert (
        call(core, PREFIX + "list_rehearsal_analyses", scope, ok=False)["code"]
        == "REHEARSAL_ANALYSIS_CORRUPT"
    )


def test_final_head_guard_rejects_selector_reply(rehearsal, monkeypatch):
    core, scope = rehearsal
    method = RehearsalReviewInventory._current
    calls = []

    def changed(self, *args):
        calls.append(True)
        if len(calls) > 1:
            from presenter_core.errors import CoreDomainError

            raise CoreDomainError("PRESENTATION_REVISION_CONFLICT", "Injected selector head change")
        return method(self, *args)

    monkeypatch.setattr(RehearsalReviewInventory, "_current", changed)
    assert (
        call(core, PREFIX + "list_rehearsal_review_messages", scope, ok=False)["code"]
        == "PRESENTATION_REVISION_CONFLICT"
    )


def import_recorded_native(core, project, path):
    document = call(
        core, "source.import", {"project_id": project, "kind": "presentation", "path": str(path)}
    )["document"]["id"]
    opened = call(core, PREFIX + "open_source", {"project_id": project, "document_id": document})
    call(
        core,
        PREFIX + "close_source",
        {"project_id": project, "document_id": document, "session_id": opened["session_id"]},
    )
    scope = {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": opened["revision_id"],
    }
    call(core, PREFIX + "select_for_rehearsal", scope)
    session = call(
        core,
        "session.start",
        {
            "project_id": project,
            "mode": "run",
            "current_slide_start": 1,
            "expected_presentation_revision_id": opened["revision_id"],
        },
    )["session"]["id"]
    call(core, "session.stop", {"project_id": project, "session_id": session})
    return scope | {"session_id": session}


def test_long_original_message_requires_explicit_unique_span(rehearsal, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    core, old_scope = rehearsal
    text = "Opening context. " + "Detailed recorded source. " * 100 + "Unique closing decision."
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text = text
    path = tmp_path / "long-message.pptx"
    deck.save(path)
    scope = import_recorded_native(core, old_scope["project_id"], path)
    listed = messages(core, scope)
    item = listed["messages"][0]
    assert listed["total"] == 1 and item["text"] == text
    assert not item["whole_message_selectable"] and len(item["text"]) > 2000
    chosen = scope | {"slide_id": item["slide_id"], "object_id": item["object_id"]}
    assert (
        call(core, PREFIX + "prepare_rehearsal_coverage_review", chosen, ok=False)["code"]
        == "INVALID_REQUEST"
    )
    call(
        core,
        PREFIX + "prepare_rehearsal_coverage_review",
        chosen | {"quote": "Detailed recorded source."},
        ok=False,
    )
    prepared = call(
        core,
        PREFIX + "prepare_rehearsal_coverage_review",
        chosen | {"quote": "Unique closing decision."},
    )
    assert prepared["point"]["quote"] == "Unique closing decision."
    assert prepared["point"]["object_sha256"] == item["object_sha256"]
    assert prepared["total_windows"] == 0


def test_native_chart_and_table_messages_keep_complete_original_representation(rehearsal, tmp_path):
    from test_native_semantic_inventory import fixture

    core, old_scope = rehearsal
    scope = import_recorded_native(
        core, old_scope["project_id"], fixture(tmp_path / "graphics.pptx")
    )
    listed = messages(core, scope)
    by_kind = {item["source_kind"]: item for item in listed["messages"]}
    assert set(by_kind) == {"native_text", "native_chart_cache", "native_table_cells"}
    assert by_kind["native_chart_cache"]["text"] == "Chart: Pilot rates\nRevenue: Q1: 0.2; Q2: 0.4"
    assert by_kind["native_table_cells"]["text"] == "Metric\tValue\nPilot\t20%"
    for kind in ("native_chart_cache", "native_table_cells"):
        item = by_kind[kind]
        assert not item["current_text_editable"] and item["current_span_status"] == "unchanged"
        prepared = call(
            core,
            PREFIX + "prepare_rehearsal_coverage_review",
            scope | {"slide_id": item["slide_id"], "object_id": item["object_id"]},
        )
        assert (
            prepared["point"]["quote"] == item["text"] and prepared["point"]["source_kind"] == kind
        )
    assert listed["excluded_objects"] == 0


def test_selector_rejects_tampered_original_native_snapshot(rehearsal):
    core, scope = rehearsal
    with core._storage.project_database(scope["project_id"]) as db:
        relative = db.execute(
            "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?",
            (scope["expected_revision_id"],),
        ).fetchone()[0]
    path = core._storage.paths.snapshot_path(scope["project_id"], relative, require_exists=True)
    path.write_bytes(path.read_bytes() + b"tampered")
    call(core, PREFIX + "list_rehearsal_review_messages", scope, ok=False)
