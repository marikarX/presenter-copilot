"""Reviewed picture brightness/contrast/grayscale edits."""

import zipfile
from pathlib import Path

import pytest
from lxml import etree
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService
from presenter_core.presentation.fidelity import DRAWING_NS, PRESENTATION_NS

A = f"{{{DRAWING_NS}}}"
P = f"{{{PRESENTATION_NS}}}"
NAME = "Workspace duplicate test picture"


class _Picture:
    def __init__(self, tmp_path: Path) -> None:
        self.core = CoreService(data_root=tmp_path / "data")
        self.project = _project_id(self.core)
        self.document = _import_presentation(self.core, self.project)["id"]
        self.revision = _open(self.core, self.project, self.document)["revision_id"]

    def target(self) -> tuple[dict, str]:
        inventory = _call(
            self.core,
            "targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": self.project,
                "document_id": self.document,
                "revision_id": self.revision,
                "slide_number": 1,
            },
        )
        assert inventory["ok"], inventory
        target = next(o for o in inventory["result"]["objects"] if o["name"] == NAME)
        return target, inventory["result"]["slide_id"]

    def adjust(self, replacement: dict, expected: dict | None = None) -> dict:
        target, slide_id = self.target()
        reply = _call(
            self.core,
            "adjust",
            "presentation.workspace.adjust_picture",
            {
                "project_id": self.project,
                "document_id": self.document,
                "expected_revision_id": self.revision,
                "slide_id": slide_id,
                "object_id": target["object_id"],
                "expected_geometry": target["geometry"],
                "expected_adjustment": expected or target["picture_adjustment"],
                "replacement_adjustment": replacement,
            },
        )
        if reply["ok"]:
            self.revision = reply["result"]["revision_id"]
        return reply

    def blip(self) -> etree._Element:
        with self.core._storage.project_database(self.project) as connection:
            row = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?",
                (self.revision,),
            ).fetchone()
        path = self.core._storage.paths.snapshot_path(
            self.project, row["local_snapshot_path"], require_exists=True
        )
        with zipfile.ZipFile(path) as archive:
            root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        picture = next(
            pic
            for pic in root.iter(f"{P}pic")
            if pic.find(f"{P}nvPicPr/{P}cNvPr").get("name") == NAME
        )
        return picture.find(f"{P}blipFill/{A}blip")


def test_adjustments_apply_read_back_and_reset(tmp_path: Path) -> None:
    picture = _Picture(tmp_path)
    try:
        target = picture.target()[0]
        assert target["picture_adjustment"] == {
            "brightness_percent": 0,
            "contrast_percent": 0,
            "grayscale": False,
        }
        crop_before = target["picture_crop"]
        wanted = {"brightness_percent": 20, "contrast_percent": -30, "grayscale": True}
        reply = picture.adjust(wanted)
        assert reply["ok"] and reply["result"]["changed"], reply
        blip = picture.blip()
        assert [child.tag.removeprefix(A) for child in blip] == ["grayscl", "lum"]
        assert blip.find(f"{A}lum").attrib == {"bright": "20000", "contrast": "-30000"}
        after = picture.target()[0]
        assert after["picture_adjustment"] == wanted
        assert after["picture_crop"] == crop_before
        reset = {"brightness_percent": 0, "contrast_percent": 0, "grayscale": False}
        assert picture.adjust(reset)["ok"]
        assert len(picture.blip()) == 0
        same = picture.adjust(reset)
        assert same["ok"] and same["result"]["changed"] is False
    finally:
        picture.core.close()


def test_stale_expected_adjustment_is_rejected(tmp_path: Path) -> None:
    picture = _Picture(tmp_path)
    try:
        reply = picture.adjust(
            {"brightness_percent": 10},
            expected={"brightness_percent": 5, "contrast_percent": 0, "grayscale": False},
        )
        assert not reply["ok"] and reply["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
    finally:
        picture.core.close()


@pytest.mark.parametrize(
    "replacement",
    [
        {"brightness_percent": 101},
        {"contrast_percent": -101},
        {"brightness_percent": 1.5},
        {"grayscale": "yes"},
        {"saturation": 10},
    ],
)
def test_invalid_adjustments_are_rejected(tmp_path: Path, replacement: dict) -> None:
    picture = _Picture(tmp_path)
    try:
        reply = picture.adjust(replacement)
        assert not reply["ok"] and reply["error"]["code"] == "INVALID_REQUEST"
    finally:
        picture.core.close()


def test_full_bleed_picture_outside_slide_bounds_can_be_adjusted(tmp_path: Path) -> None:
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image = io.BytesIO()
    Image.new("RGB", (32, 24), (200, 80, 40)).save(image, "PNG")
    image.seek(0)
    source = tmp_path / "bleed.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    picture = slide.shapes.add_picture(image, Inches(-0.5), Inches(0), Inches(11), Inches(7.5))
    picture.name = "Bleed"
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        inventory = _call(
            core,
            "targets",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": document,
                "revision_id": revision,
                "slide_number": 1,
            },
        )["result"]
        target = next(o for o in inventory["objects"] if o["name"] == "Bleed")
        assert target["geometry_editable"] is False
        reply = _call(
            core,
            "adjust",
            "presentation.workspace.adjust_picture",
            {
                "project_id": project,
                "document_id": document,
                "expected_revision_id": revision,
                "slide_id": inventory["slide_id"],
                "object_id": target["object_id"],
                "expected_geometry": target["geometry"],
                "expected_adjustment": target["picture_adjustment"],
                "replacement_adjustment": {"brightness_percent": -20, "grayscale": True},
            },
        )
        assert reply["ok"] and reply["result"]["changed"], reply.get("error")
    finally:
        core.close()
