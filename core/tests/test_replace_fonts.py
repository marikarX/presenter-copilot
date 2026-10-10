"""Replace Fonts: list the deck's typefaces and rename one everywhere."""

import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _fixture(path: Path) -> None:
    deck = Presentation()
    for title in ("Intro", "Detail"):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = f"{title} body"
        run.font.name = "Comic Sans MS"
    deck.slides[0].notes_slide.notes_text_frame.text = "Notes"
    deck.save(path)
    # Give the first run PANOSE hints that belong to the old font.
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["ppt/slides/slide1.xml"] = parts["ppt/slides/slide1.xml"].replace(
        b'<a:latin typeface="Comic Sans MS"/>',
        b'<a:latin typeface="Comic Sans MS" panose="030F0702030302020204" charset="0"/>',
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def test_replace_fonts_lists_and_renames_everywhere(tmp_path: Path) -> None:
    _fixture(tmp_path / "fonts.pptx")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, tmp_path / "fonts.pptx")["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def fonts() -> dict:
            return {
                item["name"]: item
                for item in _call(
                    core,
                    "f",
                    "presentation.workspace.list_fonts",
                    {**scope, "expected_revision_id": revision},
                )["result"]["fonts"]
            }

        before = fonts()
        assert before["Comic Sans MS"]["uses"] == 2
        assert before["Comic Sans MS"]["equation_uses"] == 0
        assert before["Calibri"]["theme"] is True
        assert not any(name.startswith("+") for name in before)

        def replace(font: str, uses: int, replacement: str) -> dict:
            return _call(
                core,
                "r",
                "presentation.workspace.replace_fonts",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "font": font,
                    "expected_uses": uses,
                    "replacement_font": replacement,
                },
            )

        stale = replace("Comic Sans MS", 3, "Georgia")
        assert stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        assert replace("Comic Sans MS", 2, "Bad<Font")["error"]["code"] == "INVALID_REQUEST"

        reply = replace("Comic Sans MS", 2, "Georgia")
        assert reply["ok"], reply
        assert reply["result"]["replaced_uses"] == 2 and reply["result"]["parts"] == 2
        revision = reply["result"]["revision_id"]
        after = fonts()
        assert "Comic Sans MS" not in after and after["Georgia"]["uses"] == 2

        with core._storage.project_database(project) as connection:
            snapshot = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, snapshot)) as archive:
            xml = archive.read("ppt/slides/slide1.xml").decode()
        assert '<a:latin typeface="Georgia"/>' in xml and "panose" not in xml

        # Replacing a theme font changes text that follows the theme.
        calibri = after["Calibri"]["uses"]
        themed = replace("Calibri", calibri, "Aptos")
        assert themed["ok"], themed
        revision = themed["result"]["revision_id"]
        theme = _call(
            core,
            "t",
            "presentation.workspace.get_theme_fonts",
            {**scope, "expected_revision_id": revision},
        )["result"]
        assert theme["body_font"] == "Aptos"

        unchanged = replace("Aptos", fonts()["Aptos"]["uses"], "Aptos")
        assert unchanged["ok"] and unchanged["result"]["changed"] is False
    finally:
        core.close()
