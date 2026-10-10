"""Alt text and the decorative flag on existing objects."""

import io
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def test_alt_text_set_decorative_clear_and_guards(tmp_path: Path) -> None:
    source = tmp_path / "alt.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    image = io.BytesIO()
    Image.new("RGB", (32, 24), (40, 120, 200)).save(image, "PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Inches(1), Inches(1), Inches(3), Inches(2)).name = "Chart"
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def target() -> dict:
            inventory = _call(
                core,
                "t",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]
            obj = next(o for o in inventory["objects"] if o["name"] == "Chart")
            return {"slide_id": inventory["slide_id"], **obj}

        def apply(text: str, decorative: bool, expected: tuple | None = None) -> dict:
            nonlocal revision
            obj = target()
            reply = _call(
                core,
                "s",
                "presentation.workspace.set_alt_text",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": obj["slide_id"],
                    "object_id": obj["object_id"],
                    "expected_alt_text": (expected or (obj["alt_text"], obj["decorative"]))[0],
                    "expected_decorative": (expected or (obj["alt_text"], obj["decorative"]))[1],
                    "alt_text": text,
                    "decorative": decorative,
                },
            )
            if reply["ok"]:
                revision = reply["result"]["revision_id"]
            return reply

        initial = target()
        assert initial["decorative"] is False and isinstance(initial["alt_text"], str)
        assert apply("Revenue rose 12% in the pilot", False)["ok"]
        assert (target()["alt_text"], target()["decorative"]) == (
            "Revenue rose 12% in the pilot",
            False,
        )
        reply = apply("", True)
        assert reply["ok"], reply["error"]
        assert (target()["alt_text"], target()["decorative"]) == ("", True)
        with core._storage.project_database(project) as connection:
            path = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, path)) as archive:
            root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        assert root.xpath("//*[local-name()='decorative']/@val") == ["1"]
        assert apply("Back again", False)["ok"]
        assert target()["decorative"] is False
        stale = apply("x", False, expected=("old", False))
        assert not stale["ok"] and stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        bad = apply("a" * 2001, False)
        assert not bad["ok"] and bad["error"]["code"] == "INVALID_REQUEST"
    finally:
        core.close()


def test_benign_nonvisual_extensions_keep_objects_groupable() -> None:
    from presenter_core.presentation.fidelity import has_unsafe_extensions

    def shape(uri: str) -> etree._Element:
        return etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr>'
            f'<p:cNvPr id="2" name="x"><a:extLst><a:ext uri="{uri}"/></a:extLst></p:cNvPr>'
            "</p:nvSpPr></p:sp>"
        )

    assert not has_unsafe_extensions(shape("{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}"))
    assert not has_unsafe_extensions(shape("{C183D7F6-B498-43B3-948B-1728B52AA6E4}"))
    assert has_unsafe_extensions(shape("{00000000-0000-0000-0000-000000000000}"))
