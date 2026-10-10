"""Reusable components: named object selections that insert into any deck in the project."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from test_object_paste import save_plan
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _deck(path: Path, shapes: bool) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if shapes:
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
        )
        badge.name = "Badge"
        badge.text_frame.text = "Confidential"
        slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(4), Inches(1), Inches(1), Inches(1)
        ).name = "Dot"
    deck.save(path)


def test_component_saves_lists_inserts_and_deletes(tmp_path: Path) -> None:
    _deck(tmp_path / "brand.pptx", True)
    _deck(tmp_path / "talk.pptx", False)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        brand = _import_presentation(core, project, tmp_path / "brand.pptx")["id"]
        talk = _import_presentation(core, project, tmp_path / "talk.pptx")["id"]
        brand_revision = _open(core, project, brand)["revision_id"]
        talk_revision = _open(core, project, talk)["revision_id"]
        inventory = _call(
            core,
            "l",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": brand,
                "revision_id": brand_revision,
                "slide_number": 1,
            },
        )["result"]
        ids = [o["object_id"] for o in inventory["objects"] if o["name"] in {"Badge", "Dot"}]
        assert len(ids) == 2
        save = {
            "project_id": project,
            "document_id": brand,
            "revision_id": brand_revision,
            "slide_id": inventory["slide_id"],
            "object_ids": ids,
            "name": "  Confidential badge ",
        }
        saved = _call(core, "s", "presentation.workspace.save_component", save)
        assert saved["ok"], saved
        component = saved["result"]["component"]
        assert component["name"] == "Confidential badge"
        assert component["names"] == ["Badge", "Dot"]

        for bad, code in (
            ({"name": " "}, "INVALID_REQUEST"),
            ({"object_ids": [inventory["slide_id"]]}, "PRESENTATION_EDIT_TARGET_NOT_FOUND"),
            ({"object_ids": [ids[0], ids[0]]}, "INVALID_REQUEST"),
        ):
            reply = _call(core, "x", "presentation.workspace.save_component", {**save, **bad})
            assert not reply["ok"] and reply["error"]["code"] == code, (bad, reply)

        listed = _call(
            core, "l", "presentation.workspace.list_components", {"project_id": project}
        )["result"]["components"]
        assert listed == [component]

        # Inserting reuses paste_objects with the component's immutable source revision.
        talk_slide = _call(
            core,
            "t",
            "presentation.workspace.list_text_targets",
            {
                "project_id": project,
                "document_id": talk,
                "revision_id": talk_revision,
                "slide_number": 1,
            },
        )["result"]["slide_id"]
        pasted = save_plan(
            core,
            {"project_id": project, "document_id": talk, "expected_revision_id": talk_revision},
            {
                "type": "paste_objects",
                "params": {
                    "slide_id": talk_slide,
                    "source_document_id": component["document_id"],
                    "source_revision_id": component["revision_id"],
                    "source_slide_id": component["slide_id"],
                    "object_ids": component["object_ids"],
                },
            },
        )
        assert len(pasted["operations"][0]["created_object_ids"]) == 2
        names = {
            o["name"]
            for o in _call(
                core,
                "t",
                "presentation.workspace.list_text_targets",
                {
                    "project_id": project,
                    "document_id": talk,
                    "revision_id": pasted["revision_id"],
                    "slide_number": 1,
                },
            )["result"]["objects"]
        }
        assert {"Badge copy", "Dot copy"} <= names

        deleted = _call(
            core,
            "d",
            "presentation.workspace.delete_component",
            {"project_id": project, "component_id": component["component_id"]},
        )
        assert deleted["ok"] and deleted["result"]["deleted"]
        again = _call(
            core,
            "d",
            "presentation.workspace.delete_component",
            {"project_id": project, "component_id": component["component_id"]},
        )
        assert again["error"]["code"] == "PRESENTATION_COMPONENT_NOT_FOUND"
        assert (
            _call(core, "l", "presentation.workspace.list_components", {"project_id": project})[
                "result"
            ]["components"]
            == []
        )
    finally:
        core.close()
