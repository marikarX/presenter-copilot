"""Lines and arrows: straight line shapes with arrowheads that move like other shapes."""

import zipfile
from pathlib import Path

from pptx import Presentation
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def test_horizontal_arrow_is_created_styled_and_moved(tmp_path: Path) -> None:
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(tmp_path / "deck.pptx")
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, tmp_path / "deck.pptx")["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def inventory() -> dict:
            return _call(
                core,
                "l",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]

        slide_id = inventory()["slide_id"]
        flat = {"x_emu": 914400, "y_emu": 2286000, "width_emu": 3657600, "height_emu": 0}
        created = _call(
            core,
            "c",
            "presentation.workspace.create_shape",
            {
                **scope,
                "expected_revision_id": revision,
                "slide_id": slide_id,
                "shape_type": "line",
                "geometry": flat,
                "fill_color": "#FFFFFF",
                "line_color": "#204A87",
            },
        )
        assert created["ok"], created.get("error")
        revision = created["result"]["revision_id"]
        line = next(o for o in inventory()["objects"] if o["kind"] == "sp")
        style = line["shape_style"]
        assert line["geometry"] == flat and line["geometry_editable"]
        assert style["shape_geometry"] == "line" and style["fill_color"] is None
        assert (style["line_start"], style["line_end"]) == ("none", "none")

        def restyle(replacement: dict) -> dict:
            current = next(o for o in inventory()["objects"] if o["object_id"] == line["object_id"])
            return _call(
                core,
                "s",
                "presentation.workspace.format_shape_style",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": slide_id,
                    "object_id": line["object_id"],
                    "expected_style_fingerprint": current["shape_style"]["fingerprint"],
                    "replacement_style": replacement,
                },
            )

        arrow = restyle({"line_end": "triangle", "line_start": "oval", "line_dash": "dash"})
        assert arrow["ok"], arrow.get("error")
        revision = arrow["result"]["revision_id"]
        style = next(o for o in inventory()["objects"] if o["object_id"] == line["object_id"])[
            "shape_style"
        ]
        assert (style["line_start"], style["line_end"], style["line_dash"]) == (
            "oval",
            "triangle",
            "dash",
        )
        with core._storage.project_database(project) as connection:
            snapshot = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, snapshot)) as archive:
            xml = archive.read("ppt/slides/slide1.xml").decode()
        # Schema order inside a:ln: fill, dash, headEnd, tailEnd.
        assert xml.index("<a:prstDash") < xml.index("<a:headEnd") < xml.index("<a:tailEnd")

        moved = _call(
            core,
            "t",
            "presentation.workspace.transform_object",
            {
                **scope,
                "expected_revision_id": revision,
                "slide_id": slide_id,
                "object_id": line["object_id"],
                "expected_geometry": flat,
                "replacement_geometry": {**flat, "y_emu": 3200400},
            },
        )
        assert moved["ok"], moved.get("error")
        revision = moved["result"]["revision_id"]

        cleared = restyle({"line_end": "none"})
        assert cleared["ok"], cleared.get("error")
        revision = cleared["result"]["revision_id"]
        style = next(o for o in inventory()["objects"] if o["object_id"] == line["object_id"])[
            "shape_style"
        ]
        assert (style["line_start"], style["line_end"]) == ("oval", "none")
        assert restyle({"line_end": "spiral"})["error"]["code"] == "INVALID_REQUEST"
    finally:
        core.close()
