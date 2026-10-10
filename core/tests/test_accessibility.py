"""Accessibility check: missing alt text and missing slide titles."""

import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def _picture(slide, name: str, alt: str) -> None:
    image = io.BytesIO()
    Image.new("RGB", (8, 8)).save(image, "PNG")
    image.seek(0)
    picture = slide.shapes.add_picture(image, Inches(1), Inches(1), Inches(1), Inches(1))
    picture.name = name
    nonvisual = picture._element.xpath("./p:nvPicPr/p:cNvPr")[0]
    nonvisual.set("descr", alt) if alt else nonvisual.attrib.pop("descr", None)


def test_accessibility_lists_missing_alt_text_and_titles(tmp_path: Path) -> None:
    source = tmp_path / "a11y.pptx"
    deck = Presentation()
    titled = deck.slides.add_slide(deck.slide_layouts[5])
    titled.shapes.title.text = "Revenue"
    _picture(titled, "No alt", "")
    _picture(titled, "Described", "Revenue chart")
    titled.shapes.add_table(2, 2, Inches(3), Inches(3), Inches(3), Inches(1))
    untitled = deck.slides.add_slide(deck.slide_layouts[6])
    _picture(untitled, "Decorative", "")
    decorative = untitled.shapes[-1]._element.xpath("./p:nvPicPr/p:cNvPr")[0]
    decorative.append(
        decorative.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}extLst", {})
    )
    extension = decorative[-1].makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}ext",
        {"uri": "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"},
    )
    decorative[-1].append(extension)
    extension.append(
        extension.makeelement(
            "{http://schemas.microsoft.com/office/drawing/2017/decorative}decorative",
            {"val": "1"},
        )
    )
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        reply = _call(
            core,
            "a",
            "presentation.workspace.check_accessibility",
            {"project_id": project, "document_id": document, "revision_id": revision},
        )
        assert reply["ok"], reply.get("error")
        issues = [
            (item["issue"], item["slide_number"], item["name"])
            for item in reply["result"]["issues"]
        ]
        assert issues == [
            ("missing_alt_text", 1, "No alt"),
            ("missing_slide_title", 2, None),
        ]
    finally:
        core.close()
