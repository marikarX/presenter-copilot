"""Provider table/chart operations preserve native package ownership and review authority."""

from __future__ import annotations

import copy
import json
from pathlib import Path

import pytest
from test_copilot_diagrams import edit_scope, operations
from test_presentation_chart_editing import imported
from test_presentation_edit_planning import call, context
from test_presentation_native_layouts import path, scope
from test_selection_actions import save

from presenter_core.ipc.core import CoreService
from presenter_core.providers.fake import DeterministicFakeReasoningProvider
from presenter_core.providers.models import ReasoningResult
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


class DataProvider(DeterministicFakeReasoningProvider):
    builder = None

    def generate(self, invocation):
        self.request_objects.append(invocation.request)
        snapshot = json.loads(invocation.request.current_slide_summary)[
            "presentation_edit_snapshot"
        ]
        return ReasoningResult(output=self.builder(snapshot))


def geometry(left=0):
    return {"x_emu": 500000 + left, "y_emu": 2200000, "width_emu": 5000000, "height_emu": 3600000}


def chart():
    return {
        "title": "Pilot revenue",
        "chart_type": "column",
        "categories": ["Current", "Pilot"],
        "series": [{"name": "Revenue", "values": [10, 12.25]}],
    }


def composed(snapshot):
    slide = snapshot["slides"][0]["slide_id"]
    cells = [["Metric", "Value"], ["Revenue", "12.25"]]
    data = chart()
    expanded = copy.deepcopy(data)
    expanded["categories"].append("Target")
    expanded["series"][0]["values"].append(15)
    expanded["series"].append({"name": "Cost", "values": [8, 9, 10]})
    entries = [
        (
            "create_authored_slide",
            {
                "slide_id": slide,
                "layout_source_slide_id": slide,
                "title": "Pilot metrics",
                "bullets": [],
                "notes": "Review metrics and evidence.",
                "template": "classic",
                "layout": "statement",
            },
        ),
        (
            "create_table",
            {
                "slide_id": {"created_by": 1},
                "geometry": geometry(),
                "header_row": True,
                "cells": cells,
            },
        ),
        ("create_chart", {"slide_id": {"created_by": 1}, "geometry": geometry(5800000), **data}),
        (
            "replace_table_cells",
            {
                "slide_id": {"created_by": 1},
                "object_id": {"created_by": 2},
                "cells": [["Metric", "Value"], ["Revenue", "12.5"]],
            },
        ),
        (
            "replace_chart_data",
            {
                "slide_id": {"created_by": 1},
                "object_id": {"created_by": 3},
                **{**data, "series": [{"name": "Revenue", "values": [10, 12.5]}]},
            },
        ),
        (
            "reshape_table",
            {
                "slide_id": {"created_by": 1},
                "object_id": {"created_by": 2},
                "cells": [
                    ["Metric", "Value", "Owner"],
                    ["Revenue", "12.5", "Pilot team"],
                    ["Cost", "9", "Finance"],
                ],
                "row_sources": [0, 1, None],
                "column_sources": [0, 1, None],
            },
        ),
        (
            "reshape_chart",
            {
                "slide_id": {"created_by": 1},
                "object_id": {"created_by": 3},
                **expanded,
                "category_sources": [0, 1, None],
                "series_sources": [0, None],
            },
        ),
    ]
    return {"summary": "Review pilot metrics", "operations": operations(entries)}


@pytest.mark.parametrize(
    "failure",
    [
        None,
        "forward",
        "unselected",
        "evidence",
        "private_fp",
        "boolean",
        "nan",
        "ragged",
        "mismatch",
        "mapping",
        "wrong_kind",
    ],
)
def test_composed_native_data_plans_bind_staged_fingerprints_and_roll_back(tmp_path: Path, failure):
    provider = DataProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        baseline = path(core, edit_scope(params)).read_bytes()

        def build(snapshot):
            output = composed(snapshot)
            ops = output["operations"]
            if failure == "forward":
                ops[3]["params"]["object_id"] = {"created_by": 9}
            if failure == "unselected":
                ops[1]["params"]["slide_id"] = snapshot["ordered_slide_ids"][1]
            if failure == "evidence":
                ops[6]["evidence_ids"] = ["invented"]
            if failure == "private_fp":
                ops[3]["params"]["expected_fingerprint"] = "0" * 64
            if failure in {"boolean", "nan"}:
                ops[2]["params"]["series"][0]["values"][0] = (
                    True if failure == "boolean" else float("nan")
                )
            if failure == "ragged":
                ops[1]["params"]["cells"][1].append("extra")
            if failure == "mismatch":
                ops[2]["params"]["series"][0]["values"].append(3)
            if failure == "mapping":
                ops[5]["params"]["row_sources"] = [0, 0, None]
            if failure == "wrong_kind":
                ops[3]["params"]["object_id"] = {"created_by": 3}
            return output

        provider.builder = build
        proposal = call(
            core, "presentation.workspace.propose_edit_plan", params, ok=failure is None
        )
        assert path(core, edit_scope(params)).read_bytes() == baseline
        if failure:
            assert not core._presentation_edit_plans._plans
            assert not list(
                core._storage.paths.safe_sources_directory(params["project_id"]).glob(
                    ".edit-*.pptx"
                )
            )
            return
        reviewed = proposal["review_operations"]
        assert reviewed[3]["previous_data"]["cells"][1][1] == "12.25"
        assert reviewed[5]["previous_data"]["cells"][1][1] == "12.5"
        assert reviewed[4]["previous_data"]["series"][0]["values"] == [10, 12.25]
        assert reviewed[6]["previous_data"]["series"][0]["values"] == [10, 12.5]
        assert (
            proposal["operations"][3]["params"]["expected_fingerprint"]
            != proposal["operations"][5]["params"]["expected_fingerprint"]
        )
        assert (
            proposal["operations"][4]["params"]["expected_fingerprint"]
            != proposal["operations"][6]["params"]["expected_fingerprint"]
        )
        s = edit_scope(params)
        save(core, s, proposal)
        tables = call(core, "presentation.workspace.list_tables", {**scope(s), "slide_number": 2})[
            "items"
        ]
        charts = call(core, "presentation.workspace.list_charts", {**scope(s), "slide_number": 2})[
            "items"
        ]
        assert tables[0]["cells"] == reviewed[5]["params"]["cells"]
        assert (
            charts[0]["data"]["categories"] == ["Current", "Pilot", "Target"]
            and len(charts[0]["data"]["series"]) == 2
        )
        call(core, "presentation.workspace.step_revision", {**s, "direction": "undo"})
        assert path(core, s).read_bytes() == baseline
    finally:
        core.close()


@pytest.mark.parametrize("mutation", [None, "linked", "extension", "stacked", "missing_cache"])
def test_selected_imported_chart_context_binds_only_supported_cached_data(tmp_path: Path, mutation):
    provider = DataProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params, catalog = imported(core, tmp_path, shared=mutation is None, mutation=mutation)
        baseline = path(core, params).read_bytes()
        seen = []

        def build(snapshot):
            seen.append(snapshot)
            assert len(snapshot["slides"]) == 1
            item = snapshot["slides"][0]["charts"][0]
            assert (
                "fingerprint" not in item
                and "chart_part" not in item
                and "workbook_rid" not in item
            )
            if mutation is None:
                assert item["data"]["series"][0]["values"] == [10, 12.25, 15] and item["editable"]
            values = copy.deepcopy(item.get("data", chart()))
            values["series"][0]["values"][0] = 11
            return {
                "summary": "Review imported chart data",
                "operations": operations(
                    [
                        (
                            "replace_chart_data",
                            {
                                "slide_id": snapshot["slides"][0]["slide_id"],
                                "object_id": item["object_id"],
                                **values,
                            },
                        )
                    ]
                ),
            }

        provider.builder = build
        proposal = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {**params, "instruction": "Review the chart data.", "slide_numbers": [1]},
            ok=mutation is None,
        )
        assert seen and path(core, params).read_bytes() == baseline
        if mutation:
            assert not seen[0]["slides"][0]["charts"][0]["editable"]
            assert not core._presentation_edit_plans._plans
        else:
            assert (
                proposal["operations"][0]["params"]["expected_fingerprint"]
                == catalog["items"][0]["fingerprint"]
            )
            assert proposal["review_operations"][0]["previous_data"] == catalog["items"][0]["data"]
            save(core, params, proposal)
            other = call(
                core, "presentation.workspace.list_charts", {**scope(params), "slide_number": 2}
            )["items"][0]
            assert other["data"]["series"][0]["values"] == [10, 12.25, 15]
    finally:
        core.close()


@pytest.mark.parametrize("mutation", [None, "merged", "mixed_runs"])
def test_existing_table_context_preserves_read_only_cells(tmp_path: Path, mutation):
    from pptx import Presentation
    from test_presentation_tables import table_request
    from test_presentation_transactions import context as direct_context

    provider = DataProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params, ops = direct_context(core)
        saved = call(core, "presentation.workspace.create_table", table_request(params, ops))
        params["expected_revision_id"] = saved["revision_id"]
        if mutation:
            deck = Presentation(path(core, params))
            table = next(s.table for s in deck.slides[0].shapes if s.has_table)
            if mutation == "merged":
                table.cell(0, 0).merge(table.cell(0, 1))
            else:
                run = table.cell(1, 1).text_frame.paragraphs[0].add_run()
                run.text = " note"
                run.font.bold = True
            file = tmp_path / "read-only-table.pptx"
            deck.save(file)
            document = call(
                core,
                "source.import",
                {"project_id": params["project_id"], "path": str(file), "kind": "presentation"},
            )["document"]["id"]
            params["document_id"] = document
            params["expected_revision_id"] = core._presentation_edits._current_revision_source(
                params["project_id"], document
            )["revision_id"]
        baseline = path(core, params).read_bytes()
        seen = []

        def build(snapshot):
            assert len(snapshot["slides"]) == 1
            item = snapshot["slides"][0]["tables"][0]
            seen.append(item)
            assert "fingerprint" not in item
            cells = copy.deepcopy(item.get("cells", [["Reviewed"]]))
            cells[-1][-1] = "Confirmed"
            return {
                "summary": "Review table cell",
                "operations": operations(
                    [
                        (
                            "replace_table_cells",
                            {
                                "slide_id": snapshot["slides"][0]["slide_id"],
                                "object_id": item["object_id"],
                                "cells": cells,
                            },
                        )
                    ]
                ),
            }

        provider.builder = build
        proposal = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {**params, "instruction": "Confirm this table value.", "slide_numbers": [1]},
            ok=mutation != "mixed_runs",
        )
        assert seen and path(core, params).read_bytes() == baseline
        if mutation == "mixed_runs":
            assert not seen[0]["editable"] and not core._presentation_edit_plans._plans
        else:
            assert seen[0]["cells"][1][1] == "$10M"
            assert proposal["review_operations"][0]["previous_data"]["cells"] == seen[0]["cells"]
            save(core, params, proposal)
    finally:
        core.close()


def test_hidden_workbook_data_is_not_sent_and_original_package_member_survives(tmp_path: Path):
    import io
    import zipfile

    from lxml import etree

    from presenter_core.presentation.fidelity import REL_NS

    provider = DataProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params, _ = imported(core, tmp_path, shared=True)
        file = tmp_path / "imported-chart.pptx"
        with zipfile.ZipFile(file) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        member = next(
            name for name in parts if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")
        )
        with zipfile.ZipFile(io.BytesIO(parts[member])) as archive:
            workbook = {name: archive.read(name) for name in archive.namelist()}
        ss = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        pkg = "http://schemas.openxmlformats.org/package/2006/relationships"
        xml = etree.fromstring(workbook["xl/workbook.xml"])
        etree.SubElement(
            xml.find(f"{{{ss}}}sheets"),
            f"{{{ss}}}sheet",
            name="Private details",
            sheetId="99",
            state="hidden",
            attrib={f"{{{REL_NS}}}id": "rId99"},
        )
        workbook["xl/workbook.xml"] = etree.tostring(xml)
        rel = etree.fromstring(workbook["xl/_rels/workbook.xml.rels"])
        etree.SubElement(
            rel,
            f"{{{pkg}}}Relationship",
            Id="rId99",
            Type=REL_NS + "/worksheet",
            Target="worksheets/sheet99.xml",
        )
        workbook["xl/_rels/workbook.xml.rels"] = etree.tostring(rel)
        types = etree.fromstring(workbook["[Content_Types].xml"])
        etree.SubElement(
            types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
            PartName="/xl/worksheets/sheet99.xml",
            ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
        )
        workbook["[Content_Types].xml"] = etree.tostring(types)
        sentinel = "Hidden workbook context must remain private"
        workbook["xl/worksheets/sheet99.xml"] = (
            f'<worksheet xmlns="{ss}"><sheetData><row r="1">'
            f'<c r="A1" t="inlineStr"><is><t>{sentinel}</t></is></c>'
            "</row></sheetData></worksheet>"
        ).encode()
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            for name, payload in workbook.items():
                archive.writestr(name, payload)
        parts[member] = buffer.getvalue()
        with zipfile.ZipFile(file, "w") as archive:
            for name, payload in parts.items():
                archive.writestr(name, payload)
        doc = call(
            core,
            "source.import",
            {"project_id": params["project_id"], "path": str(file), "kind": "presentation"},
        )["document"]["id"]
        params["document_id"] = doc
        params["expected_revision_id"] = core._presentation_edits._current_revision_source(
            params["project_id"], doc
        )["revision_id"]

        def build(snapshot):
            item = snapshot["slides"][0]["charts"][0]
            data = copy.deepcopy(item["data"])
            data["series"][0]["values"][0] = 11
            return {
                "summary": "Review cached chart numbers",
                "operations": operations(
                    [
                        (
                            "replace_chart_data",
                            {
                                "slide_id": snapshot["slides"][0]["slide_id"],
                                "object_id": item["object_id"],
                                **data,
                            },
                        )
                    ]
                ),
            }

        provider.builder = build
        proposal = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {**params, "instruction": "Review the cached chart.", "slide_numbers": [1]},
        )
        assert sentinel not in provider.request_objects[-1].serialized_input()
        save(core, params, proposal)
        with zipfile.ZipFile(path(core, params)) as archive:
            assert archive.read(member) == parts[member]
    finally:
        core.close()
