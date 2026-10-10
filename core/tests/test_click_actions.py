"""Click actions on shapes: web links, slide jumps and next/previous navigation."""

import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def test_click_action_cycles_and_cleans_relationships(tmp_path: Path) -> None:
    source = tmp_path / "nav.pptx"
    deck = Presentation()
    for number in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        if number == 0:
            slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1)
            ).name = "Button"
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def inventory(number: int) -> dict:
            return _call(
                core,
                "l",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": number},
            )["result"]

        first, third = inventory(1), inventory(3)
        button = next(o for o in first["objects"] if o["name"] == "Button")
        assert button["click_action"] == {"kind": "none", "url": None, "slide_id": None}

        def set_action(action: dict, expected: dict | None = None) -> dict:
            nonlocal revision
            current = next(o for o in inventory(1)["objects"] if o["name"] == "Button")
            reply = _call(
                core,
                "c",
                "presentation.workspace.set_click_action",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": first["slide_id"],
                    "object_id": button["object_id"],
                    "expected_click_action": expected or current["click_action"],
                    "click_action": action,
                },
            )
            if reply["ok"]:
                revision = reply["result"]["revision_id"]
            return reply

        def rels() -> str:
            with core._storage.project_database(project) as connection:
                path = connection.execute(
                    "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?",
                    (revision,),
                ).fetchone()["local_snapshot_path"]
            with zipfile.ZipFile(core._storage.paths.snapshot_path(project, path)) as archive:
                return archive.read("ppt/slides/_rels/slide1.xml.rels").decode()

        steps = [
            {"kind": "slide", "url": None, "slide_id": third["slide_id"]},
            {"kind": "url", "url": "https://example.org/agenda", "slide_id": None},
            {"kind": "next", "url": None, "slide_id": None},
            {"kind": "none", "url": None, "slide_id": None},
        ]
        for action in steps:
            assert set_action(action)["ok"]
            current = next(o for o in inventory(1)["objects"] if o["name"] == "Button")
            assert current["click_action"] == action
            text = rels()
            assert ("slide3.xml" in text) == (action["kind"] == "slide")
            assert ("example.org" in text) == (action["kind"] == "url")
        stale = set_action(steps[1], expected=steps[0])
        assert not stale["ok"] and stale["error"]["code"] == "PRESENTATION_EDIT_STALE_TARGET"
        bad = set_action({"kind": "url", "url": "javascript:alert(1)", "slide_id": None})
        assert not bad["ok"] and bad["error"]["code"] == "INVALID_REQUEST"
    finally:
        core.close()
