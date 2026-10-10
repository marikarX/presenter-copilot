"""Provider proposals keep scope, evidence, staged dependencies and review authority."""

from __future__ import annotations

import copy
import json
import uuid
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest

from presenter_core.ipc.core import CoreService
from presenter_core.providers.edit_plan_schema import PARAMETER_SCHEMAS
from presenter_core.providers.fake import DeterministicFakeReasoningProvider
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def call(core: CoreService, method: str, params: dict[str, Any], *, ok: bool = True) -> Any:
    response = core.handle_message(
        {
            "protocol_version": 1,
            "type": "request",
            "request_id": str(uuid.uuid4()),
            "method": method,
            "params": params,
        }
    )
    assert response["ok"] is ok, response
    return response["result" if ok else "error"]


def context(core: CoreService) -> dict[str, Any]:
    project = call(
        core, "project.create", {"name": "Coordinated decisions", "privacy_mode": "local_only"}
    )["project"]["id"]
    document = call(
        core,
        "presentation.workspace.create_from_outline",
        {
            "project_id": project,
            "title": "Pilot decision",
            "slides": [
                {
                    "title": "Approve the pilot",
                    "bullets": ["Review the decision owner"],
                    "notes": "Explain the tradeoffs",
                },
                {
                    "title": "Supporting evidence",
                    "bullets": ["Review the source"],
                    "notes": "Confirm the evidence",
                },
            ],
        },
    )["document"]["id"]
    revision = core._presentation_edits._current_revision_source(project, document)["revision_id"]
    return {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": revision,
        "instruction": "Clarify the pilot decision and add a speaking reminder.",
        "slide_numbers": [1],
    }


class PlanProvider(DeterministicFakeReasoningProvider):
    modifier: Any = None

    def generate(self, invocation: Any) -> Any:
        result = super().generate(invocation)
        output = copy.deepcopy(result.output)
        if self.modifier:
            self.modifier(output, invocation)
        return replace(result, output=output)


@pytest.mark.parametrize("failure", [None, "forward", "wrong_kind", "cross_slide", "deleted"])
def test_created_targets_are_bound_in_order_and_reviewed_atomically(
    tmp_path: Path, failure: str | None
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)

        def modify(output: dict[str, Any], _invocation: Any) -> None:
            slide = output["operations"][0]["params"]["slide_id"]
            style = dict.fromkeys(
                PARAMETER_SCHEMAS["format_text_shape"]["properties"]["replacement_style"][
                    "properties"
                ]
            )
            style.update(bold=True, font_size_pt=28, font_color="#123456")
            entries = [
                ("create_slide", {"slide_id": slide, "layout_source_slide_id": slide}),
                (
                    "create_text_box",
                    {"slide_id": {"created_by": 1}, "text": "New coordinated callout"},
                ),
                (
                    "format_text_shape",
                    {
                        "slide_id": {"created_by": 1},
                        "object_id": {"created_by": 2},
                        "replacement_style": style,
                    },
                ),
                (
                    "transform_object",
                    {
                        "slide_id": {"created_by": 1},
                        "object_id": {"created_by": 2},
                        "replacement_geometry": {
                            "x_emu": 100000,
                            "y_emu": 100000,
                            "width_emu": 2000000,
                            "height_emu": 500000,
                        },
                    },
                ),
            ]
            entries.append(
                (
                    "change_object_layer",
                    {
                        "slide_id": {"created_by": 1},
                        "object_id": {"created_by": 2},
                        "direction": "front",
                    },
                )
            )
            if failure == "forward":
                entries[1][1]["slide_id"] = {"created_by": 3}
            elif failure == "wrong_kind":
                entries[2][1]["object_id"] = {"created_by": 1}
            elif failure == "cross_slide":
                entries[2][1]["slide_id"] = slide
            elif failure == "deleted":
                entries.insert(2, ("delete_slide", {"slide_id": {"created_by": 1}}))
            output["operations"] = [
                {
                    "type": kind,
                    "params": values,
                    "rationale": "Review the composed callout",
                    "evidence_ids": [],
                }
                for kind, values in entries
            ]

        provider.modifier = modify
        proposal = call(
            core, "presentation.workspace.propose_edit_plan", params, ok=failure is None
        )
        current = core._presentation_edits._current_revision_source(
            params["project_id"], params["document_id"]
        )
        assert current["revision_id"] == params["expected_revision_id"]
        if failure is not None:
            assert core._presentation_edit_plans._plans == {}
            assert not list(
                core._storage.paths.safe_sources_directory(params["project_id"]).glob(
                    ".edit-*.pptx"
                )
            )
            return
        operations = proposal["operations"]
        assert operations[2]["params"]["expected_text"] == "New coordinated callout"
        assert isinstance(operations[2]["params"]["object_id"], str)
        assert operations[1]["params"]["slide_id"] == operations[3]["params"]["slide_id"]
        assert (
            operations[3]["params"]["object_id"] in operations[4]["params"]["expected_object_ids"]
        )
        scope = {key: params[key] for key in ("project_id", "document_id")}
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        assert result["sha256"] == proposal["sha256"]
        assert result["revision_number"] == 2
        targets = call(
            core,
            "presentation.workspace.list_text_targets",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "revision_id": result["revision_id"],
                "slide_number": 2,
            },
        )
        target = next(
            item for item in targets["targets"] if item["current_text"] == "New coordinated callout"
        )
        assert target["current_text"] == "New coordinated callout"
        inspected = core._presentation_edit_planning._read(
            core._presentation_edits,
            params["project_id"],
            params["document_id"],
            result["revision_id"],
        )
        obj = next(
            item for item in inspected.slides[1].objects if item.text == "New coordinated callout"
        )
        assert obj.text_style is not None and obj.text_style.bold is True
        assert obj.geometry is not None and obj.geometry.x_emu == 100000
        call(
            core,
            "presentation.workspace.step_revision",
            {**scope, "expected_revision_id": result["revision_id"], "direction": "undo"},
        )
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
    finally:
        core.close()


def test_provider_plan_is_review_only_and_keeps_successful_run_provenance(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        assert proposal["requires_user_review"] is True
        assert proposal["baseline_revision_id"] == params["expected_revision_id"]
        assert len(proposal["operations"]) == 2
        source = core._presentation_edits._current_revision_source(
            params["project_id"], params["document_id"]
        )
        assert source["revision_id"] == params["expected_revision_id"]
        request = provider.request_objects[-1]
        assert request.question == params["instruction"]
        snapshot = json.loads(request.current_slide_summary or "{}")["presentation_edit_snapshot"]
        assert len(snapshot["slides"]) == 1
        assert "part_name" not in snapshot["slides"][0]
        assert request.presentation_revision_ids == (params["expected_revision_id"],)
        assert request.source_ids == (params["document_id"],)
        with core._storage.project_database(params["project_id"]) as connection:
            run = connection.execute(
                "SELECT * FROM provider_runs WHERE id = ?", (proposal["provider_run_id"],)
            ).fetchone()
            assert run["status"] == "success"
            manifest = json.loads(run["context_manifest_json"])
            assert manifest["presentation_revision_ids"] == [params["expected_revision_id"]]
            assert params["instruction"] not in json.dumps(manifest)
            assert "Approve the pilot" not in json.dumps(manifest)
        scope = {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": proposal["plan_id"],
        }
        call(core, "presentation.workspace.open_edit_plan_preview", scope)
        applied = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {**scope, "confirm": True, "reviewed_sha256": proposal["sha256"]},
        )
        assert applied["sha256"] == proposal["sha256"]
        history = call(
            core,
            "presentation.workspace.list_revisions",
            {"project_id": params["project_id"], "document_id": params["document_id"]},
        )
        metadata = history["revisions"][0]["edit_transaction"]["operations"]
        assert all(item["provider_run_id"] == proposal["provider_run_id"] for item in metadata)
        assert "Approve the pilot" not in json.dumps(metadata)
        assert core._presentation_edit_plans._plans == {}
    finally:
        core.close()


@pytest.mark.parametrize(
    "failure",
    [
        "outside_slide",
        "outside_object",
        "invented_evidence",
        "injected_scope",
        "unknown_type",
        "style_empty",
        "huge_geometry",
        "reorder_selection",
        "malformed_notes",
        "invalid_xml",
    ],
)
def test_invalid_provider_proposal_leaves_no_private_plan_or_revision(
    tmp_path: Path, failure: str
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        second = core._presentation_workspace_sources.list_text_targets(
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "revision_id": params["expected_revision_id"],
                "slide_number": 2,
            }
        )

        def modify(output: dict[str, Any], _invocation: Any) -> None:
            operation = output["operations"][0]
            if failure == "outside_slide":
                operation["params"]["slide_id"] = second["slide_id"]
                operation["params"]["object_id"] = second["targets"][0]["object_id"]
            elif failure == "outside_object":
                operation["params"]["object_id"] = str(uuid.uuid4())
            elif failure == "invented_evidence":
                operation["evidence_ids"] = ["invented-evidence"]
            elif failure == "injected_scope":
                operation["params"]["project_id"] = str(uuid.uuid4())
            elif failure == "unknown_type":
                operation["type"] = "run_code"
            elif failure == "style_empty":
                operation["type"] = "format_text_shape"
                operation["params"].pop("replacement_text")
                operation["params"]["replacement_style"] = dict.fromkeys(
                    PARAMETER_SCHEMAS["format_text_shape"]["properties"]["replacement_style"][
                        "properties"
                    ]
                )
            elif failure == "huge_geometry":
                operation["type"] = "transform_object"
                operation["params"].pop("replacement_text")
                operation["params"]["replacement_geometry"] = {
                    "x_emu": 0,
                    "y_emu": 0,
                    "width_emu": 2_000_000_000,
                    "height_emu": 2_000_000_000,
                }
            elif failure == "reorder_selection":
                operation["type"] = "reorder_slides"
                operation["params"] = {
                    "ordered_slide_ids": [second["slide_id"], operation["params"]["slide_id"]]
                }
            elif failure == "malformed_notes":
                output["operations"][1]["params"]["replacement_text"] = None
            else:
                operation["params"]["replacement_text"] = "Invalid\x00XML"

        provider.modifier = modify
        call(core, "presentation.workspace.propose_edit_plan", params, ok=False)
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
        assert core._presentation_edit_plans._plans == {}
        directory = core._storage.paths.safe_sources_directory(params["project_id"])
        assert list(directory.glob(".edit-*.pptx")) == []
        with core._storage.project_database(params["project_id"]) as connection:
            assert (
                connection.execute(
                    "SELECT status FROM provider_runs ORDER BY started_at DESC"
                ).fetchone()[0]
                == "error"
            )
    finally:
        core.close()


def test_staged_text_then_format_and_multi_slide_structure_resolve_current_preconditions(
    tmp_path: Path,
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        params["slide_numbers"] = [1, 2]

        def modify(output: dict[str, Any], invocation: Any) -> None:
            snapshot = json.loads(invocation.request.current_slide_summary)[
                "presentation_edit_snapshot"
            ]
            target = output["operations"][0]["params"]
            style = dict.fromkeys(
                PARAMETER_SCHEMAS["format_text_shape"]["properties"]["replacement_style"][
                    "properties"
                ]
            )
            style["bold"] = True
            output["operations"].insert(
                1,
                {
                    "type": "format_text_shape",
                    "params": {
                        "slide_id": target["slide_id"],
                        "object_id": target["object_id"],
                        "replacement_style": style,
                    },
                    "rationale": "Emphasize the coordinated title.",
                    "evidence_ids": [],
                },
            )
            output["operations"].append(
                {
                    "type": "reorder_slides",
                    "params": {"ordered_slide_ids": list(reversed(snapshot["ordered_slide_ids"]))},
                    "rationale": "Review this proposed narrative order.",
                    "evidence_ids": [],
                }
            )

        provider.modifier = modify
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        assert (
            proposal["operations"][1]["params"]["expected_text"]
            == proposal["operations"][0]["params"]["replacement_text"]
        )
        scope = {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": proposal["plan_id"],
        }
        call(core, "presentation.workspace.open_edit_plan_preview", scope)
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {**scope, "confirm": True, "reviewed_sha256": proposal["sha256"]},
        )
        assert result["revision_number"] == 2
        assert len(result["operations"]) == 5
    finally:
        core.close()


@pytest.mark.parametrize("phase", ["generation", "finalization"])
def test_stale_during_generation_or_finalization_discards_the_proposal(
    tmp_path: Path, phase: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        inventory = core._presentation_workspace_sources.list_text_targets(
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "revision_id": params["expected_revision_id"],
                "slide_number": 1,
            }
        )

        def edit() -> None:
            core._presentation_edits.replace_speaker_notes(
                {
                    "project_id": params["project_id"],
                    "document_id": params["document_id"],
                    "expected_revision_id": params["expected_revision_id"],
                    "slide_id": inventory["slide_id"],
                    "expected_text": "Explain the tradeoffs",
                    "replacement_text": "Changed concurrently",
                }
            )

        if phase == "generation":
            provider.modifier = lambda _output, _invocation: edit()
        else:
            original = core._provider_execution.execute

            def execute(**kwargs: Any) -> Any:
                result = original(**kwargs)
                edit()
                return result

            monkeypatch.setattr(core._provider_execution, "execute", execute)
        error = call(core, "presentation.workspace.propose_edit_plan", params, ok=False)
        assert error["code"] == "PRESENTATION_REVISION_CONFLICT"
        assert core._presentation_edit_plans._plans == {}
        assert (
            list(
                core._storage.paths.safe_sources_directory(params["project_id"]).glob(
                    ".edit-*.pptx"
                )
            )
            == []
        )
    finally:
        core.close()


@pytest.mark.parametrize("acknowledged", [False, True])
def test_remote_provider_is_blocked_before_receiving_the_snapshot_in_local_only(
    tmp_path: Path, acknowledged: bool
) -> None:
    provider = PlanProvider(locality="remote")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        if acknowledged:
            call(core, "project.acknowledge_remote_reasoning", {"project_id": params["project_id"]})
        error = call(core, "presentation.workspace.propose_edit_plan", params, ok=False)
        assert error["code"] == "PRIVACY_LOCAL_ONLY_REMOTE_BLOCKED"
        assert provider.request_objects == []
    finally:
        core.close()


def test_supplied_evidence_is_reviewable_and_retained_as_content_free_ids(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)

        def cite(output: dict[str, Any], invocation: Any) -> None:
            assert invocation.request.evidence
            output["operations"][0]["evidence_ids"] = [
                invocation.request.evidence[0]["evidence_id"]
            ]

        provider.modifier = cite
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        ids = proposal["review_operations"][0]["evidence_ids"]
        assert ids[0] in {item["evidence_id"] for item in proposal["evidence"]}
        scope = {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": proposal["plan_id"],
        }
        call(core, "presentation.workspace.open_edit_plan_preview", scope)
        call(
            core,
            "presentation.workspace.apply_edit_plan",
            {**scope, "confirm": True, "reviewed_sha256": proposal["sha256"]},
        )
        with core._storage.project_database(params["project_id"]) as connection:
            metadata = json.loads(
                connection.execute(
                    "SELECT operation_metadata_json FROM presentation_edit_transactions"
                ).fetchone()[0]
            )
        assert metadata[0]["evidence_ids"] == ids
    finally:
        core.close()


def test_provider_can_propose_an_authored_slide_with_private_review(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)

        def author(output: dict[str, Any], invocation: Any) -> None:
            slide_id = output["operations"][0]["params"]["slide_id"]
            output["operations"] = [
                {
                    "type": "create_authored_slide",
                    "params": {
                        "slide_id": slide_id,
                        "layout_source_slide_id": slide_id,
                        "title": "An authored decision",
                        "bullets": ["Ask for the pilot"],
                        "notes": "Review the rationale with the team",
                        "template": "paper",
                        "layout": "statement",
                    },
                    "rationale": "Add the requested decision slide",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = author
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        assert proposal["review_operations"][0]["type"] == "create_authored_slide"
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
        scope = {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": proposal["plan_id"],
        }
        preview = call(core, "presentation.workspace.open_edit_plan_preview", scope)
        assert preview["fidelity_report"]["slide_count"] == 3
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {**scope, "confirm": True, "reviewed_sha256": proposal["sha256"]},
        )
        assert result["revision_number"] == 2
        with core._storage.project_database(params["project_id"]) as connection:
            metadata = connection.execute(
                "SELECT operation_metadata_json FROM presentation_edit_transactions"
            ).fetchone()[0]
        assert "create_authored_slide" in metadata
        assert "An authored decision" not in metadata
        assert "Review the rationale" not in metadata
    finally:
        core.close()


def test_missing_successful_run_prevents_later_apply(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        scope = {
            "project_id": params["project_id"],
            "document_id": params["document_id"],
            "plan_id": proposal["plan_id"],
        }
        call(core, "presentation.workspace.open_edit_plan_preview", scope)
        with core._storage.project_database(params["project_id"]) as connection:
            connection.execute(
                "DELETE FROM provider_runs WHERE id = ?", (proposal["provider_run_id"],)
            )
            connection.commit()
        error = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {**scope, "confirm": True, "reviewed_sha256": proposal["sha256"]},
            ok=False,
        )
        assert error["code"] == "PRESENTATION_EDIT_PLAN_PROVENANCE_INVALID"
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
        call(core, "presentation.workspace.reject_edit_plan", scope)
    finally:
        core.close()


def test_complete_edit_snapshot_is_never_silently_truncated_to_fit(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        document = call(
            core,
            "presentation.workspace.create_from_outline",
            {
                "project_id": params["project_id"],
                "title": "Large complete draft",
                "slides": [
                    {
                        "title": f"Slide {number}",
                        "bullets": ["A supported sentence " * 8],
                        "notes": "A complete speaking explanation. " * 100,
                    }
                    for number in range(20)
                ],
            },
        )["document"]["id"]
        revision = core._presentation_edits._current_revision_source(
            params["project_id"], document
        )["revision_id"]
        error = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {
                "project_id": params["project_id"],
                "document_id": document,
                "expected_revision_id": revision,
                "instruction": "Review the whole narrative.",
            },
            ok=False,
        )
        assert error["code"] == "PRESENTATION_PLAN_CONTEXT_TOO_LARGE"
        assert provider.request_objects == []
        assert core._presentation_edit_plans._plans == {}
    finally:
        core.close()


@pytest.mark.parametrize("failure", [None, "unattached", "duplicate", "stale", "project_library"])
def test_attached_picture_is_core_bound_without_provider_image_bytes(
    tmp_path: Path, failure: str | None
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {
                "project_id": params["project_id"],
                "document_id": params["document_id"],
                "revision_id": params["expected_revision_id"],
                "slide_number": 1,
            },
        )
        encoded = (
            "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEklEQVR4nGMU"
            "VDJmYGBgYgADAATQAGoNqSJLAAAAAElFTkSuQmCC"
        )
        picture = call(
            core,
            "presentation.workspace.create_picture",
            {key: params[key] for key in ("project_id", "document_id", "expected_revision_id")}
            | {
                "slide_id": inventory["slide_id"],
                "geometry": {
                    "x_emu": 10000,
                    "y_emu": 10000,
                    "width_emu": 100000,
                    "height_emu": 100000,
                },
                "alt_text": "Source",
                "image_base64": encoded,
            },
        )
        old_revision = params["expected_revision_id"]
        params["expected_revision_id"] = picture["revision_id"]
        assets = call(
            core,
            "presentation.workspace.list_picture_assets",
            {key: params[key] for key in ("project_id", "document_id", "expected_revision_id")},
        )["assets"]
        assert len(assets) == 1
        asset = assets[0]
        project_library = failure == "project_library"
        if project_library:
            imported = call(
                core,
                "presentation.workspace.import_project_picture",
                {
                    "project_id": params["project_id"],
                    "name": "Reusable visual.png",
                    "image_base64": encoded,
                },
            )
            asset = imported["asset"]
            failure = None

        params["asset_ids"] = [asset["asset_id"]]
        params["instruction"] = "Insert the attached picture at the lower right."
        captured = []

        def modify(output: dict[str, Any], invocation: Any) -> None:
            payload = invocation.request.current_slide_summary
            assert (
                "image_base64" not in payload
                and "thumbnail_base64" not in payload
                and encoded not in payload
            )
            snapshot = json.loads(payload)["presentation_edit_snapshot"]
            assert snapshot["assets"][0]["asset_id"] == asset["asset_id"]
            captured.append(snapshot)
            if failure == "unattached":
                output["operations"][0]["params"]["asset_id"] = "0" * 64

        provider.modifier = modify
        if failure == "duplicate":
            params["asset_ids"] *= 2
        if failure == "stale":
            params["expected_revision_id"] = old_revision
        proposal = call(
            core, "presentation.workspace.propose_edit_plan", params, ok=failure is None
        )
        if failure:
            assert not core._presentation_edit_plans._plans
            assert (
                core._presentation_edits._current_revision_source(
                    params["project_id"], params["document_id"]
                )["revision_id"]
                == picture["revision_id"]
            )
            if failure in {"duplicate", "stale"}:
                assert not captured
            return
        assert captured
        if project_library:
            call(
                core,
                "presentation.workspace.delete_project_picture",
                {
                    "project_id": params["project_id"],
                    "asset_id": asset["asset_id"],
                    "expected_sha256": asset["sha256"],
                    "confirm": True,
                },
            )
        assert "image_base64" not in json.dumps(proposal)
        assert proposal["operations"][0]["params"]["asset_id"] == asset["asset_id"]
        scope = {key: params[key] for key in ("project_id", "document_id")}
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        assert result["sha256"] == proposal["sha256"] and result["revision_number"] == 3
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": result["revision_id"], "slide_number": 1},
        )
        assert sum(obj["kind"] == "pic" for obj in inventory["objects"]) == 2
        with core._storage.project_database(params["project_id"]) as connection:
            audit = connection.execute(
                "SELECT operation_metadata_json FROM presentation_edit_transactions"
            ).fetchone()[0]
            assert encoded not in audit and "image_base64" not in audit
        restored = call(
            core,
            "presentation.workspace.step_revision",
            {**scope, "expected_revision_id": result["revision_id"], "direction": "undo"},
        )
        assert restored["revision_id"] == picture["revision_id"]
    finally:
        core.close()


def test_provider_picture_adjustment_binds_current_adjustment_and_applies(
    tmp_path: Path,
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        scope = {key: params[key] for key in ("project_id", "document_id")}
        slide_id = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": params["expected_revision_id"], "slide_number": 1},
        )["slide_id"]
        picture = call(
            core,
            "presentation.workspace.create_picture",
            {**scope, "expected_revision_id": params["expected_revision_id"]}
            | {
                "slide_id": slide_id,
                "geometry": {
                    "x_emu": 10000,
                    "y_emu": 10000,
                    "width_emu": 100000,
                    "height_emu": 100000,
                },
                "alt_text": "Source",
                "image_base64": (
                    "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEklEQVR4nGMU"
                    "VDJmYGBgYgADAATQAGoNqSJLAAAAAElFTkSuQmCC"
                ),
            },
        )
        params["expected_revision_id"] = picture["revision_id"]

        def objects(revision: str) -> list[dict[str, Any]]:
            return call(
                core,
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["objects"]

        target = next(obj for obj in objects(picture["revision_id"]) if obj["kind"] == "pic")
        wanted = {"brightness_percent": 15, "contrast_percent": 25, "grayscale": True}

        def modify(output: dict[str, Any], _invocation: Any) -> None:
            output["operations"] = [
                {
                    "type": "adjust_picture",
                    "params": {
                        "slide_id": slide_id,
                        "object_id": target["object_id"],
                        "replacement_adjustment": wanted,
                    },
                    "rationale": "Mute the picture behind the text",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = modify
        params["instruction"] = "Make the picture grayscale and a little brighter."
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        bound = proposal["operations"][0]["params"]
        assert bound["expected_adjustment"] == target["picture_adjustment"]
        assert bound["expected_geometry"] == target["geometry"]
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        adjusted = next(
            obj for obj in objects(result["revision_id"]) if obj["object_id"] == target["object_id"]
        )
        assert adjusted["picture_adjustment"] == wanted
    finally:
        core.close()


def test_provider_rotation_binds_current_rotation_and_applies(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        scope = {key: params[key] for key in ("project_id", "document_id")}
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": params["expected_revision_id"], "slide_number": 1},
        )
        target = next(obj for obj in inventory["objects"] if obj["rotation"] is not None)
        wanted = {"rotation_degrees": 15.5, "flip_horizontal": False, "flip_vertical": True}

        def modify(output: dict[str, Any], _invocation: Any) -> None:
            output["operations"] = [
                {
                    "type": "rotate_object",
                    "params": {
                        "slide_id": inventory["slide_id"],
                        "object_id": target["object_id"],
                        "replacement_rotation": wanted,
                    },
                    "rationale": "Tilt the callout",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = modify
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        bound = proposal["operations"][0]["params"]
        assert bound["expected_rotation"] == target["rotation"]
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        rotated = next(
            obj
            for obj in call(
                core,
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": result["revision_id"], "slide_number": 1},
            )["objects"]
            if obj["object_id"] == target["object_id"]
        )
        assert rotated["rotation"] == wanted
        assert rotated["geometry"] == target["geometry"]
    finally:
        core.close()


def test_provider_slide_background_binds_current_background_and_applies(
    tmp_path: Path,
) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        scope = {key: params[key] for key in ("project_id", "document_id")}
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": params["expected_revision_id"], "slide_number": 1},
        )

        def modify(output: dict[str, Any], _invocation: Any) -> None:
            output["operations"] = [
                {
                    "type": "set_slide_background",
                    "params": {"slide_id": inventory["slide_id"], "replacement_color": "#0B1F33"},
                    "rationale": "Darken the title slide",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = modify
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        bound = proposal["operations"][0]["params"]
        assert bound["expected_background"] == {
            "state": inventory["background"]["state"],
            "color": inventory["background"]["color"],
        }
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        after = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": result["revision_id"], "slide_number": 1},
        )
        assert after["background"]["color"] == "#0B1F33"
    finally:
        core.close()


def test_provider_text_link_binds_current_link_and_applies(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        scope = {key: params[key] for key in ("project_id", "document_id")}
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": params["expected_revision_id"], "slide_number": 1},
        )
        target = next(obj for obj in inventory["objects"] if obj["hyperlink"] is not None)

        def modify(output: dict[str, Any], _invocation: Any) -> None:
            output["operations"] = [
                {
                    "type": "set_text_hyperlink",
                    "params": {
                        "slide_id": inventory["slide_id"],
                        "object_id": target["object_id"],
                        "replacement_url": "https://example.com/pilot",
                        "replacement_slide_id": None,
                    },
                    "rationale": "Link the pilot brief",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = modify
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        bound = proposal["operations"][0]["params"]
        assert bound["expected_url"] is None and bound["expected_text"]
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        linked = next(
            obj
            for obj in call(
                core,
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": result["revision_id"], "slide_number": 1},
            )["objects"]
            if obj["object_id"] == target["object_id"]
        )
        assert linked["hyperlink"] == {"url": "https://example.com/pilot", "slide_id": None}
    finally:
        core.close()


def test_planning_snapshot_is_compact_and_uses_schema_units(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        captured: list[str] = []

        def modify(_output: dict[str, Any], invocation: Any) -> None:
            captured.append(invocation.request.current_slide_summary)

        provider.modifier = modify
        call(core, "presentation.workspace.propose_edit_plan", params)
        snapshot = json.loads(captured[0])["presentation_edit_snapshot"]
        objects = [obj for slide in snapshot["slides"] for obj in slide["objects"]]
        text = json.dumps(objects)
        assert '"fingerprint"' not in text and ": null" not in text and ": false" not in text
        rotated = [obj["rotation"] for obj in objects if "rotation" in obj]
        assert rotated and all("rotation_degrees" in item or item == {} for item in rotated)
    finally:
        core.close()


def test_provider_animations_see_sequence_bind_fingerprint_and_apply(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        scope = {key: params[key] for key in ("project_id", "document_id")}
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": params["expected_revision_id"], "slide_number": 1},
        )
        targets = [obj["object_id"] for obj in inventory["objects"]][:2]
        captured: list[dict[str, Any]] = []

        def modify(output: dict[str, Any], invocation: Any) -> None:
            snapshot = json.loads(invocation.request.current_slide_summary)
            captured.append(snapshot["presentation_edit_snapshot"]["slides"][0])
            output["operations"] = [
                {
                    "type": "set_slide_animations",
                    "params": {
                        "slide_id": inventory["slide_id"],
                        "effects": [
                            {"object_id": targets[0], "effect": "fade", "start": "on_click"},
                            {"object_id": targets[1], "effect": "wipe", "start": "after_previous"},
                        ],
                    },
                    "rationale": "Reveal the decision before its detail",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = modify
        proposal = call(core, "presentation.workspace.propose_edit_plan", params)
        assert captured[0]["animations_editable"] is True
        assert captured[0]["animations"] == []
        state = call(
            core,
            "presentation.workspace.get_slide_animations",
            {
                **scope,
                "expected_revision_id": params["expected_revision_id"],
                "slide_id": inventory["slide_id"],
            },
        )
        assert proposal["operations"][0]["params"]["expected_fingerprint"] == state["fingerprint"]
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        after = call(
            core,
            "presentation.workspace.get_slide_animations",
            {
                **scope,
                "expected_revision_id": result["revision_id"],
                "slide_id": inventory["slide_id"],
            },
        )
        assert [(e["object_id"], e["effect"], e["start"]) for e in after["effects"]] == [
            (targets[0], "fade", "on_click"),
            (targets[1], "wipe", "after_previous"),
        ]
    finally:
        core.close()


def test_provider_deck_theme_and_sections_need_whole_deck_and_apply(tmp_path: Path) -> None:
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        params = context(core)
        scope = {key: params[key] for key in ("project_id", "document_id")}
        sections = call(
            core,
            "presentation.workspace.get_slide_sections",
            {**scope, "expected_revision_id": params["expected_revision_id"]},
        )
        first, second = sections["slide_ids"]
        captured: list[dict[str, Any]] = []
        operations = [
            {
                "type": "set_theme_fonts",
                "params": {"heading_font": "Georgia", "body_font": "Segoe UI"},
            },
            {
                "type": "set_theme_colors",
                "params": {
                    "colors": {
                        slot: "#1F3A5F" if slot == "accent1" else None
                        for slot in PARAMETER_SCHEMAS["set_theme_colors"]["properties"]["colors"][
                            "properties"
                        ]
                    }
                },
            },
            {
                "type": "set_slide_sections",
                "params": {
                    "sections": [
                        {"section_id": None, "name": "Decision", "first_slide_id": first},
                        {"section_id": None, "name": "Evidence", "first_slide_id": second},
                    ]
                },
            },
        ]

        def modify(output: dict[str, Any], invocation: Any) -> None:
            captured.append(
                json.loads(invocation.request.current_slide_summary)["presentation_edit_snapshot"]
            )
            output["operations"] = [
                {**item, "rationale": "Restyle the deck", "evidence_ids": []} for item in operations
            ]

        provider.modifier = modify
        partial = call(core, "presentation.workspace.propose_edit_plan", params, ok=False)
        assert "entire presentation" in partial["message"]
        assert "theme_fonts" not in captured[0]
        proposal = call(
            core, "presentation.workspace.propose_edit_plan", {**params, "slide_numbers": [1, 2]}
        )
        assert {"theme_fonts", "theme_colors", "sections"} <= set(captured[1])
        bound = [item["params"] for item in proposal["operations"]]
        assert bound[1]["colors"] == {"accent1": "#1F3A5F"}
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        current = {**scope, "expected_revision_id": result["revision_id"]}
        fonts = call(core, "presentation.workspace.get_theme_fonts", current)
        assert (fonts["heading_font"], fonts["body_font"]) == ("Georgia", "Segoe UI")
        colors = call(core, "presentation.workspace.get_theme_colors", current)
        assert colors["colors"]["accent1"] == "#1F3A5F"
        after = call(core, "presentation.workspace.get_slide_sections", current)
        assert [(s["name"], s["slide_ids"]) for s in after["sections"]] == [
            ("Decision", [first]),
            ("Evidence", [second]),
        ]
    finally:
        core.close()


def test_provider_alt_text_sees_missing_flag_binds_and_applies(tmp_path: Path) -> None:
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    from test_presentation_workspace_source import _import_presentation, _open

    source = tmp_path / "alt.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Revenue"
    image = io.BytesIO()
    Image.new("RGB", (8, 8)).save(image, "PNG")
    image.seek(0)
    picture = slide.shapes.add_picture(image, Inches(1), Inches(2), Inches(2), Inches(2))
    picture.name = "Chart image"
    picture._element.xpath("./p:nvPicPr/p:cNvPr")[0].attrib.pop("descr", None)
    deck.save(source)
    provider = PlanProvider(locality="local")
    core = CoreService(
        data_root=tmp_path / "data",
        reasoning_provider=provider,
        embedding_adapter=DeterministicEmbeddingAdapter(dimension=2),
    )
    try:
        project = call(core, "project.create", {"name": "Alt text", "privacy_mode": "local_only"})[
            "project"
        ]["id"]
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}
        inventory = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": revision, "slide_number": 1},
        )
        target = next(o for o in inventory["objects"] if o["name"] == "Chart image")
        captured: list[dict[str, Any]] = []

        def modify(output: dict[str, Any], invocation: Any) -> None:
            snapshot = json.loads(invocation.request.current_slide_summary)
            captured.append(snapshot["presentation_edit_snapshot"]["slides"][0])
            output["operations"] = [
                {
                    "type": "set_alt_text",
                    "params": {
                        "slide_id": inventory["slide_id"],
                        "object_id": target["object_id"],
                        "alt_text": "Revenue rose 12% during the pilot",
                        "decorative": False,
                    },
                    "rationale": "Describe the chart for screen readers",
                    "evidence_ids": [],
                }
            ]

        provider.modifier = modify
        proposal = call(
            core,
            "presentation.workspace.propose_edit_plan",
            {
                **scope,
                "expected_revision_id": revision,
                "instruction": "Add alt text to the chart.",
                "slide_numbers": [1],
            },
        )
        flagged = [o for o in captured[0]["objects"] if o.get("alt_text_missing")]
        assert [o["object_id"] for o in flagged] == [target["object_id"]]
        bound = proposal["operations"][0]["params"]
        assert (bound["expected_alt_text"], bound["expected_decorative"]) == ("", False)
        call(
            core,
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": proposal["plan_id"]},
        )
        result = call(
            core,
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": proposal["plan_id"],
                "confirm": True,
                "reviewed_sha256": proposal["sha256"],
            },
        )
        after = call(
            core,
            "presentation.workspace.list_text_targets",
            {**scope, "revision_id": result["revision_id"], "slide_number": 1},
        )
        described = next(o for o in after["objects"] if o["name"] == "Chart image")
        assert described["alt_text"] == "Revenue rose 12% during the pilot"
    finally:
        core.close()
