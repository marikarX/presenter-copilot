"""Semantic graphic representations do not alter editing targets or native data."""

import hashlib
import uuid
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from presenter_core.presentation.fidelity import inspect_pptx_snapshot


def fixture(path: Path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(5), Inches(1)).text = "Pilot decision"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Revenue", [0.2, 0.4])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0), Inches(1), Inches(4), Inches(3), data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Pilot rates"
    table = slide.shapes.add_table(2, 2, Inches(4), Inches(1), Inches(4), Inches(2)).table
    for r, row in enumerate([["Metric", "Value"], ["Pilot", "20%"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    deck.save(path)
    return path


def read(path: Path):
    content = path.read_bytes()
    with path.open("rb") as stream:
        return inspect_pptx_snapshot(
            stream,
            document_id=str(uuid.uuid4()),
            expected_sha256=hashlib.sha256(content).hexdigest(),
            byte_size=len(content),
        )


def test_supported_graphic_semantics_are_separate_from_editing_targets(tmp_path: Path):
    path = fixture(tmp_path / "native.pptx")
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    document = read(path)
    inventory = document.text_targets(1)
    assert len(inventory["targets"]) == 1 and document.text_edit_candidate_count == 1
    by_kind = {target["source_kind"]: target for target in inventory["semantic_targets"]}
    assert set(by_kind) == {"native_text", "native_chart_cache", "native_table_cells"}
    assert by_kind["native_text"]["decision_text"] == "Pilot decision"
    assert (
        by_kind["native_chart_cache"]["decision_text"]
        == "Chart: Pilot rates\nRevenue: Q1: 0.2; Q2: 0.4"
    )
    assert by_kind["native_table_cells"]["decision_text"] == "Metric\tValue\nPilot\t20%"
    assert inventory["semantic_has_more"] is False
    assert inventory["semantic_excluded_objects"] == 0
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_semantic_paging_does_not_truncate_the_editing_inventory(tmp_path: Path):
    document = read(fixture(tmp_path / "paged.pptx"))
    inventory = document.text_targets(1, limit=1)
    assert inventory["has_more"] is False
    assert inventory["semantic_has_more"] is True
    assert len(inventory["targets"]) == len(inventory["semantic_targets"]) == 1


def test_external_chart_reference_has_no_semantic_cache_target(tmp_path: Path):
    path = fixture(tmp_path / "external.pptx")
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    name = "ppt/slides/_rels/slide1.xml.rels"
    root = etree.fromstring(parts[name])
    for relation in root:
        if relation.get("Type", "").endswith("/chart"):
            relation.set("TargetMode", "External")
            relation.set("Target", "https://invalid.example/chart.xml")
    parts[name] = etree.tostring(root)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    inventory = read(path).text_targets(1)
    assert {target["source_kind"] for target in inventory["semantic_targets"]} == {
        "native_text",
        "native_table_cells",
    }
    assert inventory["semantic_excluded_objects"] == 1
