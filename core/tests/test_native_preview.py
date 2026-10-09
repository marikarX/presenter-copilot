from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from presenter_core.errors import CoreDomainError
from presenter_core.presentation.native_preview import export_native_preview, run_preview_worker


def _job(tmp_path: Path) -> dict[str, Any]:
    source = tmp_path / "revision.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(12.8), Inches(7.2)
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    return {
        "snapshot_path": str(source),
        "output_path": str(tmp_path / "slide.png"),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "slide_number": 2,
    }


class Office:
    def __init__(self, *, fail: str = "") -> None:
        self.fail = fail
        self.closed = False
        self.exported: list[int] = []
        self.Presentations = SimpleNamespace(Open=self.open)
        self.Slides = SimpleNamespace(Count=2, Item=self.slide)
        self.PageSetup = SimpleNamespace(SlideWidth=921.6, SlideHeight=518.4)

    def __setattr__(self, key: str, value: Any) -> None:
        if key in {"AutomationSecurity", "DisplayAlerts", "WindowState", "Visible"}:
            raise AssertionError("Do not modify shared application state")
        object.__setattr__(self, key, value)

    def open(self, _: str, **kwargs: Any) -> Office:
        assert kwargs == {"ReadOnly": True, "Untitled": False, "WithWindow": False}
        return self

    def slide(self, number: int) -> Any:
        def export(path: str, format: str, width: int, height: int) -> None:
            self.exported.append(number)
            assert format == "PNG"
            if self.fail == "export":
                raise RuntimeError("token/private-path should not escape")
            Image.new("RGB", (width + (1 if self.fail == "size" else 0), height), "blue").save(path)

        return SimpleNamespace(Export=export)

    def Close(self) -> None:
        self.closed = True

    def Quit(self) -> None:
        raise AssertionError("Never quit a shared Office application")


def test_hidden_slide_index_and_source_ownership(tmp_path: Path) -> None:
    job = _job(tmp_path)
    before = Path(job["snapshot_path"]).read_bytes()
    office = Office()
    result = export_native_preview(job, application=office)
    assert result["slide_number"] == 2 and result["slide_count"] == 2
    assert (result["width"], result["height"]) == (960, 540)
    assert result["sha256"] == hashlib.sha256(Path(job["output_path"]).read_bytes()).hexdigest()
    assert office.closed and office.exported == [2]
    assert Path(job["snapshot_path"]).read_bytes() == before


@pytest.mark.parametrize(
    "delta, accepted", [(0.025, True), (0.0625, True), (0.063, False), (1, False)]
)
def test_bounded_native_page_size_rounding(tmp_path: Path, delta: float, accepted: bool) -> None:
    office = Office()
    office.PageSetup.SlideWidth += delta
    if accepted:
        assert export_native_preview(_job(tmp_path), application=office)["width"] == 960
    else:
        with pytest.raises(CoreDomainError, match="geometry"):
            export_native_preview(_job(tmp_path), application=office)
        assert not office.exported
    assert office.closed


@pytest.mark.parametrize("fail", ["export", "size"])
def test_closes_owned_deck_on_render_and_image_failure(tmp_path: Path, fail: str) -> None:
    office = Office(fail=fail)
    with pytest.raises(RuntimeError if fail == "export" else CoreDomainError):
        export_native_preview(_job(tmp_path), application=office)
    assert office.closed


@pytest.mark.parametrize("number", [True, 0, -1, 3, 5001, 1.5])
def test_invalid_index_never_opens_office(tmp_path: Path, number: Any) -> None:
    job = _job(tmp_path)
    job["slide_number"] = number
    office = Office()
    with pytest.raises(CoreDomainError):
        export_native_preview(job, application=office)
    assert not office.closed and not office.exported


def test_worker_bounds_and_sanitizes_failures(monkeypatch: Any) -> None:
    from presenter_core.presentation import native_preview

    def fail(_: Any) -> Any:
        raise RuntimeError("token/private-path")

    monkeypatch.setattr(native_preview, "export_native_preview", fail)
    output = io.StringIO()
    assert run_preview_worker(io.StringIO("{}"), output) == 1
    assert "token" not in output.getvalue() and "private-path" not in output.getvalue()
    assert json.loads(output.getvalue())["ok"] is False
    output = io.StringIO()
    assert run_preview_worker(io.StringIO(" " * 8193), output) == 1
    assert json.loads(output.getvalue())["error"]["code"] == "INVALID_REQUEST"
