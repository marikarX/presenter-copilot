from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from presenter_core.errors import CoreDomainError
from presenter_core.presentation import native_layers
from presenter_core.presentation.native_layers import (
    export_native_layers,
    recover_native_alpha,
    verify_native_composition,
)


def test_native_mattes_recover_alpha_and_color_on_arbitrary_background() -> None:
    foreground = Image.new("RGBA", (5, 3), (210, 70, 25, 165))
    black = Image.alpha_composite(Image.new("RGBA", foreground.size, "black"), foreground)
    white = Image.alpha_composite(Image.new("RGBA", foreground.size, "white"), foreground)
    layer = recover_native_alpha(black, white)
    for color in [(226, 237, 248), (20, 50, 90), (255, 20, 0)]:
        background = Image.new("RGBA", foreground.size, color)
        expected = Image.alpha_composite(background, foreground)
        assert verify_native_composition(expected, background, [layer]) <= 2


@pytest.mark.parametrize("alpha", [0, 255])
def test_empty_and_opaque_layers(alpha: int) -> None:
    foreground = Image.new("RGBA", (2, 2), (20, 70, 150, alpha))
    black = Image.alpha_composite(Image.new("RGBA", (2, 2), "black"), foreground)
    white = Image.alpha_composite(Image.new("RGBA", (2, 2), "white"), foreground)
    recovered = recover_native_alpha(black, white)
    assert recovered.getchannel("A").getextrema() == (alpha, alpha)
    if alpha:
        assert recovered.getpixel((0, 0)) == (20, 70, 150, 255)


def test_native_paint_order_and_overlapping_transparency_are_verified() -> None:
    background = Image.new("RGBA", (3, 3), "white")
    red = Image.new("RGBA", (3, 3), (220, 20, 20, 165))
    green = Image.new("RGBA", (3, 3), (20, 180, 20, 255))
    baseline = Image.alpha_composite(Image.alpha_composite(background, red), green)
    assert verify_native_composition(baseline, background, [red, green]) == 0
    with pytest.raises(CoreDomainError, match="reproduce"):
        verify_native_composition(baseline, background, [green, red])


def test_one_bad_pixel_cannot_hide_in_a_low_average_difference() -> None:
    baseline = Image.new("RGBA", (100, 100), "white")
    background = baseline.copy()
    background.putpixel((50, 50), (250, 255, 255, 255))
    with pytest.raises(CoreDomainError, match="reproduce"):
        verify_native_composition(baseline, background, [])


def test_dimension_mismatch_and_background_dependent_colors_are_rejected() -> None:
    with pytest.raises(CoreDomainError, match="dimensions"):
        recover_native_alpha(Image.new("RGB", (2, 2)), Image.new("RGB", (3, 2)))
    with pytest.raises(CoreDomainError, match="consistent"):
        recover_native_alpha(
            Image.new("RGB", (2, 2), (100, 10, 10)), Image.new("RGB", (2, 2), (50, 250, 250))
        )
    with pytest.raises(CoreDomainError, match="coordinates"):
        verify_native_composition(Image.new("RGB", (2, 2)), Image.new("RGB", (3, 2)), [])


def job(tmp_path: Path) -> dict[str, Any]:
    source = tmp_path / "revision.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    return {
        "snapshot_path": str(source),
        "output_path": str(tmp_path / "layers.json"),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "slide_number": 1,
        "document_id": "00000000-0000-4000-8000-000000000001",
    }


class Office:
    def __init__(self, fail: bool = False) -> None:
        self.closed = False
        self.fail = fail
        self.Presentations = SimpleNamespace(Open=self.open)
        self.Slides = SimpleNamespace(Count=1, Item=lambda _: self)
        self.PageSetup = SimpleNamespace(SlideWidth=720, SlideHeight=405)
        self.Shapes = SimpleNamespace(Count=0)
        self.Background = SimpleNamespace(Fill=SimpleNamespace(Solid=lambda: None))

    def __setattr__(self, key: str, value: Any) -> None:
        if key in {"AutomationSecurity", "DisplayAlerts", "WindowState", "Visible"}:
            raise AssertionError("Do not mutate shared Office settings")
        object.__setattr__(self, key, value)

    def open(self, _: str, **kwargs: Any) -> Office:
        assert kwargs == {"ReadOnly": True, "Untitled": False, "WithWindow": False}
        return self

    def Export(self, filename: str, format: str, width: int, height: int) -> None:
        assert format == "PNG"
        if self.fail:
            raise RuntimeError("private-token-and-path")
        Image.new("RGB", (width, height), "blue").save(filename)

    def Close(self) -> None:
        self.closed = True

    def Quit(self) -> None:
        raise AssertionError("Never quit shared Office")


def test_owned_native_deck_closes_and_source_bytes_and_receipts_are_preserved(
    tmp_path: Path,
) -> None:
    request = job(tmp_path)
    before = Path(request["snapshot_path"]).read_bytes()
    office = Office()
    result = export_native_layers(request, application=office)
    assert office.closed and result["layers"] == []
    assert result["baseline_max_channel_error"] == 0
    assert (result["width"], result["height"]) == (960, 540)
    assert Path(request["snapshot_path"]).read_bytes() == before
    assert json.loads(Path(request["output_path"]).read_text(encoding="utf-8")) == result
    background = Path(request["output_path"]).parent / result["background"]["file"]
    assert hashlib.sha256(background.read_bytes()).hexdigest() == result["background"]["sha256"]


def test_native_render_failure_closes_the_owned_deck(tmp_path: Path) -> None:
    office = Office(fail=True)
    with pytest.raises(RuntimeError):
        export_native_layers(job(tmp_path), application=office)
    assert office.closed


@pytest.mark.parametrize(
    "field,value",
    [
        ("slide_number", True),
        ("slide_number", 0),
        ("slide_number", 2),
        ("document_id", "bad"),
        ("source_sha256", "x" * 64),
        ("extra", True),
    ],
)
def test_invalid_jobs_never_open_office(tmp_path: Path, field: str, value: Any) -> None:
    request = job(tmp_path)
    request[field] = value
    with pytest.raises(CoreDomainError):
        export_native_layers(request, application=object())


def test_existing_or_outside_outputs_are_rejected(tmp_path: Path) -> None:
    request = job(tmp_path)
    Path(request["output_path"]).write_text("owned receipt", encoding="utf-8")
    with pytest.raises(CoreDomainError):
        export_native_layers(request, application=object())
    request["output_path"] = str(tmp_path.parent / "outside-layers.json")
    with pytest.raises(CoreDomainError):
        export_native_layers(request, application=object())


def test_worker_errors_do_not_disclose_raw_office_failures(monkeypatch: pytest.MonkeyPatch) -> None:
    def fail(_: dict[str, Any]) -> dict[str, Any]:
        raise RuntimeError("secret-token-and-private-path")

    monkeypatch.setattr(native_layers, "export_native_layers", fail)
    output = io.StringIO()
    assert native_layers.run_layers_worker(io.StringIO("{}"), output) == 1
    response = json.loads(output.getvalue())
    assert response["error"]["code"] == "PRESENTATION_NATIVE_LAYERS_FAILED"
    assert "secret" not in output.getvalue()


@pytest.mark.parametrize("count,height", [(201, 5.625), (100, 20)])
def test_object_and_decoded_memory_budgets_reject_before_office(
    tmp_path: Path, count: int, height: float
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    request = job(tmp_path)
    source = Path(request["snapshot_path"])
    deck = Presentation(source)
    deck.slide_height = Inches(height)
    for _ in range(count):
        deck.slides[0].shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1)
        )
    deck.save(source)
    request["source_sha256"] = hashlib.sha256(source.read_bytes()).hexdigest()
    with pytest.raises(CoreDomainError, match="limit|budget"):
        export_native_layers(request, application=object())
