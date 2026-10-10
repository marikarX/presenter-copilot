"""Shapes inside a top-level group move within the group's frame."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService


def test_group_child_moves_inside_the_frame_only(tmp_path: Path) -> None:
    source = tmp_path / "group.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    first = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    first.name = "Inner A"
    group.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(3), Inches(2), Inches(2)).name = "B"
    deck.save(source)
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, source)["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def objects() -> tuple[str, dict]:
            inventory = _call(
                core,
                "t",
                "presentation.workspace.list_text_targets",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]
            return inventory["slide_id"], {o["name"]: o for o in inventory["objects"]}

        slide_id, found = objects()
        inner = found["Inner A"]
        frame = next(o for o in found.values() if o["kind"] == "grpSp")["geometry"]
        assert inner["group_child"] and inner["geometry_editable"]
        assert not inner["delete_editable"] and not inner["group_editable"]
        assert inner["geometry"] == {
            "x_emu": Inches(1),
            "y_emu": Inches(1),
            "width_emu": Inches(2),
            "height_emu": Inches(1),
        }

        def move(geometry: dict) -> dict:
            return _call(
                core,
                "m",
                "presentation.workspace.transform_object",
                {
                    **scope,
                    "expected_revision_id": revision,
                    "slide_id": slide_id,
                    "object_id": inner["object_id"],
                    "expected_geometry": inner["geometry"],
                    "replacement_geometry": geometry,
                },
            )

        outside = move({**inner["geometry"], "x_emu": int(Inches(5))})
        assert not outside["ok"] and "group" in outside["error"]["message"]
        inside = {**inner["geometry"], "x_emu": int(Inches(2)), "width_emu": int(Inches(1))}
        reply = move(inside)
        assert reply["ok"], reply.get("error")
        revision = reply["result"]["revision_id"]
        _, after = objects()
        assert after["Inner A"]["geometry"] == inside
        assert next(o for o in after.values() if o["kind"] == "grpSp")["geometry"] == frame
    finally:
        core.close()
