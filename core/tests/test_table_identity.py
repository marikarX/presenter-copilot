"""Tables carrying PowerPoint's row/column identity extensions stay editable."""

import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from test_presentation_workspace_source import _call, _import_presentation, _open, _project_id

from presenter_core.ipc.core import CoreService

ROW = "{0D108BD9-81ED-4DB2-BD59-A6C34878D82A}"
COLUMN = "{9D8B030D-6E8A-4147-A177-3AD203B41FA5}"
A16 = "http://schemas.microsoft.com/office/drawing/2014/main"


def _deck(tmp_path: Path) -> Path:
    source = tmp_path / "plain.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = f"R{r}C{c}"
    deck.save(source)
    tagged = tmp_path / "tagged.pptx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(tagged, "w") as output:
        for info in original.infolist():
            data = original.read(info)
            if info.filename == "ppt/slides/slide1.xml":
                xml = data.decode()
                counter = iter(range(1000, 2000))
                xml = re.sub(
                    r"<a:gridCol w=\"(\d+)\"/>",
                    lambda m, ids=counter: f'<a:gridCol w="{m.group(1)}"><a:extLst>'
                    f'<a:ext uri="{COLUMN}"><a16:colId xmlns:a16="{A16}" val="{next(ids)}"/>'
                    "</a:ext></a:extLst></a:gridCol>",
                    xml,
                )
                xml = xml.replace(
                    "</a:tr>",
                    f'<a:extLst><a:ext uri="{ROW}"><a16:rowId xmlns:a16="{A16}" val="ROWID"/>'
                    "</a:ext></a:extLst></a:tr>",
                )
                for value in (3000, 3001):
                    xml = xml.replace('val="ROWID"', f'val="{value}"', 1)
                xml = re.sub(
                    r'(<p:cNvPr id="\d+" name="Table[^"]*")\s*/>',
                    r'\1><a:extLst><a:ext uri="{FF2B5EF4-FFF2-40B4-BE49-F238E27FC236}">'
                    rf'<a16:creationId xmlns:a16="{A16}" '
                    r'id="{AAAAAAAA-0000-0000-0000-000000000001}"/>'
                    r"</a:ext></a:extLst></p:cNvPr>",
                    xml,
                )
                assert "colId" in xml and "rowId" in xml and "creationId" in xml
                data = xml.encode()
            output.writestr(info, data)
    return tagged


def test_identity_tagged_table_edits_and_new_rows_get_fresh_ids(tmp_path: Path) -> None:
    core = CoreService(data_root=tmp_path / "data")
    try:
        project = _project_id(core)
        document = _import_presentation(core, project, _deck(tmp_path))["id"]
        revision = _open(core, project, document)["revision_id"]
        scope = {"project_id": project, "document_id": document}

        def table() -> tuple[str, dict]:
            result = _call(
                core,
                "l",
                "presentation.workspace.list_tables",
                {**scope, "revision_id": revision, "slide_number": 1},
            )["result"]
            return result["slide_id"], result["items"][0]

        slide_id, item = table()
        assert item["editable"] and item["structure_editable"], item.get("reason")
        edited = _call(
            core,
            "e",
            "presentation.workspace.replace_table_cells",
            {
                **scope,
                "expected_revision_id": revision,
                "slide_id": slide_id,
                "object_id": item["object_id"],
                "expected_fingerprint": item["fingerprint"],
                "cells": [["Header", "R0C1"], ["R1C0", "R1C1"]],
            },
        )
        assert edited["ok"], edited
        revision = edited["result"]["revision_id"]
        slide_id, item = table()
        plan = _call(
            core,
            "p",
            "presentation.workspace.prepare_edit_plan",
            {
                **scope,
                "expected_revision_id": revision,
                "summary": "Add a row",
                "operations": [
                    {
                        "type": "reshape_table",
                        "params": {
                            "slide_id": slide_id,
                            "object_id": item["object_id"],
                            "expected_fingerprint": item["fingerprint"],
                            "cells": [*item["cells"], ["New", ""]],
                            "row_sources": [0, 1, None],
                            "column_sources": [0, 1],
                        },
                    }
                ],
            },
        )
        assert plan["ok"], plan
        _call(
            core,
            "o",
            "presentation.workspace.open_edit_plan_preview",
            {**scope, "plan_id": plan["result"]["plan_id"]},
        )
        applied = _call(
            core,
            "a",
            "presentation.workspace.apply_edit_plan",
            {
                **scope,
                "plan_id": plan["result"]["plan_id"],
                "reviewed_sha256": plan["result"]["sha256"],
                "confirm": True,
            },
        )
        assert applied["ok"], applied
        revision = applied["result"]["revision_id"]
        with core._storage.project_database(project) as connection:
            path = connection.execute(
                "SELECT local_snapshot_path FROM presentation_revisions WHERE id=?", (revision,)
            ).fetchone()["local_snapshot_path"]
        with zipfile.ZipFile(core._storage.paths.snapshot_path(project, path)) as archive:
            xml = archive.read("ppt/slides/slide1.xml").decode()
        rows = re.findall(r'rowId[^>]*val="(\d+)"', xml)
        assert rows[:2] == ["3000", "3001"] and len(set(rows)) == 3
        assert re.findall(r'colId[^>]*val="(\d+)"', xml) == ["1000", "1001"]
        assert table()[1]["cells"][2] == ["New", ""]
    finally:
        core.close()
