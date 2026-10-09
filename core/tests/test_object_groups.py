from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from test_object_paste import inventory
from test_presentation_native_layouts import path
from test_presentation_transactions import call, context
from test_selection_actions import save

from presenter_core.errors import CoreDomainError
from presenter_core.ingestion.parsers.pptx import PptxParser
from presenter_core.ipc.core import CoreService
from presenter_core.presentation.fidelity import _NS, inspect_pptx_snapshot
from presenter_core.retrieval.embeddings import DeterministicEmbeddingAdapter


def operation(core, params, kind, values):
    return {
        **params,
        "summary": "Native group review",
        "operations": [{"type": kind, "params": values}],
    }


def group_request(core, params, objects):
    return operation(
        core,
        params,
        "group_objects",
        {
            "slide_id": inventory(core, params, 1)["slide_id"],
            "objects": [
                {"object_id": obj["object_id"], "expected_geometry": obj["geometry"]}
                for obj in objects
            ],
        },
    )


def prepare(core, request):
    return call(core, "presentation.workspace.prepare_edit_plan", request)


def ungroup_request(core, params, group):
    return operation(
        core,
        params,
        "ungroup_object",
        {
            "slide_id": inventory(core, params, 1)["slide_id"],
            "object_id": group["object_id"],
            "expected_geometry": group["geometry"],
        },
    )


def mixed_fixture(core, directory: Path, picture_format="PNG"):
    image = directory / f"picture.{picture_format.lower()}"
    Image.new("RGB", (80, 50), "purple").save(image, format=picture_format)
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[6])
    for number in range(2):
        text = first.shapes.add_textbox(
            Inches(0.5 + number * 2), Inches(0.5), Inches(1.8), Inches(0.5)
        )
        text.text = f"Decision evidence {number}"
    picture = first.shapes.add_picture(str(image), Inches(0.5), Inches(2), width=Inches(1.5))
    picture.crop_left = 0.1
    values = CategoryChartData()
    values.categories = ["Baseline", "Pilot"]
    values.add_series("Revenue", [10, 12])
    frame = first.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3), Inches(2), Inches(5), Inches(3), values
    )
    frame._element.find("p:nvGraphicFramePr/p:cNvGraphicFramePr/a:graphicFrameLocks", _NS).set(
        "noGrp", "0"
    )
    first.notes_slide.notes_text_frame.text = "Preserved group notes"
    other = deck.slides.add_slide(deck.slide_layouts[6])
    other.shapes.add_textbox(
        Inches(1), Inches(1), Inches(2), Inches(1)
    ).text = "Other slide unchanged"
    source = directory / "mixed.pptx"
    deck.save(source)
    project = call(core, "project.create", {"name": "Grouping acceptance"})["project"]["id"]
    document = call(
        core, "source.import", {"project_id": project, "path": str(source), "kind": "presentation"}
    )["document"]["id"]
    current = core._presentation_edits._current_revision_source(project, document)
    params = {
        "project_id": project,
        "document_id": document,
        "expected_revision_id": current["revision_id"],
    }
    return params, inventory(core, params, 1)


@pytest.mark.parametrize("nested", [False, True])
@pytest.mark.parametrize("picture_format", ["PNG", "JPEG"])
def test_native_group_preserves_parts_identities_retrieval_and_history(
    tmp_path: Path, nested, picture_format
):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first = mixed_fixture(core, tmp_path, picture_format)
        original = path(core, params).read_bytes()
        original_revision = params["expected_revision_id"]
        objects = [first["objects"][0], first["objects"][2], first["objects"][3]]
        assert all(obj["group_editable"] for obj in objects)
        result = save(core, params, prepare(core, group_request(core, params, objects)))
        group = next(obj for obj in inventory(core, params, 1)["objects"] if obj["kind"] == "grpSp")
        assert result["operations"][0]["selected_object_ids"] == [group["object_id"]]
        assert group["geometry_editable"] and group["group_editable"]
        grouped = path(core, params).read_bytes()
        with (
            zipfile.ZipFile(io.BytesIO(original)) as before,
            zipfile.ZipFile(io.BytesIO(grouped)) as after,
        ):
            assert before.namelist() == after.namelist()
            assert [
                name for name in before.namelist() if before.read(name) != after.read(name)
            ] == ["ppt/slides/slide1.xml"]
            root = etree.fromstring(after.read("ppt/slides/slide1.xml"))
            children = root.findall("p:cSld/p:spTree/p:grpSp/*", _NS)
            native_ids = [
                node.xpath("./*/p:cNvPr/@id", namespaces=_NS)[0]
                for node in children
                if node.tag.endswith(("}sp", "}pic", "}graphicFrame"))
            ]
            assert native_ids == [obj["ooxml_shape_id"] for obj in objects]
        grouped_targets = inventory(core, params, 1)
        with path(core, params).open("rb") as handle:
            fidelity = inspect_pptx_snapshot(
                handle,
                document_id=params["document_id"],
                expected_sha256=result["sha256"],
                byte_size=result["byte_size"],
            )
        assert any("Revenue" in obj.text for obj in fidelity.slides[0].objects)
        assert all(
            not obj["geometry_editable"]
            for obj in grouped_targets["objects"]
            if obj["object_id"] in {o["object_id"] for o in objects}
        )
        parsed = PptxParser().parse(path(core, params))
        assert "Decision evidence 0" in parsed[0].text and "Revenue" in parsed[0].text
        call(
            core,
            "source.reindex",
            {"project_id": params["project_id"], "document_id": params["document_id"]},
        )
        if nested:
            remaining = next(
                obj
                for obj in grouped_targets["objects"]
                if obj["name"] == first["objects"][1]["name"]
            )
            save(core, params, prepare(core, group_request(core, params, [group, remaining])))
            outer = next(
                obj
                for obj in inventory(core, params, 1)["objects"]
                if obj["kind"] == "grpSp" and obj["geometry_editable"]
            )
            save(core, params, prepare(core, ungroup_request(core, params, outer)))
            group = next(
                obj
                for obj in inventory(core, params, 1)["objects"]
                if obj["object_id"] == group["object_id"]
            )
        old = group["geometry"]
        replacement = {
            **old,
            "x_emu": old["x_emu"] + 91440,
            "y_emu": old["y_emu"] + 182880,
            "width_emu": old["width_emu"] * 3 // 5,
            "height_emu": old["height_emu"] * 3 // 4,
        }
        save(
            core,
            params,
            prepare(
                core,
                operation(
                    core,
                    params,
                    "transform_object",
                    {
                        "slide_id": first["slide_id"],
                        "object_id": group["object_id"],
                        "expected_geometry": old,
                        "replacement_geometry": replacement,
                    },
                ),
            ),
        )
        group["geometry"] = replacement
        saved = save(core, params, prepare(core, ungroup_request(core, params, group)))
        assert set(saved["operations"][0]["selected_object_ids"]) == {
            obj["object_id"] for obj in objects
        }
        fresh = inventory(core, params, 1)
        assert {obj["object_id"] for obj in fresh["objects"]} == {
            obj["object_id"] for obj in first["objects"]
        }
        assert all(obj["geometry_editable"] and obj["group_editable"] for obj in fresh["objects"])
        history = call(
            core,
            "presentation.workspace.list_revisions",
            {"project_id": params["project_id"], "document_id": params["document_id"]},
        )
        for _ in range(len(history["revisions"]) - 1):
            undone = call(
                core, "presentation.workspace.step_revision", {**params, "direction": "undo"}
            )
            params["expected_revision_id"] = undone["revision_id"]
        assert (
            params["expected_revision_id"] == original_revision
            and path(core, params).read_bytes() == original
        )
        redone = call(core, "presentation.workspace.step_revision", {**params, "direction": "redo"})
        params["expected_revision_id"] = redone["revision_id"]
        assert path(core, params).read_bytes() == grouped
    finally:
        core.close()


@pytest.mark.parametrize(
    "failure",
    [
        "duplicate",
        "missing",
        "stale_geometry",
        "stale_revision",
        "placeholder",
        "table",
        "cancel",
        "late",
    ],
)
def test_group_rejection_and_discard_are_atomic(tmp_path: Path, failure):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, _ = context(core)
        if failure == "table":
            from test_presentation_tables import table_request

            first = inventory(core, params, 1)
            request = table_request(params, [{"params": {"slide_id": first["slide_id"]}}])
            created = call(core, "presentation.workspace.create_table", request)
            params["expected_revision_id"] = created["revision_id"]
        first = inventory(core, params, 1)
        objects = [obj for obj in first["objects"] if obj["group_editable"]][:2]
        request = group_request(core, params, objects)
        selected = request["operations"][0]["params"]["objects"]
        if failure == "duplicate":
            selected[1] = selected[0]
        if failure == "missing":
            selected[1]["object_id"] = "11111111-1111-4111-8111-111111111111"
        if failure == "stale_geometry":
            selected[1]["expected_geometry"] = {**selected[1]["expected_geometry"], "x_emu": 1}
        if failure == "stale_revision":
            request["expected_revision_id"] = "11111111-1111-4111-8111-111111111111"
        if failure in {"placeholder", "table"}:
            obj = next(
                obj
                for obj in first["objects"]
                if not obj["group_editable"]
                and (
                    obj["kind"] == "graphicFrame"
                    if failure == "table"
                    else obj["name"] == "Slide title"
                )
            )
            selected[1] = {"object_id": obj["object_id"], "expected_geometry": obj["geometry"]}
        if failure == "late":
            request["operations"].append(
                {
                    "type": "delete_object",
                    "params": {
                        "slide_id": first["slide_id"],
                        "object_id": objects[0]["object_id"],
                        "expected_geometry": objects[0]["geometry"],
                    },
                }
            )
        before = path(core, params).read_bytes()
        if failure == "cancel":
            proposed = prepare(core, request)
            call(
                core,
                "presentation.workspace.reject_edit_plan",
                {
                    "project_id": params["project_id"],
                    "document_id": params["document_id"],
                    "plan_id": proposed["plan_id"],
                },
            )
        else:
            with pytest.raises(CoreDomainError) as caught:
                core._presentation_edit_plans.prepare(request)
            assert caught.value.code in {
                "INVALID_REQUEST",
                "PRESENTATION_EDIT_TARGET_NOT_FOUND",
                "PRESENTATION_EDIT_STALE_TARGET",
                "PRESENTATION_REVISION_CONFLICT",
                "PRESENTATION_EDIT_UNSUPPORTED",
            }
        assert path(core, params).read_bytes() == before
        assert (
            core._presentation_edits._current_revision_source(
                params["project_id"], params["document_id"]
            )["revision_id"]
            == params["expected_revision_id"]
        )
        assert not list(path(core, params).parent.glob(".edit-stage-*.pptx"))
    finally:
        core.close()


def test_gif_grouping_is_explicitly_unsupported(tmp_path: Path):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        _, first = mixed_fixture(core, tmp_path, "GIF")
        picture = next(obj for obj in first["objects"] if obj["kind"] == "pic")
        assert not picture["group_editable"] and "PNG and JPEG" in picture["group_reason"]
    finally:
        core.close()


@pytest.mark.parametrize(
    "unsupported",
    [
        "rotation",
        "flip",
        "zero_child_extent",
        "effect",
        "hidden",
        "lock",
        "animation",
        "child_outside",
        "unknown_child",
    ],
)
def test_imported_group_dependencies_stay_read_only(tmp_path: Path, unsupported):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        params, first = mixed_fixture(core, tmp_path)
        save(core, params, prepare(core, group_request(core, params, first["objects"][:2])))
        with zipfile.ZipFile(path(core, params)) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        root = etree.fromstring(parts["ppt/slides/slide1.xml"])
        group = root.find("p:cSld/p:spTree/p:grpSp", _NS)
        transform = group.find("p:grpSpPr/a:xfrm", _NS)
        if unsupported == "unknown_child":
            etree.SubElement(group, "{urn:unsupported-test}opaque")
        if unsupported == "rotation":
            transform.set("rot", "5400000")
        if unsupported == "flip":
            transform.set("flipH", "1")
        if unsupported == "zero_child_extent":
            transform.find("a:chExt", _NS).set("cx", "0")
        if unsupported == "effect":
            etree.SubElement(group.find("p:grpSpPr", _NS), "{" + _NS["a"] + "}effectLst")
        if unsupported == "hidden":
            group.find("p:nvGrpSpPr/p:cNvPr", _NS).set("hidden", "1")
        if unsupported == "lock":
            etree.SubElement(
                group.find("p:nvGrpSpPr/p:cNvGrpSpPr", _NS),
                "{" + _NS["a"] + "}grpSpLocks",
                noUngrp="1",
            )
        if unsupported == "child_outside":
            group.find("p:sp/p:spPr/a:xfrm/a:off", _NS).set("x", "-1")
        if unsupported == "animation":
            timing = etree.SubElement(root, "{" + _NS["p"] + "}timing")
            etree.SubElement(
                timing, "{" + _NS["p"] + "}spTgt", spid=first["objects"][0]["ooxml_shape_id"]
            )
        parts["ppt/slides/slide1.xml"] = etree.tostring(root)
        source = tmp_path / "unsupported-group.pptx"
        with zipfile.ZipFile(source, "w") as archive:
            for name, value in parts.items():
                archive.writestr(name, value)
        imported = call(
            core,
            "source.import",
            {"project_id": params["project_id"], "path": str(source), "kind": "presentation"},
        )
        params["document_id"] = imported["document"]["id"]
        params["expected_revision_id"] = core._presentation_edits._current_revision_source(
            params["project_id"], params["document_id"]
        )["revision_id"]
        obj = next(obj for obj in inventory(core, params, 1)["objects"] if obj["kind"] == "grpSp")
        assert not obj["geometry_editable"] and not obj["group_editable"] and obj["group_reason"]
        original = path(core, params).read_bytes()
        with pytest.raises(CoreDomainError) as caught:
            core._presentation_edit_plans.prepare(ungroup_request(core, params, obj))
        assert caught.value.code == "PRESENTATION_EDIT_UNSUPPORTED"
        assert path(core, params).read_bytes() == original
    finally:
        core.close()


@pytest.mark.parametrize("kind", ["imported", "authored"])
def test_chart_grouping_respects_import_locks_and_authored_charts_can_group(tmp_path: Path, kind):
    core = CoreService(
        data_root=tmp_path / "data", embedding_adapter=DeterministicEmbeddingAdapter(dimension=2)
    )
    try:
        if kind == "imported":
            from test_presentation_chart_editing import imported

            params, _ = imported(core, tmp_path)
            chart = next(
                obj
                for obj in inventory(core, params, 1)["objects"]
                if obj["kind"] == "graphicFrame"
            )
            assert not chart["group_editable"] and "locks" in chart["group_reason"]
        else:
            from test_presentation_charts import request

            params, ops = context(core)
            created = call(core, "presentation.workspace.create_chart", request(params, ops))
            params["expected_revision_id"] = created["revision_id"]
            first = inventory(core, params, 1)
            chart = next(obj for obj in first["objects"] if obj["kind"] == "graphicFrame")
            assert chart["group_editable"]
            shape = next(
                obj for obj in first["objects"] if obj["kind"] == "sp" and obj["group_editable"]
            )
            save(core, params, prepare(core, group_request(core, params, [shape, chart])))
    finally:
        core.close()
